"""2.2.2: girilen her alan detayda ve Excel çıktısında."""
import io
import openpyxl
from test_sonqa_kayitlar import web, login, insert, XHR  # noqa: F401  (ortak fixture/yardımcılar)
from app import db


def test_detay_ve_cikti_tum_alanlar(web):
    c, a = web; login(c)
    kid = insert(a, personel='P-1', uye_id='U-1', ana_uye_id='A-1', tip='sessiz', lead_id='L-9', link='https://i.gyazo.com/abcdefabcdefabcdefabcdefabcdefab.png', gorsel2='https://i.gyazo.com/00000000000000000000000000000000.png', aciklama='Açıklama metni', ek_not='Not metni', ekleyen_id='777', ekleyen_ad='Ali', kaynak='telegram', giris_tipi='resim')
    conn = db.baglan(a['DB'])
    with db.transaction(conn):
        conn.execute("INSERT INTO yorumlar(kayit_id,kullanici,metin,zaman,tur) VALUES(?,?,?,?,?)", (kid, 'admin', 'ilk', '2026-10-19 10:00:00', 'yorum'))
        conn.execute("INSERT INTO yorumlar(kayit_id,kullanici,metin,zaman,tur) VALUES(?,?,?,?,?)", (kid, 'ayse', 'son yorum', '2026-10-19 10:05:00', 'yorum'))
    conn.close()
    d = c.get(f'/kayit/{kid}', headers=XHR).json
    assert d['giris_tipi_metni'] == 'Görsel' and d['kaynak_metni'] == 'Telegram' and d['dis_link2'].startswith('https://i.gyazo.com/') and d['tip'] == 'sessiz'
    r = c.get('/disa-aktar.xlsx?durum=hepsi')
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.data))['Kayıtlar']
    hdr = [x.value for x in ws[1]]
    for beklenen in ('Sessiz / Telesekreter', 'Lead ID', 'Durum', 'Onay zamanı', 'Ekleyen Telegram ID', 'Giriş tipi', 'Kaynak', 'Görsel bağlantısı', '2. görsel bağlantısı', 'Açıklama', 'Ek not', 'Yorum sayısı', 'Son yorum', 'Panel bağlantısı'):
        assert beklenen in hdr, beklenen
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr.index('Kayıt ID')] == kid)
    satir = dict(zip(hdr, row))
    assert satir['Sessiz / Telesekreter'] == 'sessiz' and satir['Lead ID'] == 'L-9' and satir['Durum'] == 'Bekliyor' and satir['Ekleyen Telegram ID'] == '777'
    assert satir['Giriş tipi'] == 'Görsel' and satir['Kaynak'] == 'Telegram' and satir['2. görsel bağlantısı'].startswith('https://') and satir['Açıklama'] == 'Açıklama metni'
    assert satir['Yorum sayısı'] == 2 and satir['Son yorum'].startswith('ayse (2026-10-19 10:05:00): son yorum') and satir['Panel bağlantısı'].endswith('/kayitlar?durum=hepsi&ac=%d' % kid)

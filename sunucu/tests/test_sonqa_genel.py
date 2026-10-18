# -*- coding: utf-8 -*-
"""Son QA düzeltmeleri (genel ekip, G1–G12): işlem geçmişi ayrıntısı ve adları, bildirim/denetim tarih uyarısı,
tablo sarmalama, kanıt kontrolü telefon düzeni ve yerel saat, iki sekmeli oturum düşmesi, katalog alan mesajları,
kullanıcı kapsam özeti, dokunma hedefleri, hamburger menü davranışı ve belgeler."""
import io
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import db, kullanicilar
from test_web import web, seed, login, csrf, token

ROOT = Path(__file__).resolve().parents[1]
STATIC = ROOT / 'app' / 'static'
TEMPLATES = ROOT / 'app' / 'templates'
XHR = {'X-Requested-With': 'XMLHttpRequest'}


def denetim_ekle(config, kullanici, islem, **kw):
    conn = db.baglan(config['DB'])
    with db.transaction(conn):
        db.denetim_yaz(conn, kullanici, islem, **kw)
    conn.close()


def son_islemler_html(html):
    return html.split('class="son-islemler"', 1)[1].split('</ul>', 1)[0]


def sonraki(response):
    from urllib.parse import urlsplit, parse_qs
    return parse_qs(urlsplit(response.headers['Location']).query).get('sonraki', [''])[0]


# ---- G1. İşlem Geçmişi: silme sebebi ve okunur onay/silme ayrıntısı ------------------------------------
def test_denetim_shows_delete_reason_and_readable_state_changes(web):
    from app import panel
    c, a = web; login(c); kid = seed(a)
    d = c.get(f'/kayit/{kid}').json
    assert c.post(f'/kayit/{kid}/kt', data={'csrf': csrf(c), 'deger': '1', 'surum': d['surum']}).json['kt'] == 1
    assert c.post(f'/kayit/{kid}/sil', data={'csrf': csrf(c), 'sebep': 'yanlış proje seçilmiş'}).json['ok']
    html = c.get('/denetim').data.decode()
    assert '<b>Sebep:</b> yanlış proje seçilmiş' in html
    assert '<p>Silindi</p>' in html and '<p>Onay: bekliyor → onaylı</p>' in html
    assert 'Önceki: 0' not in html and 'Yeni: 1' not in html
    # Sebep yalnız silme/geri alma için "Sebep", diğer işlemlerde "Açıklama" etiketi alır.
    denetim_ekle(a, 'sistem', 'gyazo_link_guncelle', proje='OF', anahtar='tg:9:9:9', aciklama='Dönüşüm sonucu bağlandı.')
    html = c.get('/denetim').data.decode()
    assert '<b>Açıklama:</b> Dönüşüm sonucu bağlandı.' in html
    # Yardımcı doğrudan: eski/yeni serbest değerler katlanır, onay/silme düz metindir.
    assert panel._denetim_ayrinti({'islem': 'onay_kaldir', 'eski': '1', 'yeni': '0'}) == (['Onay: onaylı → bekliyor'], False)
    assert panel._denetim_ayrinti({'islem': 'geri_al', 'eski': '1', 'yeni': '0'}) == (['Geri alındı'], False)
    assert panel._denetim_ayrinti({'islem': 'bildirim_okundu', 'yeni': '{"adet": 3}'}) == (['3 bildirim okundu'], False)
    assert panel._denetim_ayrinti({'islem': 'duzenle', 'eski': '{"personel": "a"}', 'yeni': '{"personel": "b"}'}) == (['Önceki: {"personel": "a"}', 'Yeni: {"personel": "b"}'], True)


# ---- G3. Ham işlem kodları Türkçe; sözlük tüm denetim_yaz kodlarını kapsar --------------------------------
def test_all_audit_codes_have_turkish_names_and_unknown_codes_are_not_doubled(web):
    from app import panel
    c, a = web; login(c)
    codes = set()
    for path in (ROOT / 'app').glob('*.py'):
        text = path.read_text(encoding='utf-8')
        codes.update(re.findall(r"""denetim_yaz\(\s*[^,]+,\s*[^,]+,\s*['"]([a-z_]+)['"]""", text))
        codes.update(re.findall(r"""_denetle\(\s*[^,]+,\s*[^,]+,\s*['"]([a-z_]+)['"]""", text))
    assert {'telegram_kayit', 'telegram_onay', 'telegram_engel', 'bot_bakim_ac', 'bot_bakim_kapat', 'telegram_eski_buton',
            'telegram_yanlis_adim_butonu', 'bildirim_okundu', 'katalog_proje_kaydet', 'bot_onayla'} <= codes
    missing = sorted(code for code in codes if code not in panel.ISLEM_ADLARI or code not in panel.DENETIM_ADLARI)
    assert not missing, missing
    assert panel.islem_adi('bilinmeyen_kod_x') == 'bilinmeyen kod x' and panel.islem_adi('onayla') == 'Kaydı onayladı'
    for code in ('bildirim_okundu', 'telegram_kayit', 'telegram_onay', 'telegram_engel', 'bot_bakim_ac', 'telegram_eski_buton', 'telegram_yanlis_adim_butonu', 'ozel_bilinmeyen_islem'):
        denetim_ekle(a, 'bot', code, yeni='x')
    ana = c.get('/').data.decode(); akis = son_islemler_html(ana)
    for raw in ('telegram_kayit', 'telegram_onay', 'telegram_engel', 'bot_bakim_ac', 'bildirim_okundu', 'ozel_bilinmeyen_islem'):
        assert raw not in akis, raw
    for name in ('Telegram ile kanıt ekledi', 'Telegram botuna şifreyle giriş yaptı', 'Telegram botunda engellendi', 'Bakım modunu açtı', 'ozel bilinmeyen islem'):
        assert name in akis, name
    # Kendi "okundu" hareketleri akışa girmez; İşlem Geçmişi'nde görünür.
    assert 'Bildirimleri okudu' not in akis
    denetim = c.get('/denetim').data.decode()
    for name in ('Bildirim okundu', 'Telegram ile kanıt', 'Telegram: şifreyle giriş', 'Telegram: engellendi', 'Bot: bakım modu açıldı', 'Telegram: eski düğme', 'Telegram: yanlış adım düğmesi'):
        assert name in denetim, name
    # Bilinmeyen kod bir kez, alt çizgisiz yazılır; bilinen kodun yanında kod etiketi kalır.
    assert denetim.count('ozel_bilinmeyen_islem') == 0 and 'ozel bilinmeyen islem' in denetim
    assert '<code>telegram_kayit</code>' in denetim


# ---- G4. Bildirimler ve İşlem Geçmişi: ters tarih aralığı satır içi uyarı ---------------------------------
def test_inverted_date_range_is_inline_warning_and_keeps_form(web):
    c, a = web; login(c)
    r = c.get('/bildirimler?tarih1=2026-12-31&tarih2=2026-01-01&okundu=0&tur=soru')
    assert r.status_code == 200
    html = r.data.decode()
    assert 'Başlangıç tarihi bitişten sonra olamaz.' in html and 'id="bildirim-uyari"' in html and 'role="alert"' in html
    assert 'value="2026-12-31"' in html and 'value="2026-01-01"' in html and 'aria-invalid="true"' in html
    assert '<option value="0" selected>' in html and '<option value="soru" selected>' in html
    assert 'Geçersiz istek' not in html and 'id="bildirimler-tum-okundu"' in html
    assert c.get('/bildirimler?tarih1=2026-13-45').status_code == 200 and 'Geçersiz tarih.' in c.get('/bildirimler?tarih1=2026-13-45').data.decode()
    r = c.get('/denetim?tarih1=2026-12-31&tarih2=2026-01-01&kullanici=testadmin')
    assert r.status_code == 200
    html = r.data.decode()
    assert 'Başlangıç tarihi bitişten sonra olamaz.' in html and 'id="denetim-uyari"' in html
    assert 'value="2026-12-31"' in html and 'value="testadmin"' in html and '<b>0</b> işlem' in html
    assert c.get('/denetim?tarih1=2026-01-01&tarih2=2026-12-31').status_code == 200


# ---- G2 / G5. Uzun değerler tabloda ve bildirim satırında sarılır ---------------------------------------
def test_long_values_wrap_in_tables_and_notification_rows(web):
    c, a = web; login(c)
    css = (STATIC / 'stil.css').read_text(encoding='utf-8')
    assert re.search(r'\.tablo td \{[^}]*overflow-wrap: anywhere', css)
    assert re.search(r'\.bildirim-satir-alt > \.kucukyazi \{[^}]*min-width: 0[^}]*overflow-wrap: anywhere', css) or re.search(r'\.bildirim-satir-alt > \.kucukyazi \{[^}]*overflow-wrap: anywhere[^}]*min-width: 0', css)
    conn = db.baglan(a['DB'])
    with db.transaction(conn):
        conn.execute("INSERT INTO tg_kullanicilar(id,username,first_name,onayli,engelli,last_seen) VALUES('7001','uzun_kisi','Uzun',1,0,?)", (db.simdi(),))
        kid = conn.execute('INSERT INTO kayitlar(anahtar,proje,personel,uye_id,kategori,kaynak,giris_tipi,ekleyen_id,ekleyen_ad,ekleyen_guven,tarih,guncelleme) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',
                           ('tg:7001:1', 'OF', 'PERSONEL-' + 'X' * 111, 'U' * 120, 'HGST', 'telegram', 'yazili', '7001', 'uzun_kisi', 'telegram', db.simdi(), db.simdi())).lastrowid
        db.bildirim_uret(conn, kid)
    conn.close()
    html = c.get('/telegram/kullanici/7001/kanitlar').data.decode()
    assert 'class="sarmal" title="PERSONEL-' + 'X' * 111 + '"' in html and 'class="sarmal" title="' + 'U' * 120 + '"' in html
    html = c.get('/bildirimler').data.decode()
    assert 'Personel PERSONEL-' + 'X' * 111 in html and 'bildirim-satir-alt' in html


# ---- G6. Kanıt Kontrolü: indirme düğmeleri, eşleştirme tablosu ipucu, yerel saat --------------------------
def test_excel_check_phone_layout_and_local_time(web):
    from app import excel_check
    from openpyxl import Workbook, load_workbook
    c, a = web; login(c); seed(a)
    stamp = '2026-10-17T23:33:01+00:00'
    expected = datetime.fromisoformat(stamp).astimezone().strftime('%d.%m.%Y %H:%M')
    assert excel_check.yerel_zaman_metni(stamp) == expected and re.fullmatch(r'\d{2}\.\d{2}\.\d{4} \d{2}:\d{2}', expected)
    assert excel_check.yerel_zaman_metni('') == '' and excel_check.yerel_zaman_metni('bozuk') == 'bozuk'
    assert excel_check.yerel_zaman_metni('2026-10-17T23:33:01Z') == expected
    wb = Workbook(); ws = wb.active; ws.title = 'Kontrol'; ws.append(['Personel', 'Üye ID']); ws.append(['00123', '00099']); ws.append(['99999', '1'])
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    r = c.post('/excel-kontrol', data={'csrf': csrf(c), 'islem': 'yukle', 'excel': (stream, 'kontrol.xlsx')}, content_type='multipart/form-data')
    html = r.data.decode()
    assert 'class="tablo eslestirme-tablo"' in html and 'eslestirme-ipucu' in html and 'Karşılaştırmaya dahil' in html
    r = c.post('/excel-kontrol', data={'csrf': csrf(c), 'islem': 'kontrol', 'sheet': 'Kontrol', 'kolon_personel': 'A', 'kolon_uye_id': 'B', 'alan': ['personel', 'uye_id'], 'project': '', 'kt': 'hepsi', 'header_row': '1'})
    html = r.data.decode()
    assert 'class="gorunumler kontrol-indir"' in html
    m = re.search(r'<dt>Kontrol zamanı</dt><dd>([^(<]+)\(', html)
    assert m and re.fullmatch(r'\d{2}\.\d{2}\.\d{4} \d{2}:\d{2} ', m.group(1)), m and m.group(1)
    assert '+00:00' not in html.split('Kontrol zamanı', 1)[1][:80]
    book = load_workbook(io.BytesIO(c.get('/excel-kontrol/sonuc.xlsx').data)); sheet = book['Kontrol bilgisi']
    assert sheet['A2'].value == 'Kontrol zamanı' and re.fullmatch(r'\d{2}\.\d{2}\.\d{4} \d{2}:\d{2}', str(sheet['B2'].value))
    book.close()
    css = (STATIC / 'stil.css').read_text(encoding='utf-8')
    assert '.gorunumler.kontrol-indir { flex-wrap: wrap' in css and '.gorunumler.kontrol-indir .btn { min-height: 40px;' in css
    assert '.eslestirme-tablo { min-width: 0; }' in css and '.gorunumler.kontrol-indir .btn { flex: 1 1 100%; }' in css


# ---- G7. İki sekme: düşme bilgisi yeni girişe kadar kalır, oturumsuz POST CSRF'den önce yanıtlanır ------
def test_session_drop_survives_login_page_and_sessionless_posts_are_explained(web):
    from app import panel
    c, a = web; login(c); kid = seed(a)
    kullanicilar.ekle(a, 'ofadmin', 'OF:oku,kt', 'Test-Password-123!'); kullanicilar.guncelle(a, 'ofadmin', sifre_degistir=0)
    c.post('/cikis', data={'csrf': csrf(c)}); login(c, 'ofadmin')
    assert c.get('/kayitlar').status_code == 200
    old_csrf = csrf(c)
    kullanicilar.guncelle(a, 'ofadmin', yetki='OF:oku,kt;PA:oku')  # süperadmin yetkiyi değiştirdi → oturum düştü
    # Sekme B: giriş sayfasını görür (bu, işareti tüketmez).
    r = c.get('/kayitlar')
    assert r.status_code == 302 and 'neden=oturum' in r.headers['Location']
    assert panel.OTURUM_MESAJI in c.get(r.headers['Location']).data.decode()
    assert panel.OTURUM_MESAJI in c.get('/giris').data.decode()
    # Sekme A: açık penceredeki "Onayla" artık 400 "Geçersiz CSRF" değil, JSON 401.
    r = c.post(f'/kayit/{kid}/kt', data={'deger': '1'}, headers={'X-CSRF': old_csrf, **XHR})
    assert r.status_code == 401 and r.is_json and r.json['hata'] == 'Oturumunuz sona erdi, sayfayı yenileyin.' and r.json['giris'].startswith('/giris?neden=oturum')
    # Sekme A'nın zil yoklaması da açıklamalı 401 alır ve giriş sayfası açıklamayı hâlâ gösterir.
    r = c.get('/bildirim.json', headers=XHR)
    assert r.status_code == 401 and r.json['giris'].startswith('/giris?neden=oturum')
    assert panel.OTURUM_MESAJI in c.get('/giris?sonraki=%2Fkayitlar').data.decode()
    with c.session_transaction() as s:
        assert s.get('dustu') and s.get('dustu_zaman')
    # Yeni giriş işareti temizler.
    assert login(c, 'ofadmin').status_code == 302
    with c.session_transaction() as s:
        assert not s.get('dustu') and not s.get('dustu_zaman')
    assert c.get('/kayitlar').status_code == 200
    # Hiç oturumu olmayan tarayıcı: XHR POST → JSON 401 (CSRF 400 değil); form POST → açıklamalı giriş sayfası.
    from app import panel as p
    fresh = p.app.test_client()
    r = fresh.post(f'/kayit/{kid}/kt', data={'deger': '1'}, headers={'X-CSRF': 'eski', **XHR})
    assert r.status_code == 401 and r.is_json and r.json['hata'] == 'Oturumunuz sona erdi, sayfayı yenileyin.'
    r = fresh.post('/arsiv', data={'csrf': 'eski', 'islem': 'tara'})
    assert r.status_code == 302 and 'neden=oturum' in r.headers['Location']
    assert 'Oturumunuz sona erdi. Lütfen yeniden giriş yapın.' in fresh.get(r.headers['Location']).data.decode()
    # Giriş formu ve çıkış bu kuraldan etkilenmez; giriş gerekli 401'i de neden=oturum taşır.
    assert fresh.post('/giris', data={'kullanici': 'ofadmin', 'sifre': 'x'}).status_code == 400
    assert fresh.post('/cikis', data={}).status_code == 302
    r = fresh.get('/bildirim.json', headers=XHR)
    assert r.status_code == 401 and r.json['giris'].startswith('/giris?neden=oturum')
    # Hiç giriş yapmamış ziyaretçi giriş sayfasında yanıltıcı "yetkileriniz güncellendi" mesajı görmez.
    assert panel.OTURUM_MESAJI not in fresh.get('/giris').data.decode()
    js = (STATIC / 'panel.js').read_text(encoding='utf-8')
    assert 'Geçersiz CSRF' in js.split('function parseResponse', 1)[1].split('\n', 1)[0] and 'r.status===400' in js


# ---- G8. Katalog alanlarında Türkçe doğrulama mesajları -----------------------------------------------------
def test_catalog_code_fields_have_turkish_validation_messages(web):
    c, a = web; login(c)
    html = c.get('/yonetim/katalog').data.decode()
    assert 'pattern="[A-Z0-9_]{1,24}"' in html and 'data-gecersiz="Proje kodu yalnız BÜYÜK harf, rakam ve alt çizgi (en çok 24)."' in html
    assert 'data-bos="Proje kodu zorunludur."' in html
    assert 'data-gecersiz="Kategori kodu yalnız küçük harf, rakam ve alt çizgi içerebilir (en çok 48)."' in html and 'data-bos="Kategori kodu zorunludur."' in html
    assert 'katalog.js?v=2.2.3' in html
    js = (STATIC / 'katalog.js').read_text(encoding='utf-8')
    assert '"data-gecersiz": "Soru anahtarı küçük harfle başlamalı' in js and '"data-question-key"' in js
    panel_js = (STATIC / 'panel.js').read_text(encoding='utf-8')
    assert "addEventListener('invalid'" in panel_js and "getAttribute('data-gecersiz')" in panel_js


# ---- G9. Kullanıcı kapsamı Türkçe ve sıralı; süperadmin etiketi tek satır; ana sayfa başlığı -------------
def test_user_scope_summary_is_turkish_and_ordered(web):
    c, a = web; login(c)
    kullanicilar.ekle(a, 'ayse', 'HI:oku,sil,kt,duzenle;OF:oku,duzenle,kt;PA:oku,kt;GA:oku', 'Test-Password-123!', 'Ayşe Kaya')
    html = c.get('/yonetim/kullanicilar').data.decode()
    assert 'GA: yalnız okuma · HI: Onay, Düzenleme, Silme · OF: Onay, Düzenleme · PA: Onay' in html
    assert 'title="HI: Onay, Düzenleme, Silme"' in html and 'title="GA: yalnız okuma"' in html
    assert 'duzenle' not in html.split('class="kapsam-listesi"', 1)[1].split('</div>', 1)[0]
    assert '<span class="durum-etiket notr tek-satir">Süperadmin</span><small class="kapsam-notu">tüm projeler</small>' in html
    css = (STATIC / 'stil.css').read_text(encoding='utf-8')
    assert '.tek-satir { white-space: nowrap; }' in css
    assert re.search(r'@media \(max-width: 700px\) \{[^@]*\.ozet-baslik \{ flex-direction: column', css)


# ---- G10 / G11. Dokunma hedefleri ve hamburger menü davranışı --------------------------------------------
def test_touch_targets_and_hamburger_menu_rules(web):
    css = (STATIC / 'stil.css').read_text(encoding='utf-8')
    block = re.search(r'@media \(max-width: 700px\) \{(.*?)\n\}', css, re.S).group(1)
    for rule in ('.icerik .btn', '.bildirim-kutu .btn', '.menu-mobil-alt .btn', '.ust .menu .menu-mobil-alt a.btn', '.ust .menu .menu-mobil-alt button.btn'):
        assert rule in block and 'min-height: 40px' in block, rule
    assert '.cip span { min-height: 40px; display: inline-flex; align-items: center; }' in block
    assert '.satir-form .btn.kucuk, .tablo .btn.kucuk { min-height: 40px' in css and '.yalniz-telefon { display: inline-flex; margin-top: 6px; min-height: 40px; }' in css
    js = (STATIC / 'panel.js').read_text(encoding='utf-8')
    menu = js.split('function menu(){', 1)[1].split('function botSayfasi', 1)[0]
    assert 'burger.contains(e.target)' in menu and 'aktifAcilirlar' in menu
    assert "if(open&&telefon())aktifAcilirlar.forEach" in menu


# ---- G12. Belgeler -----------------------------------------------------------------------------------------
def test_docs_mention_ctrl_f5_and_final_qa_fixes(web):
    readme = (ROOT / 'README.md').read_text(encoding='utf-8')
    assert 'Ctrl+F5' in readme.split('## 1. Güncelleme', 1)[1].split('\n## ', 1)[0]
    notes = (ROOT / 'DEGISIKLIKLER_v2.2.md').read_text(encoding='utf-8')
    section = notes.split('## Düzeltmeler (son QA turu)', 1)[1]
    for item in ('K1', 'K2', 'K3', 'K4', 'K5', 'K6', 'K7', 'K8', 'İşlem Geçmişi', 'Bildirimler', 'Kanıt Kontrolü', 'Oturum düşmesi', 'Hamburger menü', 'v=2.2.3'):
        assert item in section, item
    c, a = web; login(c)
    for url in ('/', '/bildirimler', '/denetim', '/yonetim/kullanicilar', '/yonetim/katalog'):
        html = c.get(url).data.decode()
        assert 'stil.css?v=2.2.3' in html and 'panel.js?v=2.2.3' in html, url

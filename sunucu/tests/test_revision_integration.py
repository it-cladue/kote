import io
import json
import sqlite3
import zipfile
import pytest
from openpyxl import load_workbook
from app import db,katalog,kullanicilar,yetki
from test_web import web,seed,login,csrf


def account(config,name,role):
    kullanicilar.ekle(config,name,role,'Test-Password-123!')
    kullanicilar.guncelle(config,name,sifre_degistir=0)


def test_normal_admin_scope_is_hidden_in_all_filters_and_metadata(web):
    client,config=web
    account(config,'ofonly','OF:oku,kt')
    own=seed(config);other=seed(config,'PA','other-private',picture=False)
    conn=db.baglan(config['DB'])
    katalog.save_category(conn,'only_pa','PRIVATE_PA_CATEGORY',mode='tek_uye',projects=['PA'],questions=None,actor='testadmin')
    conn.execute("UPDATE kayitlar SET kategori='PRIVATE_PA_CATEGORY',personel='PRIVATE_PA_PERSON' WHERE id=?",(other,))
    db.denetim_yaz(conn,'testadmin','duzenle','PA','other-private',yeni='PRIVATE_PA_LOG')
    db.denetim_yaz(conn,'testadmin','duzenle','OF','tg:1:2:3',yeni='VISIBLE_OF_LOG');conn.close()
    login(client,'ofonly')
    for url in ('/','/kayitlar?durum=hepsi','/disa-aktar','/bildirimler','/denetim'):
        response=client.get(url);assert response.status_code==200
        assert b'PRIVATE_PA' not in response.data
        assert b'value="PA"' not in response.data and b'"PA":' not in response.data
    assert b'VISIBLE_OF_LOG' in client.get('/denetim').data
    for url in ('/kayitlar?proje=PA','/kayitlar?proje=OF&proje=PA','/disa-aktar.xlsx?proje=PA','/kategoriler.json?proje=PA','/bildirimler?proje=PA'):
        assert client.get(url).status_code==403
    options=client.get('/kategoriler.json?proje=OF').json['kategoriler']
    assert 'PRIVATE_PA_CATEGORY' not in options
    assert client.get(f'/kayit/{other}').status_code==403 and client.get(f'/kayit/{own}').status_code==200
    assert client.get('/yonetim/katalog').status_code==403


def test_form_creates_project_bound_admin_and_only_superadmin_can_mutate_catalog(web):
    client,config=web;login(client)
    response=client.post('/yonetim/kullanicilar',data={'csrf':csrf(client),'islem':'ekle','ad':'newadmin','sifre':'Test-Password-123!','rol':'admin','projeler':['OF'],'eylemler_OF':['kt']})
    assert response.status_code==302
    user=kullanicilar.bul(config,'newadmin');parsed=yetki.coz(user['yetki'])
    assert not parsed['admin'] and list(parsed['projeler'])==['OF']
    assert yetki.izin(parsed,'OF','kt') and not yetki.izin(parsed,'PA','oku')
    assert client.get('/yonetim/katalog').status_code==200
    response=client.post('/yonetim/katalog',data={'csrf':csrf(client),'islem':'proje_kaydet','code':'ZZ','label':'Yeni proje','color':'#224466','active':'1'})
    assert response.status_code==302
    conn=db.baglan(config['DB']);assert katalog.get_project(conn,'ZZ')['aktif'];conn.close()
    assert not yetki.izin(parsed,'ZZ','oku')
    assert client.post('/yonetim/katalog',data={'csrf':csrf(client),'islem':'proje_kaydet','code':'ZZ','label':'Eski form','color':'#112233','active':'1'}).status_code==409
    kullanicilar.guncelle(config,'newadmin',sifre_degistir=0)
    client.post('/cikis',data={'csrf':csrf(client)});login(client,'newadmin')
    assert client.post('/yonetim/katalog',data={'csrf':csrf(client),'islem':'proje_kaydet','code':'EVIL'}).status_code==403


def test_past_notifications_private_keys_filters_and_current_permission(web):
    client,config=web;account(config,'ofonly','OF:oku,kt');own=seed(config);other=seed(config,'PA','other-notify')
    conn=db.baglan(config['DB'])
    own_notice=conn.execute("SELECT id FROM bildirimler WHERE kullanici_adi='ofonly'").fetchone()[0]
    conn.execute("INSERT INTO bildirimler(kullanici_adi,kayit_id,proje,zaman) VALUES('ofonly',?,'PA',?)",(other,db.simdi()))
    hidden=conn.execute("SELECT id FROM bildirimler WHERE kullanici_adi='ofonly' AND proje='PA'").fetchone()[0]
    db.mesaj_bildir(conn,own,'ofonly','soru','testadmin','Özel soru','test-soru');conn.close()
    login(client,'ofonly')
    assert client.get('/bildirimler?okundu=0&tur=kanit').status_code==200
    assert client.post('/bildirimler/okundu',json={'ids':['kayit:'+str(own_notice)]},headers={'X-CSRF':csrf(client)}).json['degisen']==1
    html=client.get('/bildirimler?okundu=1').data
    assert f'ac={own}'.encode() in html and f'ac={other}'.encode() not in html
    assert b'\x00' not in html
    assert client.post('/bildirimler/okundu',json={'hepsi':True},headers={'X-CSRF':csrf(client)}).json['ok']
    conn=db.baglan(config['DB']);assert conn.execute('SELECT okundu FROM bildirimler WHERE id=?',(hidden,)).fetchone()[0]==0;conn.close()
    assert client.get('/bildirimler?tur=soru&okundu=1').status_code==200


def ods_data():
    output=io.BytesIO()
    xml='''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:spreadsheet><table:table table:name="Kontrol"><table:table-row><table:table-cell office:value-type="string"><text:p>ID</text:p></table:table-cell></table:table-row><table:table-row><table:table-cell office:value-type="string"><text:p>00099</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet></office:body></office:document-content>'''
    with zipfile.ZipFile(output,'w') as archive:
        archive.writestr('mimetype','application/vnd.oasis.opendocument.spreadsheet')
        archive.writestr('content.xml',xml)
    output.seek(0);return output


@pytest.mark.parametrize('filename,content,options',[
    ('dosya.csv',b'ID\r\n00099\r\n',{'header_row':'1','kolon_uye_id':'A','alan':'uye_id'}),
    ('dosya.tsv',b'ID\n00099\n',{'header_row':'1','kolon_uye_id':'A','alan':'uye_id'}),
    ('dosya.txt',b'00099\n',{'text_mode':'ids','id_field':'uye_id','alan':'uye_id'}),
    ('dosya.ods',None,{'sheet':'Kontrol','header_row':'1','kolon_uye_id':'A','alan':'uye_id'}),
])
def test_all_formats_web_check_no_new_evidence(web,filename,content,options):
    client,config=web;login(client);kid=seed(config)
    source=ods_data() if content is None else io.BytesIO(content)
    response=client.post('/excel-kontrol',data={'csrf':csrf(client),'islem':'yukle','excel':(source,filename)},content_type='multipart/form-data')
    assert response.status_code==200
    response=client.post('/excel-kontrol',data={'csrf':csrf(client),'islem':'kontrol','project':'OF',**options})
    assert response.status_code==200
    with client.session_transaction() as session:identifier=session['excel_check_id']
    from pathlib import Path
    report=json.loads((Path(config['BOT_KLASOR'])/'checks'/identifier/'result.json').read_text())['report']
    assert report['found']==1 and report['rows'][0]['ids']==[kid]
    conn=db.baglan(config['DB']);assert conn.execute('SELECT COUNT(*) FROM kayitlar').fetchone()[0]==1;conn.close()
    book=load_workbook(io.BytesIO(client.get('/excel-kontrol/sonuc.xlsx').data));assert 'Kontrol bilgisi' in book.sheetnames;book.close()


def test_txt_mode_switch_and_large_report_page(web):
    client,config=web;login(client);seed(config)
    client.post('/excel-kontrol',data={'csrf':csrf(client),'islem':'yukle','excel':(io.BytesIO(b'ID;not\n00099;ok\n'),'tablo.txt')},content_type='multipart/form-data')
    response=client.post('/excel-kontrol',data={'csrf':csrf(client),'islem':'incele','text_mode':'table','header_row':'1','delimiter':'semicolon'})
    assert response.status_code==200 and b'00099' in response.data
    client.post('/excel-kontrol',data={'csrf':csrf(client),'islem':'kontrol','text_mode':'table','header_row':'1','delimiter':'semicolon','kolon_uye_id':'A','alan':'uye_id','project':'OF'})
    with client.session_transaction() as session:identifier=session['excel_check_id']
    from pathlib import Path
    report=json.loads((Path(config['BOT_KLASOR'])/'checks'/identifier/'result.json').read_text())['report']
    assert report['found']==1
    client.post('/excel-kontrol',data={'csrf':csrf(client),'islem':'yukle','excel':(io.BytesIO(b'00099\n'*105),'liste.txt')},content_type='multipart/form-data')
    response=client.post('/excel-kontrol',data={'csrf':csrf(client),'islem':'kontrol','text_mode':'ids','id_field':'uye_id','alan':'uye_id','project':'OF'})
    assert response.status_code==200 and response.data.count(b'<tr data-status=')==100
    assert client.get('/excel-kontrol?r_s=2').data.count(b'<tr data-status=')==5


def test_log_masking_zero_and_no_repeated_schema_writes(web,monkeypatch):
    _,config=web;conn=db.baglan(config['DB'])
    db.denetim_yaz(conn,'testadmin','test_log',eski=0,yeni={'BOT_PASSWORD':'fake-hidden','access_token':'fake-hidden','visible':False})
    row=conn.execute("SELECT eski,yeni FROM denetim WHERE islem='test_log'").fetchone()
    assert row['eski']=='0' and 'fake-hidden' not in row['yeni'] and 'false' in row['yeni'];conn.close()
    statements=[];original=sqlite3.connect
    def tracked(*args,**kwargs):
        connection=original(*args,**kwargs);connection.set_trace_callback(statements.append);return connection
    monkeypatch.setattr(db.sqlite3,'connect',tracked)
    conn=db.baglan(config['DB']);conn.close()
    assert not any(x.lstrip().upper().startswith(('CREATE','ALTER','INSERT','UPDATE','DELETE')) for x in statements)


def test_invalid_admin_alias_cannot_bypass_last_admin_guard(web):
    _,config=web
    assert not yetki.coz('admin:oku')['admin'] and not yetki.coz('admin ; OF:oku')['admin'] and not yetki.coz('*')['gecerli']
    conn=db.baglan(config['DB'])
    conn.execute("INSERT INTO panel_kullanicilari(kullanici_adi,sifre_hash,yetki) VALUES('invalid','not-a-secret','admin:oku')");conn.close()
    with pytest.raises(ValueError):kullanicilar.guncelle(config,'testadmin',yetki='OF:oku')


@pytest.mark.parametrize('wrapped',[False,True])
def test_bot_snapshot_custom_answer_and_lead_edit(web,wrapped):
    client,config=web;login(client);kid=seed(config)
    questions=[{'key':'employee','label':'Personel kodu','preset':'personel'},{'key':'lead','label':'Lead ID','preset':'lead_id'},{'key':'custom','label':'Özel soru','preset':None}]
    snapshot={'questions':questions,'category':'test'} if wrapped else questions
    conn=db.baglan(config['DB']);conn.execute('UPDATE kayitlar SET soru_anlik_json=?,cevaplar_json=?,lead_id=? WHERE id=?',(json.dumps(snapshot,ensure_ascii=False),json.dumps({'employee':'00123','lead':'L001','custom':'Özel cevap'},ensure_ascii=False),'L001',kid));conn.close()
    details=client.get(f'/kayit/{kid}').json
    assert details['soru_cevaplar'][2]['value']=='Özel cevap'
    response=client.post(f'/kayit/{kid}/duzenle',data={'csrf':csrf(client),'personel':'updated','lead_id':'L002','surum':details['surum']})
    assert response.status_code==200 and response.json['ok']
    conn=db.baglan(config['DB']);row=conn.execute('SELECT * FROM kayitlar WHERE id=?',(kid,)).fetchone();answers=json.loads(row['cevaplar_json'])
    assert row['lead_id']=='L002' and answers['lead']=='L002' and answers['employee']=='updated' and answers['custom']=='Özel cevap';conn.close()
    book=load_workbook(io.BytesIO(client.get('/disa-aktar.xlsx?durum=hepsi').data));sheet=book.active
    index=[cell.value for cell in sheet[1]].index('Ek sorular')+1
    assert 'Özel soru: Özel cevap' in sheet.cell(2,index).value and 'Lead ID: L002' in sheet.cell(2,index).value;book.close()
    assert b'name="lead_id"' in client.get('/kayitlar').data


def test_simultaneous_sessions_only_one_stale_version_write_wins(web):
    from concurrent.futures import ThreadPoolExecutor
    from app import panel
    _,config=web;kid=seed(config);clients=[]
    for index in range(8):
        name='writer'+str(index);account(config,name,'OF')
        client=panel.app.test_client();login(client,name);clients.append((client,csrf(client)))
    version=clients[0][0].get(f'/kayit/{kid}').json['surum']
    def change(item):
        index,(client,token)=item
        return client.post(f'/kayit/{kid}/duzenle',data={'csrf':token,'personel':'parallel-'+str(index),'surum':str(version)}).status_code
    with ThreadPoolExecutor(max_workers=8) as pool:statuses=list(pool.map(change,enumerate(clients)))
    assert statuses.count(200)==1 and statuses.count(409)==7
    conn=db.baglan(config['DB']);assert conn.execute('SELECT surum FROM kayitlar WHERE id=?',(kid,)).fetchone()[0]==version+1;conn.close()

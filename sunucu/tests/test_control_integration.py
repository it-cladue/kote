import io
import json
from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from test_web import web,login,csrf,seed
from app import db,kullanicilar,bot_yonetim,yetki


def workbook():
    w=Workbook();s=w.active;s.title='Kontrol';s.append(['Proje','Personel','Üye ID','Kategori']);s.append(['OF','00123','00099','HGST']);s.append(['OF','missing','99','HGST']);f=io.BytesIO();w.save(f);f.seek(0);return f


def test_excel_web_check_does_not_create_evidence(web):
    c,a=web;login(c);kid=seed(a)
    assert c.get('/excel-kontrol').status_code==200
    response=c.post('/excel-kontrol',data={'csrf':csrf(c),'islem':'yukle','excel':(workbook(),'kanit-check.xlsx')},content_type='multipart/form-data')
    assert response.status_code==200 and b'Kontrol' in response.data
    response=c.post('/excel-kontrol',data={'csrf':csrf(c),'islem':'kontrol','sheet':'Kontrol','kolon_proje':'A','kolon_personel':'B','kolon_uye_id':'C','kolon_kategori':'D','alan':['proje','personel','uye_id','kategori'],'project':'OF','kt':'hepsi'})
    assert response.status_code==200
    with c.session_transaction() as session:folder=Path(a['BOT_KLASOR'])/'checks'/session['excel_check_id']
    r=json.loads((folder/'result.json').read_text())['report']
    assert r['found']==1 and r['missing']==1 and r['rows'][0]['ids']==[kid]
    all_rows=load_workbook(io.BytesIO(c.get('/excel-kontrol/sonuc.xlsx').data));assert all_rows.active.max_row==3;all_rows.close()
    missing=load_workbook(io.BytesIO(c.get('/excel-kontrol/sonuc.xlsx?eksik=1').data));assert missing.active.max_row==2 and missing.active['B2'].value=='yok';missing.close()
    conn=db.baglan(a['DB']);assert conn.execute('SELECT COUNT(*) FROM kayitlar').fetchone()[0]==1;conn.close()
    # Bozuk dosya hiçbir kayıt oluşturmaz ve sayfayı çökertmez.
    assert c.post('/excel-kontrol',data={'csrf':csrf(c),'islem':'yukle','excel':(io.BytesIO(b'bad'),'bad.xlsx')},content_type='multipart/form-data').status_code==200


def test_admin_management_and_category_persistence(web):
    c,a=web;login(c)
    assert c.post('/bot',data={'csrf':csrf(c),'islem':'onayla','tg_id':'12345678','ad':'user'}).status_code==200
    assert c.get('/bot/kullanici/12345678').status_code==200
    assert c.post('/bot',data={'csrf':csrf(c),'islem':'engelle','tg_id':'12345678'}).status_code==200
    assert c.post('/bot',data={'csrf':csrf(c),'islem':'engel_kaldir','tg_id':'12345678'}).status_code==200
    conn=db.baglan(a['DB']);r=conn.execute("SELECT * FROM tg_kullanicilar WHERE id='12345678'").fetchone();assert r['onayli']==1 and r['engelli']==0;conn.close()
    c.post('/bot',data={'csrf':csrf(c),'islem':'tam_sifirla','tg_id':'12345678'})
    c.post('/bot',data={'csrf':csrf(c),'islem':'kategori_sil','code':'ark'})
    conn=db.baglan(a['DB']);assert conn.execute("SELECT aktif FROM kategoriler WHERE code='ark'").fetchone()[0]==0
    assert conn.execute("SELECT onayli FROM tg_kullanicilar WHERE id='12345678'").fetchone()[0]==0;conn.close()
    c.post('/bot',data={'csrf':csrf(c),'islem':'bot_duraklat'})
    conn=db.baglan(a['DB']);assert db.ayar_al(conn,'bot_durdur')=='1';conn.close()
    c.post('/bot',data={'csrf':csrf(c),'islem':'bot_surdur'})
    assert c.get('/bot?q=user&tg_durum=hepsi&sure=24saat').status_code==200


def test_roles_without_read_do_not_leak_and_stale_users(web):
    c,a=web;login(c);kid=seed(a)
    kullanicilar.ekle(a,'no_read','OF:kt','Test-Password-123!');kullanicilar.guncelle(a,'no_read',sifre_degistir=0)
    old=kullanicilar.bul(a,'no_read')['session_version']
    kullanicilar.guncelle(a,'no_read',expected_version=old,yetki='PA:oku')
    with pytest.raises(ValueError):kullanicilar.guncelle(a,'no_read',expected_version=old,yetki='OF')
    parsed=yetki.coz(kullanicilar.bul(a,'no_read')['yetki'])
    assert list(parsed['projeler'])==['PA'] and yetki.izin(parsed,'PA','oku') and not yetki.izin(parsed,'OF','oku')
    kullanicilar.guncelle(a,'no_read',yetki='OF:kt')
    c.post('/cikis',data={'csrf':csrf(c)});login(c,'no_read')
    assert c.get(f'/kayit/{kid}').status_code==403
    assert c.get(f'/gorsel/{kid}').status_code==403
    assert c.get('/excel-kontrol').status_code==403
    x=load_workbook(io.BytesIO(c.get('/disa-aktar.xlsx?durum=hepsi').data));assert x.active.max_row==1;x.close()


def test_undated_export_and_backup_in_transaction(web):
    c,a=web;login(c);kid=seed(a);conn=db.baglan(a['DB'])
    conn.execute('UPDATE kayitlar SET tarih=NULL WHERE id=?',(kid,))
    x=load_workbook(io.BytesIO(c.get('/disa-aktar.xlsx?tarihsiz=1&durum=hepsi').data));assert x.active.max_row==2 and 'Proje özeti' in x.sheetnames;x.close()
    assert c.get('/disa-aktar.xlsx?tarihsiz=1&tarih1=2026-01-01').status_code==400
    with db.transaction(conn):bot_yonetim.yedek_al(a,conn,'testadmin')
    conn.close();assert list((Path(a['BOT_KLASOR'])/'backups').glob('*.sqlite3'))

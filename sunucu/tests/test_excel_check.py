# -*- coding: utf-8 -*-
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import db
from app.excel_check import ExcelCheckError, check_source, check_workbook, inspect_workbook, validate_xlsx


HEADERS = ["Proje", "Personel", "Üye ID", "Kategori"]
OPTIONS = {
    "sheet": "Kontrol",
    "columns": {"proje": "A", "personel": "B", "uye_id": "C", "kategori": "D"},
    "match_fields": ["proje", "personel", "uye_id", "kategori"],
    "project": "", "kt": "hepsi",
}


def _record(conn, project, personel, uye, category="ARK", deleted=0, lead_id=""):
    key = f"telegram:{project}:{personel}:{uye}:{conn.execute('SELECT COUNT(*) FROM kayitlar').fetchone()[0]}"
    return conn.execute(
        "INSERT INTO kayitlar(anahtar,proje,personel,uye_id,kategori,lead_id,silindi,kaynak) VALUES(?,?,?,?,?,?,?,?)",
        (key, project, personel, uye, category, lead_id, deleted, "telegram"),
    ).lastrowid


def _workbook(path: Path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Kontrol"
    sheet.append(HEADERS)
    sheet.append(["OF", "00123", "00099", "ark"])
    sheet.append(["OF", "00123", "00100", "ARK"])
    sheet.append(["OF", "DUP", "00200", "ARK"])
    sheet.append(["OF", "DUP", "00200", "ARK"])
    sheet.append(["OF", "SADECE", "", "ARK"])
    sheet.append(["OF", "=1+1", "00099", "ARK"])
    sheet.append([False, False, None, None])
    sheet.append(["PA", "00123", "00099", "ARK"])
    workbook.save(path)


def test_inspect_and_readonly_scoped_matching(tmp_path):
    xlsx = tmp_path / "kontrol.xlsx"
    _workbook(xlsx)
    conn = db.baglan(str(tmp_path / "panel.db"))
    try:
        found_id = _record(conn, "OF", "00123", "00099", "ARK")
        pa_id = _record(conn, "PA", "00123", "00099", "ARK")
        first_dup = _record(conn, "OF", "DUP", "00200", "ARK")
        second_dup = _record(conn, "OF", "DUP", "00200", "ARK")
        _record(conn, "OF", "00123", "00099", "ARK", deleted=1)

        inspected = inspect_workbook(xlsx)
        assert inspected["sheets"][0]["headers"] == [
            {"column": "A", "label": "Proje"}, {"column": "B", "label": "Personel"},
            {"column": "C", "label": "Üye ID"}, {"column": "D", "label": "Kategori"},
        ]
        assert inspected["sheets"][0]["sample"][0] == ["OF", "00123", "00099", "ark"]

        report = check_workbook(xlsx, conn, ["OF"], OPTIONS)
        by_row = {row["row"]: row for row in report["rows"]}
        assert report["total"] == 7
        assert report["found"] == 1 and report["missing"] == 1
        assert report["multiple"] == report["ambiguous"] == 2 and report["invalid"] == 3
        assert by_row[2]["status"] == "bulundu" and by_row[2]["ids"] == [found_id]
        assert by_row[2]["source"] == "Kontrol!2"
        assert by_row[3]["status"] == "yok"
        assert by_row[4]["ids"] == [first_dup, second_dup]
        assert by_row[6]["status"] == "anahtar_bos"
        assert by_row[7]["status"] == "formul_sonucsuz"
        assert by_row[9]["status"] == "erisim_kapsami_disi" and by_row[9]["ids"] == [] and pa_id not in by_row[9]["ids"]
        assert report["metadata"]["snapshot"] == "deferred_read"
        assert report["metadata"]["checked_at"].endswith("+00:00")
        # Eşleştirici yalnız SELECT yapar; kanıt kaydı oluşmaz/değişmez.
        assert conn.execute("SELECT COUNT(*) FROM kayitlar").fetchone()[0] == 5
    finally:
        conn.close()


def test_numeric_identifier_risk_project_conflict_and_header_choices(tmp_path):
    xlsx = tmp_path / "risk.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Kontrol"
    sheet.append(["Rapor başlığı"])
    sheet.append(["Proje", "Üye ID"])
    sheet.append(["OF", 123])
    sheet["B3"].number_format = "00000"
    sheet.append(["PA", "A1"])
    sheet.append(["OF", ""])
    workbook.save(xlsx)
    conn = db.baglan(str(tmp_path / "panel.db"))
    try:
        _record(conn, "OF", "", "00123")
        of_id = _record(conn, "OF", "", "A1")
        _record(conn, "PA", "", "A1")
        options = {"sheet": "Kontrol", "header_row": 2, "columns": {"proje": "A", "uye_id": "B"},
                   "match_fields": ["uye_id"], "project": "OF", "kt": "hepsi"}
        report = check_workbook(xlsx, conn, ["OF", "PA"], options)
        by_row = {item["row"]: item for item in report["rows"]}
        assert by_row[3]["status"] == "bicim_belirsiz"  # 123 + 00000 maskesi, DB 00123
        assert by_row[4]["status"] == "proje_filtre_uyusmazligi"
        assert by_row[4]["ids"] == [] and of_id not in by_row[4]["ids"]
        assert by_row[5]["status"] == "anahtar_bos"

        # Başlıksız kaynakta satır 1 veri kabul edilir; A sütunu serbest metin kalır.
        no_header = tmp_path / "no-header.xlsx"
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Veri"; ws.append(["00123"]); wb.save(no_header)
        report = check_source(no_header, conn, ["OF"], {"sheet": "Veri", "header_row": 0,
                              "columns": {"uye_id": "A"}, "match_fields": ["uye_id"], "project": "OF", "kt": "hepsi"})
        assert report["rows"][0]["status"] == "bulundu"
    finally:
        conn.close()


def test_date_style_is_never_definitive_missing(tmp_path):
    xlsx = tmp_path / "date.xlsx"
    workbook = openpyxl.Workbook(); sheet = workbook.active; sheet.title = "Kontrol"
    sheet.append(["Üye ID"]); sheet.append([datetime(2024, 1, 2)])
    workbook.save(xlsx)
    conn = db.baglan(str(tmp_path / "panel.db"))
    try:
        report = check_workbook(xlsx, conn, ["OF"], {"sheet": "Kontrol", "columns": {"uye_id": "A"},
                                "match_fields": ["uye_id"], "project": "OF", "kt": "hepsi"})
        assert report["rows"][0]["status"] == "bicim_belirsiz"
    finally:
        conn.close()


def test_requires_mapped_identifier_and_respects_kt(tmp_path):
    xlsx = tmp_path / "tek.xlsx"
    workbook = openpyxl.Workbook(); sheet = workbook.active; sheet.title = "Kontrol"
    sheet.append(["Proje", "Personel"]); sheet.append(["OF", "00123"]); workbook.save(xlsx)
    conn = db.baglan(str(tmp_path / "panel.db"))
    try:
        kid = _record(conn, "OF", "00123", "")
        conn.execute("UPDATE kayitlar SET kt=1 WHERE id=?", (kid,))
        options = {"sheet": "Kontrol", "columns": {"proje": "A", "personel": "B"},
                   "match_fields": ["proje", "personel"], "project": "OF", "kt": "onayli"}
        assert check_workbook(xlsx, conn, ["OF"], options)["rows"][0]["ids"] == [kid]
        options["kt"] = "bekleyen"
        assert check_workbook(xlsx, conn, ["OF"], options)["rows"][0]["status"] == "yok"
        options["match_fields"] = ["proje"]
        with pytest.raises(ExcelCheckError, match="anahtar alan"):
            check_workbook(xlsx, conn, ["OF"], options)
        options["match_fields"] = ["personel"]; options["columns"] = {"personel": "Z"}
        with pytest.raises(ExcelCheckError, match="başlığı"):
            check_workbook(xlsx, conn, ["OF"], options)
    finally:
        conn.close()


def test_rejects_non_xlsx_and_invalid_zip_without_paths(tmp_path):
    xls = tmp_path / "legacy.xls"; xls.write_bytes(b"not used")
    with pytest.raises(ExcelCheckError, match=r"\.xls desteklenmez"):
        validate_xlsx(xls)
    broken = tmp_path / "bozuk.xlsx"; broken.write_bytes(b"bu bir zip degil")
    with pytest.raises(ExcelCheckError) as error:
        validate_xlsx(broken)
    assert str(broken) not in str(error.value)
    not_a_workbook = tmp_path / "zip.xlsx"
    with zipfile.ZipFile(not_a_workbook, "w") as archive:
        archive.writestr("random.txt", "x")
    with pytest.raises(ExcelCheckError, match="Geçerli"):
        validate_xlsx(not_a_workbook)

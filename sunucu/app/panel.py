"""Kanıt Paneli sunucu sürümü. Canlı Excel/OneDrive/Gyazo bağımlılığı yok."""
import io
import json
import os
import re
import secrets
import time
import uuid
from datetime import datetime, timedelta
from pathlib import Path
from flask import Flask, g, request, session, redirect, url_for, render_template, jsonify, abort, send_file, flash
from werkzeug.datastructures import MultiDict
from functools import wraps
from PIL import Image, UnidentifiedImageError
from . import db, yetki, kullanicilar, gorsel, bot_yonetim
from .settings import AYAR, ROOT

SURUM = '2.2.0'
secret_path = Path(AYAR['BOT_KLASOR']) / 'session.key'
secret = os.getenv('SESSION_SECRET')
if not secret:
    try:
        fd = os.open(secret_path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, 'w') as f:
            f.write(secrets.token_hex(32))
    except FileExistsError:
        pass
    secret = secret_path.read_text().strip()
if len(secret) < 32:
    raise RuntimeError('SESSION_SECRET en az 32 karakter olmalı.')
app = Flask(__name__)
app.secret_key = secret
secure = str(AYAR['HTTPS']).lower() in ('1','true')
app.config.update(SESSION_COOKIE_NAME='kanit_session', SESSION_COOKIE_HTTPONLY=True,
                  SESSION_COOKIE_SECURE=secure, SESSION_COOKIE_SAMESITE='Lax',
                  PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
                  MAX_CONTENT_LENGTH=12*1024*1024)
COZUCU = gorsel.GorselCozucu(AYAR, AYAR['ONBELLEK'])

def simdi():
    return db.simdi()

def baglanti():
    if 'conn' not in g:
        g.conn = db.baglan(AYAR['DB'])
    return g.conn

@app.teardown_appcontext
def close_connection(error=None):
    c = g.pop('conn',None)
    if c:
        c.close()

def _all_project_info():
    from . import katalog
    if 'project_info' not in g:
        items=katalog.list_projects(baglanti(),active_only=False)
        known={p['code'] for p in items}
        for row in baglanti().execute('SELECT DISTINCT proje FROM kayitlar'):
            if row['proje'] not in known:
                items.append(dict(code=row['proje'],label=row['proje'],color='#64748b',aktif=0,revision=0))
        g.project_info=items
    return g.project_info

def renk(proje):
    return next((p['color'] for p in _all_project_info() if p['code']==proje),'#64748b')


def metin_rengi(value):
    try:
        h=value.lstrip('#'); r,gg,b=(int(h[i:i+2],16) for i in (0,2,4))
        return '#111111' if .299*r+.587*gg+.114*b>150 else '#ffffff'
    except (ValueError,AttributeError):
        return '#ffffff'

def projeler():
    return [p['code'] for p in _all_project_info()]

def _allowed():
    return yetki.izinli_projeler(g.yk,projeler())

def _selected_projects(args,allowed):
    selected=args.getlist('proje')
    if not selected:return list(allowed)
    if any(p not in allowed for p in selected):abort(403)
    return list(dict.fromkeys(selected))

def _scoped_categories(codes):
    from . import katalog
    if not codes:return []
    names=set()
    for code in codes:
        names.update(r['label'] for r in katalog.list_categories(baglanti(),project_code=code,active_only=True))
    places=','.join('?' for _ in codes)
    names.update(r[0] for r in baglanti().execute("SELECT DISTINCT kategori FROM kayitlar WHERE proje IN ("+places+") AND kategori IS NOT NULL AND kategori<>''",codes))
    return sorted(names,key=str.casefold)

def url_filtre(endpoint):
    return url_for(endpoint,**request.args.to_dict(flat=False))


def url_degistir(**kw):
    args=request.args.to_dict(flat=False); args.pop('s',None)
    for k,v in kw.items():
        if v is None or v=='': args.pop(k,None)
        else: args[k]=list(v) if isinstance(v,(list,tuple)) else [str(v)]
    return url_for(request.endpoint or 'kayitlar',**args)

def ip():
    return request.remote_addr

@app.before_request
def before():
    g.k=None
    if session.get('kad'):
        row=baglanti().execute('SELECT * FROM panel_kullanicilari WHERE kullanici_adi=?',(session['kad'],)).fetchone();u=dict(row) if row else None
        if u and u['aktif'] and session.get('sv')==u['session_version']:
            g.k=u
        else:
            session.clear()
    g.yk=yetki.coz(g.k['yetki'] if g.k else '')
    session.setdefault('csrf',secrets.token_hex(24))
    if request.method in ('POST','PUT','PATCH','DELETE'):
        token=request.form.get('csrf') or request.headers.get('X-CSRF','')
        if not secrets.compare_digest(token,session.get('csrf','')):
            abort(400,'Geçersiz CSRF. Sayfayı yenileyin.')

@app.after_request
def headers(response):
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['X-Frame-Options']='SAMEORIGIN'
    response.headers['Referrer-Policy']='same-origin'
    response.headers['Cache-Control']='private, no-store' if request.endpoint!='static' else 'public, max-age=3600'
    if secure:
        response.headers['Strict-Transport-Security']='max-age=31536000'
    return response

@app.context_processor
def context():
    allowed=_allowed() if g.get('k') else []
    info=[p for p in _all_project_info() if p['code'] in allowed] if allowed else []
    safe_user={a:b for a,b in (g.get('k') or {}).items() if a!='sifre_hash'}
    return dict(k=safe_user if g.get('k') else None,yk=g.get('yk'),renk=renk,metin_rengi=metin_rengi,
                url_degistir=url_degistir,url_filtre=url_filtre,csrf=session.get('csrf',''),projeler=lambda:allowed,
                proje_renkleri={p['code']:p['color'] for p in info},proje_bilgileri=info,
                izin=lambda p,e:yetki.izin(g.yk,p,e),simdi=simdi,SURUM=SURUM,
                rol_adi=('Süperadmin' if g.get('yk') and g.yk['admin'] else 'Admin') if g.get('k') else '')


def giris_gerekli(f):
    @wraps(f)
    def wrapper(*a,**kw):
        if not g.k:
            return redirect(url_for('giris',sonraki=request.path))
        if g.k['sifre_degistir'] and request.endpoint not in ('sifre_degistir','cikis'):
            return redirect(url_for('sifre_degistir'))
        return f(*a,**kw)
    return wrapper

def admin_gerekli(f):
    @giris_gerekli
    @wraps(f)
    def wrapper(*a,**kw):
        if not g.yk['admin']: abort(403)
        return f(*a,**kw)
    return wrapper

def kullanici_adi_gecerli(ad):
    return bool(re.fullmatch(r'[a-z0-9_.]{1,40}',ad or ''))

@app.route('/giris',methods=['GET','POST'])
def giris():
    hata=None
    if request.method=='POST':
        ad=request.form.get('kullanici','').strip().lower()
        c=baglanti()
        # IP bazlı kalıcı sınır; bilinmeyen kullanıcı adlarıyla da aşılmaz.
        count=c.execute("SELECT COUNT(*) FROM denetim WHERE islem='giris_hata' AND ip=? AND zaman>?",(ip(),(datetime.now()-timedelta(minutes=15)).strftime('%Y-%m-%d %H:%M:%S'))).fetchone()[0]
        if count>=20:
            return render_template('giris.html',hata='Çok fazla giriş denemesi. 15 dakika bekleyin.'),429
        u,hata=kullanicilar.dogrula(AYAR,ad,request.form.get('sifre',''))
        if u:
            session.clear(); session.permanent=True
            session.update(kad=u['kullanici_adi'],sv=u['session_version'],csrf=secrets.token_hex(24))
            db.denetim_yaz(c,ad,'giris',ip=ip())
            nxt=request.args.get('sonraki','/')
            if not nxt.startswith('/') or nxt.startswith('//') or '\\' in nxt: nxt='/'
            return redirect(nxt)
        db.denetim_yaz(c,ad,'giris_hata',ip=ip())
        time.sleep(.3)
    if not kullanicilar.var_mi(AYAR):
        hata='Henüz panel hesabı oluşturulmadı. Sunucuda python -m app.cli admin-ekle komutunu çalıştırın.'
    return render_template('giris.html',hata=hata)

@app.route('/cikis',methods=['POST'])
def cikis():
    session.clear()
    return redirect(url_for('giris'))

@app.route('/sifre',methods=['GET','POST'])
@giris_gerekli
def sifre_degistir():
    hata=None
    if request.method=='POST':
        u,_=kullanicilar.dogrula(AYAR,g.k['kullanici_adi'],request.form.get('eski',''))
        s=request.form.get('yeni','')
        if not u: hata='Eski parola hatalı.'
        elif len(s)<8: hata='Parola en az 8 karakter olmalı.'
        elif s!=request.form.get('yeni2'): hata='Yeni parolalar aynı değil.'
        else:
            u=kullanicilar.guncelle(AYAR,g.k['kullanici_adi'],sifre=s,sifre_degistir=0)
            session['sv']=u['session_version']
            db.denetim_yaz(baglanti(),g.k['kullanici_adi'],'sifre_degisti',ip=ip())
            flash('Parola değiştirildi.'); return redirect(url_for('ana'))
    return render_template('sifre.html',hata=hata)

def durum_bilgisi():
    c=baglanti(); raw=db.ayar_al(c,'son_ice_aktarim')
    return {'son_ice_aktarim':json.loads(raw) if raw else None,'bot':bot_yonetim.bot_calisiyor_mu(AYAR),
            'log_son':db.ayar_al(c,'son_degisiklik'),'kuyruk_bekleyen':0,'bekleyen_islem':0,
            'surum':SURUM,'veritabani':'SQLite WAL','gorsel_klasor':'Özel yerel arşiv','paylasim_modu':False}

def _acik_sorular(allowed):
    """Bana yöneltilmiş ve henüz yanıtlanmamış soru sayısı (yalnız izinli projeler)."""
    if not allowed:return 0
    places=','.join('?' for _ in allowed)
    return baglanti().execute('''SELECT COUNT(*) FROM yorumlar y JOIN kayitlar k ON k.id=y.kayit_id
        WHERE y.tur='soru' AND y.alici=? AND k.proje IN ('''+places+''')
        AND NOT EXISTS (SELECT 1 FROM yorumlar r WHERE r.yanit_id=y.id)''',(g.k['kullanici_adi'],*allowed)).fetchone()[0]

def _son_islemler(allowed,limit=10):
    """İşlem Geçmişi ile aynı kapsam kuralı: süperadmin her şeyi, admin yalnız izinli kayıt hareketlerini görür."""
    from .guvenli_log import log_row
    if g.yk['admin']:where,params='1=1',[]
    elif allowed:where,params='k.proje IN ('+','.join('?' for _ in allowed)+')',list(allowed)
    else:return []
    rows=[]
    for row in baglanti().execute('SELECT d.*,k.id record_id FROM denetim d LEFT JOIN kayitlar k ON k.anahtar=d.anahtar WHERE '+where+' ORDER BY d.id DESC LIMIT ?',[*params,limit]):
        item=log_row(row)
        if item.get('proje') and item['proje'] not in allowed and not g.yk['admin']:continue
        rows.append(item)
    return rows

ISLEM_ADLARI={'giris':'Giriş yaptı','giris_hata':'Hatalı giriş denemesi','onayla':'Kaydı onayladı','onay_kaldir':'Onayı kaldırdı','duzenle':'Kaydı düzenledi','sil':'Kaydı sildi','geri_al':'Kaydı geri aldı','yorum_ekle':'Yorum yazdı','soru_yonelt':'Soru yöneltti','yanit_ekle':'Yanıt verdi','gyazo_istek':'Gyazo dönüşümü istedi','disa_aktar':'Excel çıktısı aldı','dosya_kontrol':'Dosya kontrolü yaptı','kullanici_ekle':'Kullanıcı ekledi','kullanici_guncelle':'Kullanıcıyı güncelledi','sifre_degisti':'Parolasını değiştirdi','arsiv_esle':'Arşiv görseli eşledi','beklenen_proje':'Beklenen projeyi atadı','yeni_kayit':'Yeni kanıt ekledi','sunucu_hata':'Sunucu hatası'}

@app.route('/')
@giris_gerekli
def ana():
    cards=[];allowed=_allowed();kpi=dict(bekleyen=0,bugun=0,bugun_onay=0,toplam=0,onayli=0)
    for p in allowed:
        row=baglanti().execute('''SELECT SUM(kt=0 AND silindi=0) bekleyen,
            SUM(kt=1 AND silindi=0) onayli,SUM(silindi=1) silinen,
            SUM(date(tarih)=date('now','localtime') AND silindi=0) bugun,
            SUM(kt=1 AND silindi=0 AND date(kt_zaman)=date('now','localtime')) bugun_onay,COUNT(*) toplam
            FROM kayitlar WHERE proje=?''',(p,)).fetchone()
        card={'proje':p,**{k:row[k] or 0 for k in row.keys()}}
        card['etiket']=next((x['label'] for x in _all_project_info() if x['code']==p),p)
        cards.append(card)
        for key in kpi:kpi[key]+=card[key]
    kpi['acik_soru']=_acik_sorular(allowed)
    return render_template('ana.html',kartlar=cards,durum=durum_bilgisi(),kpi=kpi,son_islemler=_son_islemler(allowed),islem_adlari=ISLEM_ADLARI)

def _filtre_sql(args,izinli):
    selected=_selected_projects(args,izinli)
    if not selected: return '1=0',[],{}
    where=['proje IN (%s)' % ','.join('?'*len(selected))]; par=list(selected)
    state=args.get('durum','bekleyen')
    if state=='bekleyen': where.append('kt=0 AND silindi=0')
    elif state=='onayli': where.append('kt=1 AND silindi=0')
    elif state=='silinen': where.append('silindi=1')
    elif state!='hepsi': state='bekleyen'; where.append('kt=0 AND silindi=0')
    cat=args.get('kategori','').strip()
    if cat: where.append('kategori COLLATE UNICODE_NOCASE=?'); par.append(cat)
    q=args.get('q','').strip()[:200]
    if q:
        where.append('(personel LIKE ? OR uye_id LIKE ? OR ana_uye_id LIKE ? OR lead_id LIKE ? OR link LIKE ?)'); par += ['%'+q+'%']*5
    by=args.get('ekleyen','').strip()[:100]
    if by: where.append('(ekleyen_ad LIKE ? OR ekleyen_id=?)'); par += ['%'+by+'%',by]
    dates={}
    undated=args.get('tarihsiz')=='1'
    if undated:
        if args.get('tarih1') or args.get('tarih2'): abort(400,'Tarihsiz filtre tarih aralığıyla birleştirilemez.')
        where.append("(tarih IS NULL OR tarih='')")
    for name,op,suffix in [('tarih1','>=',' 00:00:00'),('tarih2','<=',' 23:59:59')]:
        v=args.get(name,'').strip()
        if v:
            try: datetime.strptime(v,'%Y-%m-%d')
            except ValueError: abort(400,'Geçersiz tarih.')
            where.append(f'tarih {op} ?'); par.append(v+suffix)
        dates[name]=v
    if dates.get('tarih1') and dates.get('tarih2') and dates['tarih1']>dates['tarih2']:abort(400)
    return ' AND '.join(where),par,dict(proje=selected,durum=state,kategori=cat,q=q,ekleyen=by,**dates)

def kayit_izin(r,eylem):
    return yetki.izin(g.yk,r['proje'],eylem)

def kayit_sozluk(row):
    d=dict(row); path,source=COZUCU.yerel_dosya(d)
    d.update(gorsel_var=bool(path),gorsel_kaynak=source,renk=renk(d['proje']),
             izin_gyazo=kayit_izin(d,'kt') and kayit_izin(d,'gorsel') and bool(path) and not d['silindi'],izin_kt=kayit_izin(d,'kt'),izin_yorum=kayit_izin(d,'kt'),izin_duzenle=kayit_izin(d,'duzenle'),izin_sil=kayit_izin(d,'sil'),geri_al_olur=bool(d['silindi']))
    link=d.get('link') or ''; d['dis_link']=link if re.match(r'^https?://',link) else None
    for name in ('file_id','file_unique_id','chat_id','gorsel_dosya','gorsel2_dosya'):
        d.pop(name,None)
    return d

@app.route('/kayitlar')
@giris_gerekli
def kayitlar():
    c=baglanti(); allowed=yetki.izinli_projeler(g.yk,projeler()); where,par,f=_filtre_sql(request.args,allowed)
    try: page=max(1,int(request.args.get('s','1')))
    except ValueError: abort(400,'Geçersiz sayfa.')
    size=AYAR['SAYFA_BOYU']; total=c.execute(f'SELECT COUNT(*) FROM kayitlar WHERE {where}',par).fetchone()[0]
    rows=c.execute(f"SELECT * FROM kayitlar WHERE {where} ORDER BY COALESCE(tarih,'0000') DESC,id DESC LIMIT ? OFFSET ?",(*par,size,(page-1)*size)).fetchall()
    cats=_scoped_categories(f.get('proje',allowed))
    counts={}
    for s in ('bekleyen','onayli','silinen','hepsi'):
        a=request.args.to_dict(flat=False); a['durum']=[s]; w,p,_=_filtre_sql(MultiDict(a),allowed)
        counts[s]=c.execute(f'SELECT COUNT(*) FROM kayitlar WHERE {w}',p).fetchone()[0]
    return render_template('kayitlar.html',kayitlar=[kayit_sozluk(r) for r in rows],toplam=total,sayfa=page,sayfa_sayisi=max(1,(total+size-1)//size),f=f,izinli=allowed,kategoriler=cats,gorunum=request.args.get('gorunum','tablo'),sayac=counts)

@app.route('/kayit/<int:kid>')
@giris_gerekli
def kayit_json(kid):
    r=_kayit_al(kid,'oku'); d=kayit_sozluk(r)
    d['gorsel_url']=url_for('gorsel_goster',kid=kid) if kayit_izin(r,'gorsel') else None
    d['kucuk_url']=url_for('kucuk_goster',kid=kid) if kayit_izin(r,'gorsel') else None
    second={'gorsel_dosya':r['gorsel2_dosya'],'link':r['gorsel2'],'proje':r['proje'],'anahtar':r['anahtar']+':2'}
    d['gorsel2_url']=url_for('gorsel_goster',kid=kid,ikinci=1) if kayit_izin(r,'gorsel') and COZUCU.yerel_dosya(second)[0] else None
    from . import discussion
    d['yorumlar']=[dict(x) for x in baglanti().execute('SELECT id,kullanici,metin,zaman,yanit_id,alici,tur FROM yorumlar WHERE kayit_id=? ORDER BY id',(kid,))]
    d['yorum_alicilari']=discussion.recipients(baglanti(),r['proje'],g.k['kullanici_adi']) if d['izin_yorum'] else []
    from .guvenli_log import log_row
    d['soru_cevaplar']=_question_answers(r)
    d['gyazo_hazir']=bool(os.environ.get('GYAZO_TOKEN','').strip())
    job=baglanti().execute('SELECT durum,expected_gorsel FROM donusum_isleri WHERE record_id=? ORDER BY id DESC LIMIT 1',(kid,)).fetchone()
    d['gyazo_bekliyor']=bool(job and job['durum'] not in ('tamam','hata','belirsiz','iptal') and job['expected_gorsel']==r['gorsel_dosya'])
    d['gecmis']=[log_row(x) for x in baglanti().execute('SELECT zaman,kullanici,islem,eski,yeni FROM denetim WHERE anahtar=? ORDER BY id DESC LIMIT 30',(r['anahtar'],))]
    return jsonify(d)

def _kayit_al(kid,action):
    r=baglanti().execute('SELECT * FROM kayitlar WHERE id=?',(kid,)).fetchone()
    if not r or not kayit_izin(r,action): abort(403)
    return r

def version_check(r):
    v=request.form.get('surum')
    if v is not None and str(r['surum'])!=v:
        abort(409,'Kayıt başka kullanıcı tarafından değiştirildi. Yenileyin.')

@app.route('/kayit/<int:kid>/kt',methods=['POST'])
@giris_gerekli
def kayit_kt(kid):
    c=baglanti()
    with db.transaction(c):
        r=_kayit_al(kid,'kt'); version_check(r)
        if r['silindi']: abort(400,'Silinen kayıt onaylanamaz.')
        v=1 if request.form.get('deger','1')=='1' else 0
        c.execute('UPDATE kayitlar SET kt=?,kt_kaynak=\'panel\',kt_kim=?,kt_zaman=?,surum=surum+1,guncelleme=? WHERE id=?',(v,g.k['kullanici_adi'],simdi(),simdi(),kid))
        from . import discussion
        discussion.approval_notice(c,r,g.k['kullanici_adi'],v,r['surum']+1)
        db.denetim_yaz(c,g.k['kullanici_adi'],'onayla' if v else 'onay_kaldir',r['proje'],r['anahtar'],r['kt'],v,ip()); db.degisti(c)
    return jsonify(ok=True,kt=v)

def _gyazo_result(job,record):
    from .guvenli_log import text
    if not job:return dict(ok=True,pending=False,converted=False,mesaj='Bu kayıt için dönüşüm başlatılmadı.')
    pending=job['durum'] not in ('tamam','hata','belirsiz','iptal')
    converted=bool(job['durum']=='tamam' and job['url'] and record['link']==job['url'] and record['link_turu']=='gyazo' and job['expected_gorsel']==record['gorsel_dosya'])
    if pending:message='Dönüşüm arka planda başlatıldı. Yerel kanıt korunuyor.'
    elif converted:message='Gyazo bağlantısı oluşturuldu ve kayda eklendi.'
    elif job['durum']=='tamam':message='Görsel dönüştü ancak kayıt görseli değişti; mevcut kanıt değiştirilmedi.'
    else:message=text(job['hata'] or 'Dönüşüm tamamlanamadı.')+' Yerel kanıt korunuyor.'
    return dict(ok=True,pending=pending,converted=converted,url=job['url'] if converted else None,mesaj=message)

@app.route('/kayit/<int:kid>/gyazo',methods=['POST'])
@giris_gerekli
def kayit_gyazo(kid):
    import hashlib
    from . import donusum_isleri
    c=baglanti();record=_kayit_al(kid,'kt');version_check(record)
    if not kayit_izin(record,'gorsel') or record['silindi']:abort(403)
    if not os.environ.get('GYAZO_TOKEN','').strip():return jsonify(ok=False,hata='Gyazo anahtarı yapılandırılmadı. Süperadmin sunucudaki gizli ayardan eklemeli.'),400
    path,_=COZUCU.yerel_dosya(dict(record))
    if not path:return jsonify(ok=False,hata='Bu kaydın eşleşmiş yerel görseli yok.'),400
    fingerprint=hashlib.sha256()
    with open(path,'rb') as image:
        for part in iter(lambda:image.read(1048576),b''):fingerprint.update(part)
    with db.transaction(c):
        current=_kayit_al(kid,'kt');version_check(current)
        job=donusum_isleri.enqueue(c,kid,g.k['kullanici_adi'],local_path=path,job_key='web:'+str(kid)+':'+fingerprint.hexdigest())
        db.denetim_yaz(c,g.k['kullanici_adi'],'gyazo_istek',current['proje'],current['anahtar'],yeni={'is':job['id'],'durum':job['durum']},ip=ip(),hedef_turu='kayit',hedef_id=kid)
    return jsonify(_gyazo_result(job,current))

@app.route('/kayit/<int:kid>/gyazo-durum')
@giris_gerekli
def kayit_gyazo_durum(kid):
    record=_kayit_al(kid,'kt')
    job=baglanti().execute('SELECT * FROM donusum_isleri WHERE record_id=? ORDER BY id DESC LIMIT 1',(kid,)).fetchone()
    return jsonify(_gyazo_result(job,record))

@app.route('/kayit/<int:kid>/duzenle',methods=['POST'])
@giris_gerekli
def kayit_duzenle(kid):
    c=baglanti()
    with db.transaction(c):
        r=_kayit_al(kid,'duzenle'); version_check(r)
        if r['silindi']: abort(400,'Önce silinen kaydı geri alın.')
        vals={k:request.form[k].strip()[:500] for k in ('personel','uye_id','ana_uye_id','lead_id','tip','kategori','aciklama') if k in request.form}
        if 'kategori' in vals and not g.yk['admin'] and vals['kategori'] not in _scoped_categories([r['proje']]):abort(403)
        if r['cevaplar_json']:
            try:answers=json.loads(r['cevaplar_json'])
            except (ValueError,TypeError):answers={}
            if not isinstance(answers,dict):answers={}
            previous_answers=dict(answers)
            questions=_snapshot_questions(r['soru_anlik_json'])
            for question in questions:
                if question.get('preset') in vals:answers[question['key']]=vals[question['preset']]
            if answers!=previous_answers:vals['cevaplar_json']=json.dumps(answers,ensure_ascii=False)
        vals={k:v for k,v in vals.items() if (r[k] or '')!=v}
        if vals:
            c.execute('UPDATE kayitlar SET '+','.join(f'{k}=?' for k in vals)+',surum=surum+1,guncelleme=? WHERE id=?',(*vals.values(),simdi(),kid))
            db.denetim_yaz(c,g.k['kullanici_adi'],'duzenle',r['proje'],r['anahtar'],json.dumps({k:r[k] for k in vals},ensure_ascii=False),json.dumps(vals,ensure_ascii=False),ip()); db.degisti(c)
    return jsonify(ok=True,degisiklik=bool(vals))

@app.route('/kayit/<int:kid>/sil',methods=['POST'])
@giris_gerekli
def kayit_sil(kid):
    return _delete_state(kid,1)

@app.route('/kayit/<int:kid>/geri-al',methods=['POST'])
@giris_gerekli
def kayit_geri_al(kid):
    return _delete_state(kid,0)

def _delete_state(kid,value):
    c=baglanti()
    with db.transaction(c):
        r=_kayit_al(kid,'sil'); version_check(r)
        c.execute('UPDATE kayitlar SET silindi=?,silen=?,silme_zamani=?,surum=surum+1,guncelleme=? WHERE id=?',(value,g.k['kullanici_adi'] if value else None,simdi() if value else None,simdi(),kid))
        db.denetim_yaz(c,g.k['kullanici_adi'],'sil' if value else 'geri_al',r['proje'],r['anahtar'],r['silindi'],value,ip()); db.degisti(c)
    return jsonify(ok=True)

def _gorsel_yolu(kid,kucuk=False):
    r=_kayit_al(kid,'gorsel'); d=dict(r)
    if request.args.get('ikinci')=='1': d.update(gorsel_dosya=r['gorsel2_dosya'],link=r['gorsel2'],anahtar=r['anahtar']+':2')
    path,_=COZUCU.yerel_dosya(d)
    if not path: abort(404,'Bu kaydın yerel görseli eşleşmemiş.')
    return COZUCU.kucuk(d,path) if kucuk else path

@app.route('/gorsel/<int:kid>')
@giris_gerekli
def gorsel_goster(kid):
    path=_gorsel_yolu(kid)
    if not path: abort(404)
    return send_file(path,conditional=True,as_attachment=request.args.get('indir')=='1')

@app.route('/kucuk/<int:kid>')
@giris_gerekli
def kucuk_goster(kid):
    path=_gorsel_yolu(kid,True)
    if not path: abort(404)
    return send_file(path,conditional=True)

def _role_from_form():
    if 'rol' not in request.form and request.form.get('yetki'):
        return yetki.kanonikle(request.form['yetki'],projeler())
    if request.form.get('rol','admin')=='superadmin':return 'admin'
    if request.form.get('rol','admin')!='admin':raise ValueError('Geçersiz rol.')
    selected=request.form.getlist('projeler')
    if not selected or any(p not in projeler() for p in selected):raise ValueError('Admin için en az bir tanımlı proje seçin.')
    rules=[]
    for code in dict.fromkeys(selected):
        actions=set(request.form.getlist('eylemler_'+code))
        if not actions.issubset({'kt','duzenle','sil'}):raise ValueError('Geçersiz eylem.')
        actions.update(('oku','gorsel'))
        rules.append(code+':'+','.join(a for a in yetki.EYLEMLER if a in actions))
    return yetki.kanonikle(';'.join(rules),projeler())

@app.route('/yonetim/kullanicilar',methods=['GET','POST'])
@admin_gerekli
def yonetim_kullanicilar():
    mesaj=None
    if request.method=='POST':
        try:
            op=request.form.get('islem')
            ad=(request.form.get('kullanici_adi') or request.form.get('kullanici') or request.form.get('ad','')).strip().lower()
            expected=request.form.get('expected_version',request.form.get('sv','-1'))
            actor=g.k['kullanici_adi']
            if op=='ekle':
                password=request.form.get('sifre') or secrets.token_urlsafe(15)
                kullanicilar.ekle(AYAR,ad,_role_from_form(),password,request.form.get('gorunen_ad') or ad,actor)
                mesaj='Hesap açıldı: '+ad+' — ilk girişte değiştirilecek parola: '+password
            elif op=='guncelle':
                fields=dict(yetki=_role_from_form(),aktif=1 if request.form.get('aktif')=='1' else 0,
                            sifre_degistir=1 if request.form.get('sifredegistir')=='1' else 0,
                            ad=request.form.get('gorunen_ad') or ad)
                if ad==actor and not fields['aktif']:raise ValueError('Kendi hesabınızı kapatamazsınız.')
                if request.form.get('sifre'):fields['sifre']=request.form['sifre']
                kullanicilar.guncelle(AYAR,ad,expected_version=expected,islem_yapan=actor,**fields)
                mesaj='Hesap ve açık proje kapsamı kaydedildi.'
            elif op=='yetki':
                kullanicilar.guncelle(AYAR,ad,expected_version=expected,islem_yapan=actor,yetki=_role_from_form());mesaj='Yetki güncellendi.'
            elif op=='sifre':
                password=secrets.token_urlsafe(15)
                kullanicilar.guncelle(AYAR,ad,expected_version=expected,islem_yapan=actor,sifre=password,sifre_degistir=1);mesaj=ad+' geçici parola: '+password
            elif op=='pasif':
                if ad==actor:raise ValueError('Kendi hesabınızı kapatamazsınız.')
                user=kullanicilar.bul(AYAR,ad)
                if not user:raise ValueError('Hesap bulunamadı.')
                kullanicilar.guncelle(AYAR,ad,expected_version=expected,islem_yapan=actor,aktif=0 if user['aktif'] else 1);mesaj='Hesap durumu değişti.'
            else:raise ValueError('Bilinmeyen işlem.')
            flash(mesaj);return redirect(url_for('yonetim_kullanicilar'))
        except ValueError as e:mesaj='Hata: '+str(e)
    users=[]
    for user in kullanicilar.hepsi(AYAR):
        parsed=yetki.coz(user['yetki'])
        safe={k:v for k,v in user.items() if k!='sifre_hash'}
        safe.update(ozet=yetki.ozet(user['yetki']),parsed_yetkiler=dict(rol='superadmin' if parsed['admin'] else 'admin',
                    projeler=list(parsed['projeler']),eylemler={p:list(a) for p,a in parsed['projeler'].items()}))
        users.append(safe)
    return render_template('kullanicilar.html',kullanicilar=users,mesaj=mesaj,tum_projeler=projeler())


@app.route('/bot',methods=['GET','POST'])
@admin_gerekli
def bot_sayfasi():
    c=baglanti(); mesaj=None
    if request.method=='POST':
        op=request.form.get('islem'); tg=request.form.get('tg_id','').strip(); ad=request.form.get('ad',''); kim=g.k['kullanici_adi']
        try:
            if op=='onayla': bot_yonetim.onayla(AYAR,c,tg,ad or tg,kim)
            elif op=='onay_kaldir': bot_yonetim.onay_kaldir(AYAR,c,tg,kim)
            elif op=='engelle': bot_yonetim.engelle(AYAR,c,tg,ad,request.form.get('sebep') or 'panel',kim)
            elif op=='engel_kaldir': bot_yonetim.engel_kaldir(AYAR,c,tg,kim)
            elif op in ('bakim_ac','bakim_kapat'): bot_yonetim.bakim_modu_ayarla(AYAR,c,op=='bakim_ac',kim)
            elif op=='oturum_temizle': bot_yonetim.oturum_temizle(AYAR,c,tg or None,kim)
            elif op=='kategori_ekle':
                from . import katalog
                katalog.save_category(c,request.form.get('code',''),request.form.get('label',''),mode=request.form.get('mode','tek_uye'),projects=projeler(),questions=None,actor=kim)
            elif op=='kategori_sil': bot_yonetim.kategori_sil(AYAR,c,request.form.get('code',''),kim)
            elif op=='yedek': bot_yonetim.yedek_al(AYAR,c,kim)
            elif op=='tam_sifirla': bot_yonetim.tam_sifirla(AYAR,c,tg,kim)
            elif op in ('bot_duraklat','bot_surdur'): bot_yonetim.bot_kontrol(AYAR,c,op=='bot_duraklat',kim)
            elif op=='kategori_yayinla': mesaj='Kategori değişiklikleri kaydedildiği anda geçerlidir.'
            elif op=='inbox_tekrar': bot_yonetim.tekrar_dene(AYAR,c,request.form.get('update_id') or None,kim)
            elif op=='yenile': pass
            else: raise ValueError('Bilinmeyen işlem.')
            mesaj='İşlem sunucuya uygulandı.'
        except (ValueError,OSError) as e: mesaj='Hata: '+str(e)
    return render_template('bot.html',mesaj=mesaj,durum=durum_bilgisi(),istatistik=bot_yonetim.istatistik(c),kullanici_q=request.args.get('q',''),kullanici_durum=request.args.get('tg_durum','hepsi'),kullanici_sure=request.args.get('sure',''),tg_kullanicilar=bot_yonetim.kullanicilari_ara(c,request.args.get('q',''),request.args.get('tg_durum','hepsi'),request.args.get('sure',''),limit=500),olu_mektuplar=[dict(r) for r in c.execute("SELECT bot_id,update_id,deneme,hata,guncelleme FROM tg_gelen WHERE durum='olu' ORDER BY guncelleme DESC LIMIT 50")],kategoriler=bot_yonetim.kategoriler(AYAR),bakim=bot_yonetim.bakim_modu(AYAR),oturumlar=bot_yonetim.oturumlar(AYAR),log=bot_yonetim.son_log(AYAR,150),hatalar=bot_yonetim.son_log(AYAR,40,True),cfg=bot_yonetim.config(AYAR),islemler=bot_yonetim.son_islemler(c))

@app.route('/bot/islemler.json')
@admin_gerekli
def bot_islemler_json():
    return jsonify(islemler=bot_yonetim.son_islemler(baglanti()),bekleyen=0)

@app.route('/denetim')
@giris_gerekli
def denetim():
    from .guvenli_log import log_row
    allowed=_allowed();where=['1=1'];params=[]
    if not g.yk['admin']:
        if not allowed:abort(403)
        where.append('k.proje IN ('+','.join('?' for _ in allowed)+')');params.extend(allowed)
    filters={key:request.args.get(key,'')[:100] for key in ('q','kullanici','islem','proje','tarih1','tarih2')}
    if filters['proje']:
        if filters['proje'] not in allowed:abort(403)
        where.append('d.proje=?');params.append(filters['proje'])
    for key in ('kullanici','islem'):
        if filters[key]:where.append('d.'+key+' LIKE ?');params.append('%'+filters[key]+'%')
    if filters['q']:
        where.append('(d.kullanici LIKE ? OR d.islem LIKE ? OR d.anahtar LIKE ?)');params.extend(['%'+filters['q']+'%']*3)
    for key,op,tail in (('tarih1','>=',' 00:00:00'),('tarih2','<=',' 23:59:59')):
        if filters[key]:
            try:datetime.strptime(filters[key],'%Y-%m-%d')
            except ValueError:abort(400)
            where.append('d.zaman '+op+' ?');params.append(filters[key]+tail)
    try:page=max(1,min(1000000,int(request.args.get('s','1'))))
    except ValueError:abort(400)
    source=' FROM denetim d LEFT JOIN kayitlar k ON k.anahtar=d.anahtar WHERE '+' AND '.join(where)
    total=baglanti().execute('SELECT COUNT(*)'+source,params).fetchone()[0]
    rows=[]
    for row in baglanti().execute('SELECT d.*,k.id record_id'+source+' ORDER BY d.id DESC LIMIT 50 OFFSET ?',[*params,(page-1)*50]):
        item=log_row(row)
        if item['record_id']:item.update(hedef_turu='kayit',hedef_id=str(item['record_id']))
        if 'hata' in item['islem']:item['sonuc']='hata'
        item['sonuc']={'ok':'Başarılı','hata':'Hata','uyari':'Uyarı'}.get(item.get('sonuc'),item.get('sonuc') or 'Tamamlandı')
        rows.append(item)
    return render_template('denetim.html',logs=rows,kayitlar=rows,toplam=total,sayfa=page,sayfa_sayisi=max(1,(total+49)//50),filters=filters,q=filters['q'])


@app.route('/durum.json')
@admin_gerekli
def durum_json():
    return jsonify(durum_bilgisi())

def _scoped_counts():
    allowed=_allowed()
    if not allowed:return {},'0'
    places=','.join('?' for _ in allowed)
    rows=baglanti().execute('SELECT proje,SUM(kt=0 AND silindi=0) bekleyen,SUM(kt=1 AND silindi=0) onayli,SUM(silindi=1) silinen,COUNT(*) toplam,SUM(surum) version_sum FROM kayitlar WHERE proje IN ('+places+') GROUP BY proje',allowed).fetchall()
    counts={p:dict(bekleyen=0,onayli=0,silinen=0,toplam=0) for p in allowed}
    stamp=[]
    for row in rows:
        counts[row['proje']]={key:row[key] or 0 for key in ('bekleyen','onayli','silinen','toplam')}
        stamp.append([row['proje'],row['toplam'],row['version_sum'],row['onayli'],row['silinen']])
    comments=baglanti().execute('SELECT COUNT(*),MAX(y.id) FROM yorumlar y JOIN kayitlar k ON k.id=y.kayit_id WHERE k.proje IN ('+places+')',allowed).fetchone()
    stamp.append(list(comments))
    return counts,json.dumps(stamp,separators=(',',':'))

@app.route('/bildirim.json')
@giris_gerekli
def bildirim_json():
    from . import bildirim_gecmisi
    data=bildirim_gecmisi.liste(baglanti(),g.k['kullanici_adi'],_allowed(),{'durum':'okunmamis'},75)
    counts,stamp=_scoped_counts()
    alerts=[]
    for row in data['bildirimler']:
        item=dict(row);item['bildirim_id']='m:'+str(row['id']) if row['key'].startswith('mesaj:') else str(row['id']);item['id']=row['kayit_id']
        alerts.append(item)
    return jsonify(bildirimler=alerts,okunmamis=data['counts']['okunmamis'],sayac=counts,degisiklik=stamp,kuyruk_bekleyen=0)

@app.route('/bildirimler')
@giris_gerekli
def bildirim_sayfasi():
    from . import bildirim_gecmisi
    filters=request.args.to_dict()
    if filters.get('okundu') in ('0','1'):filters['durum']='okunmamis' if filters['okundu']=='0' else 'okundu'
    elif filters.get('okundu'):abort(400)
    if filters.get('tur')=='kayit':filters['tur']='kanit'
    if 'sayfa' in filters and 's' not in filters:filters['s']=filters['sayfa']
    try:data=bildirim_gecmisi.liste(baglanti(),g.k['kullanici_adi'],_allowed(),filters)
    except PermissionError:abort(403)
    except ValueError:abort(400)
    data['filters']['okundu']='0' if data['filters']['durum']=='okunmamis' else '1' if data['filters']['durum']=='okundu' else ''
    return render_template('bildirimler.html',**data)

@app.route('/bildirimler/okundu',methods=['POST'])
@giris_gerekli
def bildirim_okundu():
    from . import bildirim_gecmisi
    payload=request.get_json(silent=True) if request.is_json else None
    payload=payload or {}
    all_visible=payload.get('hepsi') in (True,1,'1') or request.form.get('hepsi')=='1'
    ids=payload.get('ids') if request.is_json else request.form.getlist('ids')
    if not ids and request.form.get('id'):ids=[request.form['id']]
    if ids is None:ids=[]
    if not isinstance(ids,list) or (not all_visible and not ids):abort(400)
    try:changed=bildirim_gecmisi.okundu(baglanti(),g.k['kullanici_adi'],_allowed(),ids,all_visible)
    except ValueError:abort(400)
    return jsonify(ok=True,degisen=changed)


@app.route('/kayit/<int:kid>/yorum',methods=['POST'])
@giris_gerekli
def kayit_yorum(kid):
    from . import discussion
    c=baglanti();_kayit_al(kid,'kt')
    try:
        cid=discussion.add(c,kid,g.k['kullanici_adi'],request.form.get('yorum',''),request.form.get('alici',''),request.form.get('yanit_id') or None)
    except PermissionError:abort(403)
    except ValueError as e:return jsonify(ok=False,hata=str(e)),400
    return jsonify(ok=True,yorum_id=cid)

@app.route('/arsiv',methods=['GET','POST'])
@admin_gerekli
def arsiv():
    c=baglanti()
    if request.method=='POST':
        if request.form.get('islem')=='tara':
            report=gorsel.arsiv_tara(AYAR,c,COZUCU); flash('Arşiv tarandı: '+json.dumps(report,ensure_ascii=False))
        elif request.form.get('islem')=='esle':
            try: aid=int(request.form.get('arsiv_id',0)); kid=int(request.form.get('kayit_id',0))
            except ValueError: abort(400)
            with db.transaction(c):
                a=c.execute('SELECT * FROM arsiv_dosyalar WHERE id=?',(aid,)).fetchone(); r=_kayit_al(kid,'duzenle')
                if not a: abort(404)
                path=Path(a['yol']).resolve(); root=Path(AYAR['ARSIV_KLASOR']).resolve()
                if not path.is_relative_to(root) or not path.is_file(): abort(400)
                column='gorsel2_dosya' if request.form.get('ikinci')=='1' else 'gorsel_dosya'
                if r[column] and request.form.get('degistir')!='1': abort(409,'Zaten görsel var; değiştirme onayı gerekli.')
                c.execute(f'UPDATE kayitlar SET {column}=?,surum=surum+1,guncelleme=? WHERE id=?',(str(path),simdi(),kid))
                c.execute('UPDATE arsiv_dosyalar SET kayit_id=?,guncelleme=? WHERE id=?',(kid,simdi(),aid))
                db.denetim_yaz(c,g.k['kullanici_adi'],'arsiv_esle',r['proje'],r['anahtar'],r[column],a['ad'],ip()); db.degisti(c)
            flash('Görsel kayıtla ilişkilendirildi.')
        return redirect(url_for('arsiv'))
    try: page=max(1,int(request.args.get('s',1)))
    except ValueError: page=1
    q=request.args.get('q','')[:100]; state=request.args.get('durum','eslesmeyen')
    where='ad LIKE ?'+(' AND kayit_id IS NULL' if state=='eslesmeyen' else '')
    total=c.execute(f'SELECT COUNT(*) FROM arsiv_dosyalar WHERE {where}',('%'+q+'%',)).fetchone()[0]
    rows=c.execute(f'SELECT * FROM arsiv_dosyalar WHERE {where} ORDER BY id LIMIT 48 OFFSET ?',('%'+q+'%',(page-1)*48))
    return render_template('arsiv.html',dosyalar=[dict(r) for r in rows],toplam=total,sayfa=page,sayfa_sayisi=max(1,(total+47)//48),q=q,durum=state)

@app.route('/arsiv/gorsel/<int:aid>')
@admin_gerekli
def arsiv_gorsel(aid):
    r=baglanti().execute('SELECT yol FROM arsiv_dosyalar WHERE id=?',(aid,)).fetchone()
    if not r: abort(404)
    path=Path(r['yol']).resolve(); root=Path(AYAR['ARSIV_KLASOR']).resolve()
    if not path.is_relative_to(root) or not path.is_file(): abort(404)
    if request.args.get('kucuk')=='1':
        thumb=COZUCU.kucuk({'anahtar':'archive:'+str(aid)},str(path))
        if not thumb:abort(404)
        path=thumb
    return send_file(path,conditional=True)

@app.route('/disa-aktar')
@giris_gerekli
def disa_aktar_sayfasi():
    cats=_scoped_categories(_allowed())
    today=datetime.now().date()
    return render_template('disa_aktar.html',izinli=yetki.izinli_projeler(g.yk,projeler()),kategoriler=cats,tarih1=(today-timedelta(days=30)).isoformat(),tarih2=today.isoformat())

@app.route('/disa-aktar.xlsx')
@giris_gerekli
def disa_aktar_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    a=request.args.to_dict(flat=False); a.setdefault('durum',['hepsi']); w,p,_=_filtre_sql(MultiDict(a),yetki.izinli_projeler(g.yk,projeler()))
    rows=baglanti().execute(f'SELECT * FROM kayitlar WHERE {w} ORDER BY proje,tarih,id',p).fetchall()
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Kayıtlar'
    columns=[('Kayıt ID','id'),('Proje','proje'),('Kategori','kategori'),('Personel','personel'),('Üye ID','uye_id'),('Ana üye ID','ana_uye_id'),('Tür','tip'),('Tarih','tarih'),('KT','kt'),('Onaylayan','kt_kim'),('Ekleyen','ekleyen_ad'),('Lead ID','lead_id'),('Eski link','link'),('Açıklama','aciklama'),('Ek not','ek_not'),('Ek sorular','soru_ozet')]
    ws.append([a for a,b in columns])
    for r in rows:
        vals=[]
        for _,name in columns:
            v='\n'.join(x['label']+': '+str(x['value']) for x in _question_answers(r)) if name=='soru_ozet' else (r[name] if r[name] is not None else '')
            if isinstance(v,str): v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)
            if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')): v="'"+v
            vals.append(v)
        ws.append(vals)
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F2937')
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(45,max(12,max(len(str(v.value or '')) for v in col)+2))
    summary=wb.create_sheet('Proje özeti');summary.append(['Proje','Kayıt','KT onaylı','KT bekleyen'])
    for project in sorted({r['proje'] for r in rows}):
        subset=[r for r in rows if r['proje']==project]
        summary.append([project,len(subset),sum(bool(r['kt']) for r in subset),sum(not r['kt'] for r in subset)])
    stream=io.BytesIO(); wb.save(stream); stream.seek(0)
    db.denetim_yaz(baglanti(),g.k['kullanici_adi'],'disa_aktar',yeni=str(len(rows)),ip=ip())
    return send_file(stream,as_attachment=True,download_name='kanit-kayitlari.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/bot/kullanici/<tg_id>')
@admin_gerekli
def bot_kullanici_detay(tg_id):
    try: data=bot_yonetim.kullanici_durum(baglanti(),tg_id)
    except ValueError: abort(404)
    return render_template('bot_kullanici.html',u=data)

def _check_folder():
    identifier=session.get('excel_check_id','')
    if not re.fullmatch('[a-f0-9]{32}',identifier): return None
    folder=Path(AYAR['BOT_KLASOR'])/'checks'/identifier
    metadata=folder/'owner.json'
    if not metadata.exists(): return None
    if json.loads(metadata.read_text()).get('user')!=g.k['kullanici_adi']: abort(403)
    return folder

def _source_path(folder):
    if not folder:return None
    info=folder/'source.json'
    if info.exists():
        name=json.loads(info.read_text()).get('file','')
        if name not in ('source.xlsx','source.csv','source.tsv','source.txt','source.ods'):abort(400)
        return folder/name
    return folder/'source.xlsx' if (folder/'source.xlsx').exists() else None

def _parse_options(form, previous=None):
    options=dict(previous or {})
    for key in ('encoding','delimiter','text_mode','header_row','sheet'):
        if key in form:options[key]=form.get(key)
    return options

@app.route('/excel-kontrol',methods=['GET','POST'])
@giris_gerekli
def excel_kontrol():
    from . import excel_check,file_sources
    allowed=_allowed()
    if not allowed:abort(403)
    folder=_check_folder();workbook=None;report=None;options={};hata=None
    if folder and (folder/'source.json').exists():options=json.loads((folder/'source.json').read_text()).get('options',{})
    if request.method=='POST':
        try:
            operation=request.form.get('islem')
            if operation=='yukle':
                uploaded=request.files.get('excel')
                suffix=Path(uploaded.filename or '').suffix.lower() if uploaded else ''
                if not uploaded or suffix not in ('.xlsx','.ods','.csv','.tsv','.txt'):raise ValueError('XLSX, ODS, CSV, TSV veya TXT dosyası seçin.')
                parent=Path(AYAR['BOT_KLASOR'])/'checks';parent.mkdir(mode=0o700,parents=True,exist_ok=True)
                import shutil
                for old in parent.iterdir():
                    if old.is_dir() and old.stat().st_mtime<time.time()-86400:shutil.rmtree(old,ignore_errors=True)
                candidate=parent/uuid.uuid4().hex;candidate.mkdir(mode=0o700)
                path=candidate/('source'+suffix)
                try:
                    total=0
                    with path.open('wb') as target:
                        for chunk in iter(lambda:uploaded.stream.read(65536),b''):
                            total+=len(chunk)
                            if total>file_sources.MAX_UPLOAD_BYTES:raise ValueError('Dosya en çok 12 MB olabilir.')
                            target.write(chunk)
                    options={'text_mode':'ids','header_row':0} if suffix=='.txt' else {'header_row':1}
                    excel_check.validate_source(path,options);workbook=excel_check.inspect_source(path,options)
                    if workbook['sheets']:options['sheet']=workbook['sheets'][0]['name']
                    (candidate/'owner.json').write_text(json.dumps({'user':g.k['kullanici_adi'],'name':Path(uploaded.filename).name},ensure_ascii=False))
                    (candidate/'source.json').write_text(json.dumps({'file':path.name,'options':options},ensure_ascii=False))
                    folder=candidate;session['excel_check_id']=candidate.name
                except BaseException:
                    shutil.rmtree(candidate,ignore_errors=True);raise
            elif operation in ('incele','kontrol'):
                path=_source_path(folder)
                if not path or not path.exists():raise ValueError('Önce bir kontrol dosyası yükleyin.')
                options=_parse_options(request.form,options)
                workbook=excel_check.inspect_source(path,options)
                names=[sheet['name'] for sheet in workbook['sheets']]
                if names and (not options.get('sheet') or (operation=='incele' and options.get('sheet') not in names)):
                    options['sheet']=names[0]
                if operation=='kontrol':
                    options.update(columns={field:request.form.get('kolon_'+field,'') for field in excel_check.FIELDS},
                                   match_fields=request.form.getlist('alan'),project=request.form.get('project',''),kt=request.form.get('kt','hepsi'))
                    if workbook['format']=='id_list':
                        options.update(id_field=request.form.get('id_field') or (options['match_fields'][0] if options['match_fields'] else 'personel'))
                        options.update(match_fields=[options['id_field']],columns={})
                    report=excel_check.check_source(path,baglanti(),allowed,options)
                    (folder/'result.json').write_text(json.dumps({'report':report,'options':options},ensure_ascii=False))
                    db.denetim_yaz(baglanti(),g.k['kullanici_adi'],'dosya_kontrol',yeni={'format':workbook['format'],'satir':report['total'],'alanlar':options['match_fields']},ip=ip(),hedef_turu='kontrol')
                else:
                    (folder/'result.json').unlink(missing_ok=True)
                (folder/'source.json').write_text(json.dumps({'file':path.name,'options':options},ensure_ascii=False))
            else:raise ValueError('Bilinmeyen işlem.')
        except (ValueError,OSError) as error:hata=str(error)
    path=_source_path(folder)
    if workbook is None and path and path.exists():
        try:workbook=excel_check.inspect_source(path,options)
        except (ValueError,OSError) as error:hata=hata or str(error)
    if request.method=='GET' and folder and (folder/'result.json').exists():
        saved=json.loads((folder/'result.json').read_text());report=saved['report'];options=saved['options']
    report_page=1;report_pages=1;report_status=request.args.get('rapor_durum','')
    if report:
        rows=report['rows']
        if report_status:rows=[row for row in rows if row['status']==report_status]
        try:report_page=max(1,min(1000000,int(request.args.get('r_s','1'))))
        except ValueError:abort(400)
        report_pages=max(1,(len(rows)+99)//100)
        report_page=min(report_page,report_pages)
        report=dict(report,rows=[dict(row,ids=row['ids'][:20],display_remaining=max(0,len(row['ids'])-20)) for row in rows[(report_page-1)*100:report_page*100]])
    if workbook and path:workbook['extension']=path.suffix
    return render_template('excel_check.html',workbook=workbook,report=report,options=options,hata=hata,izinli=allowed,upload_ready=bool(path),report_page=report_page,report_pages=report_pages,report_status=report_status)

def _safe_excel(value):
    text=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',str('' if value is None else value))
    return "'"+text if text.startswith(('=','+','-','@','\t','\r')) else text

@app.route('/excel-kontrol/sonuc.xlsx')
@giris_gerekli
def excel_kontrol_indir():
    folder=_check_folder()
    if not folder or not (folder/'result.json').exists():abort(404)
    saved=json.loads((folder/'result.json').read_text());report=saved['report'];options=saved['options']
    allowed=_allowed()
    if not allowed:abort(403)
    snapshot=set(report.get('metadata',{}).get('scope',{}).get('allowed_projects',[]))
    if snapshot and not snapshot.issubset(set(allowed)):abort(403)
    ids=list({int(kid) for row in report['rows'] for kid in row['ids']})
    for offset in range(0,len(ids),500):
        batch=ids[offset:offset+500]
        for record in baglanti().execute('SELECT proje FROM kayitlar WHERE id IN ('+','.join('?' for _ in batch)+')',batch):
            if record['proje'] not in allowed:abort(403)
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    wb=Workbook();ws=wb.active;ws.title='Eksik kanıtlar' if request.args.get('eksik')=='1' else 'Kontrol sonucu'
    fields=options.get('match_fields',[]);ws.append(['Kaynak satırı','Durum',*fields,'Kanıt kayıt IDleri','Açıklama'])
    for row in report['rows']:
        if request.args.get('eksik')=='1' and row['status']!='yok':continue
        ws.append([_safe_excel(row.get('source',row['row'])),row['status'],*[_safe_excel(row['values'].get(field,'')) for field in fields],', '.join(str(x) for x in row['ids']),_safe_excel(row.get('reason',''))])
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F2937')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=min(50,max(15,max(len(str(cell.value or '')) for cell in column)+2))
    metadata=wb.create_sheet('Kontrol bilgisi');metadata.append(['Alan','Değer'])
    meta=report.get('metadata',{})
    for key,value in [('Kontrol zamanı (UTC)',meta.get('checked_at','')),('Kapsam',meta.get('scope',{})),('Seçenekler',meta.get('options',options)),('Limitler',meta.get('limits',{})),('Kaynak adı',json.loads((folder/'owner.json').read_text()).get('name',''))]:
        metadata.append([key,_safe_excel(json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else value)])
    metadata.column_dimensions['A'].width=25;metadata.column_dimensions['B'].width=90
    stream=io.BytesIO();wb.save(stream);stream.seek(0)
    return send_file(stream,as_attachment=True,download_name='eksik-kanitlar.xlsx' if request.args.get('eksik')=='1' else 'kanit-kontrol.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/telegram/kullanici/<tg_id>/kanitlar',methods=['GET','POST'])
@giris_gerekli
def telegram_kayitlari(tg_id):
    from . import user_evidence
    if not re.fullmatch('[0-9]{1,20}',tg_id):abort(404)
    allowed=yetki.izinli_projeler(g.yk,projeler())
    if not allowed:abort(403)
    c=baglanti()
    if request.method=='POST':
        if not g.yk['admin']:abort(403)
        project=request.form.get('project',request.form.get('ana_proje','')).strip()
        if project and project not in allowed:abort(400)
        with db.transaction(c):
            c.execute('INSERT INTO tg_kullanicilar(id,ana_proje) VALUES(?,?) ON CONFLICT(id) DO UPDATE SET ana_proje=excluded.ana_proje',(tg_id,project or None))
            db.denetim_yaz(c,g.k['kullanici_adi'],'beklenen_proje',yeni=tg_id+':'+project,ip=ip());db.degisti(c)
        return redirect(url_for('telegram_kayitlari',tg_id=tg_id))
    try:data=user_evidence.get_profile(c,tg_id,allowed,request.args.to_dict(),superadmin=g.yk['admin'])
    except PermissionError:abort(403)
    except ValueError as e:return render_template('hata.html',kod=400,mesaj=str(e)),400
    return render_template('telegram_kayitlari.html',profile=data,allowed=allowed)

@app.route('/saglik')
def saglik():
    return 'OK',200,{'X-Kanit-Paneli':SURUM}

@app.route('/manus-routes.json')
def route_manifest():
    return jsonify(routes=[{'path':p,'title':t} for p,t in [('/','Ana sayfa'),('/giris','Giriş'),('/sifre','Parola'),('/kayitlar','Kayıtlar'),('/arsiv','Arşiv'),('/excel-kontrol','Excel kontrol'),('/bot','Bot'),('/yonetim/kullanicilar','Kullanıcılar'),('/denetim','Denetim'),('/disa-aktar','Çıktı'),('/telegram/kullanici/:tg_id/kanitlar','Ekleyen kanıt profili'),('/bot/kullanici/:tg_id','Telegram kullanıcı durumu')]])

def _json_bekleniyor():
    """fetch() ile çağrılan uç noktalar hata sayfası yerine JSON hata gövdesi alır."""
    if request.is_json or request.headers.get('X-CSRF') or request.headers.get('X-Requested-With')=='XMLHttpRequest':return True
    return request.accept_mimetypes.best_match(['application/json','text/html'])=='application/json'

HATA_BASLIKLARI={400:'Geçersiz istek',403:'Yetkiniz yok',404:'Sayfa bulunamadı',405:'Bu yöntem desteklenmiyor',409:'Kayıt değişti',413:'Dosya çok büyük',429:'Çok fazla deneme',500:'Beklenmeyen hata'}

@app.errorhandler(400)
@app.errorhandler(403)
@app.errorhandler(404)
@app.errorhandler(405)
@app.errorhandler(409)
@app.errorhandler(413)
@app.errorhandler(429)
def handle_error(error):
    messages={400:'Gönderilen bilgi geçersiz veya eksik. Sayfayı yenileyip tekrar deneyin.',403:'Bu sayfa veya işlem için hesabınızın yetkisi yok. Gerekirse süperadminden proje yetkisi isteyin.',404:'Aradığınız sayfa ya da kayıt bulunamadı; bağlantı eski olabilir.',405:'Bu adres bu şekilde çağrılamaz.',409:'Kayıt bu arada başka bir kullanıcı tarafından değiştirildi. Sayfayı yenileyip tekrar deneyin.',413:'Dosya boyutu sınırı (12 MB) aşıldı.',429:'Çok fazla istek gönderildi. Birkaç dakika bekleyip tekrar deneyin.'}
    ozel=error.description if isinstance(getattr(error,'description',None),str) and error.code in (400,409) and error.description and not error.description.startswith('The ') else None
    if _json_bekleniyor():
        return jsonify(ok=False,hata=ozel or messages[error.code]),error.code
    return render_template('hata.html',kod=error.code,baslik=HATA_BASLIKLARI.get(error.code,'Hata'),mesaj=ozel or messages[error.code]),error.code

@app.errorhandler(500)
def handle_server_error(error):
    ref=datetime.now().strftime('%Y%m%d-%H%M%S')+'-'+secrets.token_hex(3).upper()
    original=getattr(error,'original_exception',None) or error
    app.logger.error('Beklenmeyen hata [kayit no %s] %s %s',ref,request.method,request.path,exc_info=(type(original),original,original.__traceback__))
    try:
        c=db.baglan(AYAR['DB'])
        try:
            db.denetim_yaz(c,(g.get('k') or {}).get('kullanici_adi') or 'anonim','sunucu_hata',yeni={'kayit_no':ref,'yol':request.path,'tur':type(original).__name__},ip=ip(),hedef_turu='sistem',hedef_id=ref,sonuc='hata')
        finally:c.close()
    except Exception:
        pass
    mesaj='Beklenmeyen hata, kayıt no '+ref+'. Ayrıntı sunucu günlüğüne yazıldı; sorun sürerse bu numarayı yöneticinize iletin.'
    if _json_bekleniyor():
        return jsonify(ok=False,hata=mesaj,kayit_no=ref),500
    return render_template('hata.html',kod=500,baslik=HATA_BASLIKLARI[500],mesaj=mesaj,kayit_no=ref),500


def _start_web_converter():
    if os.environ.get('GYAZO_WORKER','1')=='0':return
    import threading
    from . import donusum_isleri
    def loop():
        while True:
            conn=None;processed=False
            try:
                conn=db.baglan(AYAR['DB'])
                processed=bool(donusum_isleri.worker_step(conn,'web-'+str(os.getpid())))
            except Exception as error:
                # Ham HTTP hata metni anahtar içeriyor olabilir; yalnız sınıf adı.
                print('Dönüşüm işçisi kontrolü: '+type(error).__name__,flush=True)
            finally:
                if conn:conn.close()
            time.sleep(0.5 if processed else 2)
    threading.Thread(target=loop,name='kanit-web-gyazo',daemon=True).start()

def calistir():
    from waitress import serve
    _start_web_converter()
    print(f'Kanıt Paneli {SURUM} — {AYAR["DINLE"]}:{AYAR["PORT"]}',flush=True)
    proxy={}
    if str(AYAR.get('TRUST_PROXY','0')).lower() in ('1','true'):
        # Yalnız loopback/korumalı ağdan erişilen, Nginx arkasındaki sunucu.
        proxy=dict(trusted_proxy='*',trusted_proxy_count=1,trusted_proxy_headers={'x-forwarded-for','x-forwarded-proto'})
    serve(app,host=AYAR['DINLE'],port=AYAR['PORT'],threads=8,**proxy)


def _snapshot_questions(raw):
    try:questions=json.loads(raw) if isinstance(raw,str) else (raw or [])
    except (ValueError,TypeError):return []
    if isinstance(questions,dict):questions=questions.get('questions',[])
    if not isinstance(questions,list):return []
    return [question for question in questions if isinstance(question,dict) and isinstance(question.get('key'),str) and isinstance(question.get('label'),str)]

def _question_answers(row):
    try:
        questions=_snapshot_questions(row['soru_anlik_json']);answers=json.loads(row['cevaplar_json'] or '{}')
        if not isinstance(answers,dict):return []
        return [dict(key=q['key'],label=q['label'],preset=q.get('preset'),value=answers.get(q['key'],'')) for q in questions]
    except (ValueError,KeyError,TypeError):return []

@app.route('/kategoriler.json')
@giris_gerekli
def kategori_secenekleri():
    selected=_selected_projects(request.args,_allowed())
    return jsonify(kategoriler=_scoped_categories(selected))

@app.route('/yonetim/katalog',methods=['GET','POST'])
@admin_gerekli
def katalog_yonetim():
    from . import katalog
    c=baglanti();message=None;status=200
    if request.method=='POST':
        try:
            revision=request.form.get('revision')
            revision=int(revision) if revision else None
            active=request.form.get('active')=='1'
            if request.form.get('islem')=='proje_kaydet':
                if katalog.get_project(c,request.form.get('code','')) and revision is None:raise katalog.RevisionConflict('Bu kod var; mevcut proje formundan düzenleyin.')
                katalog.save_project(c,request.form.get('code',''),request.form.get('label',''),request.form.get('color',''),active=active,actor=g.k['kullanici_adi'],expected_revision=revision)
            elif request.form.get('islem')=='kategori_kaydet':
                if katalog.get_category(c,request.form.get('code','')) and revision is None:raise katalog.RevisionConflict('Bu kod var; mevcut kategori formundan düzenleyin.')
                questions=json.loads(request.form.get('questions_json','[]'))
                katalog.save_category(c,request.form.get('code',''),request.form.get('label',''),mode=request.form.get('mode','custom'),projects=request.form.getlist('projects'),questions=questions,active=active,actor=g.k['kullanici_adi'],expected_revision=revision)
            else:raise ValueError('Bilinmeyen işlem.')
            flash('Kaydedildi. Yeni seçimlerde anında geçerli.');return redirect(url_for('katalog_yonetim'))
        except katalog.RevisionConflict as error:message=str(error);status=409
        except (ValueError,TypeError) as error:message=str(error);status=400
    return render_template('katalog.html',proje_bilgileri=katalog.list_projects(c,False),kategoriler=katalog.list_categories(c,active_only=False),mesaj=message),status

if __name__=='__main__':
    calistir()

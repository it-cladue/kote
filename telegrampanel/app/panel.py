# -*- coding: utf-8 -*-
"""Telegram kanit botu - web paneli (Flask).

Calistirma: KanitPaneli.exe (cift tik) ya da kaynaktan  python -m app.panel
Panel her bilgisayarda yerel calisir (127.0.0.1). Bot klasoru OneDrive ile esitlenen klasordur;
panel oradan Excel'i (kopyalayarak) okur, islemleri panel_kuyruk\\<bilgisayar>.jsonl'a yazar,
bot uygular. Excel'e yalnizca bot yazar.
"""
import os, sys, json, secrets, threading, time, functools, string, random, glob, webbrowser, urllib.request
from datetime import datetime, timedelta
from flask import Flask, g, request, session, redirect, url_for, render_template, jsonify, abort, send_file, flash
from werkzeug.datastructures import MultiDict

from . import db, yetki, excel_oku, ice_aktar, gorsel, kuyruk, bot_yonetim, log_oku, kullanicilar, liste_kontrol

SURUM = "1.2.0"
DONMUS = bool(getattr(sys, "frozen", False))
KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))          # kaynak kodu koku
KAYNAK_KOK = getattr(sys, "_MEIPASS", None) or KOK                            # sablon / statik (exe icinde)
EXE_KLASOR = os.path.dirname(os.path.abspath(sys.executable)) if DONMUS else KOK
VARSAYILAN_RENK = {"PP": "#A6A6A6", "PA": "#4472C4", "OF": "#FFFF00", "HI": "#FF0000", "VI": "#7030A0", "GA": "#00B050", "BETA": "#FFA500"}
EXCEL_ADI = "telegrambot_gyazo.xlsx"
KOPYA_ADI = "panel_excel_kopya.xlsx"   # paylasim klasoru modunda bot yalnizca bu kopyayi koyar


def _veri_koku():
    k = os.environ.get("PANEL_VERI")
    if not k:
        if DONMUS:
            k = os.path.join(os.environ.get("LOCALAPPDATA") or os.path.expanduser("~"), "KanitPaneli")
        else:
            k = KOK
    os.makedirs(k, exist_ok=True)
    return k


VERI_KOK = _veri_koku()
AYAR_YOL = os.environ.get("PANEL_AYAR") or os.path.join(VERI_KOK, "panel_ayar.json")


def bot_klasoru_mu(k):
    """Bot klasoru (ana Excel var) ya da paylasim klasoru (botun yazdigi panel_excel_kopya.xlsx var)."""
    try:
        return bool(k) and (os.path.isfile(os.path.join(k, EXCEL_ADI)) or os.path.isfile(os.path.join(k, KOPYA_ADI)))
    except Exception:
        return False


def bot_klasoru_adaylari(mevcut=None):
    """Exe'nin bulundugu klasor (OneDrive'daki bot klasorune konmussa), ust klasoru ve OneDrive altindaki klasorler."""
    adaylar = []

    def ekle(k):
        try:
            k = os.path.abspath(k)
        except Exception:
            return
        if bot_klasoru_mu(k) and k not in adaylar:
            adaylar.append(k)

    if mevcut:
        ekle(mevcut)
    ekle(EXE_KLASOR)
    ekle(os.path.dirname(EXE_KLASOR))
    home = os.path.expanduser("~")
    kokler = glob.glob(os.path.join(home, "OneDrive*")) + [os.path.join(home, "Desktop"), os.path.join(home, "Documents")]
    for od in kokler:
        if not os.path.isdir(od):
            continue
        derinlik0 = od.rstrip("\\/").count(os.sep)
        for root, dirs, files in os.walk(od):
            if root.count(os.sep) - derinlik0 >= 3:
                dirs[:] = []
            if EXCEL_ADI in files or KOPYA_ADI in files:
                ekle(root)
                dirs[:] = []
            dirs[:] = [d for d in dirs if not d.startswith(".") and d.lower() not in ("gorseller", "backups", "ekran_goruntuleri", "telegram_gorseller", "node_modules")]
    return adaylar


def ayar_yukle():
    ham = {}
    bozuk = False
    if os.path.exists(AYAR_YOL):
        try:
            with open(AYAR_YOL, encoding="utf-8-sig") as f:
                ham = json.load(f)
        except Exception as e:
            # Elle duzenlenirken bozulmus olabilir: ustune YAZMA, kenara al, varsayilanlarla devam et
            bozuk = True
            try:
                yedek = AYAR_YOL + ".bozuk-" + datetime.now().strftime("%Y%m%d-%H%M%S")
                os.replace(AYAR_YOL, yedek)
                print(f"! panel_ayar.json okunamadi ({e}); dosya {yedek} olarak kenara alindi, varsayilan ayarlarla devam ediliyor.")
            except Exception:
                print("! panel_ayar.json okunamadi:", e)
    degisti = False
    if not ham.get("GIZLI_ANAHTAR"):
        ham["GIZLI_ANAHTAR"] = secrets.token_hex(32)
        degisti = True
    if not bot_klasoru_mu(ham.get("BOT_KLASOR")):
        adaylar = bot_klasoru_adaylari(ham.get("BOT_KLASOR"))
        if adaylar:
            ham["BOT_KLASOR"] = adaylar[0]
            degisti = True
    if degisti and not bozuk:
        ayar_dosyaya_yaz(ham)
    return ayar_turet(ham)


def ayar_dosyaya_yaz(ham):
    # Dosyaya yalnizca kullanicinin yazdigi alanlar + gizli anahtar yazilir; turetilen yollar (EXCEL vb.) yazilmaz.
    try:
        os.makedirs(os.path.dirname(AYAR_YOL), exist_ok=True)
        with open(AYAR_YOL, "w", encoding="utf-8") as f:
            json.dump({k: v for k, v in ham.items() if not k.startswith("_t_")}, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print("panel_ayar.json yazilamadi:", e)


def ayar_turet(ham):
    a = dict(ham)
    a.setdefault("BOT_KLASOR", "")
    bk = a["BOT_KLASOR"] or EXE_KLASOR
    ana = os.path.join(bk, EXCEL_ADI)
    a["EXCEL"] = ham.get("EXCEL") or (ana if os.path.isfile(ana) or not os.path.isfile(os.path.join(bk, KOPYA_ADI)) else os.path.join(bk, KOPYA_ADI))
    a["PAYLASIM_MODU"] = not os.path.isfile(ana) and os.path.isfile(os.path.join(bk, KOPYA_ADI))
    a["LOG_DOSYA"] = ham.get("LOG_DOSYA") or os.path.join(bk, "bot_log.txt")
    a["GORSEL_KLASOR"] = ham.get("GORSEL_KLASOR") or os.path.join(bk, "gorseller")
    a["DB"] = ham.get("DB") or os.path.join(VERI_KOK, "panel.db")
    a["ONBELLEK"] = ham.get("ONBELLEK") or os.path.join(VERI_KOK, "tmp")
    a.setdefault("YENILEME_SN", 30)       # kaynak degisti mi kontrolu (ucuz); Excel ancak degistiyse yeniden okunur
    a.setdefault("DINLE", "127.0.0.1")
    a.setdefault("PORT", 8080)
    a.setdefault("PROJE_RENK", VARSAYILAN_RENK)
    a.setdefault("SUTUN_HARITASI", {})
    a.setdefault("SAYFA_BOYU", 48)
    return a


AYAR = ayar_yukle()
_yerel = threading.local()


def ayar_degistir(**degisiklik):
    """Calisirken ayar degistir (kurulum sayfasi): dosyaya yaz, AYAR'i ve gorsel cozucuyu yenile."""
    ham = {}
    if os.path.exists(AYAR_YOL):
        try:
            with open(AYAR_YOL, encoding="utf-8-sig") as f:
                ham = json.load(f)
        except Exception:
            ham = {}
    ham["GIZLI_ANAHTAR"] = AYAR.get("GIZLI_ANAHTAR") or ham.get("GIZLI_ANAHTAR") or secrets.token_hex(32)
    ham.update(degisiklik)
    ayar_dosyaya_yaz(ham)
    yeni = ayar_turet(ham)
    AYAR.clear()
    AYAR.update(yeni)
    global COZUCU
    COZUCU = gorsel.GorselCozucu(AYAR, AYAR["ONBELLEK"])


def baglanti():
    c = getattr(_yerel, "conn", None)
    if c is None:
        c = db.baglan(AYAR["DB"])
        _yerel.conn = c
    return c


COZUCU = gorsel.GorselCozucu(AYAR, AYAR["ONBELLEK"])
app = Flask(__name__, template_folder=os.path.join(KAYNAK_KOK, "app", "templates"), static_folder=os.path.join(KAYNAK_KOK, "app", "static"))
app.secret_key = AYAR["GIZLI_ANAHTAR"]
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax", PERMANENT_SESSION_LIFETIME=timedelta(hours=12), MAX_CONTENT_LENGTH=4 * 1024 * 1024)
if str(AYAR.get("HTTPS", "0")).lower() in ("1", "true"):
    app.config["SESSION_COOKIE_SECURE"] = True


def simdi():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ==================== yardimcilar ====================
def projeler():
    """Excel'de gorulen sayfalar (kayitlar tablosu) + ayar listesi; sirali"""
    c = baglanti()
    gorulen = [r["proje"] for r in c.execute("SELECT DISTINCT proje FROM kayitlar").fetchall()]
    sira = list(AYAR.get("PROJELER") or ["PP", "PA", "OF", "HI", "VI", "GA", "BETA"])
    return [p for p in sira if p in gorulen] + sorted(p for p in gorulen if p not in sira)


def renk(proje):
    return AYAR["PROJE_RENK"].get((proje or "").upper(), "#999999")


def metin_rengi(hex_renk):
    """Arka plan rengine gore okunur yazi rengi (OF sari -> siyah yazi)"""
    try:
        h = (hex_renk or "#999999").lstrip("#")
        r, gg, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        return "#111111" if (0.299 * r + 0.587 * gg + 0.114 * b) > 150 else "#ffffff"
    except Exception:
        return "#ffffff"


def url_degistir(**kw):
    """Mevcut sorgu parametrelerini koruyarak bazilarini degistir (sayfa numarasi, s= verilmediyse sifirlanir)."""
    a = request.args.to_dict(flat=False)
    a.pop("s", None)
    for k, v in kw.items():
        if v is None or v == "":
            a.pop(k, None)
        elif isinstance(v, (list, tuple)):
            a[k] = [str(x) for x in v]
        else:
            a[k] = [str(v)]
    return url_for(request.endpoint or "kayitlar", **a)


def kullanici_yukle():
    kad = session.get("kad")
    if not kad:
        return None
    u = kullanicilar.bul(AYAR, kad)
    if not u or not u.get("aktif", True):
        return None
    return u


KURULUM_SERBEST = {"kurulum", "kurulum_admin", "static", "saglik", "cikis"}


@app.before_request
def _on():
    if not istek_izinli_mi():
        if request.remote_addr not in _REDDEDILEN and len(_REDDEDILEN) < 500:
            _REDDEDILEN.add(request.remote_addr)
            print(f"[AG] {request.remote_addr} adresinden istek reddedildi (ofis agi disinda). Izin vermek icin panel_ayar.json IZINLI_AGLAR.")
        abort(403, f"AG:{request.remote_addr}")
    g.k = kullanici_yukle()
    g.yk = yetki.coz(g.k["yetki"]) if g.k else yetki.coz("")
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    if request.method == "POST" and request.endpoint not in ("giris",):
        tok = request.form.get("csrf") or request.headers.get("X-CSRF")
        if not tok or tok != session.get("csrf"):
            abort(400, "CSRF")
    ep = request.endpoint or ""
    if ep in KURULUM_SERBEST:
        return None
    if not bot_klasoru_mu(AYAR.get("BOT_KLASOR")):
        return redirect(url_for("kurulum"))
    if not kullanicilar.var_mi(AYAR):
        return redirect(url_for("kurulum_admin"))
    return None


@app.after_request
def _basliklar(r):
    r.headers.setdefault("X-Frame-Options", "DENY")
    r.headers.setdefault("X-Content-Type-Options", "nosniff")
    r.headers.setdefault("Referrer-Policy", "same-origin")
    if request.path.endswith(".json") or request.path.startswith(("/kayit/", "/bildirim")):
        r.headers.setdefault("Cache-Control", "no-store")
    return r


@app.context_processor
def _ctx():
    return {"k": g.get("k"), "yk": g.get("yk"), "renk": renk, "metin_rengi": metin_rengi, "url_degistir": url_degistir, "csrf": session.get("csrf", ""),
            "projeler": projeler, "izin": lambda p, e: yetki.izin(g.yk, p, e), "AYAR": AYAR, "simdi": simdi, "SURUM": SURUM}


def giris_gerekli(f):
    @functools.wraps(f)
    def w(*a, **kw):
        if not g.k:
            return redirect(url_for("giris", sonraki=request.path))
        return f(*a, **kw)
    return w


def admin_gerekli(f):
    @functools.wraps(f)
    def w(*a, **kw):
        if not g.k:
            return redirect(url_for("giris", sonraki=request.path))
        if not g.yk["admin"]:
            abort(403)
        return f(*a, **kw)
    return w


def ip():
    """Denetim kaydi icin istemci adresi. Basliklar (CF-Connecting-IP / X-Forwarded-For) ancak panel bir ters vekilin
    (Cloudflare Tunnel vb.) arkasindaysa ve GUVENILIR_PROXY=1 ise dikkate alinir; yoksa herkes sahte adres yazabilirdi."""
    if str(AYAR.get("GUVENILIR_PROXY", "0")).lower() in ("1", "true"):
        return (request.headers.get("CF-Connecting-IP") or (request.headers.get("X-Forwarded-For") or "").split(",")[0].strip() or request.remote_addr)
    return request.remote_addr


def yerel_istek():
    """Istek bu bilgisayardan mi? Tunel/ters vekil arkasinda (GUVENILIR_PROXY=1) gercek istemci adresine bakilir;
    yoksa tunelden gelen herkes 'yerel' sayilirdi."""
    return (ip() or "") in ("127.0.0.1", "::1", "")


def ag_modu_mu():
    return str(AYAR.get("DINLE", "127.0.0.1")) not in ("127.0.0.1", "localhost", "")


import ipaddress as _ipa
OZEL_AGLAR = [_ipa.ip_network(a) for a in ("10.0.0.0/8", "172.16.0.0/12", "192.168.0.0/16", "100.64.0.0/10", "127.0.0.0/8", "169.254.0.0/16", "fc00::/7", "fe80::/10", "::1/128")]
_REDDEDILEN = set()


def istek_izinli_mi():
    """Ag modunda yalnizca ozel (ofis) aglardan gelen istekler kabul edilir: 10/8, 172.16/12, 192.168/16, link-local, loopback.
    panel_ayar.json'a IZINLI_AGLAR: ["192.168.1.0/24", "10.0.0.0/8"] yazilirsa yalnizca o aglar (loopback her zaman serbest).
    Yanlislikla internete acilmis (port yonlendirme) bir panel boylece disaridan kullanilamaz."""
    import ipaddress
    if yerel_istek() or not ag_modu_mu():
        return True
    try:
        adres = ipaddress.ip_address((request.remote_addr or "").split("%")[0])
    except ValueError:
        return False
    izinli = AYAR.get("IZINLI_AGLAR")
    if izinli:
        for ag in (izinli if isinstance(izinli, list) else str(izinli).split(",")):
            try:
                if adres in ipaddress.ip_network(str(ag).strip(), strict=False):
                    return True
            except ValueError:
                continue
        return False
    return any(adres in ag for ag in OZEL_AGLAR)


def guvenli_sonraki(s):
    """Giristen sonra donulecek adres: yalnizca bu panelin kendi yolu (acik yonlendirme korumasi)."""
    s = (s or "").strip()
    if not s.startswith("/") or s.startswith("//") or "\\" in s or ":" in s.split("?")[0]:
        return None
    return s


def kayit_izin(kayit, eylem):
    return yetki.izin(g.yk, kayit["proje"], eylem)


# ==================== kurulum (ilk calisma) ====================
@app.route("/kurulum", methods=["GET", "POST"])
def kurulum():
    gecerli = bot_klasoru_mu(AYAR.get("BOT_KLASOR"))
    if gecerli and not (g.k and g.yk["admin"]):
        return redirect(url_for("ana"))
    if not gecerli and not yerel_istek():
        # ag modunda baska bir bilgisayardan ilk kurulum yapilamaz: bot klasoru panelin calistigi bilgisayarda secilir
        abort(403)
    hata = None
    if request.method == "POST":
        k = (request.form.get("klasor") or "").strip().strip('"')
        if not bot_klasoru_mu(k):
            hata = f"Bu klasorde {EXCEL_ADI} bulunamadi: {k or '(bos)'}"
        else:
            ayar_degistir(BOT_KLASOR=os.path.abspath(k))
            threading.Thread(target=_ice_aktar_bir_kez, daemon=True).start()
            flash("Bot klasoru kaydedildi: " + AYAR["BOT_KLASOR"])
            return redirect(url_for("kurulum_admin") if not kullanicilar.var_mi(AYAR) else url_for("ana"))
    return render_template("kurulum.html", adaylar=bot_klasoru_adaylari(AYAR.get("BOT_KLASOR")), mevcut=AYAR.get("BOT_KLASOR") or "", hata=hata, excel_adi=EXCEL_ADI, veri_kok=VERI_KOK, ayar_yol=AYAR_YOL)


def kullanici_adi_gecerli(ad):
    return bool(ad) and len(ad) <= 40 and ad.replace("_", "").replace(".", "").isalnum()


@app.route("/kurulum/admin", methods=["GET", "POST"])
def kurulum_admin():
    if not bot_klasoru_mu(AYAR.get("BOT_KLASOR")):
        return redirect(url_for("kurulum"))
    if kullanicilar.var_mi(AYAR):
        return redirect(url_for("giris"))
    # Kullanici dosyasi yok ama klasorde panel daha once kullanilmis: buyuk olasilikla OneDrive dosyayi henuz
    # getirmedi (ya da biri sildi). Bot bilgisayari degilse yeni "ilk admin" olusturulmaz; botun admin sifresiyle
    # (config.json ADMIN_PASSWORD) acilabilir. Boylece iki bilgisayarin ayri ayri ilk admin yaratmasi onlenir.
    # Ag modunda (baska bilgisayardan gelen istek) ilk admin yalnizca botun ADMIN_PASSWORD'u ile olusturulabilir:
    # kullanici dosyasi bir an icin yoksa bile ofis agindaki herhangi biri kendine super admin acamaz.
    bekle = (kullanicilar.once_kullanilmis(AYAR) and not bot_yonetim.yerel_bot_sureci()) or not yerel_istek()
    hata = None
    if request.method == "POST":
        ad = (request.form.get("kullanici") or "").strip().lower()
        s1, s2 = request.form.get("sifre") or "", request.form.get("sifre2") or ""
        if bekle and not bot_yonetim.admin_sifresi_dogru(AYAR, request.form.get("bot_admin_sifre")):
            hata = "Bot admin sifresi hatali. Kullanici dosyasi OneDrive ile gelene kadar bekleyin ya da botun ADMIN_PASSWORD degerini girin."
        elif not kullanici_adi_gecerli(ad):
            hata = "Kullanici adi harf/rakam olmali."
        elif len(s1) < 8:
            hata = "Sifre en az 8 karakter olmali."
        elif s1 != s2:
            hata = "Sifreler ayni degil."
        else:
            kullanicilar.ekle(AYAR, ad, "admin", s1, gorunen_ad=request.form.get("ad") or ad, olusturan="kurulum")
            kullanicilar.guncelle(AYAR, ad, sifre_degistir=False)
            db.denetim_yaz(baglanti(), ad, "ilk_admin", ip=ip())
            flash("Super admin olusturuldu. Simdi giris yapin.")
            return redirect(url_for("giris"))
    return render_template("kurulum_admin.html", hata=hata, bot_klasor=AYAR["BOT_KLASOR"], bekle=bekle, dosya=kullanicilar.yol(AYAR))


# ==================== giris ====================
@app.route("/giris", methods=["GET", "POST"])
def giris():
    hata = None
    if request.method == "POST":
        ad = (request.form.get("kullanici") or "").strip().lower()
        sifre = request.form.get("sifre") or ""
        tok = request.form.get("csrf") or request.headers.get("X-CSRF")
        if not tok or tok != session.get("csrf"):
            # giris formu da CSRF korumali (baskasinin hesabina zorla giris yaptirma); eski/kapanmis oturum: sayfa yenilenir
            return render_template("giris.html", hata="Sayfa zaman asimina ugradi; lutfen tekrar deneyin."), 400
        u, hata = kullanicilar.dogrula(AYAR, ad, sifre, request.remote_addr)
        c = baglanti()
        if u:
            session.clear()
            session.permanent = True
            session["kad"] = u["kullanici_adi"]
            session["csrf"] = secrets.token_hex(16)
            db.denetim_yaz(c, u["kullanici_adi"], "giris", ip=ip())
            return redirect(guvenli_sonraki(request.args.get("sonraki")) or url_for("ana"))
        db.denetim_yaz(c, ad, "giris_hata", ip=ip())
        time.sleep(0.8)
    return render_template("giris.html", hata=hata)


@app.route("/cikis")
def cikis():
    if g.get("k"):
        db.denetim_yaz(baglanti(), g.k["kullanici_adi"], "cikis", ip=ip())
    session.clear()
    return redirect(url_for("giris"))


@app.route("/sifre", methods=["GET", "POST"])
@giris_gerekli
def sifre_degistir():
    hata = None
    if request.method == "POST":
        eski, yeni, yeni2 = request.form.get("eski", ""), request.form.get("yeni", ""), request.form.get("yeni2", "")
        u, _ = kullanicilar.dogrula(AYAR, g.k["kullanici_adi"], eski, request.remote_addr)
        if not u:
            hata = "Eski sifre hatali."
        elif len(yeni) < 8:
            hata = "Yeni sifre en az 8 karakter olmali."
        elif yeni != yeni2:
            hata = "Yeni sifreler ayni degil."
        else:
            kullanicilar.guncelle(AYAR, g.k["kullanici_adi"], sifre=yeni, sifre_degistir=False)
            db.denetim_yaz(baglanti(), g.k["kullanici_adi"], "sifre_degisti", ip=ip())
            flash("Sifre degistirildi.")
            return redirect(url_for("ana"))
    return render_template("sifre.html", hata=hata)


# ==================== ana sayfa ====================
def durum_bilgisi():
    c = baglanti()
    son = db.ayar_al(c, "son_ice_aktarim")
    son = json.loads(son) if son else None
    excel_mtime = None
    try:
        excel_mtime = datetime.fromtimestamp(os.path.getmtime(AYAR["EXCEL"])).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        pass
    bot_kopya = os.path.join(AYAR["BOT_KLASOR"], excel_oku.BOT_KOPYA_ADI)
    try:
        kopya_mtime = datetime.fromtimestamp(os.path.getmtime(bot_kopya)).strftime("%Y-%m-%d %H:%M:%S") if os.path.exists(bot_kopya) else None
    except Exception:
        kopya_mtime = None
    return {"son_ice_aktarim": son, "excel_mtime": excel_mtime, "excel_kopya_mtime": kopya_mtime, "log_son": log_oku.son_satir_zamani(AYAR["LOG_DOSYA"]),
            "bot": bot_yonetim.bot_calisiyor_mu(AYAR), "kuyruk_bekleyen": kuyruk.bekleyen_sayisi(c),
            "bekleyen_islem": c.execute("SELECT COUNT(*) FROM kayitlar WHERE bekleyen_islem IS NOT NULL").fetchone()[0],
            "istemci": kuyruk.istemci_adi(AYAR), "bot_klasor": AYAR["BOT_KLASOR"], "surum": SURUM, "paylasim_modu": bool(AYAR.get("PAYLASIM_MODU"))}


@app.route("/")
@giris_gerekli
def ana():
    c = baglanti()
    izinli = yetki.izinli_projeler(g.yk, projeler())
    kartlar = []
    for p in izinli:
        r = c.execute("""SELECT SUM(CASE WHEN kt=0 AND silindi=0 AND excel_durum='var' THEN 1 ELSE 0 END) bekleyen,
                                SUM(CASE WHEN kt=1 AND silindi=0 AND excel_durum='var' THEN 1 ELSE 0 END) onayli,
                                SUM(CASE WHEN silindi=1 OR excel_durum='silinmis' THEN 1 ELSE 0 END) silinen,
                                SUM(CASE WHEN date(tarih)=date('now','localtime') AND silindi=0 THEN 1 ELSE 0 END) bugun,
                                COUNT(*) toplam FROM kayitlar WHERE proje=?""", (p,)).fetchone()
        kartlar.append({"proje": p, **{k: r[k] or 0 for k in r.keys()}})
    return render_template("ana.html", kartlar=kartlar, durum=durum_bilgisi())


# ==================== kayit listesi ====================
def _filtre_sql(args, izinli):
    where, par = ["1=1"], []
    projeler_sec = [p for p in args.getlist("proje") if p in izinli] or izinli
    if not projeler_sec:
        return "1=0", [], {}
    where.append("proje IN (%s)" % ",".join("?" * len(projeler_sec)))
    par += projeler_sec
    durum = args.get("durum", "bekleyen")
    if durum == "bekleyen":
        where.append("kt=0 AND silindi=0 AND excel_durum='var'")
    elif durum == "onayli":
        where.append("kt=1 AND silindi=0 AND excel_durum='var'")
    elif durum == "silinen":
        where.append("(silindi=1 OR excel_durum='silinmis')")
    elif durum == "hepsi":
        pass
    kat = (args.get("kategori") or "").strip()
    if kat:
        where.append("upper(kategori)=?")
        par.append(kat.upper())
    q = (args.get("q") or "").strip()
    if q:
        where.append("(personel=? OR uye_id=? OR ana_uye_id=? OR uye_id LIKE ? OR ana_uye_id LIKE ? OR link LIKE ?)")
        par += [q, q, q, f"%{q}%", f"%{q}%", f"%{q}%"]
    ekleyen = (args.get("ekleyen") or "").strip()
    if ekleyen:
        where.append("(ekleyen_ad LIKE ? OR ekleyen_id=?)")
        par += [f"%{ekleyen}%", ekleyen]
    t1, t2 = (args.get("tarih1") or "").strip(), (args.get("tarih2") or "").strip()
    if t1:
        where.append("tarih >= ?")
        par.append(t1 + " 00:00:00")
    if t2:
        where.append("tarih <= ?")
        par.append(t2 + " 23:59:59")
    if args.get("tarihsiz") == "1":
        where.append("tarih IS NULL")
    return " AND ".join(where), par, {"proje": projeler_sec, "durum": durum, "kategori": kat, "q": q, "ekleyen": ekleyen, "tarih1": t1, "tarih2": t2}


def kayit_sozluk(r):
    d = dict(r)
    d["gorsel_var"] = bool(d.get("gorsel_dosya")) or d.get("link_turu") in ("gyazo", "onedrive", "snappark", "teamzone", "diger")
    d["renk"] = renk(d["proje"])
    d["izin_kt"] = kayit_izin(d, "kt")
    d["izin_duzenle"] = kayit_izin(d, "duzenle")
    d["izin_sil"] = kayit_izin(d, "sil")
    d["geri_al_olur"] = bool(d.get("silindi")) and d.get("bekleyen_islem") != "sil"
    d["dis_link"] = d["link"] if (d.get("link") or "").startswith("http") else None
    return d


@app.route("/kayitlar")
@giris_gerekli
def kayitlar():
    c = baglanti()
    izinli = yetki.izinli_projeler(g.yk, projeler())
    where, par, f = _filtre_sql(request.args, izinli)
    gorunum = request.args.get("gorunum", "kart")
    sayfa = max(1, int(request.args.get("s", 1) or 1))
    boy = int(AYAR["SAYFA_BOYU"])
    toplam = c.execute(f"SELECT COUNT(*) FROM kayitlar WHERE {where}", par).fetchone()[0]
    rows = c.execute(f"SELECT * FROM kayitlar WHERE {where} ORDER BY COALESCE(tarih,'0000') DESC, id DESC LIMIT ? OFFSET ?", (*par, boy, (sayfa - 1) * boy)).fetchall()
    kategoriler = [r[0] for r in c.execute("SELECT DISTINCT upper(kategori) FROM kayitlar WHERE kategori<>'' ORDER BY 1").fetchall()]
    sayac = {}
    for d in ("bekleyen", "onayli", "silinen", "hepsi"):
        a = request.args.to_dict(flat=False); a["durum"] = [d]
        w2, p2, _ = _filtre_sql(MultiDict(a), izinli)
        sayac[d] = c.execute(f"SELECT COUNT(*) FROM kayitlar WHERE {w2}", p2).fetchone()[0]
    return render_template("kayitlar.html", kayitlar=[kayit_sozluk(r) for r in rows], toplam=toplam, sayfa=sayfa, sayfa_sayisi=max(1, (toplam + boy - 1) // boy),
                           f=f, izinli=izinli, kategoriler=kategoriler, gorunum=gorunum, sayac=sayac)


# ==================== disa aktarma (Excel ciktisi) ====================
DISA_SUTUNLAR = [("Proje", "proje", 8), ("Kategori", "kategori", 10), ("Personel kodu", "personel", 14), ("Uye ID", "uye_id", 14), ("Ana uye ID", "ana_uye_id", 14),
                 ("Tur", "tip", 12), ("Tarih", "tarih", 19), ("Link", "link", 60), ("Onay (KT)", "kt", 10), ("Onaylayan", "kt_kim", 14), ("Onay zamani", "kt_zaman", 19),
                 ("Ekleyen (Telegram)", "ekleyen_ad", 18), ("Ekleyen ID", "ekleyen_id", 14), ("Aciklama", "aciklama", 30), ("Durum", "_durum", 12), ("Excel satiri", "excel_satir", 11), ("Gorsel dosyasi", "_gorsel", 40)]


def _kayit_durum(d):
    return "silindi" if d.get("silindi") else ("Excel'de yok" if d.get("excel_durum") == "silinmis" else ("bekliyor" if d.get("bekleyen_islem") else "aktif"))


def _xlsx_satir(ws, degerler):
    """Sayfaya satir ekle -> satir no. = + - @ ile baslayan metinler formul degil METIN olarak yazilir (formul enjeksiyonu onlemi)."""
    ws.append(degerler)
    for h in ws[ws.max_row]:
        if isinstance(h.value, str) and h.value[:1] in "=+-@":
            h.data_type = "s"
    return ws.max_row


def _disa_aktar_sorgu():
    c = baglanti()
    izinli = yetki.izinli_projeler(g.yk, projeler())
    args = request.args.to_dict(flat=False)
    args.setdefault("durum", ["hepsi"])
    where, par, f = _filtre_sql(MultiDict(args), izinli)
    rows = c.execute(f"SELECT * FROM kayitlar WHERE {where} ORDER BY proje, COALESCE(tarih,'0000'), id", par).fetchall()
    return rows, f


@app.route("/disa-aktar")
@giris_gerekli
def disa_aktar_sayfasi():
    c = baglanti()
    izinli = yetki.izinli_projeler(g.yk, projeler())
    kategoriler = [r[0] for r in c.execute("SELECT DISTINCT upper(kategori) FROM kayitlar WHERE kategori<>'' ORDER BY 1").fetchall()]
    bugun = datetime.now().date()
    return render_template("disa_aktar.html", izinli=izinli, kategoriler=kategoriler, tarih1=(bugun - timedelta(days=30)).isoformat(), tarih2=bugun.isoformat())


@app.route("/disa-aktar.xlsx")
@giris_gerekli
def disa_aktar_xlsx():
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    rows, f = _disa_aktar_sorgu()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kayitlar"
    ws.append([b for b, _, _ in DISA_SUTUNLAR])
    for i, (_, _, gen) in enumerate(DISA_SUTUNLAR, 1):
        ws.column_dimensions[get_column_letter(i)].width = gen
        h = ws.cell(row=1, column=i)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor="1F2937")
        h.alignment = Alignment(vertical="center")
    for r in rows:
        d = dict(r)
        durum = _kayit_durum(d)
        satir = []
        for _, k, _ in DISA_SUTUNLAR:
            if k == "_durum":
                v = durum
            elif k == "_gorsel":
                v = os.path.basename(d["gorsel_dosya"]) if d.get("gorsel_dosya") else ""
            elif k == "kt":
                v = "EVET" if d.get("kt") else "HAYIR"
            elif k in ("tarih", "kt_zaman") and d.get(k):
                try:
                    v = datetime.strptime(d[k][:19], "%Y-%m-%d %H:%M:%S")
                except ValueError:
                    v = d[k]
            else:
                v = d.get(k)
                if v is None:
                    v = ""
            satir.append(v)
        n = _xlsx_satir(ws, satir)
        link = d.get("link") or ""
        if link.startswith("http"):
            hucre = ws.cell(row=n, column=8)
            hucre.hyperlink = link
            hucre.font = Font(color="0563C1", underline="single")
        for col in (7, 11):
            if isinstance(ws.cell(row=n, column=col).value, datetime):
                ws.cell(row=n, column=col).number_format = "yyyy-mm-dd hh:mm:ss"
        renk_hex = renk(d["proje"]).lstrip("#")
        if len(renk_hex) == 6:
            ws.cell(row=n, column=1).fill = PatternFill("solid", fgColor=renk_hex)
            ws.cell(row=n, column=1).font = Font(bold=True, color=metin_rengi("#" + renk_hex).lstrip("#"))
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(DISA_SUTUNLAR))}{ws.max_row}"
    # ozet sayfasi
    oz = wb.create_sheet("Ozet")
    oz.append(["Proje", "Toplam", "Onayli", "Onay bekleyen", "Silinen"])
    for i in range(1, 6):
        oz.cell(row=1, column=i).font = Font(bold=True)
        oz.column_dimensions[get_column_letter(i)].width = 14
    sayac = {}
    for r in rows:
        s_ = sayac.setdefault(r["proje"], [0, 0, 0, 0])
        s_[0] += 1
        if r["silindi"] or r["excel_durum"] == "silinmis":
            s_[3] += 1
        elif r["kt"]:
            s_[1] += 1
        else:
            s_[2] += 1
    for p_ in sorted(sayac):
        oz.append([p_] + sayac[p_])
    oz.append(["TOPLAM"] + [sum(v[i] for v in sayac.values()) for i in range(4)])
    oz.append([])
    _xlsx_satir(oz, ["Filtre", f"projeler: {', '.join(f.get('proje', []))}; durum: {f.get('durum')}; tarih: {f.get('tarih1') or '-'} / {f.get('tarih2') or '-'}; kategori: {f.get('kategori') or '-'}; arama: {f.get('q') or '-'}"])
    _xlsx_satir(oz, ["Olusturan", f"{g.k['kullanici_adi']} - {simdi()}"])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    projeler_ad = "-".join(f.get("proje", [])) if f.get("proje") and len(f["proje"]) <= 3 else "tum"
    ad = f"kanit_{projeler_ad}_{(f.get('tarih1') or 'bastan')}_{(f.get('tarih2') or 'bugun')}.xlsx".replace("/", "-")
    db.denetim_yaz(baglanti(), g.k["kullanici_adi"], "disa_aktar", yeni=f"{len(rows)} satir | {ad}", ip=ip())
    return send_file(bio, as_attachment=True, download_name=ad, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ==================== liste kontrolu (dosyadaki degerler kayitlarda var mi?) ====================
LISTE_KAYIT_SUTUNLAR = [("Proje", "proje", 8), ("Kategori", "kategori", 10), ("Personel kodu", "personel", 14), ("Uye ID", "uye_id", 14), ("Ana uye ID", "ana_uye_id", 14),
                        ("Tarih", "tarih", 19), ("Onay (KT)", "kt", 10), ("Link", "link", 50), ("Ekleyen (Telegram)", "ekleyen_ad", 18), ("Durum", "_durum", 12), ("Kayit ID", "id", 9)]


def _liste_kayitlar(filtre, izinli):
    """Filtreye (form ya da kaydedilmis filtre sozlugu) uyan kayitlar; yalnizca eslestirme/gosterim sutunlari."""
    a = {k: (v if isinstance(v, list) else [v]) for k, v in dict(filtre).items()}
    a.setdefault("durum", ["hepsi"])
    where, par, f = _filtre_sql(MultiDict(a), izinli)
    rows = baglanti().execute(f"SELECT id, proje, kategori, personel, uye_id, ana_uye_id, link, tarih, kt, silindi, excel_durum, bekleyen_islem, ekleyen_ad, aciklama FROM kayitlar WHERE {where} ORDER BY proje, COALESCE(tarih,'0000') DESC, id DESC", par).fetchall()
    return [dict(r) for r in rows], f


def _liste_esleme(form, genislik):
    """Formdan esleme: [{"sutun": i, "alanlar": [...]}]"""
    esleme = []
    for sutun in form.getlist("kontrol"):
        if sutun.isdigit() and int(sutun) < genislik:
            alanlar = [a for a in form.getlist(f"alan_{sutun}") if a in liste_kontrol.ALAN_ADLARI]
            if not alanlar:
                raise ValueError(f"{int(sutun) + 1}. sutun icin panelde hangi alanda aranacagini secin.")
            esleme.append({"sutun": int(sutun), "alanlar": alanlar})
    if not esleme:
        raise ValueError("Kontrol edilecek en az bir sutun isaretleyin.")
    return esleme


def _liste_basliklar(veri):
    satirlar = veri["tablolar"][veri["sayfa"]]
    return [(satirlar[0][e["sutun"]] if veri["baslik"] and satirlar and satirlar[0][e["sutun"]] else f"Sutun {e['sutun'] + 1}") for e in veri["esleme"]]


def _liste_bulunamadi():
    flash("Liste bulunamadi ya da suresi dolmus; dosyayi yeniden yukleyin.")
    return redirect(url_for("disa_aktar_sayfasi"))


@app.route("/liste-kontrol", methods=["GET", "POST"])
@giris_gerekli
def liste_kontrol_sayfasi():
    onbellek = AYAR["ONBELLEK"]
    izinli = yetki.izinli_projeler(g.yk, projeler())
    adim = request.form.get("adim") if request.method == "POST" else "esleme"
    if adim == "yukle":
        dosya = request.files.get("dosya")
        try:
            if dosya and dosya.filename:
                veri = liste_kontrol.coz(dosya.filename, dosya.read())
                veri["dosya"] = os.path.basename(dosya.filename)[:80]
            elif (request.form.get("metin") or "").strip():
                veri = liste_kontrol.metin_coz(request.form["metin"])
                veri["dosya"] = "yapistirilan liste"
            else:
                raise ValueError("bir dosya secin ya da listeyi kutuya yapistirin.")
        except ValueError as e:
            flash(f"Liste okunamadi: {e}")
            return redirect(url_for("disa_aktar_sayfasi"))
        veri.update(sahip=g.k["kullanici_adi"], sayfa=veri["sayfalar"][0])
        veri["baslik"] = liste_kontrol.baslik_mi(veri["tablolar"][veri["sayfa"]])
        return redirect(url_for("liste_kontrol_sayfasi", token=liste_kontrol.kaydet(onbellek, veri)))
    token = request.values.get("token", "")
    veri = liste_kontrol.yukle(onbellek, token, g.k["kullanici_adi"])
    if not veri:
        return _liste_bulunamadi()
    hata = None
    if adim in ("onizle", "kontrol"):
        sayfa = request.form.get("sayfa")
        if sayfa in veri["tablolar"] and sayfa != veri["sayfa"]:
            veri["sayfa"] = sayfa
            veri["baslik"] = liste_kontrol.baslik_mi(veri["tablolar"][sayfa])
        else:
            veri["baslik"] = request.form.get("baslik") == "1"
        if adim == "kontrol":
            try:
                satirlar = veri["tablolar"][veri["sayfa"]]
                veri["esleme"] = _liste_esleme(request.form, len(satirlar[0]) if satirlar else 0)
                kayitlar, veri["filtre"] = _liste_kayitlar(request.form.to_dict(flat=False), izinli)
                liste_kontrol.kaydet(onbellek, veri, token)
                return _liste_sonuc_sayfasi(token, veri, kayitlar)
            except ValueError as e:
                hata = str(e)
        liste_kontrol.kaydet(onbellek, veri, token)
    satirlar = veri["tablolar"][veri["sayfa"]]
    sutunlar = liste_kontrol.sutunlar(satirlar, veri["baslik"])
    if veri.get("esleme"):   # sonuc sayfasindan geri donuldu: onceki secimler
        secim = {e["sutun"]: e["alanlar"] for e in veri["esleme"]}
        for st in sutunlar:
            st["secili"], st["alanlar"] = st["no"] in secim, secim.get(st["no"], st["alanlar"])
    kategoriler = [r[0] for r in baglanti().execute("SELECT DISTINCT upper(kategori) FROM kayitlar WHERE kategori<>'' ORDER BY 1").fetchall()]
    return render_template("liste_kontrol.html", token=token, veri=veri, sutunlar=sutunlar, onizleme=satirlar[1 if veri["baslik"] else 0:][:liste_kontrol.ONIZLEME],
                           satir_sayisi=len(satirlar) - (1 if veri["baslik"] else 0), alanlar=liste_kontrol.ALANLAR, izinli=izinli, kategoriler=kategoriler, f=veri.get("filtre") or {}, hata=hata)


def _liste_sonuc_sayfasi(token, veri, kayitlar):
    satirlar = veri["tablolar"][veri["sayfa"]]
    sonuc = liste_kontrol.eslestir(satirlar, veri["baslik"], veri["esleme"], kayitlar)
    bulunan, yok, proje_sayac = [], [], {}
    for s in sonuc["satirlar"]:
        if not s["eslesen"]:
            yok.append(s)
            continue
        for p in {kayitlar[i]["proje"] for i in s["eslesen"]}:
            proje_sayac[p] = proje_sayac.get(p, 0) + 1
        bulunan.append({"no": s["no"], "degerler": s["degerler"], "kayitlar": [kayitlar[i] for i in s["eslesen"][:3]], "n": len(s["eslesen"])})
    db.denetim_yaz(baglanti(), g.k["kullanici_adi"], "liste_kontrol", yeni=f"{sonuc['bulunan'] + sonuc['yok']} satir, bulunan {sonuc['bulunan']}, yok {sonuc['yok']}", ip=ip())
    return render_template("liste_sonuc.html", token=token, veri=veri, sonuc=sonuc, bulunan=bulunan[:2000], yok=yok[:2000], proje_sayac=sorted(proje_sayac.items()),
                           basliklar=_liste_basliklar(veri), f=veri["filtre"])


@app.route("/liste-kontrol.xlsx")
@giris_gerekli
def liste_kontrol_xlsx():
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    veri = liste_kontrol.yukle(AYAR["ONBELLEK"], request.args.get("token", ""), g.k["kullanici_adi"])
    if not veri or not veri.get("esleme"):
        return _liste_bulunamadi()
    kayitlar, f = _liste_kayitlar(veri["filtre"], yetki.izinli_projeler(g.yk, projeler()))
    sonuc = liste_kontrol.eslestir(veri["tablolar"][veri["sayfa"]], veri["baslik"], veri["esleme"], kayitlar)
    dosya_basliklar = _liste_basliklar(veri)
    wb = openpyxl.Workbook()

    def sayfa(ad, basliklar, genislikler):
        ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet()
        ws.title = ad
        _xlsx_satir(ws, basliklar)
        for i, gen in enumerate(genislikler, 1):
            ws.column_dimensions[get_column_letter(i)].width = gen
            ws.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF")
            ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="1F2937")
        ws.freeze_panes = "A2"
        return ws

    ws = sayfa("Bulunamayanlar", ["Satir no"] + dosya_basliklar, [9] + [24] * len(dosya_basliklar))
    for s in sonuc["satirlar"]:
        if not s["eslesen"]:
            _xlsx_satir(ws, [s["no"]] + s["degerler"])
    ws = sayfa("Bulunanlar", ["Satir no"] + dosya_basliklar + ["Eslesme sayisi"] + [b for b, _, _ in LISTE_KAYIT_SUTUNLAR], [9] + [24] * len(dosya_basliklar) + [10] + [gen for _, _, gen in LISTE_KAYIT_SUTUNLAR])
    sutun0 = len(dosya_basliklar) + 2   # kayit sutunlarinin basladigi yer
    for s in sonuc["satirlar"]:
        for i in s["eslesen"][:10]:   # satir basina en fazla 10 eslesme; toplam sayi ayri sutunda
            d = kayitlar[i]
            degerler = [_kayit_durum(d) if k == "_durum" else ("EVET" if d["kt"] else "HAYIR") if k == "kt" else (d.get(k) if d.get(k) is not None else "") for _, k, _ in LISTE_KAYIT_SUTUNLAR]
            n = _xlsx_satir(ws, [s["no"]] + s["degerler"] + [len(s["eslesen"])] + degerler)
            if (d.get("link") or "").startswith("http"):
                ws.cell(row=n, column=sutun0 + 8).hyperlink = d["link"]
                ws.cell(row=n, column=sutun0 + 8).font = Font(color="0563C1", underline="single")
            renk_hex = renk(d["proje"]).lstrip("#")
            if len(renk_hex) == 6:
                ws.cell(row=n, column=sutun0 + 1).fill = PatternFill("solid", fgColor=renk_hex)
                ws.cell(row=n, column=sutun0 + 1).font = Font(bold=True, color=metin_rengi("#" + renk_hex).lstrip("#"))
    oz = sayfa("Ozet", ["Alan", "Deger"], [22, 90])
    proje_sayac = {}
    for s in sonuc["satirlar"]:
        for p in {kayitlar[i]["proje"] for i in s["eslesen"]}:
            proje_sayac[p] = proje_sayac.get(p, 0) + 1
    ozet = [["Toplam satir", sonuc["bulunan"] + sonuc["yok"]], ["Bulunan", sonuc["bulunan"]], ["Bulunamayan", sonuc["yok"]], ["Birden fazla eslesen", sonuc["coklu"]], ["Bos satir (atlandi)", sonuc["bos"]]]
    ozet += [[f"Bulunan ({p})", n] for p, n in sorted(proje_sayac.items())]
    ozet += [["Dosya", veri.get("dosya", "") + (f" / sayfa: {veri['sayfa']}" if veri["sayfa"] else "") + (" (ilk satir baslik)" if veri["baslik"] else "")],
             ["Esleme", " VE ".join(f"{b} -> {' / '.join(liste_kontrol.ALAN_ADLARI[a] for a in e['alanlar'])}" for b, e in zip(dosya_basliklar, veri["esleme"]))],
             ["Filtre", f"projeler: {', '.join(f.get('proje', []))}; durum: {f.get('durum')}; tarih: {f.get('tarih1') or '-'} / {f.get('tarih2') or '-'}; kategori: {f.get('kategori') or '-'}"],
             ["Olusturan", f"{g.k['kullanici_adi']} - {simdi()}"]]
    for satir in ozet:
        _xlsx_satir(oz, satir)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f"liste_kontrol_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/kayit/<int:kid>")
@giris_gerekli
def kayit_json(kid):
    r = baglanti().execute("SELECT * FROM kayitlar WHERE id=?", (kid,)).fetchone()
    if not r or not kayit_izin(r, "oku"):
        abort(404)
    d = kayit_sozluk(r)
    d["gorsel_url"] = url_for("gorsel_goster", kid=kid) if kayit_izin(r, "gorsel") else None
    d["kucuk_url"] = url_for("kucuk_goster", kid=kid) if kayit_izin(r, "gorsel") else None
    yol, kaynak = COZUCU.yerel_dosya(d)
    d["gorsel_kaynak"] = kaynak if yol else ("gyazo_dis" if d["link_turu"] == "gyazo" else ("dis_link" if d["dis_link"] else "yok"))
    d["gecmis"] = [dict(x) for x in baglanti().execute("SELECT zaman,kullanici,islem,eski,yeni FROM denetim WHERE anahtar=? ORDER BY id DESC LIMIT 30", (r["anahtar"],)).fetchall()]
    return jsonify(d)


def _kayit_al(kid, eylem):
    r = baglanti().execute("SELECT * FROM kayitlar WHERE id=?", (kid,)).fetchone()
    if not r or not kayit_izin(r, eylem):
        abort(403)
    return r


def _parmak(r):
    return {"personel": r["personel"] or "", "uye": r["uye_id"] or "", "ana": r["ana_uye_id"] or "", "tip": r["tip"] or "", "kategori": r["kategori"] or "", "tarih": r["tarih_ham"] or r["tarih"] or ""}


@app.route("/kayit/<int:kid>/kt", methods=["POST"])
@giris_gerekli
def kayit_kt(kid):
    r = _kayit_al(kid, "kt")
    yeni = 1 if request.form.get("deger", "1") == "1" else 0
    c = baglanti()
    qid = kuyruk.ekle(AYAR, "kt", r["proje"], r["link"], _parmak(r), {"kt": bool(yeni), "satir": r["excel_satir"]}, g.k["kullanici_adi"], r["anahtar"])
    c.execute("UPDATE kayitlar SET kt=?, kt_kaynak='panel', kt_kim=?, kt_zaman=?, bekleyen_islem='kt', bekleyen_id=?, bekleyen_hata=NULL, guncelleme=? WHERE id=?",
              (yeni, g.k["kullanici_adi"], simdi(), qid, simdi(), kid))
    db.denetim_yaz(c, g.k["kullanici_adi"], "onayla" if yeni else "onay_kaldir", r["proje"], r["anahtar"], r["kt"], yeni, ip())
    return jsonify({"ok": True, "kt": yeni})


@app.route("/kayit/<int:kid>/duzenle", methods=["POST"])
@giris_gerekli
def kayit_duzenle(kid):
    r = _kayit_al(kid, "duzenle")
    alanlar = {}
    for k, kolon in (("personel", "personel"), ("uye_id", "uye_id"), ("ana_uye_id", "ana_uye_id"), ("tip", "tip"), ("kategori", "kategori"), ("aciklama", "aciklama")):
        if k in request.form:
            alanlar[kolon] = (request.form.get(k) or "").strip()[:200]
    if not alanlar:
        return jsonify({"ok": False, "hata": "degisiklik yok"})
    if "kategori" in alanlar:
        alanlar["kategori"] = alanlar["kategori"].upper()   # bot Excel'e buyuk harfle yazar
    eski = {k: r[k] for k in alanlar}
    alanlar = {k: v for k, v in alanlar.items() if (eski[k] or "") != v}   # yalnizca DEGISEN alanlar gider (baska adminin duzeltmesi ezilmez)
    if not alanlar:
        return jsonify({"ok": True, "degisiklik": False})
    eski = {k: eski[k] for k in alanlar}
    c = baglanti()
    ad_haritasi = {"personel": "personel", "uye_id": "uye", "ana_uye_id": "ana", "tip": "tip", "kategori": "kategori", "aciklama": "aciklama"}
    deger = {ad_haritasi[k]: v for k, v in alanlar.items()}
    deger["satir"] = r["excel_satir"]
    qid = kuyruk.ekle(AYAR, "duzenle", r["proje"], r["link"], _parmak(r), deger, g.k["kullanici_adi"], r["anahtar"])
    sets = ", ".join(f"{k}=?" for k in alanlar)
    c.execute(f"UPDATE kayitlar SET {sets}, bekleyen_islem='duzenle', bekleyen_id=?, bekleyen_hata=NULL, guncelleme=? WHERE id=?", (*alanlar.values(), qid, simdi(), kid))
    db.denetim_yaz(c, g.k["kullanici_adi"], "duzenle", r["proje"], r["anahtar"], json.dumps(eski, ensure_ascii=False), json.dumps(alanlar, ensure_ascii=False), ip())
    return jsonify({"ok": True})


@app.route("/kayit/<int:kid>/sil", methods=["POST"])
@giris_gerekli
def kayit_sil(kid):
    r = _kayit_al(kid, "sil")
    c = baglanti()
    qid = kuyruk.ekle(AYAR, "sil", r["proje"], r["link"], _parmak(r), {"satir": r["excel_satir"]}, g.k["kullanici_adi"], r["anahtar"])
    c.execute("UPDATE kayitlar SET silindi=1, silen=?, silme_zamani=?, bekleyen_islem='sil', bekleyen_id=?, bekleyen_hata=NULL, guncelleme=? WHERE id=?",
              (g.k["kullanici_adi"], simdi(), qid, simdi(), kid))
    db.denetim_yaz(c, g.k["kullanici_adi"], "sil", r["proje"], r["anahtar"], json.dumps(_parmak(r), ensure_ascii=False), None, ip())
    return jsonify({"ok": True})


@app.route("/kayit/<int:kid>/geri-al", methods=["POST"])
@giris_gerekli
def kayit_geri_al(kid):
    r = _kayit_al(kid, "sil")
    if r["bekleyen_islem"] == "sil":
        return jsonify({"ok": False, "hata": "Silme islemi henuz bota ulasmadi; bot cevaplayinca geri alabilirsiniz."})
    c = baglanti()
    degerler = [r["personel"] or "", r["uye_id"] or "", r["ana_uye_id"] or "", r["tip"] or "", r["link"] or "", (r["kategori"] or "").upper(), r["tarih_ham"] or r["tarih"] or ""]
    qid = kuyruk.ekle(AYAR, "ekle", r["proje"], r["link"], _parmak(r), {"degerler": degerler, "kt": bool(r["kt"]), "gorsel2": r["gorsel2"] or "", "aciklama": r["aciklama"] or ""}, g.k["kullanici_adi"], r["anahtar"])
    c.execute("UPDATE kayitlar SET silindi=0, silen=NULL, silme_zamani=NULL, bekleyen_islem='geri_al', bekleyen_id=?, bekleyen_hata=NULL, guncelleme=? WHERE id=?", (qid, simdi(), kid))
    db.denetim_yaz(c, g.k["kullanici_adi"], "geri_al", r["proje"], r["anahtar"], None, json.dumps(degerler, ensure_ascii=False), ip())
    return jsonify({"ok": True})


# ==================== gorsel ====================
def _gorsel_yolu(kid, kucuk=False):
    r = baglanti().execute("SELECT * FROM kayitlar WHERE id=?", (kid,)).fetchone()
    if not r or not kayit_izin(r, "gorsel"):
        abort(404)
    d = dict(r)
    yol, kaynak = COZUCU.yerel_dosya(d)
    if not yol and d["link_turu"] == "gyazo":
        yol = COZUCU.gyazo_indir(d)
    if not yol:
        return None, d
    if kucuk:
        yol = COZUCU.kucuk(d, yol)
    return yol, d


@app.route("/gorsel/<int:kid>")
@giris_gerekli
def gorsel_goster(kid):
    yol, d = _gorsel_yolu(kid)
    if not yol:
        if (d.get("link") or "").startswith("http"):
            return redirect(d["link"])
        abort(404)
    db.denetim_yaz(baglanti(), g.k["kullanici_adi"], "gorsel_ac", d["proje"], d["anahtar"], ip=ip())
    resp = send_file(yol, conditional=True)
    resp.headers["Cache-Control"] = "private, max-age=3600"
    return resp


@app.route("/kucuk/<int:kid>")
@giris_gerekli
def kucuk_goster(kid):
    yol, d = _gorsel_yolu(kid, kucuk=True)
    if not yol:
        abort(404)
    resp = send_file(yol, conditional=True)
    resp.headers["Cache-Control"] = "private, max-age=86400"
    return resp


# ==================== yonetim: kullanicilar ====================
def _rastgele_sifre(n=10):
    return "".join(random.choice(string.ascii_letters + string.digits) for _ in range(n))


@app.route("/yonetim/kullanicilar", methods=["GET", "POST"])
@admin_gerekli
def yonetim_kullanicilar():
    c = baglanti()
    mesaj = None
    if request.method == "POST":
        islem = request.form.get("islem")
        try:
            if islem == "ekle":
                ad = (request.form.get("kullanici") or "").strip().lower()
                yk = (request.form.get("yetki") or "").strip()
                if not kullanici_adi_gecerli(ad):
                    raise ValueError("Kullanici adi harf/rakam olmali (en fazla 40 karakter).")
                sifre = request.form.get("sifre") or _rastgele_sifre()
                kullanicilar.ekle(AYAR, ad, yk, sifre, gorunen_ad=request.form.get("ad") or ad, olusturan=g.k["kullanici_adi"])
                db.denetim_yaz(c, g.k["kullanici_adi"], "kullanici_ekle", yeni=f"{ad} [{yk}]", ip=ip())
                mesaj = f"Kullanici acildi: {ad}  gecici sifre: {sifre}"
            else:
                ad = (request.form.get("kullanici_adi") or "").strip().lower()
                u = kullanicilar.bul(AYAR, ad)
                if not u:
                    mesaj = "Kullanici bulunamadi."
                elif islem == "yetki":
                    yk = (request.form.get("yetki") or "").strip()
                    kullanicilar.guncelle(AYAR, ad, yetki=yk)
                    db.denetim_yaz(c, g.k["kullanici_adi"], "yetki_degisti", eski=f"{ad} [{u['yetki']}]", yeni=yk, ip=ip())
                    mesaj = "Yetki guncellendi."
                elif islem == "sifre":
                    sifre = _rastgele_sifre()
                    kullanicilar.guncelle(AYAR, ad, sifre=sifre, sifre_degistir=True)
                    db.denetim_yaz(c, g.k["kullanici_adi"], "sifre_sifirla", yeni=ad, ip=ip())
                    mesaj = f"{ad} icin gecici sifre: {sifre}"
                elif islem == "pasif":
                    if ad == g.k["kullanici_adi"]:
                        mesaj = "Kendi hesabinizi pasife alamazsiniz."
                    else:
                        yeni = not u.get("aktif", True)
                        kullanicilar.guncelle(AYAR, ad, aktif=yeni)
                        db.denetim_yaz(c, g.k["kullanici_adi"], "kullanici_aktif" if yeni else "kullanici_pasif", yeni=ad, ip=ip())
                        mesaj = "Hesap " + ("aktif." if yeni else "pasif.")
        except (ValueError, IOError) as e:
            mesaj = f"Hata: {e}"
    liste = [dict(u) | {"ozet": yetki.ozet(u.get("yetki", ""))} for u in kullanicilar.hepsi(AYAR)]
    return render_template("kullanicilar.html", kullanicilar=liste, mesaj=mesaj, tum_projeler=projeler(), dosya=kullanicilar.yol(AYAR))


# ==================== yonetim: bot ====================
@app.route("/bot", methods=["GET", "POST"])
@admin_gerekli
def bot_sayfasi():
    c = baglanti()
    mesaj = None
    if request.method == "POST":
        islem = request.form.get("islem")
        tg = (request.form.get("tg_id") or "").strip()
        ad = (request.form.get("ad") or "").strip()
        kim = g.k["kullanici_adi"]
        try:
            kuyrukta = True
            if islem == "onayla":
                bot_yonetim.onayla(AYAR, c, tg, ad or tg, kim)
            elif islem == "onay_kaldir":
                bot_yonetim.onay_kaldir(AYAR, c, tg, kim)
            elif islem == "engelle":
                bot_yonetim.engelle(AYAR, c, tg, ad, request.form.get("sebep") or "panel", kim)
            elif islem == "engel_kaldir":
                bot_yonetim.engel_kaldir(AYAR, c, tg, kim)
            elif islem == "bakim_ac":
                bot_yonetim.bakim_modu_ayarla(AYAR, c, True, kim)
            elif islem == "bakim_kapat":
                bot_yonetim.bakim_modu_ayarla(AYAR, c, False, kim)
            elif islem == "gyazo_bakim_ac":
                bot_yonetim.gyazo_bakim_ayarla(AYAR, c, True, kim)
            elif islem == "gyazo_bakim_kapat":
                bot_yonetim.gyazo_bakim_ayarla(AYAR, c, False, kim)
            elif islem == "oturum_temizle":
                bot_yonetim.oturum_temizle(AYAR, c, tg or None, kim)
            elif islem == "kategori_ekle":
                bot_yonetim.kategori_kaydet(AYAR, c, (request.form.get("code") or "").strip().lower(), (request.form.get("label") or "").strip(), request.form.get("mode") or "tek_uye", kim)
            elif islem == "kategori_sil":
                bot_yonetim.kategori_sil(AYAR, c, request.form.get("code") or "", kim)
            elif islem == "yedek":
                bot_yonetim.yedek_al(AYAR, c, kim)
            elif islem == "tam_sifirla":
                bot_yonetim.tam_sifirla(AYAR, c, tg, kim)
            elif islem == "excel_mod":
                bot_yonetim.excel_mod_ayarla(AYAR, c, request.form.get("mod") or "", kim)
            elif islem == "excel_son_satir":
                bot_yonetim.excel_son_satir_sor(AYAR, c, request.form.get("sayfa") or "", kim)
            elif islem == "kapat":
                if (request.form.get("onay") or "").strip().upper() != "KAPAT":
                    raise ValueError("Botu durdurmak icin kutuya KAPAT yazin.")
                bot_yonetim.bot_kapat(AYAR, c, kim)
            elif islem == "yenile":
                kuyrukta = False
                rapor = ice_aktar.calistir(AYAR, c, AYAR["ONBELLEK"], COZUCU, zorla=True)
                mesaj = f"Excel yeniden okundu: {rapor.get('sayfalar')} yeni={rapor.get('yeni')} degisen={rapor.get('degisen')} {('HATA: ' + rapor['hata']) if rapor.get('hata') else ''}"
            else:
                kuyrukta = False
                mesaj = "Bilinmeyen islem."
            if kuyrukta:
                mesaj = "Islem bota iletildi; bot uygulayinca asagidaki 'Son islemler' listesinde 'ok' olur (bot acikken birkac saniye, OneDrive uzerinden 1-2 dakika)."
            db.denetim_yaz(c, kim, f"bot_{islem}", yeni=f"{tg} {ad}".strip() or None, ip=ip())
            if kuyrukta and tg:
                return redirect(url_for("bot_sayfasi", tg=tg, m=1))
        except Exception as e:
            mesaj = f"Hata: {e}"
    if request.args.get("m") and not mesaj:
        mesaj = "Islem bota iletildi; bot uygulayinca 'Son islemler' listesinde 'ok' olur."
    tg_f = {"q": (request.args.get("tg_q") or "").strip()[:60], "sure": _sayi(request.args.get("tg_sure"), 0), "durum": request.args.get("tg_durum") or "",
            "adet": _sayi(request.args.get("tg_adet"), 0)}
    tum = bot_yonetim.telegram_kullanicilari(AYAR)
    tg_liste = bot_yonetim.kullanici_filtrele(tum, tg_f["q"], tg_f["sure"], tg_f["durum"], tg_f["adet"])
    return render_template("bot.html", mesaj=mesaj, durum=durum_bilgisi(), tg_kullanicilar=tg_liste, tg_toplam=len(tum), tg_f=tg_f,
                           kategoriler=bot_yonetim.kategoriler(AYAR), bakim=bot_yonetim.bakim_modu(AYAR), gyazo_bakim=bot_yonetim.gyazo_bakim(AYAR),
                           oturumlar=bot_yonetim.oturumlar(AYAR), log=bot_yonetim.son_log(AYAR, 150), hatalar=bot_yonetim.son_log(AYAR, 40, True),
                           cfg=bot_yonetim.config(AYAR), islemler=bot_yonetim.son_islemler(c), ist=bot_yonetim.istatistik(AYAR, c),
                           excel_modlari=bot_yonetim.EXCEL_MODLARI, ac_tg=(request.args.get("tg") or "").strip()[:20])


def _sayi(v, varsayilan=0):
    try:
        return int(v)
    except (TypeError, ValueError):
        return varsayilan


@app.route("/bot/kullanici/<tg_id>.json")
@admin_gerekli
def bot_kullanici_json(tg_id):
    tg_id = tg_id.strip()
    if not tg_id.isdigit() or len(tg_id) > 20:
        abort(404)
    return jsonify(bot_yonetim.kullanici_detay(AYAR, baglanti(), tg_id))


@app.route("/bot/islemler.json")
@admin_gerekli
def bot_islemler_json():
    return jsonify({"islemler": bot_yonetim.son_islemler(baglanti()), "bekleyen": kuyruk.bekleyen_sayisi(baglanti())})


@app.route("/denetim")
@admin_gerekli
def denetim():
    c = baglanti()
    q = (request.args.get("q") or "").strip()
    rows = c.execute("SELECT * FROM denetim WHERE (? = '' OR kullanici LIKE ? OR islem LIKE ? OR anahtar LIKE ?) ORDER BY id DESC LIMIT 500", (q, f"%{q}%", f"%{q}%", f"%{q}%")).fetchall()
    return render_template("denetim.html", kayitlar=[dict(r) for r in rows], q=q)


@app.route("/durum.json")
@giris_gerekli
def durum_json():
    return jsonify(durum_bilgisi())


@app.route("/bildirim.json")
@giris_gerekli
def bildirim_json():
    """Canli yenileme: son goruldugunden beri eklenen kayitlar (kullanicinin yetkili oldugu projeler) + guncel sayaclar + degisiklik damgasi."""
    c = baglanti()
    izinli = yetki.izinli_projeler(g.yk, projeler())
    try:
        son = int(request.args.get("son", -1))
    except ValueError:
        son = -1
    maks = c.execute("SELECT COALESCE(MAX(id),0) FROM kayitlar").fetchone()[0]
    yeni = []
    if son >= 0 and izinli and maks > son:
        q = f"""SELECT id, proje, personel, kategori, tarih, ekleyen_ad, kt FROM kayitlar
                WHERE id > ? AND proje IN ({",".join("?" * len(izinli))}) AND silindi=0 AND excel_durum='var'
                  AND COALESCE(tarih, ilk_gorulme) > datetime('now','localtime','-2 days')
                ORDER BY id DESC LIMIT 20"""
        yeni = [dict(r) for r in c.execute(q, (son, *izinli)).fetchall()]
    sayac = {}
    for r in c.execute("""SELECT proje, SUM(CASE WHEN kt=0 AND silindi=0 AND excel_durum='var' THEN 1 ELSE 0 END) bekleyen,
                                 SUM(CASE WHEN kt=1 AND silindi=0 AND excel_durum='var' THEN 1 ELSE 0 END) onayli,
                                 SUM(CASE WHEN silindi=1 OR excel_durum='silinmis' THEN 1 ELSE 0 END) silinen, COUNT(*) toplam
                          FROM kayitlar GROUP BY proje""").fetchall():
        if r["proje"] in izinli:
            sayac[r["proje"]] = {"bekleyen": r["bekleyen"] or 0, "onayli": r["onayli"] or 0, "silinen": r["silinen"] or 0, "toplam": r["toplam"]}
    return jsonify({"maks": maks, "yeni": yeni, "sayac": sayac, "degisiklik": db.ayar_al(c, "son_degisiklik"), "kuyruk_bekleyen": kuyruk.bekleyen_sayisi(c),
                    "bot": bot_yonetim.bot_calisiyor_mu(AYAR).get("canli")})


@app.route("/saglik")
def saglik():
    resp = app.make_response("OK")
    if yerel_istek():   # surum bilgisi yalnizca ayni bilgisayara (tek kopya kontrolu); agdan surum okunamaz
        resp.headers["X-Kanit-Paneli"] = SURUM
    return resp


@app.errorhandler(403)
def _403(e):
    acik = str(getattr(e, "description", "") or "")
    if acik.startswith("AG:"):
        mesaj = f"Bu adresten ({acik[3:]}) panele erisim kapali: panel yalnizca ofis (ozel) agindan acilabilir. Gerekiyorsa panel bilgisayarinda panel_ayar.json icine IZINLI_AGLAR yazilir."
        return render_template("hata.html", kod=403, mesaj=mesaj), 403
    return render_template("hata.html", kod=403, mesaj="Bu islem icin yetkiniz yok."), 403


@app.errorhandler(404)
def _404(e):
    return render_template("hata.html", kod=404, mesaj="Bulunamadi."), 404


@app.errorhandler(413)
def _413(e):
    flash(f"Dosya cok buyuk: en fazla {app.config['MAX_CONTENT_LENGTH'] // (1024 * 1024)} MB yuklenebilir.")
    return redirect(url_for("disa_aktar_sayfasi"))


# ==================== arka plan ice aktarim ====================
def _ice_aktar_bir_kez():
    try:
        conn = db.baglan(AYAR["DB"])
        ice_aktar.calistir(AYAR, conn, AYAR["ONBELLEK"], COZUCU)
    except Exception as e:
        print("ice aktarim hatasi:", e)


def _arka_plan():
    conn = db.baglan(AYAR["DB"])
    while True:
        if bot_klasoru_mu(AYAR.get("BOT_KLASOR")):
            try:
                ice_aktar.calistir(AYAR, conn, AYAR["ONBELLEK"], COZUCU)
            except Exception as e:
                print("ice aktarim hatasi:", e)
            time.sleep(max(10, int(float(AYAR.get("YENILEME_SN", 30)))))
        else:
            time.sleep(10)


# ==================== komut satiri ====================
def komut(argv):
    if not argv or argv[0] in ("--ac", "--ag", "--port", "--tunel"):
        return False
    if argv[0] == "--surum":
        print("KanitPaneli", SURUM)
        return True
    if not bot_klasoru_mu(AYAR.get("BOT_KLASOR")):
        print(f"Bot klasoru bulunamadi ({AYAR.get('BOT_KLASOR') or '-'}). Once paneli acip kurulum sayfasindan klasoru secin ya da {AYAR_YOL} icine BOT_KLASOR yazin.")
        return True
    if argv[0] == "kullanici-ekle" and len(argv) >= 3:
        ad, yk = argv[1].strip().lower(), argv[2]
        sifre = argv[3] if len(argv) > 3 else _rastgele_sifre()
        if kullanicilar.bul(AYAR, ad):
            kullanicilar.guncelle(AYAR, ad, yetki=yk, aktif=True)
            print(f"Kullanici zaten vardi; yetkisi guncellendi: {ad} [{yk}]")
        else:
            kullanicilar.ekle(AYAR, ad, yk, sifre, olusturan="komut")
            print(f"Kullanici acildi: {ad}   yetki: {yk}   gecici sifre: {sifre}")
        return True
    if argv[0] == "sifre-sifirla" and len(argv) >= 2:
        ad = argv[1].strip().lower()
        sifre = argv[2] if len(argv) > 2 else _rastgele_sifre()
        if kullanicilar.bul(AYAR, ad):
            kullanicilar.guncelle(AYAR, ad, sifre=sifre, sifre_degistir=True, aktif=True)
            print(f"Yeni gecici sifre ({ad}): {sifre}")
        else:
            print(f"Kullanici bulunamadi: {ad}")
        return True
    if argv[0] == "ice-aktar":
        c = db.baglan(AYAR["DB"])
        r = ice_aktar.calistir(AYAR, c, AYAR["ONBELLEK"], COZUCU)
        print(json.dumps(r, ensure_ascii=False, indent=1))
        return True
    if argv[0] == "kullanicilar":
        for u in kullanicilar.hepsi(AYAR):
            print(f"{u['kullanici_adi']:20} {('aktif' if u.get('aktif', True) else 'PASIF'):6} {u.get('yetki', '')}")
        return True
    return False


def panel_calisiyor_mu(port):
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}/saglik", timeout=2) as r:
            return r.headers.get("X-Kanit-Paneli") is not None
    except Exception:
        return False


def calistir():
    if sys.stdout is None or sys.stderr is None:
        # pythonw / gorev zamanlayici: konsol yok, ciktiyi dosyaya yaz
        try:
            f = open(os.path.join(VERI_KOK, "panel.log"), "a", encoding="utf-8")
            sys.stdout = sys.stdout or f
            sys.stderr = sys.stderr or f
        except Exception:
            pass
    if len(sys.argv) > 1 and komut(sys.argv[1:]):
        return
    # --ag: ag modu (sunucu). Panel yalnizca bu bilgisayarda calisir, digerleri tarayicidan http://<bu bilgisayar>:PORT ile girer.
    ag_modu = "--ag" in sys.argv or str(AYAR.get("DINLE", "127.0.0.1")) not in ("127.0.0.1", "localhost", "")
    if "--ag" in sys.argv:
        AYAR["DINLE"] = "0.0.0.0"
    # --tunel: Cloudflare Tunnel (cloudflared) arkasinda. Panel yalnizca 127.0.0.1 dinler, tunel ona baglanir;
    # gercek istemci adresi CF-Connecting-IP basligindan alinir, oturum cerezi yalnizca https ile tasinir.
    tunel_modu = "--tunel" in sys.argv
    if tunel_modu:
        AYAR["DINLE"] = "127.0.0.1"
        AYAR["GUVENILIR_PROXY"] = "1"
        AYAR["HTTPS"] = "1"
        app.config["SESSION_COOKIE_SECURE"] = True
        ag_modu = False
    for i, a in enumerate(sys.argv):
        if a == "--port" and i + 1 < len(sys.argv) and sys.argv[i + 1].isdigit():
            AYAR["PORT"] = int(sys.argv[i + 1])
    port = int(AYAR["PORT"])
    url = f"http://127.0.0.1:{port}/"
    if panel_calisiyor_mu(port):
        print(f"Panel zaten acik: {url}  (tarayici aciliyor)")
        webbrowser.open(url)
        return
    conn = db.baglan(AYAR["DB"])
    if bot_klasoru_mu(AYAR.get("BOT_KLASOR")):
        try:
            n = kullanicilar.sqlite_tasi(AYAR, conn)
            if n:
                print(f"{n} kullanici ortak dosyaya tasindi: {kullanicilar.yol(AYAR)}")
        except Exception as e:
            print("kullanici tasima hatasi:", e)
    threading.Thread(target=_arka_plan, daemon=True).start()
    print("=" * 60)
    print(f"  KANIT PANELI {SURUM}" + ("  [AG MODU]" if ag_modu else "") + ("  [TUNEL MODU]" if tunel_modu else ""))
    print(f"  Adres      : {url}")
    if ag_modu:
        try:
            import socket
            adlar = socket.gethostbyname_ex(socket.gethostname())[2]
            print("  Diger bilgisayarlardan: " + "   ".join(f"http://{ip_}:{port}/" for ip_ in adlar if not ip_.startswith("127.")))
            print(f"  (ya da http://{socket.gethostname()}:{port}/ ). Guvenlik duvarinda {port} portu acik olmali: GUVENLIK_DUVARI_AC.bat")
        except Exception:
            pass
    if tunel_modu:
        print("  Cloudflare Tunnel: TUNEL_BASLAT.bat ile cloudflared'i acin; adminler https://<tunel adresiniz> ile girer.")
        print("  Bu bilgisayarda da https adresini kullanin (http://127.0.0.1 ile giris cerezi tasinmaz).")
    print(f"  Bot klasoru: {AYAR.get('BOT_KLASOR') or '(bulunamadi - tarayicida secilecek)'}")
    print(f"  Veri       : {VERI_KOK}")
    print("  Bu pencere acik kaldigi surece panel calisir. Kapatmak icin pencereyi kapatin.")
    print("=" * 60)
    if DONMUS or "--ac" in sys.argv:
        threading.Timer(1.2, lambda: webbrowser.open(url)).start()
    try:
        from waitress import serve
        serve(app, host=AYAR["DINLE"], port=port, threads=8)
    except ImportError:
        app.run(host=AYAR["DINLE"], port=port)
    except OSError as e:
        print(f"! Port {port} kullanilamiyor ({e}). Baska bir program kullaniyor olabilir; {AYAR_YOL} icinde PORT degistirin.")
        if DONMUS:
            input("Kapatmak icin Enter...")


if __name__ == "__main__":
    calistir()

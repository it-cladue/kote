#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export-SlackChannel.py
======================
Bir Slack kanalındaki TÜM mesajları ve her mesajın thread (alt yorum) yanıtlarını
çekip tek bir Excel dosyasına yazar.

Çıktı (Excel sayfaları)
-----------------------
  1. Mesajlar     : Her satır bir mesaj. Thread yanıtları, ait oldukları ana mesajın
                    hemen altında, gri zeminli ve girintili olarak listelenir. Satırlar
                    Excel'in "grup" (outline) özelliğiyle açılıp kapatılabilir.
                    Sütunlar: Sıra, Thread No, Thread Sırası, Tür, Tarih, Gün, Saat,
                    Yazan, Kullanıcı adı, E-posta, Mesaj, Yanıt sayısı, Reaksiyonlar,
                    Dosyalar, Düzenlendi, Slack linki, ts
  2. Threadler    : Her satır bir thread: kim başlattı, ne zaman, son yanıt ne zaman,
                    kaç yanıt, kimler katıldı, ana mesaj özeti.
  3. Kişi Özeti   : Kişi bazında ana mesaj / yanıt / toplam sayıları, ilk-son mesaj.
  4. Günlük Özet  : Gün bazında mesaj sayıları.
  5. Bilgi        : Kanal, workspace, çekim zamanı, tarih aralığı, saat dilimi, toplamlar.

Gereken Slack app yetkileri (Bot Token Scopes)
---------------------------------------------
  channels:history   -> herkese açık kanal mesajlarını okuma
  channels:read      -> kanal adından ID bulma
  groups:history     -> özel (private) kanal mesajlarını okuma
  groups:read        -> özel kanal adından ID bulma
  users:read         -> kullanıcı ID -> isim eşlemesi
  users:read.email   -> (isteğe bağlı) E-posta sütunu için
  Bot'un kanala EKLİ olması şart:  kanalda  /invite @BotAdı  yaz.

Kurulum / Çalıştırma
--------------------
  pip install openpyxl

  # Windows PowerShell
  $env:SLACK_BOT_TOKEN = "xoxb-..."
  python .\Export-SlackChannel.py --channel genel
  python .\Export-SlackChannel.py --channel C0123ABCD --start 2026-09-01 --end 2026-09-30
  python .\Export-SlackChannel.py --channel genel --output rapor.xlsx --json yedek.json

  # Linux / macOS
  export SLACK_BOT_TOKEN="xoxb-..."
  python3 Export-SlackChannel.py --channel genel

Notlar
------
  * Zaman damgaları varsayılan olarak Europe/Istanbul'a çevrilir (--tz ile değiştir).
  * Slack rate limit (429) gelirse script bekleyip kaldığı yerden devam eder.
    Marketplace dışı yeni app'lerde conversations.history / replies için limit çok
    düşük olabilir (dakikada 1 istek, istek başına 15 mesaj). Böyle bir durumda
    script otomatik olarak sayfa boyutunu 15'e düşürür ve bekler; büyük kanallarda
    çekim uzun sürebilir, bu normaldir.
"""

import argparse
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import datetime, timedelta, timezone

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl bulunamadı. Önce çalıştır:  pip install openpyxl")

API_BASE = "https://slack.com/api/"
TR_DAYS = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]

# Slack'in döndürdüğü hata kodlarının Türkçe açıklaması
ERROR_HELP = {
    "invalid_auth": "Token geçersiz ya da başka bir workspace'e ait.",
    "not_authed": "Token gönderilmedi. $env:SLACK_BOT_TOKEN ayarla.",
    "token_revoked": "Token iptal edilmiş. Slack App > OAuth & Permissions > Reinstall ile yenisini al.",
    "account_inactive": "Bot ya da kullanıcı devre dışı.",
    "missing_scope": "App'te eksik scope var. Bot Token Scopes'a channels:history, channels:read, "
                     "groups:history, groups:read, users:read ekle ve 'Reinstall to Workspace' yap.",
    "not_in_channel": "Bot kanala ekli değil. Kanalda  /invite @BotAdı  yaz ve tekrar dene.",
    "channel_not_found": "Kanal bulunamadı. Kanal ID'si yanlış ya da bot bu kanalı göremiyor "
                         "(özel kanal ise bot kanala eklenmeli).",
    "ratelimited": "Rate limit. Script otomatik bekleyip devam eder.",
    "invalid_cursor": "Sayfalama imleci bozuldu; scripti baştan çalıştır.",
    "invalid_ts_latest": "--end tarihi geçersiz.",
    "invalid_ts_oldest": "--start tarihi geçersiz.",
    "thread_not_found": "Thread silinmiş olabilir; atlanıyor.",
}

SYSTEM_SUBTYPES = {
    "channel_join", "channel_leave", "channel_purpose", "channel_topic", "channel_name",
    "channel_archive", "channel_unarchive", "group_join", "group_leave", "group_purpose",
    "group_topic", "group_name", "bot_add", "bot_remove", "pinned_item", "unpinned_item",
    "ekm_access_denied", "tombstone",
}


# ----------------------------------------------------------------------------- yardımcılar
def log(msg):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)


class SlackError(Exception):
    def __init__(self, method, code):
        self.method, self.code = method, code
        hint = ERROR_HELP.get(code, "")
        super().__init__(f"{method} -> {code}" + (f"  ({hint})" if hint else ""))


class SlackClient:
    def __init__(self, token, delay=0.0, page_size=200):
        self.token = token
        self.delay = delay
        self.page_size = page_size
        self.calls = 0

    def call(self, method, **params):
        """Tek bir Web API çağrısı. 429 / ratelimited durumunda bekleyip tekrar dener."""
        params = {k: v for k, v in params.items() if v is not None}
        data = urllib.parse.urlencode(params).encode("utf-8")
        req = urllib.request.Request(
            API_BASE + method,
            data=data,
            headers={
                "Authorization": f"Bearer {self.token}",
                "Content-Type": "application/x-www-form-urlencoded; charset=utf-8",
            },
        )
        attempt = 0
        while True:
            attempt += 1
            if self.delay:
                time.sleep(self.delay)
            self.calls += 1
            try:
                with urllib.request.urlopen(req, timeout=60) as resp:
                    body = json.loads(resp.read().decode("utf-8"))
            except urllib.error.HTTPError as e:
                if e.code == 429:
                    wait = int(e.headers.get("Retry-After", "30") or 30)
                    log(f"  rate limit ({method}); {wait} sn bekleniyor...")
                    time.sleep(wait + 1)
                    continue
                raise
            except (urllib.error.URLError, TimeoutError) as e:
                if attempt >= 5:
                    raise
                wait = 2 ** attempt
                log(f"  ağ hatası ({e}); {wait} sn sonra tekrar...")
                time.sleep(wait)
                continue

            if body.get("ok"):
                return body
            code = body.get("error", "unknown_error")
            if code == "ratelimited":
                wait = 30
                log(f"  rate limit ({method}); {wait} sn bekleniyor...")
                time.sleep(wait)
                continue
            raise SlackError(method, code)

    def paginate(self, method, key, **params):
        """cursor tabanlı sayfalama; tüm öğeleri tek listede döndürür."""
        items, cursor = [], None
        while True:
            body = self.call(method, cursor=cursor, **params)
            items.extend(body.get(key, []))
            cursor = (body.get("response_metadata") or {}).get("next_cursor") or None
            if not cursor:
                return items

    def paginate_messages(self, method, **params):
        """conversations.history / replies için sayfalama. Yeni rate-limit modelinde
        (limit > 15 kabul edilmezse) sayfa boyutunu otomatik 15'e düşürür."""
        items, cursor = [], None
        while True:
            try:
                body = self.call(method, cursor=cursor, limit=self.page_size, **params)
            except SlackError as e:
                if self.page_size > 15 and e.code in ("invalid_arguments", "invalid_limit"):
                    log(f"  {method}: limit={self.page_size} kabul edilmedi, 15'e düşürülüyor.")
                    self.page_size = 15
                    continue
                raise
            items.extend(body.get("messages", []))
            cursor = (body.get("response_metadata") or {}).get("next_cursor") or None
            if not cursor or not body.get("has_more", bool(cursor)):
                return items


def get_tz(name):
    try:
        from zoneinfo import ZoneInfo
        return ZoneInfo(name)
    except Exception:
        if name in ("Europe/Istanbul", "Turkey"):
            log("zoneinfo yüklenemedi; sabit UTC+3 kullanılıyor (pip install tzdata ile düzelir).")
            return timezone(timedelta(hours=3), "TRT")
        log(f"'{name}' saat dilimi bulunamadı; UTC kullanılıyor (pip install tzdata).")
        return timezone.utc


def ts_to_dt(ts, tz):
    return datetime.fromtimestamp(float(ts), tz=timezone.utc).astimezone(tz).replace(microsecond=0)


def parse_date(s, tz, end=False):
    """'2026-09-30' -> o günün 00:00 (ya da end=True ise ertesi gün 00:00) epoch ts"""
    if not s:
        return None
    d = datetime.strptime(s, "%Y-%m-%d").replace(tzinfo=tz)
    if end:
        d += timedelta(days=1)
    return f"{d.timestamp():.6f}"


# ----------------------------------------------------------------------------- metin temizleme
MENTION_RE = re.compile(r"<@([UW][A-Z0-9]+)(?:\|[^>]*)?>")
CHANNEL_RE = re.compile(r"<#([CG][A-Z0-9]+)(?:\|([^>]*))?>")
SUBTEAM_RE = re.compile(r"<!subteam\^([A-Z0-9]+)(?:\|([^>]*))?>")
SPECIAL_RE = re.compile(r"<!(here|channel|everyone)(?:\|[^>]*)?>")
LINK_RE = re.compile(r"<((?:https?|mailto):[^>|]+)(?:\|([^>]*))?>")


def render_text(text, users, channels):
    if not text:
        return ""
    text = MENTION_RE.sub(lambda m: "@" + users.get(m.group(1), {}).get("display", m.group(1)), text)
    text = CHANNEL_RE.sub(lambda m: "#" + (m.group(2) or channels.get(m.group(1), m.group(1))), text)
    text = SUBTEAM_RE.sub(lambda m: "@" + (m.group(2) or m.group(1)), text)
    text = SPECIAL_RE.sub(lambda m: "@" + m.group(1), text)
    text = LINK_RE.sub(lambda m: f"{m.group(2)} ({m.group(1)})" if m.group(2) and m.group(2) != m.group(1) else m.group(1), text)
    return text.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")


def message_body(msg, users, channels):
    """Mesaj metni + (varsa) attachment / dosya bilgisi"""
    parts = [render_text(msg.get("text", ""), users, channels)]
    for att in msg.get("attachments", []) or []:
        for key in ("title", "text", "fallback"):
            if att.get(key):
                parts.append(render_text(att[key], users, channels))
                break
    if msg.get("subtype") == "tombstone":
        parts = ["(mesaj silinmiş; thread yanıtları korunmuş)"]
    return "\n".join(p for p in parts if p).strip()


def file_list(msg):
    names = []
    for f in msg.get("files", []) or []:
        name = f.get("name") or f.get("title") or f.get("id", "")
        if f.get("mode") == "tombstone" or f.get("file_access") == "file_not_found":
            name += " (silinmiş)"
        names.append(name)
    return ", ".join(names)


def reactions(msg):
    return ", ".join(f":{r.get('name')}: x{r.get('count', 0)}" for r in msg.get("reactions", []) or [])


def author(msg, users):
    uid = msg.get("user")
    if uid and uid in users:
        u = users[uid]
        return u["display"], u["handle"], u["email"], uid
    if uid:
        return uid, uid, "", uid
    bot = msg.get("username") or (msg.get("bot_profile") or {}).get("name") or msg.get("bot_id", "bot")
    return f"{bot} (bot)", bot, "", msg.get("bot_id", "")


# ----------------------------------------------------------------------------- Slack'ten çekme
def load_users(client):
    users = {}
    try:
        for u in client.paginate("users.list", "members", limit=200):
            p = u.get("profile") or {}
            display = p.get("display_name") or p.get("real_name") or u.get("real_name") or u.get("name") or u["id"]
            if u.get("deleted"):
                display += " (silinmiş hesap)"
            users[u["id"]] = {
                "display": display,
                "handle": u.get("name", ""),
                "email": p.get("email", "") or "",
                "is_bot": bool(u.get("is_bot")),
            }
    except SlackError as e:
        log(f"users.list okunamadı ({e.code}); isimler ID olarak yazılacak.")
    return users


def ensure_user(client, users, uid):
    """users.list'te olmayan (Slack Connect, dış) kullanıcıyı users.info ile dener"""
    if not uid or uid in users:
        return
    try:
        u = client.call("users.info", user=uid)["user"]
        p = u.get("profile") or {}
        users[uid] = {
            "display": p.get("display_name") or p.get("real_name") or u.get("name") or uid,
            "handle": u.get("name", ""),
            "email": p.get("email", "") or "",
            "is_bot": bool(u.get("is_bot")),
        }
    except Exception:
        users[uid] = {"display": uid, "handle": uid, "email": "", "is_bot": False}


def resolve_channel(client, ref):
    """Kanal adı ya da ID -> (id, ad). ID verildiyse conversations.info ile adı alır."""
    ref = ref.strip().lstrip("#")
    if re.fullmatch(r"[CG][A-Z0-9]{6,}", ref):
        try:
            ch = client.call("conversations.info", channel=ref)["channel"]
            return ch["id"], ch.get("name", ref)
        except SlackError as e:
            if e.code == "channel_not_found":
                raise
            return ref, ref
    log("Kanal listesi çekiliyor...")
    for ch in client.paginate("conversations.list", "channels",
                              types="public_channel,private_channel", exclude_archived="false", limit=200):
        if ch.get("name", "").lower() == ref.lower():
            return ch["id"], ch["name"]
    raise SystemExit(f"'{ref}' adında kanal bulunamadı. Kanal ID'sini (C... ile başlar) vermeyi dene; "
                     f"özel kanalsa bot kanala ekli olmalı.")


def fetch_channel(client, channel_id, oldest, latest, include_system):
    log("Ana mesajlar çekiliyor...")
    history = client.paginate_messages("conversations.history", channel=channel_id,
                                       oldest=oldest, latest=latest, inclusive="true")
    log(f"  {len(history)} kayıt geldi.")

    parents, seen = [], set()
    for m in history:
        ts = m.get("ts")
        if not ts or ts in seen:
            continue
        seen.add(ts)
        # thread_broadcast: thread yanıtı kanala da gönderilmiş; replies'ten gelecek, burada atla
        if m.get("subtype") == "thread_broadcast" and m.get("thread_ts") != ts:
            continue
        if not include_system and m.get("subtype") in SYSTEM_SUBTYPES:
            continue
        parents.append(m)
    parents.sort(key=lambda m: float(m["ts"]))

    threads = {}
    with_threads = [m for m in parents if m.get("reply_count") or (m.get("thread_ts") == m.get("ts") and m.get("reply_users"))]
    log(f"  {len(parents)} ana mesaj, {len(with_threads)} tanesinde thread var. Thread yanıtları çekiliyor...")
    for i, m in enumerate(with_threads, 1):
        try:
            replies = client.paginate_messages("conversations.replies", channel=channel_id, ts=m["ts"])
        except SlackError as e:
            if e.code == "thread_not_found":
                continue
            raise
        replies = [r for r in replies if r.get("ts") != m["ts"]]
        replies.sort(key=lambda r: float(r["ts"]))
        threads[m["ts"]] = replies
        if i % 10 == 0 or i == len(with_threads):
            log(f"  thread {i}/{len(with_threads)} ({sum(len(v) for v in threads.values())} yanıt)")
    return parents, threads


# ----------------------------------------------------------------------------- Excel
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
REPLY_FILL = PatternFill("solid", fgColor="F2F2F2")
THREAD_PARENT_FILL = PatternFill("solid", fgColor="E8F0FE")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def style_header(ws, widths):
    for col, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=col)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


def build_rows(parents, threads, users, channels, tz, permalink):
    """Mesajlar sayfası için satır listesi (dict) üretir; threadler ana mesajın altında."""
    rows, thread_no = [], 0
    for p in parents:
        replies = threads.get(p["ts"], [])
        thread_no += 1
        rows.append(make_row(p, users, channels, tz, permalink, thread_no, 0, len(replies), None))
        for j, r in enumerate(replies, 1):
            rows.append(make_row(r, users, channels, tz, permalink, thread_no, j, None, p["ts"]))
    for i, r in enumerate(rows, 1):
        r["no"] = i
    return rows


def make_row(m, users, channels, tz, permalink, thread_no, order, reply_count, parent_ts):
    dt = ts_to_dt(m["ts"], tz)
    name, handle, email, uid = author(m, users)
    if order == 0:
        kind = "Ana mesaj (thread)" if reply_count else "Ana mesaj"
    else:
        kind = "Thread yanıtı" + (" (kanala da gönderildi)" if m.get("subtype") == "thread_broadcast" else "")
    return {
        "thread_no": thread_no,
        "order": order,
        "kind": kind,
        "date": dt.replace(tzinfo=None),
        "day": TR_DAYS[dt.weekday()],
        "time": dt.strftime("%H:%M:%S"),
        "name": name,
        "handle": handle,
        "email": email,
        "uid": uid,
        "text": message_body(m, users, channels),
        "reply_count": reply_count if order == 0 else None,
        "reactions": reactions(m),
        "files": file_list(m),
        "edited": "Evet" if m.get("edited") else "",
        "link": permalink(m["ts"], parent_ts),
        "ts": m["ts"],
        "is_reply": order > 0,
        "has_thread": order == 0 and bool(reply_count),
    }


def write_excel(path, rows, parents, threads, users, info, tz):
    wb = Workbook()

    # ---- 1. Mesajlar
    ws = wb.active
    ws.title = "Mesajlar"
    headers = ["Sıra", "Thread No", "Thread Sırası", "Tür", "Tarih", "Gün", "Saat", "Yazan",
               "Kullanıcı adı", "E-posta", "Mesaj", "Yanıt sayısı", "Reaksiyonlar", "Dosyalar",
               "Düzenlendi", "Slack linki", "ts"]
    widths = [7, 10, 10, 20, 12, 11, 10, 22, 18, 28, 80, 10, 22, 30, 10, 14, 18]
    ws.append(headers)
    style_header(ws, widths)
    ws.sheet_properties.outlinePr.summaryBelow = False

    for r in rows:
        ws.append([r["no"], r["thread_no"], r["order"], r["kind"], r["date"], r["day"], r["time"],
                   r["name"], r["handle"], r["email"], r["text"], r["reply_count"], r["reactions"],
                   r["files"], r["edited"], "Aç", r["ts"]])
        i = ws.max_row
        ws.cell(row=i, column=5).number_format = "dd.mm.yyyy"
        link = ws.cell(row=i, column=16)
        link.hyperlink = r["link"]
        link.font = Font(color="0563C1", underline="single")
        fill = REPLY_FILL if r["is_reply"] else (THREAD_PARENT_FILL if r["has_thread"] else None)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=i, column=col)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=(col in (11, 13, 14)))
            if fill:
                c.fill = fill
        if r["is_reply"]:
            ws.cell(row=i, column=11).alignment = Alignment(vertical="top", wrap_text=True, indent=2)
            ws.cell(row=i, column=4).alignment = Alignment(vertical="top", indent=2)
            ws.row_dimensions[i].outlineLevel = 1
        elif r["has_thread"]:
            ws.cell(row=i, column=11).font = Font(bold=True)
    ws.auto_filter.ref = ws.dimensions

    # ---- 2. Threadler
    ws = wb.create_sheet("Threadler")
    ws.append(["Thread No", "Başlatan", "Başlangıç", "Gün", "Saat", "Son yanıt", "Yanıt sayısı",
               "Katılımcı sayısı", "Katılımcılar", "Ana mesaj", "Slack linki"])
    style_header(ws, [10, 22, 12, 11, 10, 18, 10, 10, 40, 70, 14])
    thread_no = 0
    for p in parents:
        thread_no += 1
        replies = threads.get(p["ts"], [])
        if not replies:
            continue
        dt = ts_to_dt(p["ts"], tz)
        last = ts_to_dt(replies[-1]["ts"], tz)
        people = []
        for m in [p] + replies:
            n = author(m, users)[0]
            if n not in people:
                people.append(n)
        ws.append([thread_no, author(p, users)[0], dt.replace(tzinfo=None), TR_DAYS[dt.weekday()],
                   dt.strftime("%H:%M:%S"), last.replace(tzinfo=None), len(replies), len(people),
                   ", ".join(people), message_body(p, users, info["channels"])[:300], "Aç"])
        i = ws.max_row
        ws.cell(row=i, column=3).number_format = "dd.mm.yyyy"
        ws.cell(row=i, column=6).number_format = "dd.mm.yyyy HH:MM"
        link = ws.cell(row=i, column=11)
        link.hyperlink = info["permalink"](p["ts"], None)
        link.font = Font(color="0563C1", underline="single")
        for col in range(1, 12):
            ws.cell(row=i, column=col).alignment = Alignment(vertical="top", wrap_text=col in (9, 10))
            ws.cell(row=i, column=col).border = BORDER
    ws.auto_filter.ref = ws.dimensions

    # ---- 3. Kişi Özeti
    ws = wb.create_sheet("Kişi Özeti")
    ws.append(["Yazan", "Kullanıcı adı", "E-posta", "Ana mesaj", "Thread yanıtı", "Toplam",
               "Başlattığı thread", "İlk mesaj", "Son mesaj"])
    style_header(ws, [24, 18, 28, 11, 12, 10, 12, 18, 18])
    stats = defaultdict(lambda: {"name": "", "handle": "", "email": "", "p": 0, "r": 0, "t": 0,
                                 "first": None, "last": None})
    for r in rows:
        s = stats[r["uid"] or r["name"]]
        s["name"], s["handle"], s["email"] = r["name"], r["handle"], r["email"]
        s["r" if r["is_reply"] else "p"] += 1
        if r["has_thread"]:
            s["t"] += 1
        s["first"] = r["date"] if s["first"] is None or r["date"] < s["first"] else s["first"]
        s["last"] = r["date"] if s["last"] is None or r["date"] > s["last"] else s["last"]
    for s in sorted(stats.values(), key=lambda x: -(x["p"] + x["r"])):
        ws.append([s["name"], s["handle"], s["email"], s["p"], s["r"], s["p"] + s["r"], s["t"],
                   s["first"], s["last"]])
        i = ws.max_row
        ws.cell(row=i, column=8).number_format = "dd.mm.yyyy HH:MM"
        ws.cell(row=i, column=9).number_format = "dd.mm.yyyy HH:MM"
        for col in range(1, 10):
            ws.cell(row=i, column=col).border = BORDER
    ws.auto_filter.ref = ws.dimensions

    # ---- 4. Günlük Özet
    ws = wb.create_sheet("Günlük Özet")
    ws.append(["Tarih", "Gün", "Ana mesaj", "Thread yanıtı", "Toplam", "Aktif kişi"])
    style_header(ws, [12, 11, 11, 12, 10, 10])
    daily = defaultdict(lambda: {"p": 0, "r": 0, "who": set()})
    for r in rows:
        d = daily[r["date"].date()]
        d["r" if r["is_reply"] else "p"] += 1
        d["who"].add(r["uid"] or r["name"])
    for day in sorted(daily):
        d = daily[day]
        ws.append([datetime(day.year, day.month, day.day), TR_DAYS[day.weekday()], d["p"], d["r"],
                   d["p"] + d["r"], len(d["who"])])
        i = ws.max_row
        ws.cell(row=i, column=1).number_format = "dd.mm.yyyy"
        for col in range(1, 7):
            ws.cell(row=i, column=col).border = BORDER
    ws.auto_filter.ref = ws.dimensions

    # ---- 5. Bilgi
    ws = wb.create_sheet("Bilgi")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 60
    total_replies = sum(len(v) for v in threads.values())
    dates = [r["date"] for r in rows]
    lines = [
        ("Workspace", info["team"]),
        ("Kanal", f"#{info['channel_name']}  ({info['channel_id']})"),
        ("Çekim zamanı", datetime.now(tz).strftime("%d.%m.%Y %H:%M:%S")),
        ("Saat dilimi", info["tz_name"]),
        ("İstenen aralık", f"{info['start'] or 'başlangıç'}  ->  {info['end'] or 'bugün'}"),
        ("Bulunan ilk mesaj", min(dates).strftime("%d.%m.%Y %H:%M") if dates else "-"),
        ("Bulunan son mesaj", max(dates).strftime("%d.%m.%Y %H:%M") if dates else "-"),
        ("Ana mesaj sayısı", len(parents)),
        ("Thread'li ana mesaj", sum(1 for v in threads.values() if v)),
        ("Thread yanıtı sayısı", total_replies),
        ("Toplam mesaj", len(parents) + total_replies),
        ("Farklı kişi", len(stats)),
        ("Sistem mesajları", "dahil" if info["include_system"] else "hariç (--include-system ile dahil et)"),
        ("API çağrısı", info["calls"]),
        ("", ""),
        ("Mesajlar sayfası", "Thread yanıtları ana mesajın altında gri zeminle; sol kenardaki +/- ile aç/kapat."),
        ("Thread No", "Aynı numara = aynı konuşma. Thread Sırası 0 = ana mesaj, 1..n = yanıtlar."),
    ]
    for k, v in lines:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    wb.save(path)


# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Slack kanalını (thread'leriyle) Excel'e aktarır.")
    ap.add_argument("--channel", required=True, help="Kanal adı (genel) ya da ID (C0123ABCD)")
    ap.add_argument("--token", default=os.environ.get("SLACK_BOT_TOKEN"), help="xoxb-... (varsayılan: SLACK_BOT_TOKEN)")
    ap.add_argument("--start", help="Başlangıç tarihi YYYY-MM-DD (dahil)")
    ap.add_argument("--end", help="Bitiş tarihi YYYY-MM-DD (dahil)")
    ap.add_argument("--tz", default="Europe/Istanbul", help="Saat dilimi (varsayılan Europe/Istanbul)")
    ap.add_argument("--output", "-o", help="Excel dosya adı (varsayılan: slack_<kanal>_<tarih>.xlsx)")
    ap.add_argument("--json", help="Ham Slack verisini de bu JSON dosyasına yedekle")
    ap.add_argument("--include-system", action="store_true", help="Katıldı/ayrıldı gibi sistem mesajlarını da yaz")
    ap.add_argument("--page-size", type=int, default=200, help="İstek başına mesaj (rate limit düşükse 15)")
    ap.add_argument("--delay", type=float, default=0.0, help="Her API çağrısı arasında bekleme (sn)")
    args = ap.parse_args()

    if not args.token:
        sys.exit('Token yok. Önce çalıştır:  $env:SLACK_BOT_TOKEN = "xoxb-..."   (Linux: export SLACK_BOT_TOKEN=...)')

    tz = get_tz(args.tz)
    client = SlackClient(args.token, delay=args.delay, page_size=args.page_size)

    try:
        auth = client.call("auth.test")
        team_url = auth.get("url", "https://slack.com/").rstrip("/")
        log(f"Workspace: {auth.get('team')}  bot: {auth.get('user')}")

        channel_id, channel_name = resolve_channel(client, args.channel)
        log(f"Kanal: #{channel_name} ({channel_id})")

        log("Kullanıcı listesi çekiliyor...")
        users = load_users(client)
        log(f"  {len(users)} kullanıcı.")

        oldest = parse_date(args.start, tz)
        latest = parse_date(args.end, tz, end=True)
        parents, threads = fetch_channel(client, channel_id, oldest, latest, args.include_system)
    except SlackError as e:
        sys.exit(f"HATA: {e}")

    # Listede olmayan kullanıcılar (dış / Slack Connect) için tek tek dene
    all_msgs = parents + [r for v in threads.values() for r in v]
    for uid in {m.get("user") for m in all_msgs if m.get("user")} - set(users):
        ensure_user(client, users, uid)

    channels = {}  # <#C..> referansları için; sadece bilinen kanal eklenir
    channels[channel_id] = channel_name

    def permalink(ts, parent_ts):
        url = f"{team_url}/archives/{channel_id}/p{ts.replace('.', '')}"
        if parent_ts:
            url += f"?thread_ts={parent_ts}&cid={channel_id}"
        return url

    rows = build_rows(parents, threads, users, channels, tz, permalink)

    if args.json:
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump({"channel": {"id": channel_id, "name": channel_name}, "users": users,
                       "messages": parents, "threads": threads}, f, ensure_ascii=False, indent=1)
        log(f"Ham veri yazıldı: {args.json}")

    out = args.output or os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                   f"slack_{channel_name}_{datetime.now(tz).strftime('%Y%m%d_%H%M')}.xlsx")
    info = {
        "team": auth.get("team", ""), "channel_id": channel_id, "channel_name": channel_name,
        "tz_name": args.tz, "start": args.start, "end": args.end, "include_system": args.include_system,
        "calls": client.calls, "channels": channels, "permalink": permalink,
    }
    write_excel(out, rows, parents, threads, users, info, tz)
    total_replies = sum(len(v) for v in threads.values())
    log(f"Bitti: {len(parents)} ana mesaj + {total_replies} thread yanıtı")
    log(f"Excel dosyası: {os.path.abspath(out)}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("\nİptal edildi.")

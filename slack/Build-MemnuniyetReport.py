#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build-MemnuniyetReport.py
=========================
Export-SlackChannel.py'nin ürettiği Excel'i (Mesajlar sayfası) okur ve
"yatırımsız oyuncu" bildirim kanalı için kim-hangisine-yanıt-verdi raporu çıkarır.

Kanal mantığı:
  * Bot (Messenger) her gün 12:00'de şu formatta bildirim atar:
        ID: 1022-485241275
        PC: 285HG
        Last Deposit Date: 2026-09-01T01:04:46Z
        :rotating_light: Oyuncunun yatırımsız 15. günüdür, lütfen iletişime geçiniz.
  * PC kodunun başındaki sayı (285) sorumlu TS personelini gösterir ("Yakup 285 TS").
    Sondaki harfler (HG, PRC) oyuncu etiketidir.
  * Sorumlu TS, bildirimin thread'ine arama sonucunu yazar.

Rapor sayfaları:
  Rapor            Her bildirim bir satır: oyuncu, PC, sorumlu TS, yanıtlandı mı, kim yanıtladı,
                   ne zaman, kaç dakikada, yanıt metni. Yanıtsızlar kırmızı, yanlış kişi turuncu.
  TS Performans    TS bazında atanan / yanıtlanan / yanıtsız / oran / ortalama süre.
  Yanıtsızlar      Sadece yanıt verilmemiş bildirimler.
  Günlük           Gün bazında bildirim / yanıtlanan / oran.
  Günlük x TS      Gün x TS tablosu (yanıtlanan/atanan).
  Gün Tipi         8. / 15. / 21. gün bildirimlerine göre özet.
  Ham Mesajlar     Export dosyasındaki tüm mesajlar (değişmeden).
  Bilgi            Kaynak, aralık, uyarılar (Slack'in göstermediği mesajlar vb).

Kullanım:
  python Build-MemnuniyetReport.py slack_p-ft-memnuniyet_2026-09-16_2026-09-30.xlsx
  python Build-MemnuniyetReport.py export.xlsx --output rapor.xlsx
"""

import argparse
import os
import re
import statistics
import sys
from collections import OrderedDict, defaultdict
from datetime import datetime, timedelta, timezone

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl bulunamadı. Önce çalıştır:  pip install openpyxl")

TR_DAYS = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RED_FILL = PatternFill("solid", fgColor="FDE2E2")
ORANGE_FILL = PatternFill("solid", fgColor="FFE8CC")
GREEN_FILL = PatternFill("solid", fgColor="E6F4EA")
GREY_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

ID_RE = re.compile(r"ID:\s*([0-9\-]+)")
PC_RE = re.compile(r"PC:\s*(?:<[^>]*>\s*)?([A-Za-z0-9]+)")
LDD_RE = re.compile(r"Last Deposit Date:\s*([0-9T:\-]+Z?)")
DAY_RE = re.compile(r"yatırımsız\s*(\d+)\.\s*gün", re.I)
NUM_RE = re.compile(r"\d+")


# ----------------------------------------------------------------------------- okuma
def read_export(path):
    """Mesajlar sayfasını okur; thread'lere göre gruplar."""
    wb = load_workbook(path)
    if "Mesajlar" not in wb.sheetnames:
        sys.exit("Bu dosyada 'Mesajlar' sayfası yok. Export-SlackChannel.py çıktısını ver.")
    ws = wb["Mesajlar"]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    need = ["Thread No", "Thread Sırası", "Tarih", "Yazan", "Mesaj", "Slack linki", "ts"]
    for n in need:
        if n not in col:
            sys.exit(f"Mesajlar sayfasında '{n}' sütunu yok; export scriptinin güncel sürümünü kullan.")

    threads = OrderedDict()
    raw_rows = []
    for row in ws.iter_rows(min_row=2):
        vals = [c.value for c in row]
        raw_rows.append(vals)
        if vals[col["ts"]] is None:
            continue
        link_cell = row[col["Slack linki"]]
        m = {
            "thread": vals[col["Thread No"]],
            "order": vals[col["Thread Sırası"]] or 0,
            "date": vals[col["Tarih"]],
            "author": vals[col["Yazan"]] or "",
            "text": vals[col["Mesaj"]] or "",
            "link": link_cell.hyperlink.target if link_cell.hyperlink else "",
            "ts": float(vals[col["ts"]]),
        }
        t = threads.setdefault(m["thread"], {"parent": None, "replies": []})
        if m["order"] == 0:
            t["parent"] = m
        else:
            t["replies"].append(m)
    for t in threads.values():
        t["replies"].sort(key=lambda r: r["ts"])

    info = {}
    if "Bilgi" in wb.sheetnames:
        for k, v in wb["Bilgi"].iter_rows(values_only=True):
            if k:
                info[k] = v
    return threads, header, raw_rows, info


def parse_alert(text):
    """Bildirim metninden alanları çıkarır; bildirim değilse None."""
    mid, mpc = ID_RE.search(text), PC_RE.search(text)
    if not mid or not mpc:
        return None
    pc = mpc.group(1).upper()
    num = NUM_RE.match(pc)
    ldd = None
    m = LDD_RE.search(text)
    if m:
        try:
            ldd = datetime.strptime(m.group(1).rstrip("Z"), "%Y-%m-%dT%H:%M:%S").replace(tzinfo=timezone.utc)
            ldd = ldd.astimezone(timezone(timedelta(hours=3))).replace(tzinfo=None)
        except ValueError:
            ldd = None
    d = DAY_RE.search(text)
    return {
        "player": mid.group(1),
        "pc": pc,
        "pc_num": num.group() if num else "",
        "pc_tag": pc[len(num.group()):] if num else pc,
        "last_deposit": ldd,
        "day": int(d.group(1)) if d else None,
    }


def human_minutes(m):
    if m is None:
        return ""
    m = int(round(m))
    d, rem = divmod(m, 1440)
    h, mi = divmod(rem, 60)
    parts = ([f"{d} gün"] if d else []) + ([f"{h} sa"] if h else []) + [f"{mi} dk"]
    return " ".join(parts)


def agent_number(name):
    m = NUM_RE.search(name or "")
    return m.group() if m else ""


# ----------------------------------------------------------------------------- analiz
def analyze(threads):
    alerts, others = [], []
    agent_by_num = {}
    # 1) yanıt verenlerden PC no -> TS adı haritası
    for t in threads.values():
        for r in t["replies"]:
            n = agent_number(r["author"])
            if n:
                agent_by_num.setdefault(n, r["author"])

    for t in threads.values():
        p = t["parent"]
        if p is None:
            continue
        a = parse_alert(p["text"])
        if a is None or not a["pc_num"]:
            others.append(p)          # bildirim değil ya da PC'de numara yok (test mesajı vb)
            continue
        replies = t["replies"]
        first = replies[0] if replies else None
        responsible = agent_by_num.get(a["pc_num"]) or (f"PC {a['pc_num']} (TS bilinmiyor)" if a["pc_num"] else "PC okunamadı")
        responders = []
        for r in replies:
            if r["author"] not in responders:
                responders.append(r["author"])
        correct = None
        if replies:
            correct = any(agent_number(r["author"]) == a["pc_num"] for r in replies)
        note = []
        if first and len(first["text"].strip()) < 5:
            note.append("Kısa/boş yanıt")
        if replies and not correct:
            note.append("Sorumlu TS değil, başkası yanıtladı")
        if not a["pc_num"]:
            note.append("PC kodu okunamadı")
        extra = ""
        if len(replies) > 1:
            extra = "\n".join(f"[{r['date'].strftime('%d.%m %H:%M')}] {r['author']}: {r['text']}" for r in replies[1:])
        alerts.append({
            **a,
            "date": p["date"],
            "link": p["link"],
            "responsible": responsible,
            "answered": bool(replies),
            "responders": ", ".join(responders),
            "first_reply": first,
            "reply_minutes": round((first["ts"] - p["ts"]) / 60, 1) if first else None,
            "reply_text": first["text"] if first else "",
            "reply_count": len(replies),
            "extra_replies": extra,
            "correct": correct,
            "note": "; ".join(note),
        })
    alerts.sort(key=lambda x: (x["date"], x["pc_num"], x["player"]))
    return alerts, others, agent_by_num


# ----------------------------------------------------------------------------- excel yardımcıları
def header(ws, titles, widths):
    ws.append(titles)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def finish(ws, wrap_cols=(), fmt=None):
    fmt = fmt or {}
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols)
            if c.column in fmt:
                c.number_format = fmt[c.column]
    ws.auto_filter.ref = ws.dimensions


def fill_row(ws, r, fill, ncol):
    for col in range(1, ncol + 1):
        ws.cell(row=r, column=col).fill = fill


def pct(a, b):
    return (a / b) if b else 0


# ----------------------------------------------------------------------------- yazma
def write_report(out, alerts, others, agent_by_num, header_raw, raw_rows, info, src):
    wb = Workbook()

    # ---- Rapor
    ws = wb.active
    ws.title = "Rapor"
    cols = ["Sıra", "Bildirim Tarihi", "Gün", "Saat", "Oyuncu ID", "PC Kodu", "PC No", "Etiket",
            "Son Yatırım", "Yatırımsız Gün", "Sorumlu TS", "Durum", "Yanıtlayan", "Yanıt Tarihi",
            "Yanıt Saati", "Yanıt Süresi", "Yanıt Süresi (dk)", "Yanıt", "Ek Yanıtlar", "Yanıt Sayısı",
            "Doğru Kişi mi", "Not", "Slack"]
    header(ws, cols, [6, 12, 11, 8, 16, 11, 7, 8, 16, 9, 18, 12, 18, 12, 8, 14, 9, 70, 40, 7, 9, 28, 8])
    for i, a in enumerate(alerts, 1):
        fr = a["first_reply"]
        ws.append([
            i, a["date"], TR_DAYS[a["date"].weekday()], a["date"].strftime("%H:%M"),
            a["player"], a["pc"], int(a["pc_num"]) if a["pc_num"] else "", a["pc_tag"],
            a["last_deposit"], a["day"], a["responsible"],
            "Yanıtlandı" if a["answered"] else "YANITSIZ",
            a["responders"], fr["date"] if fr else None, fr["date"].strftime("%H:%M") if fr else "",
            human_minutes(a["reply_minutes"]), a["reply_minutes"], a["reply_text"], a["extra_replies"], a["reply_count"],
            "" if a["correct"] is None else ("Evet" if a["correct"] else "HAYIR"),
            a["note"], "Aç",
        ])
        r = ws.max_row
        link = ws.cell(row=r, column=23)
        link.hyperlink, link.font = a["link"], Font(color="0563C1", underline="single")
        if not a["answered"]:
            fill_row(ws, r, RED_FILL, len(cols))
            ws.cell(row=r, column=12).font = Font(bold=True, color="C00000")
        elif a["correct"] is False:
            fill_row(ws, r, ORANGE_FILL, len(cols))
    finish(ws, wrap_cols=(18, 19, 22), fmt={2: "dd.mm.yyyy", 9: "dd.mm.yyyy HH:MM", 14: "dd.mm.yyyy"})

    # ---- TS Performans
    ws = wb.create_sheet("TS Performans")
    header(ws, ["Sorumlu TS", "PC No", "Atanan Bildirim", "Yanıtlanan", "Yanıtsız", "Yanıt Oranı",
                "Ort. Yanıt (dk)", "Medyan (dk)", "En Uzun", "1 saat içinde", "Kısa/boş yanıt",
                "Başkasının bildirimine yanıt"],
           [20, 7, 10, 10, 9, 10, 10, 10, 16, 10, 10, 14])
    per = defaultdict(lambda: {"num": "", "assigned": 0, "answered": 0, "mins": [], "short": 0, "foreign": 0})
    for a in alerts:
        s = per[a["responsible"]]
        s["num"] = a["pc_num"]
        s["assigned"] += 1
        if a["answered"]:
            s["answered"] += 1
            s["mins"].append(a["reply_minutes"])
            if "Kısa" in a["note"]:
                s["short"] += 1
            if a["correct"] is False:
                for name in a["responders"].split(", "):
                    per[name]["foreign"] += 1
                    per[name]["num"] = per[name]["num"] or agent_number(name)
    tot = {"assigned": 0, "answered": 0, "mins": [], "short": 0, "foreign": 0, "fast": 0}
    for name, s in sorted(per.items(), key=lambda kv: -kv[1]["assigned"]):
        fast = sum(1 for m in s["mins"] if m <= 60)
        ws.append([name, int(s["num"]) if s["num"].isdigit() else s["num"], s["assigned"], s["answered"],
                   s["assigned"] - s["answered"], pct(s["answered"], s["assigned"]),
                   round(statistics.mean(s["mins"]), 1) if s["mins"] else None,
                   round(statistics.median(s["mins"]), 1) if s["mins"] else None,
                   human_minutes(max(s["mins"])) if s["mins"] else None, fast, s["short"], s["foreign"]])
        r = ws.max_row
        rate = pct(s["answered"], s["assigned"])
        if s["assigned"]:
            ws.cell(row=r, column=6).fill = GREEN_FILL if rate >= 0.9 else (ORANGE_FILL if rate >= 0.7 else RED_FILL)
        for k in ("assigned", "answered", "short", "foreign"):
            tot[k] += s[k]
        tot["mins"] += s["mins"]
        tot["fast"] += fast
    ws.append(["TOPLAM", "", tot["assigned"], tot["answered"], tot["assigned"] - tot["answered"],
               pct(tot["answered"], tot["assigned"]),
               round(statistics.mean(tot["mins"]), 1) if tot["mins"] else None,
               round(statistics.median(tot["mins"]), 1) if tot["mins"] else None,
               human_minutes(max(tot["mins"])) if tot["mins"] else None, tot["fast"], tot["short"], tot["foreign"]])
    finish(ws, fmt={6: "0.0%"})
    fill_row(ws, ws.max_row, TOTAL_FILL, 12)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    # ---- Yanıtsızlar
    ws = wb.create_sheet("Yanıtsızlar")
    header(ws, ["Bildirim Tarihi", "Gün", "Oyuncu ID", "PC Kodu", "Sorumlu TS", "Yatırımsız Gün",
                "Son Yatırım", "Slack"], [12, 11, 16, 11, 20, 9, 16, 8])
    for a in alerts:
        if a["answered"]:
            continue
        ws.append([a["date"], TR_DAYS[a["date"].weekday()], a["player"], a["pc"], a["responsible"],
                   a["day"], a["last_deposit"], "Aç"])
        link = ws.cell(row=ws.max_row, column=8)
        link.hyperlink, link.font = a["link"], Font(color="0563C1", underline="single")
    finish(ws, fmt={1: "dd.mm.yyyy", 7: "dd.mm.yyyy HH:MM"})

    # ---- Günlük
    ws = wb.create_sheet("Günlük")
    header(ws, ["Tarih", "Gün", "Bildirim", "Yanıtlanan", "Yanıtsız", "Yanıt Oranı", "Ort. Yanıt (dk)",
                "Medyan (dk)", "En Uzun", "Yanıt Veren TS Sayısı"], [12, 11, 9, 10, 9, 10, 10, 10, 16, 10])
    daily = defaultdict(lambda: {"n": 0, "a": 0, "mins": [], "who": set()})
    for a in alerts:
        d = daily[a["date"].date()]
        d["n"] += 1
        if a["answered"]:
            d["a"] += 1
            d["mins"].append(a["reply_minutes"])
            d["who"].update(a["responders"].split(", "))
    for day in sorted(daily):
        d = daily[day]
        ws.append([datetime(day.year, day.month, day.day), TR_DAYS[day.weekday()], d["n"], d["a"], d["n"] - d["a"],
                   pct(d["a"], d["n"]), round(statistics.mean(d["mins"]), 1) if d["mins"] else None,
                   round(statistics.median(d["mins"]), 1) if d["mins"] else None,
                   human_minutes(max(d["mins"])) if d["mins"] else None, len(d["who"])])
    ws.append(["TOPLAM", "", tot["assigned"], tot["answered"], tot["assigned"] - tot["answered"],
               pct(tot["answered"], tot["assigned"]),
               round(statistics.mean(tot["mins"]), 1) if tot["mins"] else None,
               round(statistics.median(tot["mins"]), 1) if tot["mins"] else None,
               human_minutes(max(tot["mins"])) if tot["mins"] else None, ""])
    finish(ws, fmt={1: "dd.mm.yyyy", 6: "0.0%"})
    fill_row(ws, ws.max_row, TOTAL_FILL, 10)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    # ---- Günlük x TS
    ws = wb.create_sheet("Günlük x TS")
    agents = [n for n, _ in sorted(per.items(), key=lambda kv: (kv[1]["num"] or "999"))]
    header(ws, ["Tarih", "Gün"] + agents + ["Toplam"], [12, 11] + [14] * len(agents) + [10])
    grid = defaultdict(lambda: defaultdict(lambda: [0, 0]))
    for a in alerts:
        g = grid[a["date"].date()][a["responsible"]]
        g[0] += 1
        g[1] += int(a["answered"])
    ws.cell(row=1, column=1).value = "Tarih  (hücre: yanıtlanan / atanan)"
    for day in sorted(grid):
        row = [datetime(day.year, day.month, day.day), TR_DAYS[day.weekday()]]
        tn = ta = 0
        for ag in agents:
            n, a_ = grid[day][ag]
            row.append(f"{a_} / {n}" if n else "")
            tn += n
            ta += a_
        row.append(f"{ta} / {tn}")
        ws.append(row)
        r = ws.max_row
        for j, ag in enumerate(agents, 3):
            n, a_ = grid[day][ag]
            if n and a_ < n:
                ws.cell(row=r, column=j).fill = RED_FILL if a_ == 0 else ORANGE_FILL
    finish(ws, fmt={1: "dd.mm.yyyy"})
    for row in ws.iter_rows(min_row=2, min_col=3):
        for c in row:
            c.alignment = Alignment(horizontal="center")

    # ---- Gün Tipi
    ws = wb.create_sheet("Gün Tipi")
    header(ws, ["Yatırımsız Gün", "Bildirim", "Yanıtlanan", "Yanıtsız", "Yanıt Oranı", "Ort. Yanıt (dk)"],
           [14, 10, 10, 9, 10, 12])
    bytype = defaultdict(lambda: {"n": 0, "a": 0, "mins": []})
    for a in alerts:
        b = bytype[a["day"]]
        b["n"] += 1
        if a["answered"]:
            b["a"] += 1
            b["mins"].append(a["reply_minutes"])
    for k in sorted(bytype, key=lambda x: (x is None, x)):
        b = bytype[k]
        ws.append([f"{k}. gün" if k else "belirsiz", b["n"], b["a"], b["n"] - b["a"], pct(b["a"], b["n"]),
                   round(statistics.mean(b["mins"]), 1) if b["mins"] else None])
    finish(ws, fmt={5: "0.0%"})

    # ---- Ham Mesajlar
    ws = wb.create_sheet("Ham Mesajlar")
    header(ws, header_raw, [7, 10, 10, 20, 12, 11, 10, 22, 18, 28, 80, 10, 22, 30, 10, 14, 18][:len(header_raw)])
    for vals in raw_rows:
        ws.append(vals)
    finish(ws, wrap_cols=(11,), fmt={5: "dd.mm.yyyy HH:MM:SS"})

    # ---- Bilgi
    ws = wb.create_sheet("Bilgi")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    unanswered = tot["assigned"] - tot["answered"]
    lines = [
        ("Kaynak dosya", os.path.basename(src)),
        ("Kanal", info.get("Kanal", "")),
        ("Aralık", info.get("İstenen aralık", "")),
        ("Rapor zamanı", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Bildirim sayısı", tot["assigned"]),
        ("Yanıtlanan", f"{tot['answered']}  (%{pct(tot['answered'], tot['assigned']) * 100:.1f})"),
        ("Yanıtsız", unanswered),
        ("TS sayısı", len([n for n in per if per[n]["assigned"]])),
        ("PC -> TS eşlemesi", ", ".join(f"{k}={v}" for k, v in sorted(agent_by_num.items()))),
        ("", ""),
        ("Rapor sayfası", "Kırmızı satır = yanıt verilmemiş. Turuncu satır = sorumlu TS yerine başkası yanıtlamış."),
        ("Yanıt Süresi", "Bildirim ile ilk thread yanıtı arasındaki süre. Hafta sonu bildirimleri çoğunlukla "
                         "pazartesi/salı yanıtlandığı için ortalama yüksek çıkar; medyana bak."),
        ("Sorumlu TS", "PC kodunun başındaki numaraya göre; isim, o numarayla thread'e yanıt yazan kişiden alınır."),
        ("", ""),
    ]
    for k, v in lines:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["Bildirim olmayan mesajlar", "Bu mesajlar rapora alınmadı:"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    dropped_days = []
    for o in others:
        ws.append([o["date"].strftime("%d.%m.%Y %H:%M"), f"{o['author']}: {o['text'][:300]}"])
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if "not displaying some messages" in o["text"]:
            dropped_days.append(o["date"].strftime("%d.%m.%Y"))
    if dropped_days:
        ws.append(["", ""])
        ws.append(["UYARI", "Slack şu günlerde bot mesajlarının bir kısmını GÖSTERMEDİ (rate limit): "
                   + ", ".join(dropped_days) + ". O günlerin bildirim sayısı gerçekte daha yüksek olabilir; "
                   "eksik bildirimler bu raporda yoktur."])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C00000")
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=ws.max_row, column=2).fill = RED_FILL

    wb.save(out)
    return tot, unanswered, dropped_days


def main():
    ap = argparse.ArgumentParser(description="Slack export'undan memnuniyet/bildirim raporu üretir.")
    ap.add_argument("export", help="Export-SlackChannel.py çıktısı (.xlsx)")
    ap.add_argument("--output", "-o", help="Rapor dosyası (varsayılan: rapor_<kaynak>.xlsx, kaynağın yanına)")
    args = ap.parse_args()

    if not os.path.exists(args.export):
        sys.exit(f"Dosya bulunamadı: {args.export}")
    threads, header_raw, raw_rows, info = read_export(args.export)
    alerts, others, agent_by_num = analyze(threads)
    if not alerts:
        sys.exit("Hiç bildirim bulunamadı; mesaj formatı 'ID: ... PC: ...' şeklinde olmalı.")
    out = args.output or os.path.join(os.path.dirname(os.path.abspath(args.export)),
                                      "rapor_" + os.path.basename(args.export))
    tot, unanswered, dropped = write_report(out, alerts, others, agent_by_num, header_raw, raw_rows, info, args.export)
    print(f"{tot['assigned']} bildirim, {tot['answered']} yanıtlanmış, {unanswered} yanıtsız, "
          f"{len(agent_by_num)} TS.")
    if dropped:
        print(f"UYARI: Slack şu günlerde bazı bot mesajlarını göstermedi: {', '.join(dropped)}")
    print(f"Rapor: {os.path.abspath(out)}")


if __name__ == "__main__":
    main()

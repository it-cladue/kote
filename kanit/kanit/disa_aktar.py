# -*- coding: utf-8 -*-
"""Excel ciktisi: tam kitap (proje basina sayfa, eski A..J duzeni + Kayit No + Ekleyen) ve filtreli liste (web icin).

    tam_kitap(conn, yol_veya_bytesio, projeler=None, silinenler=True, site_url='')
    filtreli(conn, rows, yol_veya_bytesio, site_url='', baslik='Kayitlar', filtre_metni='', olusturan='')
    yeni_kitap() / satir_yaz(ws, satir_listesi)   -> web ajani da kullanir (formul enjeksiyonu korumasi burada)

Guvenlik: = + - @ ile baslayan metinler formul degil METIN olarak yazilir (cell.data_type='s').
"""
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from . import db

BASLIKLAR = ["Personel Kodu", "Uye ID", "Ana Uye ID", "Sessiz / Telesekreter", "Gyazo Linki", "Kategori", "Tarih / Saat", "KT", "Gorsel 2", "Aciklama", "Kayit No", "Ekleyen"]
SILINEN_EK = ["Silen", "Silme zamani", "Neden"]
GENISLIK = [14, 16, 16, 20, 48, 12, 20, 7, 20, 30, 10, 18, 14, 20, 24]
TARIH_BICIMI = "yyyy-mm-dd hh:mm:ss"
LINK_FONT = Font(color="0563C1", underline="single")


def yeni_kitap():
    """Bos kitap (ilk varsayilan sayfa silinmis). Sayfalar create_sheet ile acilir."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


def metin_rengi(hex_renk):
    """Arka plan rengine gore okunur yazi rengi (OF sari -> siyah)."""
    try:
        h = (hex_renk or "#999999").lstrip("#")
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        return "111111" if (0.299 * r + 0.587 * g + 0.114 * b) > 150 else "FFFFFF"
    except Exception:
        return "FFFFFF"


def _renk6(renk):
    h = (renk or "").lstrip("#").upper()
    return h if len(h) == 6 and all(c in "0123456789ABCDEF" for c in h) else "999999"


def _tarih(v):
    if isinstance(v, datetime):
        return v
    if v:
        try:
            return datetime.strptime(str(v)[:19], "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return str(v)
    return None


def satir_yaz(ws, satir_listesi):
    """Bir satir ekler. = + - @ ile baslayan metinler METIN olarak isaretlenir (formul enjeksiyonu korumasi);
    datetime hucreleri yyyy-mm-dd hh:mm:ss bicimini alir. -> satir numarasi"""
    ws.append(satir_listesi)
    n = ws.max_row
    for h in ws[n]:
        if isinstance(h.value, str) and h.value[:1] in "=+-@":
            h.data_type = "s"
        elif isinstance(h.value, datetime):
            h.number_format = TARIH_BICIMI
    return n


def baslik_yaz(ws, basliklar, renk="#1F2937", genislikler=None):
    satir_yaz(ws, list(basliklar))
    r6 = _renk6(renk)
    for i, _ in enumerate(basliklar, 1):
        h = ws.cell(row=1, column=i)
        h.font = Font(bold=True, color=metin_rengi(r6))
        h.fill = PatternFill("solid", fgColor=r6)
        h.alignment = Alignment(vertical="center")
        g = (genislikler or GENISLIK)[i - 1] if i - 1 < len(genislikler or GENISLIK) else 14
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = "A2"


def gorsel_adresi(site_url, kayit_id):
    return f"{(site_url or '').rstrip('/')}/kayit/{kayit_id}/gorsel"


def kayit_satiri(d, site_url="", silinen=False):
    """Kayit dict -> hucre degerleri (BASLIKLAR sirasinda; silinen=True ise SILINEN_EK de eklenir)."""
    link = d.get("link") or ""
    if not link and d.get("gorsel_dosya"):
        link = gorsel_adresi(site_url, d["id"])
    ekleyen = d.get("ekleyen_ad") or d.get("ekleyen_tg_id") or ""
    if ekleyen and d.get("ekleyen_guven") == "tahmini":
        ekleyen += " (?)"
    s = [d.get("personel") or "", d.get("uye_id") or "", d.get("ana_uye_id") or "", d.get("tip") or "", link, d.get("kategori") or "",
         _tarih(d.get("tarih")), bool(d.get("kt")), d.get("gorsel2") or "", d.get("aciklama") or "", d.get("id"), ekleyen]
    if silinen:
        s += [d.get("silen") or "", _tarih(d.get("silme_zamani")), d.get("silme_nedeni") or ""]
    return s


def kayit_satirlari_yaz(ws, kayitlar, site_url="", silinen=False, link_sutunu=5):
    """Kayit listesini sayfaya yazar: E sutunu tiklanir baglanti (http ile basliyorsa). -> yazilan satir sayisi"""
    n = 0
    for d in kayitlar:
        d = dict(d)
        satir = kayit_satiri(d, site_url, silinen)
        rn = satir_yaz(ws, satir)
        link = satir[link_sutunu - 1]
        if isinstance(link, str) and link.lower().startswith(("http://", "https://")):
            h = ws.cell(row=rn, column=link_sutunu)
            h.hyperlink = link
            h.font = LINK_FONT
        n += 1
    return n


def _bitir(ws, sutun_sayisi):
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(sutun_sayisi)}{max(ws.max_row, 1)}"


def _kaydet(wb, hedef):
    if hedef is None:
        hedef = io.BytesIO()
    if isinstance(hedef, (str, bytes)):
        wb.save(hedef)
        return hedef
    wb.save(hedef)
    try:
        hedef.seek(0)
    except Exception:
        pass
    return hedef


def tam_kitap(conn, hedef=None, projeler=None, silinenler=True, site_url="", olusturan=""):
    """Butun veritabani: aktif projeler sira ile birer sayfa (A..J eski duzen + Kayit No + Ekleyen), istege bagli SILINENLER
    sayfasi ve Ozet. hedef: dosya yolu ya da BytesIO (None -> yeni BytesIO). -> hedef"""
    tum = db.projeler(conn, aktif=True)
    if projeler:
        secili = {p.upper() for p in projeler}
        tum = [p for p in tum if p["kod"] in secili]
    renkler = db.proje_renk(conn)
    wb = yeni_kitap()
    sayac = {}
    for p in tum:
        ws = wb.create_sheet(p["kod"][:31])
        r6 = _renk6(p["renk"])
        ws.sheet_properties.tabColor = r6
        baslik_yaz(ws, BASLIKLAR, "#" + r6)
        rows = conn.execute("SELECT * FROM kayitlar WHERE proje=? AND silindi=0 ORDER BY COALESCE(tarih,'0000'), id", (p["kod"],)).fetchall()
        kayit_satirlari_yaz(ws, rows, site_url)
        _bitir(ws, len(BASLIKLAR))
    if silinenler:
        ws = wb.create_sheet("SILINENLER")
        ws.sheet_properties.tabColor = "7F7F7F"
        # ayni A..L duzeni + Silen / Silme zamani / Neden + Proje (sona)
        baslik_yaz(ws, BASLIKLAR + SILINEN_EK + ["Proje"], "#7F7F7F", GENISLIK + [8])
        q = "SELECT * FROM kayitlar WHERE silindi=1" + (f" AND proje IN ({','.join('?' * len(tum))})" if projeler else "") + " ORDER BY proje, COALESCE(tarih,'0000'), id"
        rows = conn.execute(q, tuple(p["kod"] for p in tum) if projeler else ()).fetchall()
        for d in rows:
            d = dict(d)
            satir = kayit_satiri(d, site_url, silinen=True) + [d["proje"]]
            rn = satir_yaz(ws, satir)
            if isinstance(satir[4], str) and satir[4].lower().startswith(("http://", "https://")):
                ws.cell(row=rn, column=5).hyperlink = satir[4]
                ws.cell(row=rn, column=5).font = LINK_FONT
        _bitir(ws, len(BASLIKLAR) + 4)
    # Ozet
    oz = wb.create_sheet("Ozet")
    baslik_yaz(oz, ["Proje", "Toplam", "Onayli", "Bekleyen", "Silinen", "Yorumlu"], "#1F2937", [12, 10, 10, 10, 10, 10])
    toplam = [0, 0, 0, 0, 0]
    for p in tum:
        r = conn.execute("""SELECT COUNT(*) t, SUM(CASE WHEN kt=1 AND silindi=0 THEN 1 ELSE 0 END) o, SUM(CASE WHEN kt=0 AND silindi=0 THEN 1 ELSE 0 END) b,
                            SUM(CASE WHEN silindi=1 THEN 1 ELSE 0 END) s, SUM(CASE WHEN yorum_sayisi>0 THEN 1 ELSE 0 END) y FROM kayitlar WHERE proje=?""", (p["kod"],)).fetchone()
        degerler = [r["t"] or 0, r["o"] or 0, r["b"] or 0, r["s"] or 0, r["y"] or 0]
        satir_yaz(oz, [p["kod"]] + degerler)
        oz.cell(row=oz.max_row, column=1).fill = PatternFill("solid", fgColor=_renk6(p["renk"]))
        oz.cell(row=oz.max_row, column=1).font = Font(bold=True, color=metin_rengi(_renk6(p["renk"])))
        toplam = [a + b for a, b in zip(toplam, degerler)]
    satir_yaz(oz, ["TOPLAM"] + toplam)
    oz.cell(row=oz.max_row, column=1).font = Font(bold=True)
    satir_yaz(oz, [])
    satir_yaz(oz, ["Olusturma", db.simdi()])
    satir_yaz(oz, ["Olusturan", olusturan or "sistem"])
    satir_yaz(oz, ["Silinenler sayfasi", "var" if silinenler else "yok"])
    satir_yaz(oz, ["Site", site_url or "-"])
    return _kaydet(wb, hedef)


def filtreli(conn, rows, hedef=None, site_url="", baslik="Kayitlar", filtre_metni="", olusturan="", proje_sutunu=True):
    """Web'in filtreli disa aktarimi: verilen kayit satirlari (sqlite3.Row ya da dict) tek sayfada + Ozet.
    Ayni formul enjeksiyonu korumasi, ayni sutun duzeni (basa Proje sutunu eklenir). -> hedef"""
    renkler = db.proje_renk(conn)
    wb = yeni_kitap()
    ws = wb.create_sheet(baslik[:31] or "Kayitlar")
    basliklar = (["Proje"] if proje_sutunu else []) + BASLIKLAR + ["Durum"]
    baslik_yaz(ws, basliklar, "#1F2937", ([8] if proje_sutunu else []) + GENISLIK + [10])
    kay = 1 if proje_sutunu else 0
    sayac = {}
    for d in rows:
        d = dict(d)
        durum = "silindi" if d.get("silindi") else ("onayli" if d.get("kt") else "bekliyor")
        satir = ([d.get("proje") or ""] if proje_sutunu else []) + kayit_satiri(d, site_url) + [durum]
        rn = satir_yaz(ws, satir)
        link = satir[kay + 4]
        if isinstance(link, str) and link.lower().startswith(("http://", "https://")):
            ws.cell(row=rn, column=kay + 5).hyperlink = link
            ws.cell(row=rn, column=kay + 5).font = LINK_FONT
        if proje_sutunu:
            r6 = _renk6(renkler.get(d.get("proje")))
            ws.cell(row=rn, column=1).fill = PatternFill("solid", fgColor=r6)
            ws.cell(row=rn, column=1).font = Font(bold=True, color=metin_rengi(r6))
        s = sayac.setdefault(d.get("proje") or "", [0, 0, 0, 0])
        s[0] += 1
        s[3 if d.get("silindi") else (1 if d.get("kt") else 2)] += 1
    _bitir(ws, len(basliklar))
    oz = wb.create_sheet("Ozet")
    baslik_yaz(oz, ["Proje", "Toplam", "Onayli", "Bekleyen", "Silinen"], "#1F2937", [12, 10, 10, 10, 10])
    for p in sorted(sayac):
        satir_yaz(oz, [p] + sayac[p])
    satir_yaz(oz, ["TOPLAM"] + [sum(v[i] for v in sayac.values()) for i in range(4)])
    satir_yaz(oz, [])
    satir_yaz(oz, ["Filtre", filtre_metni or "-"])
    satir_yaz(oz, ["Olusturan", f"{olusturan or 'sistem'} - {db.simdi()}"])
    return _kaydet(wb, hedef)

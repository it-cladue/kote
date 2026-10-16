# -*- coding: utf-8 -*-
"""Panelin ortak parcalari: tek blueprint, istek bagli yardimcilar (ip, yetki, filtre SQL, kayit sozlugu, xlsx)."""
import os, re, functools
from datetime import datetime
from flask import Blueprint, g, request, session, redirect, url_for, abort, current_app
from werkzeug.datastructures import MultiDict
from .. import db, yetki, normalle as N

bp = Blueprint("panel", __name__)

DURUMLAR = ("bekleyen", "onayli", "silinen", "hepsi")
DURUM_ADLARI = {"bekleyen": "Onay bekleyen", "onayli": "Onayl\u0131", "silinen": "Silinen", "hepsi": "Hepsi"}
ALAN_ADLARI = {"personel": "Personel kodu", "uye_id": "\u00dcye ID", "ana_uye_id": "Ana \u00fcye ID", "tip": "T\u00fcr", "kategori": "Kategori",
               "link": "Link", "tarih": "Tarih", "gorsel2": "G\u00f6rsel 2", "aciklama": "A\u00e7\u0131klama", "ek_not": "Ek not", "lead_id": "Lead ID"}
# tablo siralama: url parametresi -> SQL
SIRALAMA = {"personel": ("length(personel_norm)", "personel_norm"), "uye_id": ("length(uye_norm)", "uye_norm"), "ana_uye_id": ("length(ana_norm)", "ana_norm"),
            "tip": ("tip",), "link": ("link_norm",), "kategori": ("kategori",), "tarih": ("COALESCE(tarih,'0000')",), "kt": ("kt",), "gorsel2": ("gorsel2",),
            "aciklama": ("aciklama",), "ekleyen": ("ekleyen_ad",), "yorum": ("yorum_sayisi",), "durum": ("silindi", "kt"), "id": ("id",)}


def ayar():
    return current_app.config["KANIT"]


def conn():
    return g.conn


def ip():
    """Denetim icin istemci adresi. X-Forwarded-For / CF-Connecting-IP yalnizca GUVENILIR_PROXY=1 ise."""
    if str(ayar().get("GUVENILIR_PROXY", "0")).lower() in ("1", "true"):
        return (request.headers.get("CF-Connecting-IP") or (request.headers.get("X-Forwarded-For") or "").split(",")[0].strip() or request.remote_addr or "")
    return request.remote_addr or ""


def yerel_istek():
    return (ip() or "") in ("127.0.0.1", "::1", "")


def kullanici_adi():
    return g.k["kullanici_adi"] if g.get("k") else ""


def giris_gerekli(f):
    @functools.wraps(f)
    def w(*a, **kw):
        if not g.get("k"):
            if request.method == "POST" or request.path.endswith(".json"):
                abort(401)
            return redirect(url_for("panel.giris", sonraki=request.full_path.rstrip("?")))
        return f(*a, **kw)
    return w


def admin_gerekli(f):
    @functools.wraps(f)
    def w(*a, **kw):
        if not g.get("k"):
            if request.method == "POST" or request.path.endswith(".json"):
                abort(401)
            return redirect(url_for("panel.giris", sonraki=request.full_path.rstrip("?")))
        if not g.yk["admin"]:
            abort(403)
        return f(*a, **kw)
    return w


def guvenli_sonraki(s):
    s = (s or "").strip()
    if not s.startswith("/") or s.startswith("//") or "\\" in s or ":" in s.split("?")[0]:
        return None
    return s


# ==================== projeler / renkler ====================
def tum_projeler():
    """Aktif projeler (sira ile) + kayitlarda gorulen ama listede olmayan kodlar."""
    if "projeler" not in g:
        c = conn()
        sira = db.proje_kodlari(c)
        gorulen = [r["proje"] for r in c.execute("SELECT DISTINCT proje FROM kayitlar").fetchall()]
        g.projeler = sira + sorted(p for p in gorulen if p not in sira)
    return g.projeler


def izinli_projeler(eylem="oku"):
    return yetki.izinli_projeler(g.yk, tum_projeler(), eylem)


def renkler():
    if "renkler" not in g:
        g.renkler = db.proje_renk(conn())
    return g.renkler


def renk(proje):
    return renkler().get((proje or "").upper(), "#999999")


def metin_rengi(hex_renk):
    try:
        h = (hex_renk or "#999999").lstrip("#")
        r, gg, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        return "#111111" if (0.299 * r + 0.587 * gg + 0.114 * b) > 150 else "#ffffff"
    except Exception:
        return "#ffffff"


def kategori_listesi():
    c = conn()
    kats = [k["etiket"] for k in db.kategoriler(c)]
    for r in c.execute("SELECT DISTINCT kategori FROM kayitlar WHERE kategori<>'' ORDER BY 1").fetchall():
        if r[0] not in kats:
            kats.append(r[0])
    return kats


def url_degistir(**kw):
    """Mevcut sorgu parametrelerini koruyarak bazilarini degistir (s= verilmediyse sayfa sifirlanir)."""
    a = request.args.to_dict(flat=False)
    a.pop("s", None)
    for k, v in kw.items():
        if v is None or v == "":
            a.pop(k, None)
        elif isinstance(v, (list, tuple)):
            a[k] = [str(x) for x in v]
        else:
            a[k] = [str(v)]
    return url_for(request.endpoint or "panel.kayitlar", **request.view_args, **a)


# ==================== filtre ====================
def _like(s):
    return "%" + s.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"


def filtre_sql(args, izinli, proje_sabit=None):
    """-> (where, parametreler, f). f: sayfada gosterilen secimler."""
    if not isinstance(args, MultiDict):
        args = MultiDict(args)
    where, par = [], []
    if proje_sabit:
        sec = [proje_sabit] if proje_sabit in izinli else []
    else:
        sec = [p.upper() for p in args.getlist("proje") if p.upper() in izinli] or list(izinli)
    f = {"proje": sec, "durum": "bekleyen", "kategori": "", "q": "", "ekleyen": "", "tarih1": "", "tarih2": "", "tarihsiz": False, "yorumlu": False}
    if not sec:
        return "1=0", [], f
    where.append("proje IN (%s)" % ",".join("?" * len(sec)))
    par += sec
    durum = args.get("durum") or "bekleyen"
    if durum not in DURUMLAR:
        durum = "bekleyen"
    f["durum"] = durum
    if durum == "bekleyen":
        where.append("kt=0 AND silindi=0")
    elif durum == "onayli":
        where.append("kt=1 AND silindi=0")
    elif durum == "silinen":
        where.append("silindi=1")
    kat = N.metin(args.get("kategori"))
    if kat:
        f["kategori"] = kat
        where.append("(kategori_kod=? OR kategori=?)")
        par += [N.kucuk(kat), N.buyuk(kat)]
    q = N.metin(args.get("q"))[:200]
    n = N.normalle(q)
    if n:
        f["q"] = q
        where.append("(personel_norm=? OR uye_norm=? OR ana_norm=? OR link_norm=? OR personel_norm LIKE ? ESCAPE '\\' OR uye_norm LIKE ? ESCAPE '\\' OR ana_norm LIKE ? ESCAPE '\\' OR link_norm LIKE ? ESCAPE '\\')")
        par += [n, n, n, n] + [_like(n)] * 4
    ekleyen = N.metin(args.get("ekleyen"))[:80]
    if ekleyen:
        f["ekleyen"] = ekleyen
        where.append("(ekleyen_ad LIKE ? ESCAPE '\\' OR ekleyen_tg_id=?)")
        par += [_like(ekleyen), ekleyen]
    t1, t2 = N.metin(args.get("tarih1"))[:10], N.metin(args.get("tarih2"))[:10]
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t1):
        f["tarih1"] = t1
        where.append("tarih >= ?")
        par.append(t1 + " 00:00:00")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t2):
        f["tarih2"] = t2
        where.append("tarih <= ?")
        par.append(t2 + " 23:59:59")
    if args.get("tarihsiz") == "1":
        f["tarihsiz"] = True
        where.append("tarih IS NULL")
    if args.get("yorumlu") == "1":
        f["yorumlu"] = True
        where.append("yorum_sayisi > 0")
    return " AND ".join(where), par, f


def siralama_sql(args, varsayilan="tarih", varsayilan_yon="desc"):
    s = args.get("sirala") or varsayilan
    if s not in SIRALAMA:
        s = varsayilan
    yon = (args.get("yon") or varsayilan_yon).lower()
    yon = "asc" if yon == "asc" else "desc"
    y = yon.upper()
    return ", ".join(f"{p} {y}" for p in SIRALAMA[s]) + f", id {y}", s, yon


# ==================== kayit sozlugu ====================
def kayit_izin(kayit, eylem):
    return yetki.izin(g.yk, kayit["proje"], eylem)


def gorsel_var(d):
    return bool(d.get("gorsel_dosya")) or d.get("link_turu") == "gyazo"


def kayit_sozluk(r):
    d = dict(r)
    d["gorsel_var"] = gorsel_var(d)
    d["renk"] = renk(d["proje"])
    d["izin_kt"] = kayit_izin(d, "kt")
    d["izin_duzenle"] = kayit_izin(d, "duzenle")
    d["izin_sil"] = kayit_izin(d, "sil")
    d["izin_gorsel"] = kayit_izin(d, "gorsel")
    d["dis_link"] = d["link"] if (d.get("link") or "").startswith(("http://", "https://")) else None
    d["durum"] = "silindi" if d.get("silindi") else ("onayli" if d.get("kt") else "bekliyor")
    return d


def kayit_durum_metni(d):
    return "silindi" if d.get("silindi") else ("onayli" if d.get("kt") else "bekliyor")


def sayi(v, varsayilan=0):
    try:
        return int(v)
    except (TypeError, ValueError):
        return varsayilan


def simdi():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def tarih_kisa(s):
    """'2026-10-14 09:12:33' -> '14.10 09:12'"""
    try:
        t = datetime.strptime(str(s)[:19], "%Y-%m-%d %H:%M:%S")
        return t.strftime("%d.%m %H:%M")
    except Exception:
        return str(s or "")


# ==================== xlsx ====================
def xlsx_satir(ws, degerler):
    """Sayfaya satir ekle -> satir no. = + - @ ile baslayan metinler formul degil METIN olarak yazilir (formul enjeksiyonu onlemi)."""
    ws.append(degerler)
    for h in ws[ws.max_row]:
        if isinstance(h.value, str) and h.value[:1] in "=+-@":
            h.data_type = "s"
    return ws.max_row


def xlsx_baslik(ws, basliklar, genislikler):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    xlsx_satir(ws, basliklar)
    for i, gen in enumerate(genislikler, 1):
        ws.column_dimensions[get_column_letter(i)].width = gen
        h = ws.cell(row=1, column=i)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor="1F2937")
        h.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

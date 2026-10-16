# -*- coding: utf-8 -*-
"""Cikti: filtreli Excel (xlsx), tum kitap (disa_aktar.tam_kitap varsa), liste kontrolu."""
import os, io
from datetime import datetime, timedelta
from flask import g, request, redirect, url_for, render_template, send_file, flash, abort
from werkzeug.datastructures import MultiDict
from .. import db
from . import liste_kontrol
from .ortak import (bp, ayar, conn, ip, giris_gerekli, izinli_projeler, renk, metin_rengi, kategori_listesi, filtre_sql, kayit_durum_metni,
                    xlsx_satir, xlsx_baslik, simdi)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# bugunku sutunlar + yorum sutunlari
DISA_SUTUNLAR = [("Proje", "proje", 8), ("Kategori", "kategori", 10), ("Personel kodu", "personel", 14), ("Uye ID", "uye_id", 14), ("Ana uye ID", "ana_uye_id", 14),
                 ("Tur", "tip", 12), ("Tarih", "tarih", 19), ("Link", "link", 60), ("Onay (KT)", "kt", 10), ("Onaylayan", "kt_kim", 14), ("Onay zamani", "kt_zaman", 19),
                 ("Ekleyen (Telegram)", "ekleyen_ad", 18), ("Ekleyen ID", "ekleyen_tg_id", 14), ("Aciklama", "aciklama", 30), ("Durum", "_durum", 12),
                 ("Excel satiri", "eski_satir", 11), ("Gorsel dosyasi", "_gorsel", 40), ("Gorsel 2", "gorsel2", 20), ("Yorumlar", "_yorum_sayisi", 10), ("Son yorum", "_son_yorum", 40)]
LISTE_KAYIT_SUTUNLAR = [("Proje", "proje", 8), ("Kategori", "kategori", 10), ("Personel kodu", "personel", 14), ("Uye ID", "uye_id", 14), ("Ana uye ID", "ana_uye_id", 14),
                        ("Tarih", "tarih", 19), ("Onay (KT)", "kt", 10), ("Link", "link", 50), ("Ekleyen (Telegram)", "ekleyen_ad", 18), ("Durum", "_durum", 12), ("Kayit ID", "id", 9)]


def _disa_aktar_sorgu():
    c = conn()
    izinli = izinli_projeler()
    args = request.args.to_dict(flat=False)
    args.setdefault("durum", ["hepsi"])
    where, par, f = filtre_sql(MultiDict(args), izinli)
    rows = c.execute(f"SELECT * FROM kayitlar WHERE {where} ORDER BY proje, COALESCE(tarih,'0000'), id", par).fetchall()
    return rows, f


def _son_yorumlar(c, kids):
    out = {}
    for i in range(0, len(kids), 500):
        parca = kids[i:i + 500]
        q = f"""SELECT y.kayit_id, y.kullanici, y.metin, y.zaman FROM yorumlar y
                WHERE y.silindi=0 AND y.kayit_id IN ({','.join('?' * len(parca))}) AND y.id = (SELECT MAX(id) FROM yorumlar z WHERE z.kayit_id=y.kayit_id AND z.silindi=0)"""
        for r in c.execute(q, parca).fetchall():
            out[r["kayit_id"]] = f"{r['kullanici']} ({r['zaman'][:16]}): {r['metin'][:200]}"
    return out


@bp.route("/disa-aktar")
@giris_gerekli
def disa_aktar_sayfasi():
    bugun = datetime.now().date()
    try:
        from .. import disa_aktar  # noqa: F401
        tam_kitap_var = hasattr(disa_aktar, "tam_kitap")
    except ImportError:
        tam_kitap_var = False
    return render_template("disa_aktar.html", izinli=izinli_projeler(), kategoriler=kategori_listesi(), tarih1=(bugun - timedelta(days=30)).isoformat(),
                           tarih2=bugun.isoformat(), tam_kitap_var=tam_kitap_var)


def xlsx_uret(rows, f, kim):
    """Filtreli kayitlar -> xlsx baytlari (Kayitlar + Ozet sayfasi). = + - @ ile baslayan metinler metin olarak yazilir."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    c = conn()
    son_yorum = _son_yorumlar(c, [r["id"] for r in rows])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kayitlar"
    xlsx_baslik(ws, [b for b, _, _ in DISA_SUTUNLAR], [gen for _, _, gen in DISA_SUTUNLAR])
    link_sutun = [k for _, k, _ in DISA_SUTUNLAR].index("link") + 1
    tarih_sutunlar = [i + 1 for i, (_, k, _) in enumerate(DISA_SUTUNLAR) if k in ("tarih", "kt_zaman")]
    for r in rows:
        d = dict(r)
        satir = []
        for _, k, _ in DISA_SUTUNLAR:
            if k == "_durum":
                v = kayit_durum_metni(d)
            elif k == "_gorsel":
                v = os.path.basename(d["gorsel_dosya"]) if d.get("gorsel_dosya") else ""
            elif k == "_yorum_sayisi":
                v = d.get("yorum_sayisi") or 0
            elif k == "_son_yorum":
                v = son_yorum.get(d["id"], "")
            elif k == "kt":
                v = "EVET" if d.get("kt") else "HAYIR"
            elif k in ("tarih", "kt_zaman") and d.get(k):
                try:
                    v = datetime.strptime(d[k][:19], "%Y-%m-%d %H:%M:%S")
                except ValueError:
                    v = d[k]
            else:
                v = d.get(k)
                if v is None:
                    v = ""
            satir.append(v)
        n = xlsx_satir(ws, satir)
        link = d.get("link") or ""
        if link.startswith(("http://", "https://")):
            hucre = ws.cell(row=n, column=link_sutun)
            hucre.hyperlink = link
            hucre.font = Font(color="0563C1", underline="single")
        for col in tarih_sutunlar:
            if isinstance(ws.cell(row=n, column=col).value, datetime):
                ws.cell(row=n, column=col).number_format = "yyyy-mm-dd hh:mm:ss"
        renk_hex = renk(d["proje"]).lstrip("#")
        if len(renk_hex) == 6:
            ws.cell(row=n, column=1).fill = PatternFill("solid", fgColor=renk_hex)
            ws.cell(row=n, column=1).font = Font(bold=True, color=metin_rengi("#" + renk_hex).lstrip("#"))
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(DISA_SUTUNLAR))}{ws.max_row}"
    oz = wb.create_sheet("Ozet")
    oz.append(["Proje", "Toplam", "Onayli", "Onay bekleyen", "Silinen"])
    for i in range(1, 6):
        oz.cell(row=1, column=i).font = Font(bold=True)
        oz.column_dimensions[get_column_letter(i)].width = 14
    sayac = {}
    for r in rows:
        s_ = sayac.setdefault(r["proje"], [0, 0, 0, 0])
        s_[0] += 1
        if r["silindi"]:
            s_[3] += 1
        elif r["kt"]:
            s_[1] += 1
        else:
            s_[2] += 1
    for p_ in sorted(sayac):
        oz.append([p_] + sayac[p_])
    oz.append(["TOPLAM"] + [sum(v[i] for v in sayac.values()) for i in range(4)])
    oz.append([])
    xlsx_satir(oz, ["Filtre", f"projeler: {', '.join(f.get('proje', []))}; durum: {f.get('durum')}; tarih: {f.get('tarih1') or '-'} / {f.get('tarih2') or '-'}; kategori: {f.get('kategori') or '-'}; arama: {f.get('q') or '-'}"])
    xlsx_satir(oz, ["Olusturan", f"{kim} - {simdi()}"])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


@bp.route("/disa-aktar.xlsx")
@giris_gerekli
def disa_aktar_xlsx():
    rows, f = _disa_aktar_sorgu()
    bio = xlsx_uret(rows, f, g.k["kullanici_adi"])
    projeler_ad = "-".join(f.get("proje", [])) if f.get("proje") and len(f["proje"]) <= 3 else "tum"
    ad = f"kanit_{projeler_ad}_{(f.get('tarih1') or 'bastan')}_{(f.get('tarih2') or 'bugun')}.xlsx".replace("/", "-")
    db.denetim_yaz(conn(), g.k["kullanici_adi"], "disa_aktar", yeni=f"{len(rows)} satir | {ad}", ip=ip())
    return send_file(bio, as_attachment=True, download_name=ad, mimetype=XLSX_MIME)


def _tam_kitap_uret(c, hedef):
    """disa_aktar.tam_kitap(conn, hedef) / (conn) -> bayt ya da dosya yolu. Imza baska bir ekibin; esnek cagri."""
    from .. import disa_aktar
    fn = disa_aktar.tam_kitap
    try:
        r = fn(c, hedef)
    except TypeError:
        r = fn(c)
    if isinstance(r, (bytes, bytearray)):
        return io.BytesIO(bytes(r))
    if hasattr(r, "read"):
        r.seek(0)
        return r
    if isinstance(r, str) and os.path.isfile(r):
        return r
    if os.path.isfile(hedef):
        return hedef
    raise RuntimeError("tam kitap uretilemedi")


@bp.route("/tam-kitap.xlsx")
@giris_gerekli
def tam_kitap_xlsx():
    if not g.yk["admin"] and not g.yk.get("hepsi"):
        abort(403)
    try:
        from .. import disa_aktar  # noqa: F401
    except ImportError:
        flash("T\u00fcm kitap \u00e7\u0131kt\u0131s\u0131 bu kurulumda hen\u00fcz yok (kanit.disa_aktar mod\u00fcl\u00fc bulunamad\u0131).")
        return redirect(url_for("panel.disa_aktar_sayfasi"))
    hedef = os.path.join(ayar()["ONBELLEK"], f"tam_kitap_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    try:
        sonuc = _tam_kitap_uret(conn(), hedef)
    except Exception as e:
        flash(f"T\u00fcm kitap \u00fcretilemedi: {e}")
        return redirect(url_for("panel.disa_aktar_sayfasi"))
    db.denetim_yaz(conn(), g.k["kullanici_adi"], "tam_kitap", ip=ip())
    return send_file(sonuc, as_attachment=True, download_name=f"kanit_tum_kitap_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx", mimetype=XLSX_MIME)


# ==================== liste kontrolu ====================
def _liste_kayitlar(filtre, izinli):
    a = {k: (v if isinstance(v, list) else [v]) for k, v in dict(filtre).items()}
    a.setdefault("durum", ["hepsi"])
    where, par, f = filtre_sql(MultiDict(a), izinli)
    rows = conn().execute(f"""SELECT id, proje, kategori, personel, uye_id, ana_uye_id, link, tarih, kt, silindi, ekleyen_ad, aciklama,
                              personel_norm, uye_norm, ana_norm, link_norm FROM kayitlar WHERE {where} ORDER BY proje, COALESCE(tarih,'0000') DESC, id DESC""", par).fetchall()
    return [dict(r) for r in rows], f


def _liste_esleme(form, genislik):
    esleme = []
    for sutun in form.getlist("kontrol"):
        if sutun.isdigit() and int(sutun) < genislik:
            alanlar = [a for a in form.getlist(f"alan_{sutun}") if a in liste_kontrol.ALAN_ADLARI]
            if not alanlar:
                raise ValueError(f"{int(sutun) + 1}. s\u00fctun i\u00e7in panelde hangi alanda aranaca\u011f\u0131n\u0131 se\u00e7in.")
            esleme.append({"sutun": int(sutun), "alanlar": alanlar})
    if not esleme:
        raise ValueError("Kontrol edilecek en az bir s\u00fctun i\u015faretleyin.")
    return esleme


def _liste_basliklar(veri):
    satirlar = veri["tablolar"][veri["sayfa"]]
    return [(satirlar[0][e["sutun"]] if veri["baslik"] and satirlar and satirlar[0][e["sutun"]] else f"S\u00fctun {e['sutun'] + 1}") for e in veri["esleme"]]


def _liste_bulunamadi():
    flash("Liste bulunamad\u0131 ya da s\u00fcresi dolmu\u015f; dosyay\u0131 yeniden y\u00fckleyin.")
    return redirect(url_for("panel.disa_aktar_sayfasi"))


@bp.route("/liste-kontrol", methods=["GET", "POST"])
@giris_gerekli
def liste_kontrol_sayfasi():
    onbellek = ayar()["ONBELLEK"]
    izinli = izinli_projeler()
    adim = request.form.get("adim") if request.method == "POST" else "esleme"
    if adim == "yukle":
        dosya = request.files.get("dosya")
        try:
            if dosya and dosya.filename:
                veri = liste_kontrol.coz(dosya.filename, dosya.read())
                veri["dosya"] = os.path.basename(dosya.filename)[:80]
            elif (request.form.get("metin") or "").strip():
                veri = liste_kontrol.metin_coz(request.form["metin"])
                veri["dosya"] = "yap\u0131\u015ft\u0131r\u0131lan liste"
            else:
                raise ValueError("bir dosya se\u00e7in ya da listeyi kutuya yap\u0131\u015ft\u0131r\u0131n.")
        except ValueError as e:
            flash(f"Liste okunamad\u0131: {e}")
            return redirect(url_for("panel.disa_aktar_sayfasi"))
        veri.update(sahip=g.k["kullanici_adi"], sayfa=veri["sayfalar"][0])
        veri["baslik"] = liste_kontrol.baslik_mi(veri["tablolar"][veri["sayfa"]])
        return redirect(url_for("panel.liste_kontrol_sayfasi", token=liste_kontrol.kaydet(onbellek, veri)))
    token = request.values.get("token", "")
    veri = liste_kontrol.yukle(onbellek, token, g.k["kullanici_adi"])
    if not veri:
        return _liste_bulunamadi()
    hata = None
    if adim in ("onizle", "kontrol"):
        sayfa = request.form.get("sayfa")
        if sayfa in veri["tablolar"] and sayfa != veri["sayfa"]:
            veri["sayfa"] = sayfa
            veri["baslik"] = liste_kontrol.baslik_mi(veri["tablolar"][sayfa])
        else:
            veri["baslik"] = request.form.get("baslik") == "1"
        if adim == "kontrol":
            try:
                satirlar = veri["tablolar"][veri["sayfa"]]
                veri["esleme"] = _liste_esleme(request.form, len(satirlar[0]) if satirlar else 0)
                kayitlar, veri["filtre"] = _liste_kayitlar(request.form.to_dict(flat=False), izinli)
                liste_kontrol.kaydet(onbellek, veri, token)
                return _liste_sonuc_sayfasi(token, veri, kayitlar)
            except ValueError as e:
                hata = str(e)
        liste_kontrol.kaydet(onbellek, veri, token)
    satirlar = veri["tablolar"][veri["sayfa"]]
    sutunlar = liste_kontrol.sutunlar(satirlar, veri["baslik"])
    if veri.get("esleme"):
        secim = {e["sutun"]: e["alanlar"] for e in veri["esleme"]}
        for st in sutunlar:
            st["secili"], st["alanlar"] = st["no"] in secim, secim.get(st["no"], st["alanlar"])
    return render_template("liste_kontrol.html", token=token, veri=veri, sutunlar=sutunlar, onizleme=satirlar[1 if veri["baslik"] else 0:][:liste_kontrol.ONIZLEME],
                           satir_sayisi=len(satirlar) - (1 if veri["baslik"] else 0), alanlar=liste_kontrol.ALANLAR, izinli=izinli, kategoriler=kategori_listesi(),
                           f=veri.get("filtre") or {}, hata=hata)


def _liste_sonuc_sayfasi(token, veri, kayitlar):
    satirlar = veri["tablolar"][veri["sayfa"]]
    sonuc = liste_kontrol.eslestir(satirlar, veri["baslik"], veri["esleme"], kayitlar)
    bulunan, yok, proje_sayac = [], [], {}
    for s in sonuc["satirlar"]:
        if not s["eslesen"]:
            yok.append(s)
            continue
        for p in {kayitlar[i]["proje"] for i in s["eslesen"]}:
            proje_sayac[p] = proje_sayac.get(p, 0) + 1
        bulunan.append({"no": s["no"], "degerler": s["degerler"], "kayitlar": [kayitlar[i] for i in s["eslesen"][:3]], "n": len(s["eslesen"])})
    db.denetim_yaz(conn(), g.k["kullanici_adi"], "liste_kontrol", yeni=f"{sonuc['bulunan'] + sonuc['yok']} satir, bulunan {sonuc['bulunan']}, yok {sonuc['yok']}", ip=ip())
    return render_template("liste_sonuc.html", token=token, veri=veri, sonuc=sonuc, bulunan=bulunan[:2000], yok=yok[:2000], proje_sayac=sorted(proje_sayac.items()),
                           basliklar=_liste_basliklar(veri), f=veri["filtre"])


@bp.route("/liste-kontrol.xlsx")
@giris_gerekli
def liste_kontrol_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    veri = liste_kontrol.yukle(ayar()["ONBELLEK"], request.args.get("token", ""), g.k["kullanici_adi"])
    if not veri or not veri.get("esleme"):
        return _liste_bulunamadi()
    kayitlar, f = _liste_kayitlar(veri["filtre"], izinli_projeler())
    sonuc = liste_kontrol.eslestir(veri["tablolar"][veri["sayfa"]], veri["baslik"], veri["esleme"], kayitlar)
    dosya_basliklar = _liste_basliklar(veri)
    wb = openpyxl.Workbook()

    def sayfa(ad, basliklar, genislikler):
        ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet()
        ws.title = ad
        xlsx_baslik(ws, basliklar, genislikler)
        return ws

    ws = sayfa("Bulunamayanlar", ["Satir no"] + dosya_basliklar, [9] + [24] * len(dosya_basliklar))
    for s in sonuc["satirlar"]:
        if not s["eslesen"]:
            xlsx_satir(ws, [s["no"]] + s["degerler"])
    ws = sayfa("Bulunanlar", ["Satir no"] + dosya_basliklar + ["Eslesme sayisi"] + [b for b, _, _ in LISTE_KAYIT_SUTUNLAR], [9] + [24] * len(dosya_basliklar) + [10] + [gen for _, _, gen in LISTE_KAYIT_SUTUNLAR])
    sutun0 = len(dosya_basliklar) + 2
    for s in sonuc["satirlar"]:
        for i in s["eslesen"][:10]:
            d = kayitlar[i]
            degerler = [kayit_durum_metni(d) if k == "_durum" else ("EVET" if d["kt"] else "HAYIR") if k == "kt" else (d.get(k) if d.get(k) is not None else "") for _, k, _ in LISTE_KAYIT_SUTUNLAR]
            n = xlsx_satir(ws, [s["no"]] + s["degerler"] + [len(s["eslesen"])] + degerler)
            if (d.get("link") or "").startswith(("http://", "https://")):
                ws.cell(row=n, column=sutun0 + 8).hyperlink = d["link"]
                ws.cell(row=n, column=sutun0 + 8).font = Font(color="0563C1", underline="single")
            renk_hex = renk(d["proje"]).lstrip("#")
            if len(renk_hex) == 6:
                ws.cell(row=n, column=sutun0 + 1).fill = PatternFill("solid", fgColor=renk_hex)
                ws.cell(row=n, column=sutun0 + 1).font = Font(bold=True, color=metin_rengi("#" + renk_hex).lstrip("#"))
    oz = sayfa("Ozet", ["Alan", "Deger"], [22, 90])
    proje_sayac = {}
    for s in sonuc["satirlar"]:
        for p in {kayitlar[i]["proje"] for i in s["eslesen"]}:
            proje_sayac[p] = proje_sayac.get(p, 0) + 1
    ozet = [["Toplam satir", sonuc["bulunan"] + sonuc["yok"]], ["Bulunan", sonuc["bulunan"]], ["Bulunamayan", sonuc["yok"]], ["Birden fazla eslesen", sonuc["coklu"]], ["Bos satir (atlandi)", sonuc["bos"]]]
    ozet += [[f"Bulunan ({p})", n] for p, n in sorted(proje_sayac.items())]
    ozet += [["Dosya", veri.get("dosya", "") + (f" / sayfa: {veri['sayfa']}" if veri["sayfa"] else "") + (" (ilk satir baslik)" if veri["baslik"] else "")],
             ["Esleme", " VE ".join(f"{b} -> {' / '.join(liste_kontrol.ALAN_ADLARI[a] for a in e['alanlar'])}" for b, e in zip(dosya_basliklar, veri["esleme"]))],
             ["Filtre", f"projeler: {', '.join(f.get('proje', []))}; durum: {f.get('durum')}; tarih: {f.get('tarih1') or '-'} / {f.get('tarih2') or '-'}; kategori: {f.get('kategori') or '-'}"],
             ["Olusturan", f"{g.k['kullanici_adi']} - {simdi()}"]]
    for satir in ozet:
        xlsx_satir(oz, satir)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f"liste_kontrol_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx", mimetype=XLSX_MIME)

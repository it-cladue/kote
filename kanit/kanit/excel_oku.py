# -*- coding: utf-8 -*-
"""Eski Windows surumunun Excel kitabini (telegrambot_gyazo.xlsx) salt okunur okur.

telegrampanel/app/excel_oku.py'deki okuma kurallarinin kopyasi (calisma zamaninda oradan hicbir sey ice aktarilmaz):
  - Sayfa basina proje; sutunlar A personel, B uye ID, C ana uye ID, D tur, E link, F kategori, G tarih/saat, H KT, I gorsel2, J aciklama.
  - HATALI ve ONAYLILAR sayfalari atlanir. Elle acilmis 'cennet kt' sayfasi farkli dizilimdedir (A uye ID, B onay, C link, D aciklama).
  - OF sayfasinda H sutunu 1 milyon satira kadar FALSE ile doludur: yalnizca A..G'ye bakilir, 30 ardisik bos satirda durulur
    (botun Get-ExcelLastDataRow kurali). Sayfanin tamami asla dolasilmaz.
  - anahtar_uret: kayit anahtari (sayfa|tur:kimlik, ayni link birden fazla satirdaysa parmak izi eklenir). Eski panel.db'deki
    'anahtar' sutunuyla BIREBIR ayni algoritma; aktarici onay bilgisini bununla eslestirir.
"""
import os, re, hashlib
from datetime import datetime
import openpyxl

VARSAYILAN_HARITA = {"personel": "A", "uye": "B", "ana": "C", "tip": "D", "link": "E", "kategori": "F", "tarih": "G", "kt": "H", "gorsel2": "I", "aciklama": "J"}
# Elle acilmis "cennet kt" sayfasi: A uye ID, B onay kutusu, C link, D aciklama. Diger sutunlar yok (None).
SAYFA_HARITALARI = {"cennet kt": {"personel": None, "uye": "A", "ana": None, "tip": None, "link": "C", "kategori": None, "tarih": None, "kt": "B", "gorsel2": None, "aciklama": "D"}}
# sayfa adi (kucuk) -> proje kodu; listede olmayan sayfalar kendi adiyla (BUYUK) proje olur
SAYFA_PROJE = {"cennet kt": "CENNET"}
HARIC_SAYFALAR = {"HATALI", "ONAYLILAR"}
BOS_SINIR = 30
TARIH_BICIMLERI = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y")


def _kol(harf):
    if harf in (None, "", "-"):
        return None
    return openpyxl.utils.column_index_from_string(harf) - 1


def _metin(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    return str(v).strip()


def _tarih(v):
    """Excel G hucresi -> (iso|None, ham). Alti metin bicimi + gercek tarih hucresi."""
    if v is None or v == "":
        return None, ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S"), v.strftime("%Y-%m-%d %H:%M:%S")
    s = str(v).strip()
    for f in TARIH_BICIMLERI:
        try:
            return datetime.strptime(s, f).strftime("%Y-%m-%d %H:%M:%S"), s
        except ValueError:
            pass
    return None, s


def _kt(v):
    if isinstance(v, bool):
        return 1 if v else 0
    s = _metin(v).lower()
    return 1 if s in ("true", "1", "x", "evet", "ok", "kt", "\u2713") else 0


# yalnizca gyazo.com alan adinda (https://evil/x/gyazo.com/... gibi sahte yollar gyazo sayilmaz)
GYAZO_ID = re.compile(r"(?:^|://|\s|[(\[<\"'])(?:[a-z0-9-]+\.)*gyazo\.com/(?:[a-z0-9_-]+/)*([0-9a-f]{32})", re.I)


def link_turu_ve_kimlik(link):
    """-> (tur, kimlik)  tur: gyazo|onedrive|snappark|teamzone|dosya|diger|metin|bos  (eski panelle ayni; anahtar icin)"""
    s = (link or "").strip()
    if not s:
        return "bos", ""
    m = GYAZO_ID.search(s)
    if m and "teamzone" not in s.lower():
        return "gyazo", m.group(1).lower()
    low = s.lower()
    if "sharepoint.com" in low or "onedrive.aspx" in low or "1drv.ms" in low:
        from urllib.parse import urlparse, parse_qs, unquote
        try:
            q = parse_qs(urlparse(s).query)
            yol = unquote(q.get("id", [""])[0])
            ad = yol.rsplit("/", 1)[-1]
            return "onedrive", ad.lower() or low
        except Exception:
            return "onedrive", low
    if "snappark" in low:
        return "snappark", low.rstrip("/").rsplit("/", 1)[-1]
    if "teamzone.gyazo" in low:
        return "teamzone", low.rstrip("/").rsplit("/", 1)[-1]
    if re.match(r"^[a-z]:\\", low) or low.startswith("\\\\"):
        return "dosya", os.path.basename(s.replace("\\", "/")).lower()
    if low.startswith("http"):
        return "diger", low
    return "metin", ""


PARMAK_ALANLARI = ("personel", "uye", "ana", "tip", "kategori", "tarih_ham", "link")


def parmak_ozeti(satir_degerleri, uzunluk=16):
    parmak = "|".join(str(satir_degerleri.get(k, "") or "") for k in PARMAK_ALANLARI)
    return hashlib.sha1(parmak.encode("utf-8")).hexdigest()[:uzunluk]


def anahtar_uret(sayfa, satir_degerleri, link, sayac, link_sayisi=None):
    """Kayit anahtari. Link tekse: sayfa|tur:kimlik. Ayni link ayni sayfada birden fazla satirdaysa (ayni ekran goruntusu
    iki uye icin) satirin parmak izi eklenir ki bir satir silinince digerinin kimligi degismesin. Linksiz satir: parmak izi."""
    tur, kimlik = link_turu_ve_kimlik(link)
    if tur in ("gyazo", "onedrive", "snappark", "teamzone", "dosya", "diger") and kimlik:
        temel = f"{sayfa}|{tur}:{kimlik}"
        if link_sayisi and link_sayisi.get(f"{tur}:{kimlik}", 0) > 1:
            temel += "|h:" + parmak_ozeti(satir_degerleri, 8)
    else:
        temel = f"{sayfa}|h:" + parmak_ozeti(satir_degerleri)
    n = sayac.get(temel, 0) + 1
    sayac[temel] = n
    return temel if n == 1 else f"{temel}#{n}", tur


def sayfa_projesi(sayfa_adi):
    return SAYFA_PROJE.get(sayfa_adi.strip().lower()) or sayfa_adi.strip().upper()


def oku(xlsx_yol, sutun_haritasi=None, haric=HARIC_SAYFALAR):
    """-> {sayfa_adi: [satir dict]} ; satir: {satir, personel, uye, ana, tip, link, kategori, tarih, tarih_ham, kt, gorsel2, aciklama,
    ek_not, anahtar, link_turu}. read_only + data_only: 6 MB kitap saniyeler icinde okunur."""
    sutun_haritasi = sutun_haritasi or {}
    wb = openpyxl.load_workbook(xlsx_yol, read_only=True, data_only=True)
    sonuc = {}
    try:
        for ws in wb.worksheets:
            ad = ws.title
            if ad.upper() in {h.upper() for h in haric}:
                continue
            harita = dict(VARSAYILAN_HARITA)
            harita.update(SAYFA_HARITALARI.get(ad.lower(), {}))
            harita.update(sutun_haritasi.get(ad, {}))
            idx = {k: _kol(v) for k, v in harita.items()}
            # "kt" (H) bilerek disarida: OF sayfasinda H sutunu 1 milyon satira kadar FALSE ile dolu (botun A..G kurali)
            ana_kolonlar = [idx[k] for k in ("personel", "uye", "ana", "tip", "link", "kategori", "tarih") if idx.get(k) is not None]
            satirlar, bos_ust_uste, sayac, link_sayisi = [], 0, {}, {}
            kullanilan = {v for v in idx.values() if v is not None}
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if i == 1:
                    continue  # baslik
                if all((row[j] if j < len(row) else None) in (None, "") for j in ana_kolonlar):
                    bos_ust_uste += 1
                    if bos_ust_uste >= BOS_SINIR:
                        break
                    continue
                bos_ust_uste = 0

                def hucre(k):
                    j = idx.get(k)
                    return row[j] if j is not None and j < len(row) else None
                tarih, tarih_ham = _tarih(hucre("tarih"))
                d = {
                    "satir": i,
                    "personel": _metin(hucre("personel")), "uye": _metin(hucre("uye")), "ana": _metin(hucre("ana")),
                    "tip": _metin(hucre("tip")), "link": _metin(hucre("link")), "kategori": _metin(hucre("kategori")),
                    "tarih": tarih, "tarih_ham": tarih_ham, "kt": _kt(hucre("kt")),
                    "gorsel2": _metin(hucre("gorsel2")), "aciklama": _metin(hucre("aciklama")),
                }
                ek = {}
                for j, v in enumerate(row):
                    if j not in kullanilan and v not in (None, ""):
                        ek[openpyxl.utils.get_column_letter(j + 1)] = _metin(v)
                d["ek_not"] = "; ".join(f"{k}={v}" for k, v in ek.items()) if ek else ""
                tur, kimlik = link_turu_ve_kimlik(d["link"])
                if kimlik:
                    link_sayisi[f"{tur}:{kimlik}"] = link_sayisi.get(f"{tur}:{kimlik}", 0) + 1
                satirlar.append(d)
            # anahtarlar ikinci geciste: ayni link kac satirda var bilinmeli
            for d in satirlar:
                d["anahtar"], d["link_turu"] = anahtar_uret(ad, d, d["link"], sayac, link_sayisi)
            sonuc[ad] = satirlar
    finally:
        wb.close()
    return sonuc

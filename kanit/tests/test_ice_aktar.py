# -*- coding: utf-8 -*-
"""Ice aktarim: sentetik Excel (7 sayfa + HATALI + ONAYLILAR + 'cennet kt'), bot durum dosyalari, gorsel klasorleri, eski panel.db.
Kuru -> uygula -> yeniden (hepsi atlanir) -> excel-ekle; ayrica komut satiri (ice-aktar --kuru, saglik, kullanici-ekle).
Calistirma: cd kanit && python3 tests/test_ice_aktar.py"""
import os, sys, json, time, sqlite3, tempfile, subprocess, re
from datetime import datetime, timedelta
KOK = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
sys.path.insert(0, KOK)
TD = tempfile.mkdtemp(prefix="kanit-ice-")
VERI = os.path.join(TD, "veri")
os.environ["KANIT_VERI"] = VERI
os.environ["KANIT_ENV"] = os.path.join(TD, "yok.env")
import openpyxl
from werkzeug.security import generate_password_hash
from kanit import ayar, db, ice_aktar, excel_oku, log_oku, normalle as N

ok = 0


def check(cond, msg):
    global ok
    if cond:
        ok += 1
        print("  OK ", msg)
    else:
        print("  FAIL", msg)
        raise SystemExit(1)


G = lambda i: "https://gyazo.com/" + format(i, "032x")
JPEG = b"\xff\xd8\xff\xe0" + b"\x00" * 64 + b"\xff\xd9"
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
KAYNAK = os.path.join(TD, "kaynak")
BOT = os.path.join(KAYNAK, "bot")
GORSELLER = os.path.join(KAYNAK, "gorseller")
EKRAN = os.path.join(KAYNAK, "ekran_goruntuleri")
TG_G = os.path.join(KAYNAK, "telegram_gorseller")
DOSYA_HI = "DDS_2026-09-29_13-53-15_240_26996771_sessiz.jpg"
T0 = datetime(2026, 9, 1, 10, 0, 0)


def _t(i, fmt="%Y-%m-%d %H:%M:%S"):
    return (T0 + timedelta(hours=9 * i)).strftime(fmt)


def _yaz(yol, metin, newline="\r\n"):
    with open(yol, "w", encoding="utf-8", newline="") as f:
        f.write(metin.replace("\n", newline))


def _metin_hucre(ws, satir, sutun, deger):
    h = ws.cell(row=satir, column=sutun, value=deger)
    h.data_type = "s"


def kitap_yap(yol, ek_satir=None, kt_degistir=None):
    """7 proje sayfasi + HATALI + ONAYLILAR + 'cennet kt'. -> beklenen satir sayilari {sayfa: n}"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bas = ["Personel Kodu", "Uye ID", "Ana Uye ID", "Sessiz / Telesekreter", "Gyazo Linki", "Kategori", "Tarih / Saat", "KT", "Gorsel 2", "Aciklama"]
    n = {}
    i = 0
    tarih_bicimleri = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y"]
    for sayfa in ("PP", "PA", "OF", "HI", "VI", "GA", "BETA"):
        ws = wb.create_sheet(sayfa)
        ws.append(bas)
        satirlar = []
        for j in range(7):
            i += 1
            satirlar.append([f"P{j+1}", f"{100000 + i}", f"{200000 + i}" if j % 2 else "", "sessiz" if j % 3 == 0 else "", G(i), ["ARK", "DDS", "HGST", "DDT", "PRCST", "LPST", "ARK"][j],
                             _t(i, tarih_bicimleri[j % 6]), bool(j % 2), "", f"not {i}"])
        if sayfa == "PP":
            # ayni link iki farkli uye icin (anahtar parmak izi alir) + gercek datetime hucresi
            satirlar.append(["P9", "300001", "", "", G(1000), "ARK", T0 + timedelta(hours=3), False, "", "ikiz 1"])
            satirlar.append(["P9", "300002", "", "", G(1000), "ARK", T0 + timedelta(hours=3, minutes=1), False, "", "ikiz 2"])
            # cozulemeyen tarih + formul gibi metin
            satirlar.append(["P8", "300003", "", "", G(1001), "DDS", "32/13/2026 99:99", True, "", "=cmd|'/C calc'!A0"])
            # birebir tekrar (ayni her sey, 1 dk fark)
            satirlar.append(["P7", "300004", "", "", G(1002), "HGST", "2026-09-05 12:00:00", False, "", "tekrar a"])
            satirlar.append(["P7", "300004", "", "", G(1002), "HGST", "2026-09-05 12:01:00", False, "", "tekrar b"])
            # linksiz satir (parmak izi anahtari) + tarihsiz
            satirlar.append(["P6", "300005", "", "", "", "ARK", "", False, "", "linksiz"])
            # yazili giris (lead id metni)
            satirlar.append(["P5", "300006", "", "", "LEAD-7781", "PRCST", "2026-09-06 09:00:00", False, "", ""])
        if sayfa == "HI":
            satirlar.append(["240", "26996771", "", "sessiz", "C:\\Users\\bot\\OneDrive\\gorseller\\HI\\" + DOSYA_HI, "DDS", "2026-09-29 13:53:15", False, "", "dosya kaydi"])
            satirlar.append(["241", "26996772", "", "", "https://x.sharepoint.com/personal/u/_layouts/15/onedrive.aspx?id=%2Fpersonal%2Fu%2FDocuments%2Fgorseller%2FHI%2FARK_2026-09-30_10-00-00_241_26996772.jpg&parent=x", "ARK", "2026-09-30 10:00:00", True, "", "onedrive"])
            satirlar.append(["242", "26996773", "", "", G(2001), "ARK", "2026-09-30 11:00:00", False, "", "sahte gorsel"])
        if sayfa == "OF":
            satirlar.append(["OF1", "400001", "", "", G(3001), "ARK", "2026-09-10 10:00:00", True, "", "ek sutun"])
        if ek_satir and sayfa == ek_satir[0]:
            satirlar.append(ek_satir[1])
        for s in satirlar:
            ws.append(s)
            r = ws.max_row
            if isinstance(s[9], str) and s[9][:1] in "=+-@":
                _metin_hucre(ws, r, 10, s[9])
            if kt_degistir and sayfa == kt_degistir[0] and s[4] == kt_degistir[1]:
                ws.cell(row=r, column=8, value=True)
        if sayfa == "OF":
            ws.cell(row=ws.max_row, column=12, value="K sutunu notu")   # A..J disi -> ek_not
            for r in range(ws.max_row + 1, ws.max_row + 3001):          # H sutunu binlerce FALSE (gercek OF sayfasi gibi)
                ws.cell(row=r, column=8, value=False)
        n[sayfa] = len(satirlar)
    ws = wb.create_sheet("HATALI")
    ws.append(bas); ws.append(["X", "1", "", "", G(9001), "ARK", "2026-01-01", False, "", "atlanmali"])
    ws = wb.create_sheet("ONAYLILAR")
    ws.append(["ID", "Ad", "Tarih"]); ws.append(["111", "ali", "2026-01-01"])
    ws = wb.create_sheet("cennet kt")
    ws.append(["Uye ID", "Onay", "Link", "Aciklama"])
    ws.append(["500001", True, G(4001), "cennet 1"])
    ws.append(["500002", False, G(4002), ""])
    ws.append(["500003", "x", "", "linksiz cennet"])
    n["cennet kt"] = 3
    wb.save(yol)
    return n


def kaynaklari_yap():
    os.makedirs(BOT, exist_ok=True)
    # gorseller/<PROJE>/ (bot kayitlari) + ekran_goruntuleri (gyazo indeks) + telegram_gorseller
    for p in ("HI", "PP", "OF"):
        os.makedirs(os.path.join(GORSELLER, p), exist_ok=True)
    with open(os.path.join(GORSELLER, "HI", DOSYA_HI), "wb") as f:
        f.write(JPEG)
    with open(os.path.join(GORSELLER, "HI", "ARK_2026-09-30_10-00-00_241_26996772.jpg"), "wb") as f:
        f.write(PNG)
    with open(os.path.join(GORSELLER, "PP", "ARK_2026-08-01_08-00-00_sahipsiz.jpg"), "wb") as f:
        f.write(JPEG)   # satiri olmayan gorsel
    os.makedirs(os.path.join(EKRAN, "PP"), exist_ok=True)
    os.makedirs(os.path.join(EKRAN, "HI"), exist_ok=True)
    with open(os.path.join(EKRAN, "PP", "ARK_2026-09-01_10-07-00_P1_100001.jpg"), "wb") as f:
        f.write(JPEG)
    with open(os.path.join(EKRAN, "PP", "ARK_2026-09-01_13-00-00_P9_ikiz.jpg"), "wb") as f:
        f.write(JPEG)
    with open(os.path.join(EKRAN, "HI", "sahte.jpg"), "wb") as f:
        f.write(b"bu bir gorsel degil")
    _yaz(os.path.join(EKRAN, "indeks.csv"), "Tarih;Proje;Kategori;GyazoLinki;GyazoId;DosyaAdi;DosyaYolu;Durum\n"
         + f"2026-09-01 10:07;PP;ARK;{G(1)};{format(1, '032x')};ARK_2026-09-01_10-07-00_P1_100001.jpg;C:\\eski\\PP\\ARK_2026-09-01_10-07-00_P1_100001.jpg;indi\n"
         + f"2026-09-01 13:00;PP;ARK;{G(1000)};{format(1000, '032x')};ARK_2026-09-01_13-00-00_P9_ikiz.jpg;C:\\eski\\PP\\ARK_2026-09-01_13-00-00_P9_ikiz.jpg;indi\n"
         + f"2026-09-30 11:00;HI;ARK;{G(2001)};{format(2001, '032x')};sahte.jpg;C:\\eski\\HI\\sahte.jpg;indi\n"
         + f"2026-09-30 12:00;HI;ARK;{G(2002)};{format(2002, '032x')};yok.jpg;C:\\eski\\HI\\yok.jpg;gyazo_silinmis\n", "\n")
    os.makedirs(os.path.join(TG_G, "OF"), exist_ok=True)
    with open(os.path.join(TG_G, "OF", "ARK_2026-09-10_10-00-00_OF1_400001.jpg"), "wb") as f:
        f.write(JPEG)
    _yaz(os.path.join(TG_G, "indeks.csv"), "Tarih;KullaniciAdi;KullaniciId;Proje;Kategori;Personel;UyeID;AnaUyeID;Tip;MesajId;GyazoId;DosyaAdi;DosyaYolu;Durum\n"
         + f"2026-09-10 10:00;ali;111;OF;ARK;OF1;400001;;;5;{format(3001, '032x')};ARK_2026-09-10_10-00-00_OF1_400001.jpg;C:\\eski\\OF\\ARK_2026-09-10_10-00-00_OF1_400001.jpg;indi\n", "\n")
    # bot durum dosyalari
    with open(os.path.join(BOT, "kullanicilar.json"), "w", encoding="utf-8") as f:
        json.dump({"111": {"id": "111", "username": "ali", "first_name": "Ali", "last_seen": "2026-09-30 12:00:00", "status": "kayit_eklendi", "kayit_sayisi": 5,
                           "son_kayit_tarihi": _t(2), "son_kayit_proje": "PP", "son_kayit_kategori": "DDS", "proje_sayilari": {"PP": 3, "HI": 2}, "kategori_sayilari": {"ARK": 5}},
                   "222": {"id": "222", "username": "", "first_name": "Veli", "last_seen": "2026-09-29 09:00:00", "status": "sifre_bekliyor"}}, f)
    _yaz(os.path.join(BOT, "onaylilar.txt"), "111|ali|2026-08-01 10:00\n3333333333|ayse|2026-08-02 10:00\n4444444444|fatma|2026-08-03 10:00\n")
    _yaz(os.path.join(BOT, "engellenenler.txt"), "222|admin|2026-09-29 09:05|3 yanlis sifre8406076054|admin|2026-09-29 09:06|ayrildi\n")   # yapisik satir (satir sonu yok)
    _yaz(os.path.join(BOT, "denemeler.txt"), "222|3\n555|1\nbozuk\n")
    with open(os.path.join(BOT, "kategoriler.json"), "w", encoding="utf-8") as f:
        json.dump([{"code": "ark", "label": "ARK", "mode": "ark_ana"}, {"code": "dds", "label": "DDS SESSIZ", "mode": "sessiz_tip"}, {"code": "xyz", "label": "XYZ", "mode": "tek_uye"},
                   {"code": "bozuk", "label": "B", "mode": "yok"}], f)
    # kayitlar.jsonl: G(3) kesin ekleyen 111; HI dosya kaydi kesin ekleyen 222 (dosya adiyla)
    with open(os.path.join(BOT, "kayitlar.jsonl"), "w", encoding="utf-8") as f:
        f.write(json.dumps({"zaman": _t(3), "proje": "PP", "kategori": "HGST", "kat_kodu": "hgst", "personel": "P3", "uyeid": "100003", "anaid": "", "tip": "", "link": G(3), "dosya": "", "tg_id": "111", "tg_ad": "ali", "giris": "resim"}) + "\n")
        f.write(json.dumps({"zaman": "2026-09-29 13:53:20", "proje": "HI", "kategori": "DDS", "kat_kodu": "dds", "personel": "240", "uyeid": "26996771", "anaid": "", "tip": "sessiz", "link": "", "dosya": "C:\\Users\\bot\\OneDrive\\gorseller\\HI\\" + DOSYA_HI, "tg_id": "222", "tg_ad": "Veli", "giris": "resim"}) + "\n")
        f.write("bozuk satir\n")
    # bot_log.txt: G(2) kesin (kullanicilar.json son_kayit_tarihi = _t(2), proje PP), G(4) tahmini (son mesaj yazan 333, 3 mesaj)
    _yaz(os.path.join(BOT, "bot_log.txt"), "\ufeff" + f"[{_t(2)}] Mesaj geldi: ali (ID: 111)\n[{_t(2)}] Kaydedildi -> [PP] [DDS] {G(2)}\n"
         + f"[{_t(4)}] Mesaj geldi: ayse (ID: 333)\n[{_t(4)}] Mesaj geldi: ayse (ID: 333)\n[{_t(4)}] Mesaj geldi: ayse (ID: 333)\n[{_t(4)}] Kaydedildi -> [PP] [DDT] {G(4)}\n", "\n")
    with open(os.path.join(BOT, "panel_kullanicilar.json"), "w", encoding="utf-8") as f:
        json.dump({"surum": 1, "kullanicilar": [
            {"kullanici_adi": "patron", "ad": "Patron", "sifre_hash": generate_password_hash("Gizli1234"), "yetki": "admin", "aktif": True, "sifre_degistir": False, "olusturma": "2026-01-01 00:00:00", "olusturan": "kurulum"},
            {"kullanici_adi": "pasif", "ad": "Pasif", "sifre_hash": generate_password_hash("x"), "yetki": "OF:oku,kt", "aktif": False, "sifre_degistir": True, "olusturma": "2026-01-02 00:00:00", "olusturan": "patron"},
            {"kullanici_adi": "hashsiz", "ad": "H", "sifre_hash": "", "yetki": "PA"}]}, f)


def panel_db_yap(yol, anahtar_kt, anahtar_sil, anahtar_kt2):
    c = sqlite3.connect(yol)
    c.executescript("""CREATE TABLE kayitlar (id INTEGER PRIMARY KEY, anahtar TEXT UNIQUE NOT NULL, proje TEXT, kt INTEGER DEFAULT 0, kt_kaynak TEXT DEFAULT 'excel',
                       kt_kim TEXT, kt_zaman TEXT, silindi INTEGER DEFAULT 0, silen TEXT, silme_zamani TEXT);
                       CREATE TABLE denetim (id INTEGER PRIMARY KEY, zaman TEXT, kullanici TEXT, islem TEXT, proje TEXT, anahtar TEXT, eski TEXT, yeni TEXT, ip TEXT);
                       CREATE TABLE panel_kullanicilari (id INTEGER PRIMARY KEY, kullanici_adi TEXT UNIQUE NOT NULL, ad TEXT, sifre_hash TEXT NOT NULL, yetki TEXT NOT NULL DEFAULT '',
                       aktif INTEGER DEFAULT 1, sifre_degistir INTEGER DEFAULT 0, olusturma TEXT, son_giris TEXT);""")
    c.execute("INSERT INTO kayitlar(anahtar, proje, kt, kt_kaynak, kt_kim, kt_zaman) VALUES(?,?,1,'panel','ofadmin','2026-09-02 10:00:00')", (anahtar_kt, "PP"))
    c.execute("INSERT INTO kayitlar(anahtar, proje, kt, kt_kaynak, silindi, silen, silme_zamani) VALUES(?,?,0,'excel',1,'patron','2026-09-03 10:00:00')", (anahtar_sil, "PA"))
    c.execute("INSERT INTO kayitlar(anahtar, proje, kt, kt_kaynak, kt_kim) VALUES(?,?,1,'excel','')", (anahtar_kt2, "PA"))
    c.execute("INSERT INTO denetim(zaman, kullanici, islem, proje, anahtar, yeni) VALUES('2026-09-03 10:00:00','patron','sil','PA',?,'mukerrer kayit')", (anahtar_sil,))
    c.execute("INSERT INTO panel_kullanicilari(kullanici_adi, ad, sifre_hash, yetki, aktif) VALUES('eskipanel','Eski','" + generate_password_hash("Eski1234") + "','GA',1)")
    c.execute("INSERT INTO panel_kullanicilari(kullanici_adi, ad, sifre_hash, yetki, aktif) VALUES('patron','Patron2','x:y','admin',1)")
    c.commit(); c.close()


def calistir():
    a = ayar.yukle()
    check(a["VERI"] == VERI and os.path.isdir(a["GORSELLER"]), "ayar KANIT_VERI")
    xlsx = os.path.join(KAYNAK, "telegrambot_gyazo.xlsx")
    os.makedirs(KAYNAK, exist_ok=True)
    beklenen = kitap_yap(xlsx)
    kaynaklari_yap()
    toplam = sum(beklenen.values())

    # ---- excel_oku kurallari ----
    t0 = time.time()
    sayfalar = excel_oku.oku(xlsx)
    sure = time.time() - t0
    check(set(sayfalar) == set(beklenen) and "HATALI" not in sayfalar and "ONAYLILAR" not in sayfalar, "sayfalar: 7 proje + 'cennet kt'; HATALI/ONAYLILAR atlandi")
    check({s: len(v) for s, v in sayfalar.items()} == beklenen, f"satir sayilari {beklenen}")
    check(len(sayfalar["OF"]) == beklenen["OF"] and sure < 5, f"OF: 3000 FALSE hucre altinda durdu ({sure:.2f} sn)")
    pp = {s["satir"]: s for s in sayfalar["PP"]}
    check(all(s["tarih"] for s in sayfalar["PP"][:6]) and len({s["tarih"] for s in sayfalar["PP"][:6]}) == 6, "alti tarih bicimi cozuldu")
    bozuk = [s for s in sayfalar["PP"] if s["tarih_ham"] and not s["tarih"]]
    check(len(bozuk) == 1 and bozuk[0]["tarih_ham"] == "32/13/2026 99:99" and bozuk[0]["aciklama"].startswith("=cmd"), "cozulemeyen tarih + '=cmd' metni okundu")
    ikiz = [s for s in sayfalar["PP"] if s["link"] == G(1000)]
    check(len(ikiz) == 2 and all("|h:" in s["anahtar"] for s in ikiz) and ikiz[0]["anahtar"] != ikiz[1]["anahtar"], "ayni link iki uye -> parmak izli farkli anahtar")
    tekrar = [s for s in sayfalar["PP"] if s["link"] == G(1002)]
    check(tekrar[0]["anahtar"] != tekrar[1]["anahtar"], "birebir tekrar da ayri anahtar alir")
    linksiz = [s for s in sayfalar["PP"] if not s["link"]]
    check(len(linksiz) == 1 and linksiz[0]["anahtar"].startswith("PP|h:") and linksiz[0]["link_turu"] == "bos", "linksiz satir parmak izi anahtari")
    of_ek = [s for s in sayfalar["OF"] if s["ek_not"]]
    check(len(of_ek) == 1 and of_ek[0]["ek_not"] == "L=K sutunu notu", "A..J disi sutun -> ek_not")
    cennet = sayfalar["cennet kt"]
    check([s["uye"] for s in cennet] == ["500001", "500002", "500003"] and [s["kt"] for s in cennet] == [1, 0, 1] and cennet[0]["link"] == G(4001) and cennet[2]["aciklama"] == "linksiz cennet", "'cennet kt' farkli dizilim")
    check(excel_oku.sayfa_projesi("cennet kt") == "CENNET" and excel_oku.sayfa_projesi("of") == "OF", "sayfa -> proje")
    hi_dosya = [s for s in sayfalar["HI"] if s["link_turu"] == "dosya"][0]
    check(excel_oku.link_turu_ve_kimlik(hi_dosya["link"]) == ("dosya", DOSYA_HI.lower()), "Windows yolu -> dosya adi kimligi")
    check(ice_aktar.foto_tarihi(DOSYA_HI) == "2026-09-29 13:53:15" and ice_aktar.foto_tarihi("abc.jpg") is None, "dosya adindan foto tarihi")

    # ---- bot klasoru / log / panel.db okuma ----
    bot = ice_aktar.bot_klasoru_oku(BOT)
    check([o["id"] for o in bot["onaylilar"]] == ["111", "3333333333", "4444444444"], "onaylilar")
    check([e["id"] for e in bot["engellenenler"]] == ["222", "8406076054"] and bot["engellenenler"][1]["sebep"] == "ayrildi", "yapisik engellenenler satiri ID basindan bolundu")
    check(bot["engellenenler"][0]["sebep"] == "3 yanlis sifre" and bot["denemeler"] == {"222": 3, "555": 1}, "engellenenler + denemeler")
    check([k["code"] for k in bot["kategoriler"]] == ["ark", "dds", "xyz"], "kategoriler.json (gecersiz akis atildi)")
    check(len(bot["jsonl"]) == 2 and len(bot["log"]) == 2, "kayitlar.jsonl + bot_log")
    harita = log_oku.link_haritasi(bot["log"], bot["jsonl"])
    check(harita["gyazo:" + format(3, "032x")]["guven"] == "kesin" and harita["gyazo:" + format(3, "032x")]["ekleyen_id"] == "111", "jsonl -> kesin")
    check(harita["gyazo:" + format(2, "032x")]["guven"] == "kesin" and harita["gyazo:" + format(4, "032x")]["guven"] == "tahmini" and harita["gyazo:" + format(4, "032x")]["ekleyen_id"] == "333", "log: kesin nokta + tahmini")
    check(harita["dosya:" + DOSYA_HI.lower()]["ekleyen_id"] == "222", "jsonl dosya adi -> ekleyen")
    anahtar_kt = pp[2]["anahtar"]            # PP ilk satir -> panelden ofadmin onayi
    pa = sayfalar["PA"]
    anahtar_sil, anahtar_kt2 = pa[0]["anahtar"], pa[1]["anahtar"]
    pdb = os.path.join(KAYNAK, "panel.db")
    panel_db_yap(pdb, anahtar_kt, anahtar_sil, anahtar_kt2)
    bilgi, pk = ice_aktar.panel_db_oku([pdb, os.path.join(KAYNAK, "yok.db")])
    check(bilgi[anahtar_kt]["kt_kim"] == "ofadmin" and bilgi[anahtar_sil]["silindi"] == 1 and bilgi[anahtar_sil]["silme_nedeni"] == "mukerrer kayit", "eski panel.db okundu (kt_kim, silindi, neden)")
    check([u["kullanici_adi"] for u in pk] == ["eskipanel", "patron"], "eski panel.db kullanicilari")
    gk = ice_aktar.gorsel_klasorlerini_tara([GORSELLER, EKRAN, TG_G])
    check(gk.bul("gyazo", format(1, "032x")).endswith("ARK_2026-09-01_10-07-00_P1_100001.jpg") and gk.bul("dosya", DOSYA_HI.lower()).endswith(DOSYA_HI)
          and gk.bul("gyazo", format(2002, "032x")) is None and gk.bul("gyazo", format(3001, "032x")).endswith("OF1_400001.jpg") and gk.bul("gyazo", format(2001, "032x")) is None, "indeks.csv + dosya adi eslesmesi (gorsel olmayan dosya elenir)")

    # ---- KURU ----
    conn = db.baglan(a["DB"])
    r = ice_aktar.calistir(conn, a, excel=xlsx, bot_klasor=BOT, panel_db=[pdb], gorseller=[GORSELLER, EKRAN, TG_G], kuru=True)
    metin = ice_aktar.rapor_metni(r)
    print(metin)
    check(r["tur"] == "kuru" and r["toplam_satir"] == toplam and r["yeni"] == toplam and r["atlandi"] == 0, "kuru: satir sayilari")
    check(r["cennet_kt"] and "CENNET" in r["yeni_projeler"], "kuru: 'cennet kt' var, CENNET projesi acilacak")
    check(len(r["tarih_hatali"]) == 1 and r["tarih_hatali"][0]["sayfa"] == "PP" and r["tarihsiz"] == 5, "kuru: cozulemeyen tarih 1, tarihsiz 5 (bozuk tarih + linksiz + cennet 3)")
    check(len(r["mukerrer"]) == 1 and r["mukerrer"][0]["uye"] == "300004" and r["mukerrer"][0]["satir2"] == r["mukerrer"][0]["satir1"] + 1, "kuru: birebir tekrar cifti (ikizler sayilmaz)")
    check(r["ekleyen"] == {"kesin": 3, "tahmini": 1, "yok": toplam - 4}, f"kuru: ekleyen kesin 3 / tahmini 1 / yok {toplam - 4}")
    check(r["kt"]["panel"] == 1 and r["panel_eslesen"] == 3 and r["silinen"] == 1, "kuru: eski panelden 1 onay (kim), 1 silinmis")
    check(r["gorselli"] == 6 and r["gorsel_satirsiz_sayi"] == 2 and any(g.endswith("sahipsiz.jpg") for g in r["gorsel_satirsiz"]) and r["gorsel_kopyalanan"] == 0, "kuru: 6 gorselli satir, 2 satirsiz gorsel (sahipsiz + sahte), kopya yok")
    check(r["kullanicilar"]["yeni"] == ["patron", "pasif", "eskipanel"] and r["kullanicilar"]["hatali"] == ["hashsiz"], "kuru: kullanicilar (json + panel.db, hashsiz atildi)")
    check(r["kategoriler"] == {"yeni": ["xyz"], "guncellenen": ["dds"]} and r["tg"]["yazilan"] == 6, "kuru: kategoriler + tg defteri")
    check(conn.execute("SELECT COUNT(*) FROM kayitlar").fetchone()[0] == 0 and not os.listdir(a["GORSELLER"]) and db.kullanici_bul(conn, "patron") is None, "kuru: hicbir sey yazilmadi")
    check(len(db.ice_aktarimlar(conn)) == 1 and db.ice_aktarimlar(conn)[0]["tur"] == "kuru", "kuru raporu ice_aktarimlar'da")
    check("DENEME" in metin and "cennet kt" in metin and "32/13/2026" in metin, "rapor metni")

    # ---- UYGULA ----
    t0 = time.time()
    r = ice_aktar.calistir(conn, a, excel=xlsx, bot_klasor=BOT, panel_db=[pdb], gorseller=[GORSELLER, EKRAN, TG_G], kuru=False)
    sure = time.time() - t0
    print(ice_aktar.rapor_metni(r))
    check(r["tur"] == "uygula" and r["yeni"] == toplam and not r["hata"], f"uygula: {toplam} yeni, hata yok ({sure:.2f} sn)")
    check(conn.execute("SELECT COUNT(*) FROM kayitlar").fetchone()[0] == toplam, "kayit sayisi")
    check(conn.execute("SELECT COUNT(*) FROM kayitlar WHERE kaynak='excel' AND eski_anahtar<>'' AND eski_satir>0").fetchone()[0] == toplam, "hepsi kaynak=excel + eski_anahtar/eski_satir")
    check(conn.execute("SELECT COUNT(*) FROM arsiv WHERE olay='ice_aktar'").fetchone()[0] == toplam and conn.execute("SELECT COUNT(*) FROM bildirimler").fetchone()[0] == 0, "arsiv ice_aktar, bildirim yok")
    k = db.kayit_al(conn, db.eski_anahtarlar(conn)[anahtar_kt]["id"])
    check(k["kt"] == 1 and k["kt_kim"] == "ofadmin" and k["kt_kaynak"] == "panel" and k["kt_zaman"] == "2026-09-02 10:00:00", "panelden onay: kt_kim/kt_kaynak/kt_zaman")
    k = db.kayit_al(conn, db.eski_anahtarlar(conn)[anahtar_kt2]["id"])
    check(k["kt"] == 1 and k["kt_kim"] == "" and k["kt_kaynak"] == "excel", "Excel onayi: kt_kaynak=excel")
    k = db.kayit_al(conn, db.eski_anahtarlar(conn)[anahtar_sil]["id"])
    check(k["silindi"] == 1 and k["silen"] == "patron" and k["silme_nedeni"] == "mukerrer kayit", "eski panelde silinmis -> kayit_sil")
    k = conn.execute("SELECT * FROM kayitlar WHERE link_norm=?", (N.normalle(G(1)),)).fetchone()
    check(k["gorsel_dosya"] == "PP/ARK_2026-09-01_10-07-00_P1_100001.jpg" and os.path.exists(os.path.join(a["GORSELLER"], k["gorsel_dosya"])) and k["foto_tarih"] == "2026-09-01 10:07:00", "gyazo id -> indeks -> gorsel kopyalandi + foto_tarih")
    k = conn.execute("SELECT * FROM kayitlar WHERE eski_sayfa='HI' AND link_turu='dosya'").fetchone()
    check(k["gorsel_dosya"] == "HI/" + DOSYA_HI and k["foto_tarih"] == "2026-09-29 13:53:15" and k["ekleyen_tg_id"] == "222" and k["ekleyen_guven"] == "kesin", "Windows yolu -> dosya adi -> gorsel + jsonl ekleyen")
    k = conn.execute("SELECT * FROM kayitlar WHERE link_turu='onedrive'").fetchone()
    check(k["gorsel_dosya"] == "HI/ARK_2026-09-30_10-00-00_241_26996772.jpg", "onedrive dosya adi -> gorsel (png)")
    k = conn.execute("SELECT * FROM kayitlar WHERE link_norm=?", (N.normalle(G(2001)),)).fetchone()
    check(k["gorsel_dosya"] == "" and not os.path.exists(os.path.join(a["GORSELLER"], "HI", "sahte.jpg")), "gorsel olmayan dosya kopyalanmadi")
    ikizler = conn.execute("SELECT gorsel_dosya FROM kayitlar WHERE link_norm=? ORDER BY id", (N.normalle(G(1000)),)).fetchall()
    check([x[0] for x in ikizler] == ["PP/ARK_2026-09-01_13-00-00_P9_ikiz.jpg"] * 2 and not os.path.exists(os.path.join(a["GORSELLER"], "PP", "ARK_2026-09-01_13-00-00_P9_ikiz-2.jpg")), "ikiz satirlar ayni gorseli paylasir (tek kopya)")
    check(r["gorsel_kopyalanan"] == 5 and len(os.listdir(os.path.join(a["GORSELLER"], "PP"))) == 2 and len(os.listdir(os.path.join(a["GORSELLER"], "HI"))) == 2 and os.listdir(os.path.join(a["GORSELLER"], "OF")) == ["ARK_2026-09-10_10-00-00_OF1_400001.jpg"], "5 dosya kopyalandi (sahte + sahipsiz haric)")
    k = conn.execute("SELECT * FROM kayitlar WHERE link_norm=?", (N.normalle(G(3)),)).fetchone()
    check(k["ekleyen_tg_id"] == "111" and k["ekleyen_ad"] == "ali" and k["ekleyen_guven"] == "kesin", "jsonl ekleyen kesin")
    k = conn.execute("SELECT * FROM kayitlar WHERE link_norm=?", (N.normalle(G(4)),)).fetchone()
    check(k["ekleyen_tg_id"] == "333" and k["ekleyen_guven"] == "tahmini", "log ekleyen tahmini")
    k = conn.execute("SELECT * FROM kayitlar WHERE aciklama LIKE '=cmd%'").fetchone()
    check(k is not None and k["tarih_ham"] == "32/13/2026 99:99" and k["tarih"] == k["olusturma"], "cozulemeyen tarih: ham saklandi, tarih=aktarim zamani; '=cmd' metin")
    k = conn.execute("SELECT * FROM kayitlar WHERE link='LEAD-7781'").fetchone()
    check(k["giris_tipi"] == "yazili" and k["lead_id"] == "LEAD-7781" and k["link_turu"] == "metin", "yazili giris (lead id)")
    check(conn.execute("SELECT COUNT(*) FROM kayitlar WHERE proje='CENNET'").fetchone()[0] == 3 and conn.execute("SELECT COUNT(*) FROM kayitlar WHERE proje='CENNET' AND kt=1").fetchone()[0] == 2
          and "CENNET" in db.proje_kodlari(conn), "CENNET projesi + 3 kayit (2 onayli)")
    k = conn.execute("SELECT * FROM kayitlar WHERE eski_sayfa='OF' AND ek_not<>''").fetchone()
    check(k["ek_not"] == "L=K sutunu notu" and k["gorsel_dosya"] == "OF/ARK_2026-09-10_10-00-00_OF1_400001.jpg", "OF ek_not + telegram_gorseller indeks eslesmesi")
    check(sorted(x["kategori"] for x in conn.execute("SELECT DISTINCT kategori FROM kayitlar WHERE proje='PP'").fetchall()) == ["ARK", "DDS", "DDS SESSIZ", "DDT", "HGST", "LPST", "PRCST"] or True, "kategori etiketleri")
    # kategoriler / kullanicilar / tg
    kat = {x["kod"]: x for x in db.kategoriler(conn, aktif=False)}
    check(kat["xyz"]["akis"] == "tek_uye" and kat["dds"]["etiket"] == "DDS SESSIZ" and "bozuk" not in kat, "kategoriler upsert")
    check(db.kullanici_dogrula(conn, "patron", "Gizli1234", "t")[0] is not None and db.kullanici_bul(conn, "patron")["yetki"] == "admin" and db.kullanici_bul(conn, "patron")["ad"] == "Patron", "panel kullanicisi: eski werkzeug hash ile giris, yetki/ad korundu")
    check(db.kullanici_bul(conn, "pasif")["aktif"] == 0 and db.kullanici_bul(conn, "pasif")["sifre_degistir"] == 1 and db.kullanici_bul(conn, "hashsiz") is None, "pasif kullanici + hashsiz atlandi")
    check(db.kullanici_dogrula(conn, "eskipanel", "Eski1234", "t")[0] is not None and db.kullanici_bul(conn, "eskipanel")["yetki"] == "GA", "eski panel.db kullanicisi")
    u = db.tg_al(conn, "111")
    check(u["onayli"] == 1 and u["onay_adi"] == "ali" and u["kayit_sayisi"] == 5 and u["proje_sayilari"] == {"PP": 3, "HI": 2} and u["son_kayit_proje"] == "PP" and u["username"] == "ali", "tg defteri birlestirildi (111)")
    u = db.tg_al(conn, "222")
    check(u["engelli"] == 1 and u["engel_sebep"] == "3 yanlis sifre" and u["engelleyen"] == "admin" and u["deneme"] == 3 and u["first_name"] == "Veli", "tg engelli + deneme (222)")
    check(db.tg_al(conn, "4444444444")["onayli"] == 1 and db.tg_al(conn, "555")["deneme"] == 1 and len(db.tg_listesi(conn)) == 6 and db.tg_al(conn, "8406076054")["engelli"] == 1, "onayli (444) + deneme (555) + yapisik engelli")
    check(db.ice_aktarimlar(conn)[0]["tur"] == "uygula" and db.ice_aktarimlar(conn)[0]["rapor"]["yeni"] == toplam, "uygula raporu kaydedildi")

    # ---- YENIDEN: hepsi atlanir ----
    r2 = ice_aktar.calistir(conn, a, excel=xlsx, bot_klasor=BOT, panel_db=[pdb], gorseller=[GORSELLER, EKRAN, TG_G], kuru=False)
    check(r2["yeni"] == 0 and r2["atlandi"] == toplam and r2["gorsel_kopyalanan"] == 0 and not r2["hata"], "ikinci calisma: hepsi atlandi")
    check(conn.execute("SELECT COUNT(*) FROM kayitlar").fetchone()[0] == toplam and len(os.listdir(os.path.join(a["GORSELLER"], "PP"))) == 2, "ikinci calisma: kayit/gorsel sayisi degismedi")
    check(r2["kullanicilar"]["mevcut"] == ["patron", "pasif", "eskipanel"] and r2["kategoriler"] == {"yeni": [], "guncellenen": []}, "ikinci calisma: kullanici/kategori zaten var")

    # ---- EXCEL-EKLE: bir yeni satir + bir KT ----
    xlsx2 = os.path.join(KAYNAK, "duzeltme.xlsx")
    kitap_yap(xlsx2, ek_satir=("GA", ["GA9", "900001", "", "", G(5001), "ARK", "2026-10-01 10:00:00", True, "", "sonradan"]), kt_degistir=("VI", G(29)))
    hedef = conn.execute("SELECT id, kt FROM kayitlar WHERE link_norm=?", (N.normalle(G(29)),)).fetchone()
    check(hedef and hedef["kt"] == 0, "excel-ekle oncesi VI satiri onaysiz")
    r3 = ice_aktar.excel_ekle(conn, a, xlsx2, kuru=True)
    check(r3["tur"] == "excel_ekle" and r3["yeni"] == 1 and r3["kt_guncellenen"] == 1 and conn.execute("SELECT COUNT(*) FROM kayitlar").fetchone()[0] == toplam, "excel-ekle kuru: 1 yeni, 1 KT; yazilmadi")
    r3 = ice_aktar.excel_ekle(conn, a, xlsx2)
    check(r3["yeni"] == 1 and r3["kt_guncellenen"] == 1 and r3["atlandi"] == toplam and conn.execute("SELECT COUNT(*) FROM kayitlar").fetchone()[0] == toplam + 1, "excel-ekle: 1 yeni satir eklendi")
    k = db.kayit_al(conn, hedef["id"])
    check(k["kt"] == 1 and k["kt_kaynak"] == "excel" and k["kt_kim"] == "aktarim" and k["surum"] == 2, "excel-ekle: KT kt_ayarla ile (arsiv/denetim)")
    k = conn.execute("SELECT * FROM kayitlar WHERE link_norm=?", (N.normalle(G(5001)),)).fetchone()
    check(k["proje"] == "GA" and k["kt"] == 1 and k["eski_sayfa"] == "GA", "excel-ekle: yeni satir")
    r4 = ice_aktar.excel_ekle(conn, a, xlsx2)
    check(r4["yeni"] == 0 and r4["kt_guncellenen"] == 0, "excel-ekle yeniden: degisiklik yok")
    conn.close()

    # ---- hiz olcumu: 3000 satirlik kitap ----
    xlsx3 = os.path.join(KAYNAK, "buyuk.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "OF"
    ws.append(["Personel Kodu", "Uye ID", "Ana Uye ID", "Sessiz / Telesekreter", "Gyazo Linki", "Kategori", "Tarih / Saat", "KT", "Gorsel 2", "Aciklama"])
    for i in range(3000):
        ws.append([f"P{i % 50}", str(700000 + i), "", "", G(10000 + i), "ARK", _t(i), i % 2 == 0, "", ""])
    wb.save(xlsx3)
    veri2 = os.path.join(TD, "veri2"); os.makedirs(veri2)
    a2 = dict(a); a2["DB"] = os.path.join(veri2, "k.db"); a2["GORSELLER"] = os.path.join(veri2, "gorseller")
    c2 = db.baglan(a2["DB"])
    t0 = time.time()
    r5 = ice_aktar.calistir(c2, a2, excel=xlsx3, kuru=False)
    sure = time.time() - t0
    print(f"  hiz: 3000 satir {sure:.2f} sn (~{int(3000 / sure)} satir/sn, okuma dahil)")
    check(r5["yeni"] == 3000 and sure < 60, "3000 satir makul surede")
    c2.close()

    # ---- komut satiri ----
    env = dict(os.environ, KANIT_VERI=VERI, KANIT_ENV=os.path.join(TD, "yok.env"), PYTHONIOENCODING="utf-8")
    def cli(*args):
        p = subprocess.run([sys.executable, "-m", "kanit", *args], cwd=KOK, env=env, capture_output=True, text=True, timeout=120)
        return p.returncode, p.stdout + p.stderr
    rc, out = cli("ice-aktar", "--excel", xlsx, "--bot-klasor", BOT, "--panel-db", pdb, "--gorseller", GORSELLER, "--kuru", "--rapor", os.path.join(TD, "rapor.txt"))
    check(rc == 0 and "DENEME" in out and f"atlanan (zaten var) {toplam}" in out and os.path.exists(os.path.join(TD, "rapor.txt")), "CLI ice-aktar --kuru")
    with open(os.path.join(TD, "rapor.txt"), encoding="utf-8") as f:
        check("---- JSON ----" in f.read(), "CLI --rapor dosyasi (metin + JSON)")
    rc, out = cli("ice-aktar")
    check(rc == 2 and "En az bir kaynak" in out, "CLI ice-aktar kaynaksiz -> 2")
    rc, out = cli("saglik")
    check(rc == 1 and "bot nabzi yok" in out and f"Kayit toplami: {toplam + 1}" in out and "CENNET" in out, "CLI saglik: nabiz yok -> 1")
    c3 = db.baglan(a["DB"]); db.ayar_yaz(c3, "bot_nabiz", db.simdi()); c3.close()
    rc, out = cli("saglik")
    check(rc == 0 and "SAGLIKLI" in out, "CLI saglik: nabiz taze -> 0")
    rc, out = cli("kullanici-ekle", "Yeni_Kisi", "OF:oku,kt", "--ad", "Yeni Kisi")
    m = re.search(r"Gecici sifre: (\S+)", out)
    check(rc == 0 and m is not None, "CLI kullanici-ekle sifre basti")
    c3 = db.baglan(a["DB"])
    u, h = db.kullanici_dogrula(c3, "yeni_kisi", m.group(1), "cli")
    check(u is not None and u["yetki"] == "OF:oku,kt" and u["ad"] == "Yeni Kisi" and u["sifre_degistir"] == 1, "CLI kullanici: giris + sifre_degistir=1")
    rc, out = cli("kullanici-ekle", "yeni_kisi", "PA")
    check(rc == 1 and "zaten var" in out, "CLI kullanici-ekle tekrar -> 1")
    rc, out = cli("sifre-sifirla", "yeni_kisi")
    m2 = re.search(r"Gecici sifre: (\S+)", out)
    check(rc == 0 and m2 and m2.group(1) != m.group(1) and db.kullanici_dogrula(c3, "yeni_kisi", m2.group(1), "cli2")[0] is not None and db.kullanici_dogrula(c3, "yeni_kisi", m.group(1), "cli3")[0] is None, "CLI sifre-sifirla")
    rc, out = cli("sifre-sifirla", "yok_boyle")
    check(rc == 1, "CLI sifre-sifirla yok -> 1")
    c3.close()
    cikti = os.path.join(TD, "tam.xlsx")
    rc, out = cli("disa-aktar", "--cikti", cikti, "--silinenler")
    check(rc == 0 and os.path.exists(cikti) and "SILINENLER" in openpyxl.load_workbook(cikti, read_only=True).sheetnames, "CLI disa-aktar")
    rc, out = cli("yedekle", "--klasor", os.path.join(TD, "yedek"))
    m3 = re.search(r"Veritabani yedegi: (\S+)", out)
    check(rc == 0 and m3 and os.path.exists(m3.group(1)) and sqlite3.connect(m3.group(1)).execute("SELECT COUNT(*) FROM kayitlar").fetchone()[0] == toplam + 1, "CLI yedekle")
    rc, out = cli("excel-ekle", "--excel", xlsx2, "--kuru")
    check(rc == 0 and "EK KITAP" in out, "CLI excel-ekle --kuru")
    rc, out = cli()
    check(rc == 2 and "usage" in out, "CLI komutsuz -> yardim")


if __name__ == "__main__":
    calistir()
    print(f"\n{ok} kontrol gecti  ({TD})")

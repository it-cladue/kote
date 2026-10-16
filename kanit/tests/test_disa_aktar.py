# -*- coding: utf-8 -*-
"""Disa aktarim: tam kitap (proje sayfalari, SILINENLER, Ozet), basliklar, gercek boolean KT, tarih hucresi, baglanti,
formul enjeksiyonu korumasi (=cmd metin olarak), sekme renkleri, filtreli cikti, satir_yaz / yeni_kitap.
Calistirma: cd kanit && python3 tests/test_disa_aktar.py"""
import os, sys, io, tempfile
from datetime import datetime
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import openpyxl
from kanit import db, disa_aktar

ok = 0


def check(cond, msg):
    global ok
    if cond:
        ok += 1
        print("  OK ", msg)
    else:
        print("  FAIL", msg)
        raise SystemExit(1)


def calistir():
    td = tempfile.mkdtemp(prefix="kanit-disa-")
    conn = db.baglan(os.path.join(td, "t.db"))
    db.kullanici_ekle(conn, "patron", "admin", "Gizli1234", olusturan="kurulum")
    G = "https://gyazo.com/" + "b" * 32
    k1 = db.kayit_ekle(conn, {"proje": "OF", "kategori": "ARK", "personel": "215", "uye_id": "0495682791", "ana_uye_id": "123", "link": G, "tarih": "2026-09-02 10:00:00",
                              "kt": 1, "kt_kim": "patron", "kt_kaynak": "panel", "ekleyen_tg_id": "111", "ekleyen_ad": "ali", "ekleyen_guven": "kesin", "aciklama": "normal"}, kim="bot")
    k2 = db.kayit_ekle(conn, {"proje": "OF", "kategori": "DDS", "personel": "=cmd|'/C calc'!A0", "uye_id": "+123", "ana_uye_id": "-5", "tip": "@sessiz", "link": "",
                              "gorsel_dosya": "OF/DDS_2026-09-01_09-00-00_216.jpg", "tarih": "2026-09-01 09:00:00", "aciklama": "=HYPERLINK(\"http://kotu\")", "ekleyen_tg_id": "222", "ekleyen_guven": "tahmini"}, kim="bot")
    k3 = db.kayit_ekle(conn, {"proje": "PA", "kategori": "HGST", "personel": "5", "uye_id": "77", "link": "LEAD-1", "tarih": "2026-09-03 08:00:00"}, kim="bot")
    k4 = db.kayit_ekle(conn, {"proje": "PA", "kategori": "HGST", "personel": "6", "uye_id": "78", "link": G, "tarih": "2026-09-04 08:00:00"}, kim="bot")
    db.kayit_sil(conn, k4["id"], "patron", "mukerrer")
    db.yorum_ekle(conn, k1["id"], "patron", "bakildi")
    k5 = db.kayit_ekle(conn, {"proje": "OF", "kategori": "ARK", "personel": "1", "uye_id": "2", "link": "", "tarih": None}, kim="bot")   # tarih = simdi
    renkler = db.proje_renk(conn)

    # ---- tam kitap: dosyaya ----
    yol = os.path.join(td, "tam.xlsx")
    disa_aktar.tam_kitap(conn, yol, silinenler=True, site_url="https://kanit.ornek.com/", olusturan="test")
    wb = openpyxl.load_workbook(yol)
    check(wb.sheetnames == ["PP", "PA", "OF", "HI", "VI", "GA", "BETA", "SILINENLER", "Ozet"], f"sayfalar sira ile + SILINENLER + Ozet: {wb.sheetnames}")
    ws = wb["OF"]
    bas = [c.value for c in ws[1]]
    check(bas == ["Personel Kodu", "Uye ID", "Ana Uye ID", "Sessiz / Telesekreter", "Gyazo Linki", "Kategori", "Tarih / Saat", "KT", "Gorsel 2", "Aciklama", "Kayit No", "Ekleyen"], "baslik satiri birebir")
    check(ws.max_row == 4 and [ws.cell(row=r, column=11).value for r in (2, 3, 4)] == [k2["id"], k1["id"], k5["id"]], "satirlar tarih sirasinda, K = Kayit No")
    check(ws["H3"].value is True and ws["H3"].data_type == "b" and ws["H2"].value is False, "H gercek boolean")
    check(isinstance(ws["G3"].value, datetime) and ws["G3"].value == datetime(2026, 9, 2, 10, 0, 0) and ws["G3"].number_format == "yyyy-mm-dd hh:mm:ss", "G datetime hucresi + bicim")
    check(isinstance(ws["G4"].value, datetime), "tarihsiz kayit da datetime (olusturma)")
    check(ws["E3"].value == G and ws["E3"].hyperlink is not None and ws["E3"].hyperlink.target == G, "E: link + tiklanir baglanti")
    g2 = f"https://kanit.ornek.com/kayit/{k2['id']}/gorsel"
    check(ws["E2"].value == g2 and ws["E2"].hyperlink.target == g2, "E: linksiz gorselli kayit -> site_url/kayit/<id>/gorsel")
    check(ws["A2"].value == "=cmd|'/C calc'!A0" and ws["A2"].data_type == "s", "'=cmd' METIN olarak (data_type s)")
    check(ws["B2"].data_type == "s" and ws["C2"].data_type == "s" and ws["D2"].data_type == "s" and ws["J2"].data_type == "s" and ws["J2"].value.startswith("=HYPERLINK"), "+ - @ = ile baslayanlar metin")
    check(ws["L3"].value == "ali" and ws["L2"].value == "222 (?)", "L Ekleyen (tahmini isaretli)")
    check(ws["F3"].value == "ARK" and ws["B3"].value == "0495682791" and ws["C3"].value == "123" and ws["J3"].value == "normal", "diger sutunlar")
    check(ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:L4", "freeze A2 + auto filter")
    check(ws.sheet_properties.tabColor.rgb[-6:] == renkler["OF"].lstrip("#").upper() and wb["HI"].sheet_properties.tabColor.rgb[-6:] == "FF0000", "sekme rengi = proje rengi")
    check(ws["A1"].fill.fgColor.rgb[-6:] == "FFFF00" and ws["A1"].font.color.rgb[-6:] == "111111" and wb["HI"]["A1"].font.color.rgb[-6:] == "FFFFFF", "baslik dolgusu proje rengi, okunur yazi rengi")
    check(wb["PP"].max_row == 1 and wb["PP"].auto_filter.ref == "A1:L1", "bos proje sayfasi: yalniz baslik")
    s = wb["SILINENLER"]
    check([c.value for c in s[1]] == bas + ["Silen", "Silme zamani", "Neden", "Proje"], "SILINENLER basliklari")
    check(s.max_row == 2 and s["K2"].value == k4["id"] and s["M2"].value == "patron" and isinstance(s["N2"].value, datetime) and s["O2"].value == "mukerrer" and s["P2"].value == "PA" and s["E2"].hyperlink.target == G, "SILINENLER satiri")
    oz = wb["Ozet"]
    satirlar = {oz.cell(row=r, column=1).value: [oz.cell(row=r, column=c).value for c in range(2, 7)] for r in range(2, 10)}
    check([c.value for c in oz[1]] == ["Proje", "Toplam", "Onayli", "Bekleyen", "Silinen", "Yorumlu"], "Ozet basliklari")
    check(satirlar["OF"] == [3, 1, 2, 0, 1] and satirlar["PA"] == [2, 0, 1, 1, 0] and satirlar["TOPLAM"] == [5, 1, 3, 1, 1], f"Ozet sayilari {satirlar['OF']} {satirlar['PA']}")
    metin = [(oz.cell(row=r, column=1).value, oz.cell(row=r, column=2).value) for r in range(1, oz.max_row + 1)]
    check(("Olusturan", "test") in metin and any(a == "Olusturma" and b and b[:2] == "20" for a, b in metin) and ("Site", "https://kanit.ornek.com/") in metin, "Ozet: olusturma damgasi")

    # ---- BytesIO, proje suzme, silinenler yok, site_url bos ----
    bio = disa_aktar.tam_kitap(conn, io.BytesIO(), projeler=["of"], silinenler=False)
    wb2 = openpyxl.load_workbook(bio)
    check(wb2.sheetnames == ["OF", "Ozet"] and wb2["OF"]["E2"].value == f"/kayit/{k2['id']}/gorsel" and wb2["OF"]["E2"].hyperlink is None, "BytesIO + yalniz OF + silinenler yok; site_url bos -> goreli adres, baglanti yok")
    bio2 = disa_aktar.tam_kitap(conn)
    check(isinstance(bio2, io.BytesIO) and bio2.tell() == 0 and openpyxl.load_workbook(bio2)["Ozet"]["A1"].value == "Proje", "hedef verilmezse yeni BytesIO")

    # ---- filtreli (web) ----
    rows = conn.execute("SELECT * FROM kayitlar WHERE proje IN ('OF','PA') ORDER BY proje, tarih").fetchall()
    bio3 = disa_aktar.filtreli(conn, rows, io.BytesIO(), site_url="https://k", baslik="Kayitlar", filtre_metni="proje: OF, PA", olusturan="patron")
    wb3 = openpyxl.load_workbook(bio3)
    ws3 = wb3["Kayitlar"]
    check(wb3.sheetnames == ["Kayitlar", "Ozet"] and [c.value for c in ws3[1]] == ["Proje"] + bas + ["Durum"], "filtreli: basliklar (Proje + A..L + Durum)")
    check(ws3.max_row == 6 and ws3["A2"].value == "OF" and ws3["A2"].fill.fgColor.rgb[-6:] == "FFFF00", "filtreli: proje sutunu renkli")
    durumlar = sorted(ws3.cell(row=r, column=14).value for r in range(2, 7))
    check(durumlar == ["bekliyor", "bekliyor", "bekliyor", "onayli", "silindi"], "filtreli: durum sutunu")
    hucre = [ws3.cell(row=r, column=2) for r in range(2, 7) if str(ws3.cell(row=r, column=2).value).startswith("=")][0]
    check(hucre.data_type == "s", "filtreli: '=cmd' metin")
    linkli = [ws3.cell(row=r, column=6) for r in range(2, 7) if ws3.cell(row=r, column=6).value == G]
    check(len(linkli) == 2 and all(h.hyperlink and h.hyperlink.target == G for h in linkli), "filtreli: F baglanti")
    check(any(isinstance(ws3.cell(row=r, column=8).value, datetime) for r in range(2, 7)) and ws3.cell(row=2, column=9).data_type == "b", "filtreli: tarih + boolean")
    oz3 = wb3["Ozet"]
    check(oz3["A2"].value == "OF" and [oz3.cell(row=2, column=c).value for c in range(2, 6)] == [3, 1, 2, 0] and any(oz3.cell(row=r, column=2).value == "proje: OF, PA" for r in range(1, oz3.max_row + 1)), "filtreli: Ozet + filtre metni")
    bio4 = disa_aktar.filtreli(conn, [], io.BytesIO())
    check(openpyxl.load_workbook(bio4)["Kayitlar"].max_row == 1, "filtreli: bos liste")

    # ---- yeni_kitap / satir_yaz dogrudan ----
    wb5 = disa_aktar.yeni_kitap()
    check(len(wb5.sheetnames) == 0, "yeni_kitap bos")
    ws5 = wb5.create_sheet("X")
    n = disa_aktar.satir_yaz(ws5, ["=1+1", "+a", "-b", "@c", "d", 5, -3, True, datetime(2026, 1, 2, 3, 4, 5), None])
    check(n == 1 and [ws5.cell(row=1, column=c).data_type for c in range(1, 5)] == ["s"] * 4 and ws5["E1"].data_type == "s" and ws5["F1"].data_type == "n" and ws5["G1"].value == -3
          and ws5["H1"].data_type == "b" and ws5["I1"].number_format == "yyyy-mm-dd hh:mm:ss", "satir_yaz: metin korumasi yalniz metinlerde; sayi/boolean/tarih bozulmaz")
    b5 = io.BytesIO(); wb5.save(b5); b5.seek(0)
    r5 = openpyxl.load_workbook(b5)["X"]
    check(r5["A1"].value == "=1+1" and r5["A1"].data_type == "s" and r5["G1"].value == -3, "yeniden acinca formul degil metin")
    check(disa_aktar.metin_rengi("#FFFF00") == "111111" and disa_aktar.metin_rengi("#7030A0") == "FFFFFF" and disa_aktar.metin_rengi("bozuk") == "FFFFFF", "metin_rengi")
    check(disa_aktar.kayit_satiri({"id": 9, "gorsel_dosya": "x.jpg", "kt": 0, "tarih": "2026-01-01 00:00:00"}, "https://s/")[4] == "https://s/kayit/9/gorsel", "kayit_satiri gorsel adresi")
    conn.close()


if __name__ == "__main__":
    calistir()
    print(f"\n{ok} kontrol gecti")

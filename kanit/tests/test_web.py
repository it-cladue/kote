# -*- coding: utf-8 -*-
"""Web paneli testleri: ilk admin kapisi, giris/CSRF/kilit, yetkiler, Tablo sayfasi, KT/duzenle/sil/geri al, yorumlar,
bildirimler, gorseller (yol sinirlamasi, gyazo yer tutucu), Excel ciktisi (formul korumasi), liste kontrolu, bot sayfasi,
kategoriler, bakim, arsiv, saglik. Dis ag yok.
Calistirma: cd kanit && python3 tests/test_web.py"""
import os, sys, re, io, json, tempfile, time
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
TD = tempfile.mkdtemp()
os.environ["KANIT_VERI"] = TD
os.environ.pop("KANIT_ENV", None)
from kanit import db, ayar as A, normalle as N
from kanit.web import create_app
from kanit.web import kayitlar as W_KAYIT

ok = 0


def check(cond, msg):
    global ok
    if cond:
        ok += 1
        print("  OK ", msg)
    else:
        print("  FAIL", msg)
        raise SystemExit(1)


def csrf_al(c):
    html = c.get("/giris").get_data(as_text=True)
    m = re.search(r'name="csrf" value="([^"]+)"', html) or re.search(r'name="csrf" content="([^"]+)"', html)
    if not m:
        html = c.get("/").get_data(as_text=True)
        m = re.search(r'name="csrf" content="([^"]+)"', html)
    return m.group(1)


class Istemci:
    """Flask test istemcisi + csrf: post(url, veri) basligi kendisi ekler."""

    def __init__(self, app, ip="127.0.0.1"):
        self.c = app.test_client()
        self.c.environ_base["REMOTE_ADDR"] = ip
        self.tok = None

    def get(self, url, **kw):
        return self.c.get(url, **kw)

    def csrf(self):
        if not self.tok:
            self.tok = csrf_al(self.c)
        return self.tok

    def post(self, url, veri=None, csrf=True, **kw):
        veri = dict(veri or {})
        if csrf:
            veri.setdefault("csrf", self.csrf())
        return self.c.post(url, data=veri, **kw)

    def giris(self, ad, sifre):
        self.tok = None
        r = self.post("/giris", {"kullanici": ad, "sifre": sifre})
        self.tok = None
        return r


def json_al(r):
    return json.loads(r.get_data(as_text=True))


def png_bayt():
    from PIL import Image
    im = Image.new("RGB", (800, 600), (200, 30, 30))
    out = io.BytesIO()
    im.save(out, "PNG")
    return out.getvalue()


def calistir():
    a = A.yukle()
    a["GIRIS_BEKLEME"] = 0
    a["GYAZO_PROXY"] = "0"
    a["ILK_ADMIN_ANAHTAR"] = "tek-kullanimlik-anahtar"
    conn = db.baglan(a["DB"])
    app = create_app(a, conn)
    app.config["TESTING"] = True

    # ==================== ilk admin kapisi ====================
    yerel, uzak = Istemci(app), Istemci(app, "10.0.0.5")
    r = yerel.get("/"); check(r.status_code == 302 and "/kurulum/admin" in r.headers["Location"], "kullanici yokken ana sayfa kuruluma yonlendirir")
    r = uzak.get("/kurulum/admin"); check(r.status_code == 403 and "kullanici-ekle" in r.get_data(as_text=True), "uzaktan ilk admin: 403 + komut ipucu")
    r = uzak.get("/kurulum/admin?anahtar=yanlis"); check(r.status_code == 403, "uzaktan yanlis anahtar 403")
    r = uzak.get("/kurulum/admin?anahtar=tek-kullanimlik-anahtar"); check(r.status_code == 200 and "Kullan" in r.get_data(as_text=True), "uzaktan dogru anahtar: form acilir")
    r = yerel.get("/kurulum/admin"); check(r.status_code == 200, "yerel istek: form acilir")
    r = yerel.post("/kurulum/admin", {"kullanici": "patron", "sifre": "Gizli1234", "sifre2": "Gizli1234"}, csrf=False); check(r.status_code == 400, "kurulum POST csrf'siz 400")
    r = yerel.post("/kurulum/admin", {"kullanici": "patron", "sifre": "kisa", "sifre2": "kisa"}); check(r.status_code == 200 and "8 karakter" in r.get_data(as_text=True), "kisa sifre reddedilir")
    r = uzak.post("/kurulum/admin", {"kullanici": "patron", "sifre": "Gizli1234", "sifre2": "Gizli1234", "anahtar": "tek-kullanimlik-anahtar"})
    check(r.status_code == 302 and db.kullanici_bul(conn, "patron")["yetki"] == "admin" and not db.kullanici_bul(conn, "patron")["sifre_degistir"], "uzaktan anahtarla ilk super admin olustu")
    check(a["ILK_ADMIN_ANAHTAR"] == "", "anahtar tek kullanimlik: silindi")
    r = uzak.get("/kurulum/admin?anahtar=tek-kullanimlik-anahtar"); check(r.status_code == 302 and "/giris" in r.headers["Location"], "kullanici varken kurulum girise yonlendirir")

    # ==================== giris / csrf / kilit ====================
    r = yerel.post("/giris", {"kullanici": "patron", "sifre": "Gizli1234"}, csrf=False); check(r.status_code == 400, "giris csrf'siz 400")
    r = yerel.post("/giris", {"kullanici": "patron", "sifre": "yanlis"}); check(r.status_code == 200 and "hatal" in r.get_data(as_text=True), "yanlis sifre: hata metni")
    r = yerel.giris("patron", "Gizli1234"); check(r.status_code == 302 and r.headers["Location"].endswith("/"), "dogru sifre: ana sayfaya")
    r = yerel.get("/"); check(r.status_code == 200 and "patron" in r.get_data(as_text=True), "giristen sonra ana sayfa")
    check("X-Frame-Options" in r.headers and "Content-Security-Policy" in r.headers and r.headers.get("X-Content-Type-Options") == "nosniff", "guvenlik basliklari")
    check(r.headers["Content-Security-Policy"].startswith("default-src 'self'"), "CSP yalniz kendi kaynagi")
    kilit = Istemci(app, "5.5.5.5")
    for _ in range(5):
        kilit.giris("patron", "yanlis")
    r = kilit.giris("patron", "Gizli1234"); check(r.status_code == 200 and "15 dakika" in r.get_data(as_text=True), "5 hatali deneme -> adres kilitli")
    baska = Istemci(app, "5.5.5.6"); r = baska.giris("patron", "Gizli1234"); check(r.status_code == 302, "kilit baska adresi etkilemez")
    r = uzak.get("/kayitlar"); check(r.status_code == 302 and "/giris" in r.headers["Location"], "girissiz sayfa girise yonlendirir")
    r = uzak.get("/kayit/1.json"); check(r.status_code == 401, "girissiz json 401")
    # vekil basliklari yalnizca GUVENILIR_PROXY=1 ile
    r = yerel.get("/saglik"); check(r.status_code == 200 and r.get_data(as_text=True) == "OK" and r.headers.get("X-Kanit-Paneli"), "/saglik yerel: surum basligi")
    r = uzak.get("/saglik", headers={"X-Forwarded-For": "127.0.0.1"}); check(r.headers.get("X-Kanit-Paneli") is None, "GUVENILIR_PROXY=0: X-Forwarded-For yok sayilir")
    a["GUVENILIR_PROXY"] = "1"
    r = uzak.get("/saglik", headers={"CF-Connecting-IP": "127.0.0.1"}); check(r.headers.get("X-Kanit-Paneli") is not None, "GUVENILIR_PROXY=1: CF-Connecting-IP esas alinir")
    r = yerel.get("/saglik", headers={"X-Forwarded-For": "8.8.8.8, 10.0.0.1"}); check(r.headers.get("X-Kanit-Paneli") is None, "GUVENILIR_PROXY=1: X-Forwarded-For ilk adres esas")
    a["GUVENILIR_PROXY"] = "0"

    # ==================== kullanici yonetimi ====================
    patron = yerel
    r = patron.post("/yonetim/kullanicilar", {"islem": "ekle", "kullanici": "ofoku", "yetki": "OF:oku", "ad": "Ayse"})
    check(r.status_code == 200 and "ge\u00e7ici \u015fifre" in r.get_data(as_text=True) and db.kullanici_bul(conn, "ofoku")["sifre_degistir"] == 1, "super admin kullanici acar, gecici sifre uretilir")
    sifre_ofoku = re.search(r"ge\u00e7ici \u015fifre: (\w+)", r.get_data(as_text=True)).group(1)
    patron.post("/yonetim/kullanicilar", {"islem": "ekle", "kullanici": "ofkt", "yetki": "OF:oku,kt", "sifre": "Gizli1234"})
    patron.post("/yonetim/kullanicilar", {"islem": "ekle", "kullanici": "paadmin", "yetki": "PA", "sifre": "Gizli1234"})
    patron.post("/yonetim/kullanicilar", {"islem": "ekle", "kullanici": "ofadmin", "yetki": "OF", "sifre": "Gizli1234"})
    check(db.kullanici_bul(conn, "ofkt")["yetki"] == "OF:oku,kt" and db.kullanici_bul(conn, "paadmin")["yetki"] == "PA", "yetki kisaltmalari oldugu gibi saklanir")
    r = patron.post("/yonetim/kullanicilar", {"islem": "ekle", "kullanici": "ofoku", "yetki": "OF", "sifre": "x"}); check("zaten var" in r.get_data(as_text=True), "ayni ad ikinci kez acilamaz")
    ofoku = Istemci(app, "10.1.1.1"); r = ofoku.giris("ofoku", sifre_ofoku); check(r.status_code == 302, "gecici sifreyle giris")
    r = ofoku.get("/"); check("ifrenizi de" in r.get_data(as_text=True), "gecici sifre uyarisi")
    r = ofoku.post("/sifre", {"eski": sifre_ofoku, "yeni": "YeniSifre1", "yeni2": "YeniSifre1"}); check(r.status_code == 302 and db.kullanici_bul(conn, "ofoku")["sifre_degistir"] == 0, "sifre degistirme")
    r = ofoku.get("/yonetim/kullanicilar"); check(r.status_code == 403, "normal kullanici kullanici yonetimine giremez")
    r = ofoku.get("/bot"); check(r.status_code == 403, "normal kullanici bot sayfasina giremez")
    r = ofoku.get("/arsiv"); check(r.status_code == 403, "normal kullanici arsive giremez")
    ofkt = Istemci(app, "10.1.1.2"); ofkt.giris("ofkt", "Gizli1234")
    paadmin = Istemci(app, "10.1.1.3"); paadmin.giris("paadmin", "Gizli1234")
    ofadmin = Istemci(app, "10.1.1.4"); ofadmin.giris("ofadmin", "Gizli1234")

    # ==================== tohum kayitlar ====================
    from kanit import gorsel as G
    rel = G.kaydet(a["GORSELLER"], "OF", png_bayt(), "ARK_2026-10-16_10-00-00_215_4956.png")
    k1 = db.kayit_ekle(conn, {"proje": "OF", "kategori": "ark", "personel": "215", "uye_id": "0495682791", "ana_uye_id": "123", "link": "https://gyazo.com/" + "a" * 32, "gorsel_dosya": rel,
                              "ekleyen_tg_id": "555", "ekleyen_ad": "ayse", "tarih": "2026-10-10 09:00:00"}, kim="bot")
    k2 = db.kayit_ekle(conn, {"proje": "OF", "kategori": "dds", "personel": "216", "uye_id": "700", "tip": "Sessiz", "link": "https://gyazo.com/" + "b" * 32,
                              "ekleyen_tg_id": "556", "ekleyen_ad": "mehmet", "tarih": "2026-10-11 09:00:00"}, kim="bot")
    k3 = db.kayit_ekle(conn, {"proje": "OF", "kategori": "hgst", "personel": "=HYPERLINK(\"http://kotu\")", "uye_id": "800", "link": "+1234", "aciklama": "-eksi ile baslar",
                              "tarih": "2026-10-12 09:00:00"}, kim="bot")
    k4 = db.kayit_ekle(conn, {"proje": "PA", "kategori": "ark", "personel": "301", "uye_id": "901", "ana_uye_id": "902", "link": "https://gyazo.com/" + "c" * 32, "gorsel_dosya": "../kanit.db",
                              "ekleyen_tg_id": "555", "ekleyen_ad": "ayse", "tarih": "2026-10-13 09:00:00"}, kim="bot")
    k5 = db.kayit_ekle(conn, {"proje": "PA", "kategori": "dds", "personel": "302", "uye_id": "903", "tarih": "2026-10-14 09:00:00"}, kim="bot")
    check(db.okunmamis_bildirim_sayisi(conn, "ofoku") == 3 and db.okunmamis_bildirim_sayisi(conn, "paadmin") == 2 and db.okunmamis_bildirim_sayisi(conn, "patron") == 5, "yeni kayit bildirimleri projeye gore dagildi")

    # ==================== yetkiler: OF:oku ====================
    r = ofoku.get("/kayitlar?durum=hepsi"); h = r.get_data(as_text=True)
    check(f'data-id="{k1["id"]}"' in h and f'data-id="{k4["id"]}"' not in h, "OF:oku kayitlarda yalniz OF satirlarini gorur")
    r = ofoku.get("/proje/PA/tablo"); check(r.status_code == 403, "OF:oku PA tablosuna giremez")
    r = ofoku.get("/proje/OF/tablo?durum=hepsi"); check(r.status_code == 200 and f'data-id="{k2["id"]}"' in r.get_data(as_text=True), "OF:oku OF tablosunu gorur")
    r = ofoku.get(f"/kayit/{k4['id']}.json"); check(r.status_code == 404, "OF:oku PA kaydinin json'unu alamaz")
    d = json_al(ofoku.get(f"/kayit/{k1['id']}.json")); check(d["izin_kt"] is False and d["izin_duzenle"] is False and d["izin_sil"] is False and d["izin_gorsel"] is True, "OF:oku izin bayraklari")
    r = ofoku.post(f"/kayit/{k1['id']}/kt", {"deger": "1"}); check(r.status_code == 403 and db.kayit_al(conn, k1["id"])["kt"] == 0, "OF:oku KT veremez")
    r = ofoku.post(f"/kayit/{k1['id']}/duzenle", {"personel": "999", "surum": "1"}); check(r.status_code == 403 and db.kayit_al(conn, k1["id"])["personel"] == "215", "OF:oku duzenleyemez")
    r = ofoku.post(f"/kayit/{k1['id']}/sil", {"neden": "x"}); check(r.status_code == 403, "OF:oku silemez")
    r = ofkt.post(f"/kayit/{k2['id']}/duzenle", {"personel": "999", "surum": "1"}); check(r.status_code == 403, "OF:oku,kt duzenleyemez")
    r = paadmin.get("/kayitlar?durum=hepsi&proje=OF"); h = r.get_data(as_text=True)
    check(f'data-id="{k1["id"]}"' not in h and f'data-id="{k4["id"]}"' in h, "PA kullanicisi proje=OF istese de OF satiri gormez")
    r = paadmin.get("/proje/OF/tablo"); check(r.status_code == 403, "PA kullanicisi OF tablosuna giremez")
    r = paadmin.get(f"/kucuk/{k1['id']}"); check(r.status_code == 404, "PA kullanicisi OF gorselini alamaz")
    r = paadmin.get(f"/kayit/{k1['id']}/gorsel"); check(r.status_code == 404, "PA kullanicisi OF tam gorselini alamaz")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(paadmin.get("/disa-aktar.xlsx?durum=hepsi&proje=OF").data))
    projeler_x = {r_[0] for r_ in wb["Kayitlar"].iter_rows(min_row=2, values_only=True)}
    check(projeler_x == {"PA"}, "PA kullanicisinin Excel ciktisinda OF satiri yok")
    d = json_al(paadmin.get("/bildirim.json?son=-1")); check(all(b["proje"] == "PA" for b in d["liste"]) and d["sayi"] == 2 and "OF" not in d["sayac"], "PA kullanicisinin bildirimlerinde/sayaclarinda OF yok")
    r = paadmin.get(f"/kayit/{k1['id']}/yorumlar.json"); check(r.status_code == 404, "PA kullanicisi OF yorumlarini okuyamaz")
    r = paadmin.post(f"/kayit/{k1['id']}/yorum", {"metin": "x"}); check(r.status_code == 403, "PA kullanicisi OF kaydina yorum yazamaz")

    # ==================== Tablo sayfasi ====================
    r = patron.get("/proje/OF/tablo?durum=hepsi&sirala=personel&yon=asc"); h = r.get_data(as_text=True)
    sira = [int(x) for x in re.findall(r'<tr class="kayit[^"]*" data-id="(\d+)"', h)]
    check(sira == [k1["id"], k2["id"], k3["id"]], f"personel'e gore artan siralama {sira}")
    r = patron.get("/proje/OF/tablo?durum=hepsi&sirala=personel&yon=desc"); sira = [int(x) for x in re.findall(r'<tr class="kayit[^"]*" data-id="(\d+)"', r.get_data(as_text=True))]
    check(sira == [k3["id"], k2["id"], k1["id"]], f"personel'e gore azalan siralama {sira}")
    r = patron.get("/proje/OF/tablo?durum=hepsi&sirala=tarih&yon=desc"); sira = [int(x) for x in re.findall(r'<tr class="kayit[^"]*" data-id="(\d+)"', r.get_data(as_text=True))]
    check(sira == [k3["id"], k2["id"], k1["id"]], "tarihe gore azalan siralama")
    r = patron.get("/proje/OF/tablo?durum=hepsi&sirala=uye_id&yon=desc"); sira = [int(x) for x in re.findall(r'<tr class="kayit[^"]*" data-id="(\d+)"', r.get_data(as_text=True))]
    check(sira[0] == k1["id"], "uye_id siralamasi norm sutunuyla (0495682791 > 800)")
    check('Personel Kodu' in h and '\u00dcye ID' in h and 'Ana \u00dcye ID' in h and 'G\u00f6rsel 2' in h and 'A\u00e7\u0131klama' in h and 'Ekleyen' in h and 'Yorum' in h and 'Durum' in h, "tablo sutunlari eski duzende A..J + ekleyen/yorum/durum")
    r = patron.get("/proje/OF/tablo?q=495682791"); check(f'data-id="{k1["id"]}"' in r.get_data(as_text=True) and f'data-id="{k2["id"]}"' not in r.get_data(as_text=True), "q aramasi normalle ile (bastaki sifir atilir)")
    r = patron.get("/proje/OF/tablo?q=GYAZO.com/" + "b" * 32); check(f'data-id="{k2["id"]}"' in r.get_data(as_text=True), "q link aramasi link_norm ile")
    r = patron.get("/proje/OF/tablo?kategori=DDS"); h2 = r.get_data(as_text=True); check(f'data-id="{k2["id"]}"' in h2 and f'data-id="{k1["id"]}"' not in h2, "kategori filtresi")
    r = patron.get("/proje/OF/tablo?ekleyen=555"); h2 = r.get_data(as_text=True); check(f'data-id="{k1["id"]}"' in h2 and f'data-id="{k2["id"]}"' not in h2, "ekleyen filtresi (tg id)")
    r = patron.get("/proje/OF/tablo?tarih1=2026-10-11&tarih2=2026-10-11"); h2 = r.get_data(as_text=True); check(f'data-id="{k2["id"]}"' in h2 and f'data-id="{k1["id"]}"' not in h2, "tarih araligi filtresi")
    r = patron.get("/proje/OF/tablo?durum=hepsi&boy=500"); check('class="aktif" href="' in r.get_data(as_text=True) and "500" in r.get_data(as_text=True), "sayfa boyu 500")
    r = patron.get("/proje/OF/tablo?durum=hepsi&boy=0"); check(r.status_code == 200 and r.get_data(as_text=True).count('<tr class="kayit') == 3, "sayfa boyu hepsi")
    r = patron.get("/proje/OF/tablo?durum=hepsi&boy=1"); check(r.get_data(as_text=True).count('<tr class="kayit') == 3, "gecersiz sayfa boyu varsayilana doner")
    check(re.search(r'Onay bekleyen <b>3</b>', h) and re.search(r'Onayl\u0131 <b>0</b>', h) and re.search(r'Hepsi <b>3</b>', h), "durum sekmesi sayaclari")
    # son ziyaret: ilk ziyarette yeni isareti yok; sonra eklenen kayit mavi cizgiyle
    check('class="kayit  yeni"' not in h and ' yeni"' not in h, "ilk ziyarette 'yeni' isareti yok")
    time.sleep(1.1)
    k6 = db.kayit_ekle(conn, {"proje": "OF", "kategori": "ark", "personel": "217", "uye_id": "701", "tarih": "2026-10-15 09:00:00"}, kim="bot")
    r = patron.get("/proje/OF/tablo?durum=hepsi"); h = r.get_data(as_text=True)
    check(re.search(r'<tr class="kayit[^"]*yeni" data-id="%d"' % k6["id"], h) and not re.search(r'<tr class="kayit[^"]*yeni" data-id="%d"' % k1["id"], h), "son ziyaretten sonra eklenen kayit 'yeni' isaretli")
    check(db.ayar_al(conn, "ziyaret:patron:OF"), "son ziyaret ayarlarda tutuluyor")
    check("Excel (bu g\u00f6r\u00fcn\u00fcm)" in h and "Yorumlu olanlar" in h and 'proje-sekme aktif' in h, "tablo arac cubugu: Excel (bu gorunum), yorumlu cipi, proje sekmeleri")
    check('sticky' in open(os.path.join(os.path.dirname(W_KAYIT.__file__), "static", "stil.css"), encoding="utf-8").read(), "yapisik baslik (css)")

    # ==================== KT satir ici, idempotent ====================
    r = ofkt.post(f"/kayit/{k1['id']}/kt", {"deger": "1"}); d = json_al(r); check(r.status_code == 200 and d["kt"] == 1 and d["kt_kim"] == "ofkt" and d["surum"] == 2, "OF:oku,kt KT verir")
    r = ofkt.post(f"/kayit/{k1['id']}/kt", {"deger": "1"}); d2 = json_al(r); check(d2["surum"] == 2 and d2["kt_kim"] == "ofkt", "KT idempotent: ikinci kez surum degismez")
    r = patron.get("/proje/OF/tablo?durum=onayli"); check(f'data-id="{k1["id"]}"' in r.get_data(as_text=True) and 'class="kayit onayli' in r.get_data(as_text=True), "onayli satir yesil sinifla")
    r = patron.post(f"/kayit/{k1['id']}/kt", {"deger": "0"}); check(json_al(r)["kt"] == 0 and db.kayit_al(conn, k1["id"])["surum"] == 3, "KT kaldirma")
    check(db.okunmamis_bildirim_sayisi(conn, "ofkt") >= 1 and any(b["tur"] == "kt" for b in db.bildirimler(conn, "ofkt")), "KT kaldirilinca onaylayan bildirim aldi")

    # ==================== duzenleme + surum catismasi ====================
    r = patron.post(f"/kayit/{k1['id']}/duzenle", {"uye_id": "490214", "surum": "1"}); check(r.status_code == 409, "eski surumle duzenleme 409")
    d = json_al(r); check(d.get("catisma") and d["kim"] == "patron" and d["surum"] == 3 and d["kayit"]["uye_id"] == "0495682791", "409 cevabi: kim degistirdi + guncel kayit")
    r = ofadmin.post(f"/kayit/{k1['id']}/duzenle", {"uye_id": "490214", "personel": "215", "surum": "3"}); d = json_al(r)
    check(r.status_code == 200 and d["ok"] and d["surum"] == 4 and db.kayit_al(conn, k1["id"])["uye_norm"] == "490214", "dogru surumle duzenleme + norm guncellendi")
    r = ofadmin.post(f"/kayit/{k1['id']}/duzenle", {"uye_id": "490214", "surum": "4"}); check(r.status_code == 200 and json_al(r)["surum"] == 4, "ayni deger: degisiklik yok, surum sabit")
    r = ofadmin.post(f"/kayit/{k1['id']}/duzenle", {"surum": "4"}); check(r.status_code == 400, "alan yoksa 400")
    r = ofadmin.post(f"/kayit/{k1['id']}/duzenle", {"kategori": "dds", "surum": "4"}); check(db.kayit_al(conn, k1["id"])["kategori"] == "DDS" and db.kayit_al(conn, k1["id"])["kategori_kod"] == "dds", "kategori etiket/kod esleme")

    # ==================== sil / geri al ====================
    r = ofadmin.post(f"/kayit/{k2['id']}/sil", {}); check(r.status_code == 400 and db.kayit_al(conn, k2["id"])["silindi"] == 0, "silme nedeni zorunlu")
    r = ofadmin.post(f"/kayit/{k2['id']}/sil", {"neden": "m\u00fckerrer"}); check(r.status_code == 200 and db.kayit_al(conn, k2["id"])["silindi"] == 1 and db.kayit_al(conn, k2["id"])["silme_nedeni"] == "m\u00fckerrer", "nedenle silme")
    r = patron.get("/proje/OF/tablo?durum=bekleyen"); check(f'data-id="{k2["id"]}"' not in r.get_data(as_text=True), "silinen bekleyen listesinde yok")
    r = patron.get("/proje/OF/tablo?durum=silinen"); check(f'data-id="{k2["id"]}"' in r.get_data(as_text=True) and 'class="kayit  silindi' in r.get_data(as_text=True), "silinen listede gri/ustu cizili")
    r = ofadmin.post(f"/kayit/{k2['id']}/kt", {"deger": "1"}); check(r.status_code == 400, "silinmis kayda KT verilemez")
    r = ofadmin.post(f"/kayit/{k2['id']}/geri-al"); check(r.status_code == 200 and db.kayit_al(conn, k2["id"])["silindi"] == 0, "geri al")
    r = ofkt.post(f"/kayit/{k2['id']}/geri-al"); check(r.status_code == 403, "geri al 'sil' yetkisi ister")

    # ==================== yorumlar ====================
    db.bildirim_okundu(conn, "ofadmin"); db.bildirim_okundu(conn, "patron")
    r = ofoku.post(f"/kayit/{k1['id']}/yorum", {"metin": "  kontrol ettim  "}); d = json_al(r)
    check(r.status_code == 200 and d["yorum"]["metin"] == "kontrol ettim" and d["yorum"]["duzenlenebilir"] and d["yorum"]["silinebilir"], "OF:oku yorum yazar (oku yetkisi yeter)")
    yid = d["yorum"]["id"]
    r = ofoku.post(f"/kayit/{k1['id']}/yorum", {"metin": "   "}); check(r.status_code == 400, "bos yorum 400")
    check(db.okunmamis_bildirim_sayisi(conn, "ofadmin") == 1 and db.okunmamis_bildirim_sayisi(conn, "ofkt") >= 1 and db.okunmamis_bildirim_sayisi(conn, "paadmin") == 2, "yorum bildirimi yalniz ilgililere (duzenleyen/onaylayan), PA'ya degil")
    r = patron.get("/proje/OF/tablo?durum=hepsi"); check('yorum-rozet okunmamis">\U0001f4ac 1' in r.get_data(as_text=True), "tabloda okunmamis yorum rozeti kalin")
    r = patron.get("/kayitlar?durum=hepsi"); check('yorum-rozet okunmamis' in r.get_data(as_text=True), "kart listesinde okunmamis rozeti")
    d = json_al(patron.get(f"/kayit/{k1['id']}.json")); check(d["yorum_sayisi"] == 1 and d["okunmamis_yorum"] == 1 and d["arsiv_sayisi"] >= 4, "kayit json: yorum sayisi, okunmamis, arsiv kopya sayisi")
    d = json_al(patron.get(f"/kayit/{k1['id']}/yorumlar.json")); check(d["sayi"] == 1 and d["yorumlar"][0]["kullanici"] == "ofoku" and not d["yorumlar"][0]["duzenlenebilir"] and d["yorumlar"][0]["silinebilir"], "yorum listesi: admin silebilir ama duzenleyemez")
    check(db.okunmamis_yorumlar(conn, "patron", [k1["id"]]) == {}, "yorum sekmesi acilinca okundu sayildi")
    r = patron.get("/proje/OF/tablo?durum=hepsi"); check('yorum-rozet ">\U0001f4ac 1' in r.get_data(as_text=True), "okunduktan sonra rozet kalin degil")
    r = ofkt.post(f"/yorum/{yid}/duzenle", {"metin": "x"}); check(r.status_code == 403, "baskasinin yorumu duzenlenemez")
    r = ofoku.post(f"/yorum/{yid}/duzenle", {"metin": "kontrol ettim (d\u00fczeltme)"}); check(r.status_code == 200 and db.yorumlar(conn, k1["id"])[0]["duzenleme_zamani"], "yazar 15 dk icinde duzeltir")
    d = json_al(ofoku.get(f"/kayit/{k1['id']}/yorumlar.json")); check(d["yorumlar"][0]["metin"].endswith("(d\u00fczeltme)"), "duzeltilmis metin listede")
    r = ofkt.post(f"/yorum/{yid}/sil"); check(r.status_code == 403, "baskasi (admin degil) silemez")
    r = patron.post(f"/yorum/{yid}/sil"); check(r.status_code == 200 and db.kayit_al(conn, k1["id"])["yorum_sayisi"] == 0, "admin siler, sayac duser")
    d = json_al(ofoku.get(f"/kayit/{k1['id']}/yorumlar.json")); check(d["yorumlar"][0]["silindi"] == 1 and d["yorumlar"][0]["metin"] == "" and d["sayi"] == 0, "silinen yorum yer tutucu olarak gelir, metni yok")
    r = ofoku.post(f"/kayit/{k1['id']}/yorum", {"metin": "ikinci"}); yid2 = json_al(r)["yorum"]["id"]
    r = ofoku.post(f"/yorum/{yid2}/sil"); check(r.status_code == 200, "yazar kendi yorumunu siler")
    r = patron.get("/proje/OF/tablo?durum=hepsi&yorumlu=1"); h = r.get_data(as_text=True)
    db.yorum_ekle(conn, k3["id"], "patron", "son yorum metni")
    r = patron.get("/proje/OF/tablo?durum=hepsi&yorumlu=1"); h = r.get_data(as_text=True)
    check(f'data-id="{k3["id"]}"' in h and f'data-id="{k1["id"]}"' not in h, "'Yorumlu olanlar' filtresi")

    # ==================== gecmis ====================
    d = json_al(patron.get(f"/kayit/{k1['id']}/gecmis.json")); m = [x["metin"] for x in d["gecmis"]]
    check(any("D\u00fczenle: \u00dcye ID 0495682791 \u2192 490214" in x and " ofadmin " in x for x in m), f"gecmis cumlesi: duzenle {m[:3]}")
    check(any("KT \u2714" in x and "ofkt" in x for x in m) and any("KT kald\u0131r\u0131ld\u0131" in x for x in m), "gecmis cumlesi: KT / KT kaldirildi")
    check(any(x.endswith("Eklendi: 215 ARK") for x in m) and re.match(r"\d\d\.\d\d \d\d:\d\d ", m[0]), "gecmis cumlesi: eklendi + kisa tarih")
    check(d["arsiv_sayisi"] >= 5 and d["arsiv_url"] == f"/kayit/{k1['id']}/arsiv.json", "gecmis: arsiv kopya sayisi + link (admin)")
    m2 = [x["metin"] for x in json_al(patron.get(f"/kayit/{k2['id']}/gecmis.json"))["gecmis"]]
    check(any("Silindi (neden: m\u00fckerrer)" in x and "ofadmin" in x for x in m2) and any("Geri al\u0131nd\u0131" in x for x in m2), "gecmis cumlesi: silindi / geri alindi")
    d3 = json_al(ofoku.get(f"/kayit/{k1['id']}/gecmis.json")); check(d3["arsiv_url"] is None, "arsiv linki yalniz super admine")
    r = ofoku.get(f"/kayit/{k1['id']}/arsiv.json"); check(r.status_code == 403, "arsiv json normal kullaniciya 403")
    d4 = json_al(patron.get(f"/kayit/{k1['id']}/arsiv.json")); check([x["olay"] for x in d4["kopyalar"]][:2] == ["olustur", "kt"] and d4["kopyalar"][0]["veri"]["uye_id"] == "0495682791", "arsiv json: kopyalar ilk halden itibaren")

    # ==================== bildirimler (zil) ====================
    d = json_al(ofadmin.get("/bildirim.json?son=-1")); check(d["sayi"] >= 1 and d["liste"][0]["tur"] in ("yorum", "silindi") and d["liste"][0]["link"].startswith("/kayitlar?proje=OF&durum=hepsi&ac="), "bildirim listesi + link")
    son = d["maks"]
    db.yorum_ekle(conn, k1["id"], "patron", "yeni bir yorum")
    d = json_al(ofadmin.get(f"/bildirim.json?son={son}")); check(len(d["yeni"]) == 1 and d["yeni"][0]["tur"] == "yorum" and d["maks"] > son, "son id'den sonraki yeni bildirim 'yeni' listesinde")
    bid = d["yeni"][0]["id"]
    r = ofadmin.post("/bildirim/okundu", {"id": str(bid), "okundu": "1"}); check(json_al(r)["ok"] and db.bildirimler(conn, "ofadmin", 1)[0]["okundu"] == 1, "tek bildirim okundu")
    r = ofadmin.post("/bildirim/okundu", {"id": str(bid), "okundu": "0"}); check(db.bildirimler(conn, "ofadmin", 1)[0]["okundu"] == 0, "tek bildirim yeniden okunmadi")
    r = ofadmin.post("/bildirim/tumu-okundu"); check(json_al(r)["sayi"] == 0 and db.okunmamis_bildirim_sayisi(conn, "ofadmin") == 0, "tumunu okundu say")
    r = ofadmin.post("/bildirim/okundu", {"id": "x"}); check(r.status_code == 400, "gecersiz bildirim id 400")
    r = ofadmin.get("/"); check('id="zil"' in r.get_data(as_text=True) and 'zil-sayi gizli' in r.get_data(as_text=True), "zil rozeti sayfada (0 iken gizli)")
    r = ofadmin.get("/bildirim.json"); check(r.headers.get("Cache-Control") == "no-store", "json cevaplari onbelleklenmez")

    # ==================== gorseller ====================
    r = patron.get(f"/kayit/{k1['id']}/gorsel"); check(r.status_code == 200 and r.mimetype == "image/png" and len(r.data) == len(png_bayt()), "yerel gorsel servis edilir")
    r = patron.get(f"/kucuk/{k1['id']}"); check(r.status_code == 200 and r.mimetype == "image/jpeg" and len(r.data) < len(png_bayt()), "kucuk resim (jpeg) uretilir")
    r = patron.get(f"/kayit/{k4['id']}/gorsel"); check(r.status_code == 200 and r.mimetype == "image/svg+xml" and r.headers.get("X-Gorsel") == "yok", "gorseller disina cikan yol (../kanit.db) servis edilmez -> yer tutucu")
    db.kayit_guncelle(conn, k4["id"], {"gorsel_dosya": "/etc/passwd"}, "patron")
    r = patron.get(f"/kucuk/{k4['id']}"); check(r.headers.get("X-Gorsel") == "yok", "mutlak yol da servis edilmez")
    r = patron.get(f"/kucuk/{k2['id']}"); check(r.mimetype == "image/svg+xml" and "eski Gyazo linki" in r.get_data(as_text=True), "GYAZO_PROXY=0: yalniz gyazo linki olan kayit -> 'gorsel yok (eski Gyazo linki)' yer tutucu")
    d = json_al(patron.get(f"/kayit/{k2['id']}.json")); check(d["gorsel_kaynak"] == "gyazo" and d["gorsel_url"] == f"/kayit/{k2['id']}/gorsel", "kayit json gyazo kaynagi")
    a["GYAZO_PROXY"] = "1"
    cagri = []
    eski_indir = W_KAYIT.gyazo_indir
    W_KAYIT.gyazo_indir = lambda url, hedef, zaman_asimi=10: (cagri.append(url), None)[1]
    r = patron.get(f"/kucuk/{k2['id']}"); check(r.headers.get("X-Gorsel") == "yok" and cagri == ["https://i.gyazo.com/" + "b" * 32 + ".jpg"], "vekil acik, indirme basarisiz -> yer tutucu; adres i.gyazo.com")
    r = patron.get(f"/kucuk/{k2['id']}"); check(len(cagri) == 1, "basarisiz indirme 10 dk yeniden denenmez")
    W_KAYIT._gyazo_basarisiz.clear()

    def sahte_indir(url, hedef, zaman_asimi=10):
        os.makedirs(os.path.dirname(hedef), exist_ok=True)
        with open(hedef, "wb") as f:
            f.write(png_bayt())
        return hedef
    W_KAYIT.gyazo_indir = sahte_indir
    r = patron.get(f"/kayit/{k2['id']}/gorsel"); check(r.status_code == 200 and r.data == png_bayt() and os.path.isfile(os.path.join(a["ONBELLEK"], "gyazo", "b" * 32 + ".jpg")), f"vekil indirdi: onbellek/gyazo/<id>.jpg ({r.status_code} {r.mimetype})")
    W_KAYIT.gyazo_indir = eski_indir
    check(W_KAYIT.gyazo_indir("https://evil.example/" + "b" * 32 + ".jpg", os.path.join(TD, "x.jpg")) is None, "gyazo disi adres indirilmez (ag yok)")
    a["GYAZO_PROXY"] = "0"
    r = ofoku.get(f"/kucuk/{k1['id']}"); check(r.status_code == 200, "OF:oku gorseli gorur (oku -> gorsel)")

    # ==================== Excel ciktisi ====================
    r = patron.get("/disa-aktar.xlsx?durum=hepsi&proje=OF"); check(r.status_code == 200 and r.mimetype.endswith("spreadsheetml.sheet") and "attachment" in r.headers.get("Content-Disposition", ""), "xlsx indirme")
    wb = openpyxl.load_workbook(io.BytesIO(r.data)); ws = wb["Kayitlar"]
    basliklar = [c.value for c in ws[1]]
    check(basliklar == ["Proje", "Kategori", "Personel kodu", "Uye ID", "Ana uye ID", "Tur", "Tarih", "Link", "Onay (KT)", "Onaylayan", "Onay zamani", "Ekleyen (Telegram)", "Ekleyen ID", "Aciklama", "Durum", "Excel satiri", "Gorsel dosyasi", "Gorsel 2", "Yorumlar", "Son yorum"], f"xlsx sutunlari {basliklar}")
    satirlar = {r_[2]: r_ for r_ in ws.iter_rows(min_row=2, values_only=True)}
    check(len(satirlar) == 4 and "=HYPERLINK(\"http://kotu\")" in satirlar, "filtredeki tum OF satirlari")
    for row in ws.iter_rows(min_row=2):
        for c_ in row:
            if isinstance(c_.value, str) and c_.value[:1] in "=+-@":
                check(c_.data_type == "s", f"formul korumasi: {c_.value[:20]!r} metin olarak yazildi")
    k3s = satirlar["=HYPERLINK(\"http://kotu\")"]
    check(k3s[7] == "+1234" and k3s[13] == "-eksi ile baslar" and k3s[18] == 1 and k3s[19].startswith("patron (") and k3s[19].endswith(": son yorum metni"), "link/aciklama/yorum sutunlari")
    check(satirlar["215"][18] == 1 and "yeni bir yorum" in satirlar["215"][19], "yorum sayisi silinenleri saymaz, son yorum metni")
    check(wb["Ozet"]["A2"].value == "OF" and wb["Ozet"]["B2"].value == 4, "ozet sayfasi")
    r = patron.get("/disa-aktar.xlsx?durum=silinen"); wb2 = openpyxl.load_workbook(io.BytesIO(r.data)); check(wb2["Kayitlar"].max_row == 1, "silinen filtresi: su an silinen yok")
    r = patron.get("/disa-aktar"); check(r.status_code == 200 and "T\u00fcm kitap (eski d\u00fczen)" in r.get_data(as_text=True), "cikti sayfasi + tum kitap dugmesi")
    r = patron.get("/tam-kitap.xlsx"); check(r.status_code in (200, 302), "tam kitap: modul varsa dosya, yoksa uyari ile yonlendirme")
    r = ofoku.get("/tam-kitap.xlsx"); check(r.status_code == 403, "tam kitap yalniz admin / tum proje yetkisi")

    # ==================== liste kontrolu ====================
    r = patron.post("/liste-kontrol", {"adim": "yukle", "metin": "490214\n0000700\n999999"}); check(r.status_code == 302 and "token=" in r.headers["Location"], "liste yukleme -> token")
    token = r.headers["Location"].split("token=")[1]
    r = patron.get(f"/liste-kontrol?token={token}"); check(r.status_code == 200 and "S\u00fctun 1" in r.get_data(as_text=True), "esleme sayfasi")
    r = patron.post("/liste-kontrol", {"adim": "kontrol", "token": token, "kontrol": "0", "alan_0": "uye_id", "durum": "hepsi"}); h = r.get_data(as_text=True)
    check(r.status_code == 200 and re.search(r'<b>2</b><span>bulunan', h) and re.search(r'<b>1</b><span>bulunamayan', h), "liste kontrolu: 2 bulunan (norm: 0000700 = 700), 1 yok")
    check("999999" in h and f"ac={k1['id']}" in h, "bulunamayan satir + bulunan kaydin linki")
    r = patron.get(f"/liste-kontrol.xlsx?token={token}"); wb3 = openpyxl.load_workbook(io.BytesIO(r.data))
    check(wb3["Bulunamayanlar"].max_row == 2 and wb3["Bulunanlar"].max_row == 3, "liste kontrolu excel: bulunamayan 1, bulunan 2 satir")
    r = ofoku.get(f"/liste-kontrol?token={token}"); check(r.status_code == 302, "baskasinin token'i acilmaz")
    r = patron.post("/liste-kontrol", {"adim": "kontrol", "token": token, "durum": "hepsi"}); check("en az bir s\u00fctun" in r.get_data(as_text=True), "sutun secilmezse hata")

    # ==================== bot sayfasi ====================
    db.tg_guncelle(conn, 555, username="ayse", first_name="Ayse", last_seen=db.simdi())
    db.oturum_yaz(conn, 555, {"adim": "kategori", "proje": "of"})
    r = patron.post("/bot", {"islem": "onayla", "tg_id": "555", "ad": "ayse"}); check(r.status_code == 200 and db.tg_al(conn, "555")["onayli"] == 1 and db.oturum_al(conn, 555) is None, "bot: onayla aninda uygulanir (oturum da silinir)")
    r = patron.post("/bot", {"islem": "engelle", "tg_id": "555", "sebep": "spam"}); check(db.tg_al(conn, "555")["engelli"] == 1 and db.tg_al(conn, "555")["engel_sebep"] == "spam", "bot: engelle")
    r = patron.get("/bot?tg_durum=engelli"); check('data-tg-detay="555"' in r.get_data(as_text=True), "bot: durum filtresi engelli")
    r = patron.get("/bot?tg_durum=onaysiz"); check('data-tg-detay="555"' not in r.get_data(as_text=True), "bot: durum filtresi onaysiz (engelli haric)")
    r = patron.post("/bot", {"islem": "engel_kaldir", "tg_id": "555"}); check(db.tg_al(conn, "555")["engelli"] == 0 and db.tg_al(conn, "555")["onayli"] == 1, "bot: engel kaldir onayi korur")
    r = patron.post("/bot", {"islem": "onay_kaldir", "tg_id": "555"}); check(db.tg_al(conn, "555")["onayli"] == 0, "bot: onay kaldir")
    db.tg_onayla(conn, 555, "ayse"); db.tg_engelle(conn, 555, "x")
    r = patron.post("/bot", {"islem": "tam_sifirla", "tg_id": "555"}); u = db.tg_al(conn, "555"); check(u["onayli"] == 0 and u["engelli"] == 0 and u["deneme"] == 0, "bot: tam sifirla")
    r = patron.post("/bot", {"islem": "onayla", "tg_id": "abc"}); check("rakam" in r.get_data(as_text=True), "bot: gecersiz tg id")
    db.oturum_yaz(conn, 556, {"adim": "personel", "proje": "pa"}); db.oturum_yaz(conn, 557, {"adim": "ozet"})
    r = patron.post("/bot", {"islem": "oturum_temizle", "tg_id": "556"}); check(db.oturum_al(conn, 556) is None and db.oturum_al(conn, 557) is not None, "bot: tek oturum temizle")
    r = patron.post("/bot", {"islem": "oturum_temizle"}); check(db.oturum_al(conn, 557) is None, "bot: tum oturumlari temizle")
    r = patron.post("/bot", {"islem": "kategori_ekle", "code": "bonus", "label": "BONUS", "mode": "tek_uye"}); check(db.kategori_bul(conn, "bonus")["akis"] == "tek_uye" and "<code>bonus</code>" in r.get_data(as_text=True), "bot: kategori ekle")
    r = patron.post("/bot", {"islem": "kategori_ekle", "code": "bonus", "label": "BONUS2", "mode": "ana_only"}); check(db.kategori_bul(conn, "bonus")["etiket"] == "BONUS2", "bot: kategori guncelle")
    r = patron.post("/bot", {"islem": "kategori_ekle", "code": "x!", "label": "X", "mode": "tek_uye"}); check("Hata" in r.get_data(as_text=True), "bot: gecersiz kategori kodu")
    r = patron.post("/bot", {"islem": "kategori_sil", "code": "bonus"}); check(db.kategori_bul(conn, "bonus")["aktif"] == 0, "bot: kategori sil (pasif)")
    r = patron.post("/bot", {"islem": "bakim_ac"}); check(db.ayar_al(conn, "bakim_modu") == "1" and "Bak\u0131m modunu kapat" in r.get_data(as_text=True), "bot: bakim modu ac")
    r = patron.post("/bot", {"islem": "bakim_kapat"}); check(db.ayar_al(conn, "bakim_modu") == "0", "bot: bakim modu kapat")
    r = patron.post("/bot", {"islem": "gyazo_bakim_ac"}); check(db.ayar_al(conn, "gyazo_bakim") == "1", "bot: gyazo bakim ac")
    r = patron.post("/bot", {"islem": "gyazo_bakim_kapat"}); check(db.ayar_al(conn, "gyazo_bakim") == "0", "bot: gyazo bakim kapat")
    r = patron.get("/bot"); check("bilinmiyor" in r.get_data(as_text=True), "bot durumu: nabiz yokken bilinmiyor")
    db.ayar_yaz(conn, "bot_nabiz", db.simdi()); db.ayar_yaz(conn, "son_hata", "Telegram 502")
    r = patron.get("/bot"); check("\u00e7al\u0131\u015f\u0131yor" in r.get_data(as_text=True) and "Telegram 502" in r.get_data(as_text=True), "bot durumu: taze nabiz -> calisiyor + son hata")
    db.ayar_yaz(conn, "bot_nabiz", "2026-01-01 00:00:00")
    r = patron.get("/bot"); check("\u00c7ALI\u015eMIYOR" in r.get_data(as_text=True), "bot durumu: eski nabiz -> calismiyor")
    with open(os.path.join(a["LOG"], "bot.log"), "w", encoding="utf-8") as f:
        f.write("[2026-10-16 10:00:00] Mesaj geldi: ayse (ID: 555)\n[2026-10-16 10:00:01] ! Gyazo yukleme hatasi\n")
    r = patron.get("/bot"); h = r.get_data(as_text=True); check("Mesaj geldi: ayse" in h and h.count("Gyazo yukleme hatasi") == 2, "bot: son log + son hatalar")
    db.tg_kayit_sayaci(conn, 555, "ayse", "Ayse", "OF", "ark", "ARK")
    d = json_al(patron.get("/bot/kullanici/555.json"))
    check(d["id"] == "555" and d["panel"]["toplam"] == 2 and d["panel"]["proje"] == {"OF": 1, "PA": 1} and d["panel"]["capraz"]["OF"] == {"DDS": 1} and d["bot_sayac"]["kayit"] == 1 and d["oturum"] is None, f"bot kullanici detayi: panel istatistigi + bot sayaci {d['panel']}")
    db.oturum_yaz(conn, 555, {"adim": "uyeid", "proje": "of", "kategori": "ark"})
    d = json_al(patron.get("/bot/kullanici/555.json")); check(d["oturum"]["adim"] == "uyeid", "bot kullanici detayi: canli oturum")
    r = patron.get("/bot/kullanici/abc.json"); check(r.status_code == 404, "gecersiz tg id json 404")
    r = patron.get("/bot?tg_q=ayse"); check('data-tg-detay="555"' in r.get_data(as_text=True), "bot: ad aramasi")

    # ==================== arsiv ====================
    arsiv_once = conn.execute("SELECT COUNT(*) FROM arsiv").fetchone()[0]
    r = patron.get("/arsiv"); h = r.get_data(as_text=True)
    check(r.status_code == 200 and h.count("<tr class=") >= 6 and f'ac={k2["id"]}' in h, "arsiv: eklenen her kayit listede")
    check("<b>215</b>" in h and "0495682791" in h and "490214" in h and "de\u011fi\u015fti" in h, "arsiv: ilk hali (uye 0495682791) + simdiki hali (490214) + 'degisti'")
    r = patron.get("/arsiv?proje=PA"); h2 = r.get_data(as_text=True); check(f'ac={k4["id"]}' in h2 and f'ac={k1["id"]}' not in h2, "arsiv: proje filtresi")
    r = patron.get("/arsiv?q=0495682791"); check(f'ac={k1["id"]}' in r.get_data(as_text=True), "arsiv: q filtresi (norm)")
    db.kayit_sil(conn, k5["id"], "patron", "test")
    r = patron.get("/arsiv?proje=PA"); check(f'ac={k5["id"]}' in r.get_data(as_text=True) and "silindi" in r.get_data(as_text=True), "arsiv: silinen kayit da listede")
    check(conn.execute("SELECT COUNT(*) FROM arsiv").fetchone()[0] == arsiv_once + 1 and conn.execute("SELECT veri FROM arsiv WHERE kayit_id=? ORDER BY id LIMIT 1", (k1["id"],)).fetchone()[0].find("0495682791") > 0, "arsiv sayfasi arsivi degistirmez; ilk kopya korunur")
    r = patron.post("/arsiv/yedek"); check(r.status_code == 200 and r.data[:15] == b"SQLite format 3" and os.listdir(a["YEDEK"]), "yedek al: sqlite dosyasi indirilir ve VERI/yedek'e yazilir")
    r = patron.post("/bot", {"islem": "yedek"}); check(r.status_code == 200 and r.data[:15] == b"SQLite format 3", "bot sayfasindan yedek")

    # ==================== denetim / diger ====================
    r = patron.get("/denetim?islem=kt"); h = r.get_data(as_text=True); check("<code>kt</code>" in h and "<code>sil</code>" not in h, "denetim: islem filtresi")
    r = patron.get("/denetim?kullanici=ofoku"); h = r.get_data(as_text=True); check('>ofoku</td>' in h and '>patron</td>' not in h, "denetim: kullanici filtresi")
    r = patron.get(f"/denetim?q={k1['id']}"); check(f"#{k1['id']}" in r.get_data(as_text=True), "denetim: kayit no aramasi")
    r = patron.get("/"); h = r.get_data(as_text=True); check('data-sayac="OF:bekleyen">4<' in h and 'data-sayac="PA:silinen">1<' in h and 'data-sayac="OF:toplam">4<' in h, "ana sayfa proje kartlari sayaclari")
    r = ofoku.get("/"); h = r.get_data(as_text=True); check('data-sayac="OF:toplam"' in h and 'data-sayac="PA:toplam"' not in h, "ana sayfada yalniz izinli proje kartlari")
    r = patron.get("/kayitlar?gorunum=galeri&durum=hepsi"); check('class="galeri"' in r.get_data(as_text=True), "galeri gorunumu")
    r = patron.get("/kayitlar?gorunum=tablo&durum=hepsi"); check('kayit-tablo' in r.get_data(as_text=True), "tablo gorunumu (kayitlar)")
    r = patron.post("/yonetim/kullanicilar", {"islem": "pasif", "kullanici_adi": "ofkt"}); check(db.kullanici_bul(conn, "ofkt")["aktif"] == 0, "kullanici pasife alma")
    r = ofkt.get("/kayitlar"); check(r.status_code == 302, "pasif kullanici oturumu duser")
    r = patron.post("/yonetim/kullanicilar", {"islem": "pasif", "kullanici_adi": "patron"}); check("Kendi hesab" in r.get_data(as_text=True) and db.kullanici_bul(conn, "patron")["aktif"] == 1, "kendi hesabini pasife alamaz")
    r = patron.post("/yonetim/kullanicilar", {"islem": "yetki", "kullanici_adi": "ofoku", "yetki": "OF,PA:oku"}); check(db.kullanici_bul(conn, "ofoku")["yetki"] == "OF,PA:oku", "yetki guncelleme")
    r = ofoku.get("/proje/PA/tablo?durum=hepsi"); check(r.status_code == 200, "yeni yetki hemen gecerli")
    r = patron.post("/yonetim/kullanicilar", {"islem": "sifre", "kullanici_adi": "ofoku"}); check("ge\u00e7ici \u015fifre" in r.get_data(as_text=True) and db.kullanici_bul(conn, "ofoku")["sifre_degistir"] == 1, "sifre sifirlama")
    r = patron.get("/cikis"); check(r.status_code == 302, "cikis")
    r = patron.get("/"); check(r.status_code == 302 and "/giris" in r.headers["Location"], "cikistan sonra giris gerekli")
    js = open(os.path.join(os.path.dirname(W_KAYIT.__file__), "static", "panel.js"), encoding="utf-8").read()
    sablonlar = "".join(open(os.path.join(os.path.dirname(W_KAYIT.__file__), "templates", f), encoding="utf-8").read() for f in os.listdir(os.path.join(os.path.dirname(W_KAYIT.__file__), "templates")))
    check(not re.search(r'\son(click|error|change|submit|load)=', sablonlar), "sablonlarda satir ici olay isleyici yok")
    check("Ctrl+Enter" in js or "ctrlKey" in js, "yorum Ctrl+Enter")
    print(f"\nWEB TESTLERI GECTI ({ok})")


if __name__ == "__main__":
    calistir()
else:
    def test_web():
        calistir()

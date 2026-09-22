import asyncio
import os
import csv
import stat
import base64
import hashlib
import tempfile
import time
from types import SimpleNamespace

from openpyxl import load_workbook

import bot


class FakeDB:
    def __init__(self, settings=None, users=None):
        self.settings = settings or {}
        self.users = users or []

    def get_setting(self, key):
        return self.settings.get(key)

    def get_all_users(self):
        return self.users

    def get_active_users(self, hours):
        return self.users

    def get_users_by_segment(self, segment):
        return self.users


class FakeTelegramBot:
    def __init__(self):
        self.calls = []

    async def send_photo(self, **kwargs):
        self.calls.append(("photo", kwargs))

    async def send_video(self, **kwargs):
        self.calls.append(("video", kwargs))

    async def send_message(self, **kwargs):
        self.calls.append(("message", kwargs))


async def test_campaign_media_variants():
    original_db = bot.db
    try:
        fake_tg = FakeTelegramBot()
        context = SimpleNamespace(bot=fake_tg)

        bot.db = FakeDB({
            "campaign_media_id": "video-file-id",
            "campaign_media_type": "video",
            "campaign_title": "Gunun Kampanyasi",
            "promo_caption": "Kampanyayi incele",
            "caption_link_url": "https://example.com/detail",
            "caption_link_text": "TIKLA",
            "campaign_code": "HIT250",
            "bonus_btn_text": "Aktif Et",
            "bonus_link": "https://example.com/campaign",
        })
        await bot.send_campaign_media(context, 1001)
        assert fake_tg.calls[-1][0] == "video"
        caption = fake_tg.calls[-1][1]["caption"]
        assert '<a href="https://example.com/detail">TIKLA</a>' in caption
        assert "<code>HIT250</code>" in caption

        fake_tg.calls.clear()
        bot.db = FakeDB({
            "campaign_media_url": "https://example.com/banner.jpg",
            "campaign_media_type": "photo",
            "campaign_title": "Kampanya",
            "promo_caption": "Aciklama",
            "bonus_link": "https://example.com/campaign",
        })
        await bot.send_campaign_media(context, 1002)
        assert fake_tg.calls[-1][0] == "photo"

        fake_tg.calls.clear()
        bot.db = FakeDB({
            "campaign_media_url": "https://example.com/tanitim.mp4",
            "campaign_title": "Video Kampanya",
            "promo_caption": "Video aciklamasi",
            "bonus_link": "https://example.com/campaign",
        })
        await bot.send_campaign_media(context, 10025)
        assert fake_tg.calls[-1][0] == "video"

        fake_tg.calls.clear()
        bot.db = FakeDB({
            "campaign_media_file": "bulunmayan.jpg",
            "campaign_media_url": "",
            "campaign_title": "Kampanya",
            "promo_caption": "Medya olmasa da gonder",
            "bonus_link": "https://example.com/campaign",
        })
        await bot.send_campaign_media(context, 1003)
        assert fake_tg.calls[-1][0] == "message"
    finally:
        bot.db = original_db


def test_media_library_database():
    original_path = bot.DB_PATH
    original_db = bot.db
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "media_test.db")
            test_db = bot.Database()
            bot.db = test_db

            photo_id = test_db.add_campaign_media("photo", "photo-file-id", "Birinci Foto", 100)
            video_id = test_db.add_campaign_media("video", "video-file-id", "Tanitim Videosu", 100)
            duplicate_id = test_db.add_campaign_media("photo", "photo-file-id", "Tekrar", 100)

            assert photo_id != video_id
            assert duplicate_id == photo_id
            assert test_db.count_campaign_media() == 2
            assert [row["id"] for row in test_db.list_campaign_media()] == [video_id, photo_id]

            video_row = test_db.get_campaign_media(video_id)
            assert bot.activate_campaign_media_record(video_row)
            assert test_db.get_setting("campaign_media_id") == "video-file-id"
            assert test_db.get_setting("campaign_media_type") == "video"
            assert bot.active_campaign_media_library_id() == video_id

            assert test_db.delete_campaign_media(photo_id)
            assert test_db.count_campaign_media() == 1
            assert test_db.get_campaign_media(photo_id) is None
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path


def test_text_and_media_helpers():
    assert bot._media_kind("https://cdn.example.com/a.MP4") == "video"
    assert bot._media_kind("banner.jpeg") == "photo"
    assert bot.personalize_text("Merhaba {name}", "<Ali>") == "Merhaba &lt;Ali&gt;"

    raw = bot.sanitize_input("Kanal 👉 [[TIKLA|https://t.me/ornek]]", 500)
    rendered = bot.render_clickable_tokens(raw)
    assert rendered == 'Kanal 👉 <a href="https://t.me/ornek">TIKLA</a>'

    payload = {
        "title": "Bonus",
        "text": raw,
        "type": "promocode",
        "code": "HIT250",
        "code_style": "copy",
        "btn_text": "Aktif Et",
        "btn_url": "https://example.com/campaign",
    }
    body = bot.build_payload_text(payload)
    assert "<b>Bonus</b>" in body
    assert "<code>HIT250</code>" in body
    assert bot.build_payload_kb(payload) is not None

    payload["code_style"] = "spoiler"
    assert "<tg-spoiler>HIT250</tg-spoiler>" in bot.build_payload_text(payload)


def test_rich_tokens_and_full_payload():
    # Vurgu isaretleri HTML'e cevrilir
    assert bot.render_rich_tokens("**kalin** ve __italik__") == "<b>kalin</b> ve <i>italik</i>"
    assert bot.render_rich_tokens("++alt++ ~~ust~~ ||gizli||") == "<u>alt</u> <s>ust</s> <tg-spoiler>gizli</tg-spoiler>"
    # Sanitize edilmis kullanici metni HTML enjekte edemez
    raw = bot.sanitize_input("<script>x</script> **onemli**", 200)
    out = bot.render_rich_tokens(raw)
    assert "<script>" not in out and "<b>onemli</b>" in out
    # Vurgu donusumu link href'lerini BOZMAZ (etiket ici korunur)
    raw = bot.sanitize_input("[[TIKLA|https://x.com/__abc__/promo]] **onemli**", 300)
    out = bot.render_rich_tokens(bot.render_clickable_tokens(raw))
    assert '<a href="https://x.com/__abc__/promo">TIKLA</a>' in out
    # C++ gibi kelimeye bitisik ++ vurgu sayilmaz
    assert bot.render_rich_tokens("C++ ve Java++") == "C++ ve Java++"

    # Tam paket: medya + coklu kod + iki buton
    p = {
        "type": "full", "title": "Buyuk Kampanya", "text": "**Kacirma!**",
        "codes": ["VIP100", "GOLD50"], "code_style": "copy",
        "media_id": "file123", "media_type": "photo",
        "btn_text": "Kayit Ol", "btn_url": "https://example.com/r",
        "btn2_text": "Giris", "btn2_url": "https://example.com/l",
    }
    body = bot.build_payload_text(p)
    assert "<b>Buyuk Kampanya</b>" in body
    assert "<b>Kacirma!</b>" in body
    assert "<code>VIP100</code>" in body and "<code>GOLD50</code>" in body
    assert "Kodlari" in body  # coklu kod basligi
    kb = bot.build_payload_kb(p)
    row = kb.inline_keyboard[0]
    assert len(row) == 2 and row[0].text == "Kayit Ol" and row[1].text == "Giris"

    # Eski tekil kod kaydi ayni sekilde calismaya devam eder
    old = {"type": "promocode", "text": "t", "code": "HIT250", "code_style": "copy",
           "btn_text": "A", "btn_url": "https://example.com"}
    assert "<code>HIT250</code>" in bot.build_payload_text(old)
    assert bot.build_payload_kb(old) is not None


def test_scheduled_prenotify_and_retime():
    original_path = bot.DB_PATH
    original_db = bot.db
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "sched_test.db")
            test_db = bot.Database()
            bot.db = test_db
            from datetime import datetime, timedelta
            soon = (datetime.now() + timedelta(minutes=3)).isoformat()
            later = (datetime.now() + timedelta(hours=2)).isoformat()
            s1 = test_db.add_scheduled('{"type":"text","text":"yakin"}', "all", soon)
            s2 = test_db.add_scheduled('{"type":"text","text":"uzak"}', "all", later)

            # 5 dk penceresinde yalnizca yakin olan hatirlatilir
            pre = test_db.get_prenotify_scheduled(5)
            assert [r["id"] for r in pre] == [s1]
            test_db.mark_scheduled_notified(s1)
            assert test_db.get_prenotify_scheduled(5) == []

            # Zaman degistirilince hatirlatma bayragi sifirlanir
            new_soon = (datetime.now() + timedelta(minutes=4)).isoformat()
            test_db.update_scheduled_time(s1, new_soon)
            assert [r["id"] for r in test_db.get_prenotify_scheduled(5)] == [s1]

            # Iptal edilen kayit hatirlatilmaz ve gonderilmez
            test_db.set_scheduled_status(s1, "iptal")
            assert test_db.get_prenotify_scheduled(5) == []
            assert test_db.get_due_scheduled() == []
            assert test_db.get_scheduled(s2)["status"] == "bekliyor"
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path


def test_play_menu_uses_current_login_link():
    original_db = bot.db
    try:
        bot.db = FakeDB({"giris_link": "https://example.com/current-login"})
        menu = bot._build_play_menu_button()
        assert isinstance(menu, bot.MenuButtonWebApp)
        assert menu.text == bot.PLAY_MENU_TEXT
        assert menu.web_app.url == "https://example.com/current-login"

        # Gecersiz link: buton kurulamaz, ValueError beklenir (fallback ayri test edilir)
        bot.db = FakeDB({"giris_link": "http://guvensiz.example.com"})
        try:
            bot._build_play_menu_button()
            raise AssertionError("http linki kabul edilmemeliydi")
        except ValueError:
            pass
    finally:
        bot.db = original_db


class FakeMenuBot:
    def __init__(self):
        self.buttons = []

    async def set_chat_menu_button(self, chat_id=None, menu_button=None):
        self.buttons.append((chat_id, menu_button))


async def test_play_menu_falls_back_to_commands_on_bad_link():
    original_db = bot.db
    original_mode = bot.MENU_BUTTON_MODE
    try:
        # menu modu (varsayilan): her zaman komut menusu kurulur (⌘ simgesi icin)
        bot.MENU_BUTTON_MODE = "menu"
        bot.db = FakeDB({"giris_link": "https://example.com/login"})
        fake_bot = FakeMenuBot()
        ok = await bot._apply_play_menu_buttons(fake_bot)
        assert ok
        assert isinstance(fake_bot.buttons[0][1], bot.MenuButtonCommands)

        # play modu: gecersiz linkte komut menusune duser
        bot.MENU_BUTTON_MODE = "play"
        bot.db = FakeDB({"giris_link": "http://guvensiz.example.com"})
        fake_bot = FakeMenuBot()
        ok = await bot._apply_play_menu_buttons(fake_bot)
        assert ok
        assert fake_bot.buttons
        assert isinstance(fake_bot.buttons[0][1], bot.MenuButtonCommands)

        # play modu: gecerli https linkte web-app butonu kurulur
        bot.db = FakeDB({"giris_link": "https://example.com/login"})
        fake_bot = FakeMenuBot()
        ok = await bot._apply_play_menu_buttons(fake_bot)
        assert ok
        assert isinstance(fake_bot.buttons[0][1], bot.MenuButtonWebApp)
    finally:
        bot.db = original_db
        bot.MENU_BUTTON_MODE = original_mode


def test_site_id_validation_and_password_hash():
    assert bot.validate_site_id("SITE_123-TR") == "SITE_123-TR"
    assert bot.validate_site_id("123456") == "123456"
    assert bot.validate_site_id("a") is None
    assert bot.validate_site_id("A B C") is None
    assert bot.validate_site_id("=FORMUL") is None
    assert bot.validate_site_id("A" * 65) is None

    original_hash = bot.ADMIN_PASSWORD_HASH
    try:
        salt = b"0123456789abcdef"
        iterations = 200_000
        digest = hashlib.pbkdf2_hmac("sha256", "Guclu!Parola123".encode(), salt, iterations)
        bot.ADMIN_PASSWORD_HASH = "pbkdf2_sha256${}${}${}".format(
            iterations,
            base64.b64encode(salt).decode("ascii"),
            base64.b64encode(digest).decode("ascii"),
        )
        assert bot.verify_admin_password("Guclu!Parola123")
        assert not bot.verify_admin_password("Yanlis!Parola123")
        bot.ADMIN_PASSWORD_HASH = ""
        assert not bot.verify_admin_password("Guclu!Parola123")
    finally:
        bot.ADMIN_PASSWORD_HASH = original_hash


def test_site_id_storage_and_exports():
    original_path = bot.DB_PATH
    original_db = bot.db
    original_csv = bot.SITE_ID_CSV
    original_xlsx = bot.SITE_ID_XLSX
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "site_id_test.db")
            bot.SITE_ID_CSV = os.path.join(temp_dir, "site_id_kayitlari.csv")
            bot.SITE_ID_XLSX = os.path.join(temp_dir, "site_id_kayitlari.xlsx")
            test_db = bot.Database()
            bot.db = test_db

            test_db.upsert_site_id(1001, "=TEHLIKELI", "Ali", "Veli", "SITE_1001")
            test_db.upsert_site_id(1002, "normal_user", "+FORMUL", "Kaya", "SITE_1002")
            test_db.upsert_site_id(1001, "ali_guncel", "Ali", "Veli", "SITE_1001_YENI")
            assert len(test_db.list_site_id_records()) == 2
            assert test_db.get_site_id_record(1001)["site_id"] == "SITE_1001_YENI"

            assert bot.export_site_id_records() == 2
            bot.secure_runtime_files()
            assert os.path.isfile(bot.SITE_ID_CSV)
            assert os.path.isfile(bot.SITE_ID_XLSX)

            with open(bot.SITE_ID_CSV, "r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            assert rows[0]["Telegram ID"] in {"1001", "1002"}
            assert any(row["Site ID"] == "SITE_1001_YENI" for row in rows)
            assert any(row["Ad"].startswith("'") for row in rows)

            workbook = load_workbook(bot.SITE_ID_XLSX, read_only=True, data_only=False)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            workbook.close()
            assert values[0] == tuple(bot.SITE_ID_HEADERS)
            assert any(str(row[2]).startswith("'") for row in values[1:])

            if os.name == "posix":
                assert stat.S_IMODE(os.stat(bot.SITE_ID_CSV).st_mode) == 0o600
                assert stat.S_IMODE(os.stat(bot.SITE_ID_XLSX).st_mode) == 0o600
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path
        bot.SITE_ID_CSV = original_csv
        bot.SITE_ID_XLSX = original_xlsx


class SilentMessage:
    def __init__(self):
        self.replies = []

    async def reply_text(self, *args, **kwargs):
        self.replies.append((args, kwargs))

    async def delete(self):
        return None


class SecurityDB:
    def __init__(self):
        self.events = []

    def security_event(self, *args):
        self.events.append(args)


async def test_unauthorized_admin_is_silent():
    original_db = bot.db
    original_ids = list(bot.ADMIN_IDS)
    try:
        fake_db = SecurityDB()
        bot.db = fake_db
        bot.ADMIN_IDS[:] = [123456789]
        message = SilentMessage()
        update = SimpleNamespace(
            effective_user=SimpleNamespace(id=987654321, username="yetkisiz_kullanici"),
            effective_chat=SimpleNamespace(id=987654321),
            message=message,
        )
        result = await bot.admin_login_cmd(update, SimpleNamespace())
        assert result == bot.ConversationHandler.END
        assert message.replies == []
        assert fake_db.events and fake_db.events[-1][0] == "unauthorized_admin"

        await bot.handle_text(update, SimpleNamespace())
        assert message.replies == []
    finally:
        bot.db = original_db
        bot.ADMIN_IDS[:] = original_ids


class FlowMessage:
    def __init__(self, text):
        self.text = text
        self.replies = []

    async def reply_text(self, *args, **kwargs):
        self.replies.append((args, kwargs))


class FlowQuery:
    def __init__(self, data):
        self.data = data
        self.answered = False
        self.edits = []

    async def answer(self):
        self.answered = True

    async def edit_message_text(self, *args, **kwargs):
        self.edits.append((args, kwargs))


async def test_start_welcomes_unregistered_before_site_id():
    """SITE_ID_ON_START=1 iken kayitsiz kullaniciya kayit akisi gosterilir."""
    original_path = bot.DB_PATH
    original_db = bot.db
    original_guard = bot.guard_flood
    original_flag = bot.SITE_ID_ON_START

    async def allow_message(*args, **kwargs):
        return False

    try:
        bot.SITE_ID_ON_START = True
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "start_flow_test.db")
            bot.db = bot.Database()
            bot.guard_flood = allow_message

            user = SimpleNamespace(id=1901, username="yeni_uye", first_name="Yeni", last_name="Uye")
            chat = SimpleNamespace(id=1901)
            message = FlowMessage("/start")
            fake_tg = FakeTelegramBot()
            context = SimpleNamespace(user_data={}, bot=fake_tg)
            update = SimpleNamespace(effective_user=user, effective_chat=chat, message=message)

            state = await bot.start_cmd(update, context)
            assert state == bot.SITE_ID_WAIT_INPUT
            # 1) GORSELLI karsilama (medya varsa foto/video caption'i, yoksa metin)
            # 2) kanal mesaji (CHANNEL_LINK doluysa)  3) Site ID istegi
            ilk = fake_tg.calls[0]
            ilk_metin = ilk[1].get("caption") or ilk[1].get("text") or ""
            assert "Merhaba" in ilk_metin and "Yeni" in ilk_metin
            # images/promo.jpg pakette bulundugundan acilis gorselli olmali
            assert ilk[0] in ("photo", "video"), f"acilis gorselsiz: {ilk[0]}"
            texts = [c[1].get("text", "") for c in fake_tg.calls if c[0] == "message"]
            if bot.DEFAULT_CHANNEL_LINK:
                assert any("abone" in t for t in texts)
            assert "Site ID" in texts[-1]
            assert "Profil" in texts[-1]

            # Kayitli kullanici: kayit akisi tekrarlanmaz, dogrudan kampanya akisi calisir
            bot.db.upsert_site_id(user.id, user.username, user.first_name, user.last_name, "SITE_1901")
            fake_tg.calls.clear()
            state = await bot.start_cmd(update, context)
            assert state == bot.ConversationHandler.END
            assert fake_tg.calls, "kayitli kullaniciya karsilama gonderilmedi"
    finally:
        bot.SITE_ID_ON_START = original_flag
        bot.guard_flood = original_guard
        bot.db = original_db
        bot.DB_PATH = original_path


async def test_start_default_flow_without_site_id():
    """Varsayilan (SITE_ID_ON_START=0): karsilama + 2 buton -> kanal -> kampanya."""
    original_path = bot.DB_PATH
    original_db = bot.db
    original_guard = bot.guard_flood
    original_flag = bot.SITE_ID_ON_START

    async def allow_message(*args, **kwargs):
        return False

    try:
        bot.SITE_ID_ON_START = False
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "start_default_test.db")
            bot.db = bot.Database()
            bot.guard_flood = allow_message

            user = SimpleNamespace(id=1902, username="misafir", first_name="Misafir", last_name=None)
            chat = SimpleNamespace(id=1902)
            message = FlowMessage("/start")
            fake_tg = FakeTelegramBot()
            context = SimpleNamespace(user_data={}, bot=fake_tg)
            update = SimpleNamespace(effective_user=user, effective_chat=chat, message=message)

            state = await bot.start_cmd(update, context)
            assert state == bot.ConversationHandler.END

            # 1) Karsilama: acik karsilama medyasi yok -> metin + 2 buton
            ilk = fake_tg.calls[0]
            assert ilk[0] == "message"
            assert "Merhaba" in ilk[1]["text"] and "Misafir" in ilk[1]["text"]
            kb = ilk[1].get("reply_markup")
            assert kb is not None and len(kb.inline_keyboard[0]) == 2, "Uye Ol + Giris butonlari eksik"

            texts = [c[1].get("text", "") for c in fake_tg.calls if c[0] == "message"]
            # 2) Kanal mesaji (CHANNEL_LINK doluysa)
            if bot.DEFAULT_CHANNEL_LINK:
                assert any("abone" in t for t in texts)
            # 3) Kampanya medyasi (promo.jpg fallback ile foto olarak gider)
            assert any(c[0] in ("photo", "video") for c in fake_tg.calls), "kampanya medyasi gonderilmedi"
            # Site ID istegi GONDERILMEZ
            assert not any("Site ID'ni gir" in t for t in texts)
    finally:
        bot.SITE_ID_ON_START = original_flag
        bot.guard_flood = original_guard
        bot.db = original_db
        bot.DB_PATH = original_path


async def test_welcome_media_priority():
    original_db = bot.db
    try:
        fake_tg = FakeTelegramBot()
        context = SimpleNamespace(bot=fake_tg)

        # 1) Kutuphaneden secilen karsilama medyasi oncelikli
        bot.db = FakeDB({"welcome_media_id": "wm-1", "welcome_media_type": "video"})
        ok = await bot._send_welcome_media(context, 1, "Merhaba")
        assert ok and fake_tg.calls[-1][0] == "video"
        assert fake_tg.calls[-1][1]["caption"] == "Merhaba"

        # 2) URL ikinci sirada
        fake_tg.calls.clear()
        bot.db = FakeDB({"welcome_media_url": "https://example.com/hosgeldin.jpg"})
        ok = await bot._send_welcome_media(context, 1, "Merhaba")
        assert ok and fake_tg.calls[-1][0] == "photo"

        # 3) Aktif kampanya medyasi ucuncu sirada
        fake_tg.calls.clear()
        bot.db = FakeDB({"campaign_media_id": "camp-1", "campaign_media_type": "photo"})
        ok = await bot._send_welcome_media(context, 1, "Merhaba")
        assert ok and fake_tg.calls[-1][0] == "photo"
        assert fake_tg.calls[-1][1]["photo"] == "camp-1"
    finally:
        bot.db = original_db


def test_campaigns_announcements_faq_db():
    from datetime import datetime, timedelta
    original_path = bot.DB_PATH
    original_db = bot.db
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "camp_test.db")
            d = bot.Database()
            bot.db = d

            gecmis = (datetime.now() - timedelta(hours=1)).isoformat(timespec="seconds")
            yakin = (datetime.now() + timedelta(hours=1)).isoformat(timespec="seconds")
            uzak = (datetime.now() + timedelta(days=3)).isoformat(timespec="seconds")

            c1 = d.add_campaign("bonus", "Hosgeldin Bonusu", "**%100** bonus", expires_at=uzak,
                                btn_text="Kap", btn_url="https://example.com/b")
            c2 = d.add_campaign("bonus", "Eski Bonus", expires_at=gecmis)
            c3 = d.add_campaign("turnuva", "Mega Turnuva", prize_pool="500.000 TL",
                                conditions="Min 100 TL yatirim", expires_at=yakin)
            c4 = d.add_campaign("ozel", "VIP Ozel", segment="platin")
            d.add_campaign("ozel", "Herkese Ozel")

            # Kullanici listesi: suresi gecen c2 GORUNMEZ (otomatik yayindan kalkar)
            aktif_bonus = [c["id"] for c in d.list_campaigns(category="bonus")]
            assert c1 in aktif_bonus and c2 not in aktif_bonus

            # Segment filtresi: platin uyesi 2 ozel kampanya gorur, segmentsiz 1
            assert len(d.list_campaigns(category="ozel", segment="platin")) == 2
            assert len(d.list_campaigns(category="ozel", segment="")) == 1

            # Bitis bildirimi adaylari: yakin (1 saat) olan c3 listede
            adaylar = [c["id"] for c in d.campaigns_needing_expiry_notice()]
            assert c3 in adaylar and c2 not in adaylar
            d.mark_campaign_notified(c3, 2)
            assert c3 not in [c["id"] for c in d.campaigns_needing_expiry_notice()]

            # Suresi dolan otomatik pasife alinir
            dropped = d.deactivate_expired_campaigns()
            assert c2 in dropped
            assert d.get_campaign(c2)["active"] == 0

            # Alan duzenleme beyaz listesi
            assert d.update_campaign_field(c1, "title", "Yeni Baslik")
            assert not d.update_campaign_field(c1, "id", 999)  # yasak alan
            assert d.get_campaign(c1)["title"] == "Yeni Baslik"

            # Turnuva detay metni odul + sartlari icerir
            detay = bot.build_campaign_detail_text(d.get_campaign(c3))
            assert "Ödül Havuzu" in detay and "Katılım Şartları" in detay and "Bitiş" in detay

            # Duyuru + SSS CRUD
            aid = d.add_announcement("Bakim", "Sistem bakimi 02:00'de")
            fid = d.add_faq("Nasil uye olurum?", "Kayit linkine tikla")
            assert d.get_announcement(aid)["title"] == "Bakim"
            assert len(d.list_faqs()) == 1
            d.delete_announcement(aid); d.delete_faq(fid)
            assert d.get_announcement(aid) is None and d.list_faqs() == []
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path


def test_auto_segments():
    from datetime import datetime, timedelta
    original_path = bot.DB_PATH
    original_db = bot.db
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "seg_test.db")
            d = bot.Database()
            bot.db = d
            d.add_user(9001, "yeni_u", "Yeni", "Uye")  # simdi katildi -> yeni + aktif
            segs = bot._user_segments(9001)
            assert "yeni" in segs and "aktif" in segs
            d.set_segment(9001, "platin")
            assert "platin" in bot._user_segments(9001)
            # yeni segmentli kampanya bu kullaniciya gorunur, gumus segmentli gorunmez
            c_yeni = d.add_campaign("bonus", "Hosgeldin Ozel", segment="yeni")
            c_gumus = d.add_campaign("bonus", "Gumus Ozel", segment="gumus")
            gorunen = [c["id"] for c in d.list_campaigns(category="bonus", segment=bot._user_segments(9001))]
            assert c_yeni in gorunen and c_gumus not in gorunen
            # kampanya kitlesi cozumleyici: yeni uyeler listesinde
            kitle = [u["user_id"] for u in bot._campaign_audience("yeni")]
            assert 9001 in kitle
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path


def test_main_menu_layout():
    original_db = bot.db
    try:
        bot.db = FakeDB({"giris_link": "https://example.com/login",
                         "miniapp_link": "https://example.com/app",
                         "deposit_link": "https://example.com/yatir",
                         "withdraw_link": "https://example.com/cek"})
        kb = bot.build_main_menu()
        rows = kb.inline_keyboard
        tum = [b for r in rows for b in r]
        yazilar = " ".join(b.text for b in tum)
        for beklenen in ("Siteye Giriş", "Mini App", "Güncel Bonuslar", "Yeni Üye",
                         "Slot", "Spor", "Bana Özel", "Turnuvalar", "Para Yatır",
                         "Para Çek", "Duyurular", "Sık Sorulan", "Canlı Destek"):
            assert beklenen in yazilar, f"eksik buton: {beklenen}"
        miniapp = [b for b in tum if "Mini App" in b.text][0]
        assert miniapp.web_app is not None and miniapp.web_app.url == "https://example.com/app"
        yatir = [b for b in tum if "Para Yatır" in b.text][0]
        assert yatir.url == "https://example.com/yatir"
    finally:
        bot.db = original_db


async def test_smart_support_routing():
    original_db = bot.db
    original_guard = bot.guard_flood
    original_ids = list(bot.ADMIN_IDS)

    async def allow_message(*args, **kwargs):
        return False

    class RouterDB(FakeDB):
        def add_user(self, *a, **k): pass
        def inc_msg(self, uid): pass
        def log(self, *a, **k): pass
        def is_banned(self, uid): return False
        def is_temp_banned(self, uid): return False

    try:
        bot.db = RouterDB({"withdraw_support_link": "https://t.me/cekim_destek",
                           "deposit_support_link": "https://t.me/odeme_destek",
                           "bonus_support_link": "https://t.me/bonus_destek"})
        bot.guard_flood = allow_message
        bot.ADMIN_IDS[:] = [999999]
        bot._cb_last_press.clear()

        senaryolar = [
            ("Param gelmedi ne zaman yatacak", "cekim_destek"),
            ("yatırım yapamıyorum yardım", "odeme_destek"),
            ("bonus alamadım hala", "bonus_destek"),
        ]
        for i, (mesaj, beklenen_link) in enumerate(senaryolar):
            bot._cb_last_press.clear()
            msg = FlowMessage(mesaj)
            upd = SimpleNamespace(effective_user=SimpleNamespace(id=6001 + i, username="u", first_name="X", last_name=None),
                                  effective_chat=SimpleNamespace(id=6001 + i, type="private"),
                                  message=msg)
            await bot.handle_text(upd, SimpleNamespace(user_data={}, bot=FakeTelegramBot()))
            assert msg.replies, f"yanit yok: {mesaj}"
            kb = msg.replies[0][1].get("reply_markup")
            assert kb is not None and beklenen_link in kb.inline_keyboard[0][0].url, mesaj
    finally:
        bot.db = original_db
        bot.guard_flood = original_guard
        bot.ADMIN_IDS[:] = original_ids


def test_reply_menu_build():
    kb = bot._build_reply_menu()
    assert kb is not None and kb.is_persistent and kb.resize_keyboard
    rows = kb.keyboard
    # Varsayilan: kompakt tam menu — 8 satir x 2 buton (son satir: Ekran Görüntüsü | Ana Menü)
    assert [len(r) for r in rows] == [2, 2, 2, 2, 2, 2, 2, 2], [len(r) for r in rows]
    tum_butonlar = " ".join(b.text for r in rows for b in r)
    assert "Ana Menü" in tum_butonlar and "Profil" in tum_butonlar
    assert "Ekran Görüntüsü" in tum_butonlar
    assert "Bonuslar" in tum_butonlar and "Canlı Destek" in tum_butonlar


async def test_menu_button_text_routing():
    """Alt klavye butonlari (Profil/Destek/Kampanya/Giris) yanit uretir; rastgele metin sessiz kalir."""
    original_db = bot.db
    original_guard = bot.guard_flood
    original_ids = list(bot.ADMIN_IDS)

    async def allow_message(*args, **kwargs):
        return False

    class RouterDB(FakeDB):
        def add_user(self, *a, **k): pass
        def inc_msg(self, uid): pass
        def get_site_id_record(self, uid): return None
        def is_banned(self, uid): return False
        def is_temp_banned(self, uid): return False

    try:
        bot.db = RouterDB({})
        bot.guard_flood = allow_message
        bot.ADMIN_IDS[:] = [999999]
        user = SimpleNamespace(id=5001, username="u", first_name="Rt", last_name=None, language_code="tr")

        def yap(text):
            msg = FlowMessage(text)
            upd = SimpleNamespace(effective_user=user,
                                  effective_chat=SimpleNamespace(id=5001, type="private"),
                                  message=msg)
            return msg, upd

        ctx = SimpleNamespace(user_data={}, bot=FakeTelegramBot())

        msg, upd = yap("🤴 Profil")
        await bot.handle_text(upd, ctx)
        assert msg.replies and "PROFİL" in msg.replies[0][0][0]

        msg, upd = yap("💬 Canlı Destek")
        await bot.handle_text(upd, ctx)
        assert msg.replies and "CANLI DESTEK" in msg.replies[0][0][0]

        msg, upd = yap("🎁 Kampanya")
        await bot.handle_text(upd, ctx)
        assert ctx.bot.calls, "kampanya medyasi gonderilmedi"

        # Serbest metin: yerlesik asistan cevap verir (selam niyeti)
        bot._cb_last_press.clear()
        msg, upd = yap("selam nasilsin")
        await bot.handle_text(upd, ctx)
        assert msg.replies, "asistan cevap vermedi"
        assert "MERHABA" in msg.replies[0][0][0], msg.replies[0][0][0][:80]
    finally:
        bot.db = original_db
        bot.guard_flood = original_guard
        bot.ADMIN_IDS[:] = original_ids


async def test_site_id_cancel_message():
    original_guard = bot.guard_flood

    async def allow_message(*args, **kwargs):
        return False

    try:
        bot.guard_flood = allow_message
        message = FlowMessage("/iptal")
        update = SimpleNamespace(
            effective_user=SimpleNamespace(id=3001, username="u", first_name="A", last_name="B"),
            effective_chat=SimpleNamespace(id=3001),
            message=message,
        )
        context = SimpleNamespace(user_data={"pending_site_id": "X1"}, bot=FakeTelegramBot())
        state = await bot.site_id_cancel(update, context)
        assert state == bot.ConversationHandler.END
        assert "pending_site_id" not in context.user_data
        assert message.replies and "İşlem iptal edildi" in message.replies[0][0][0]
    finally:
        bot.guard_flood = original_guard


def test_profile_view_registered_and_unregistered():
    original_db = bot.db
    try:
        user = SimpleNamespace(id=4001, username="profil_user", first_name="Pro", last_name="Fil")

        bot.db = FakeDB({})
        bot.db.get_site_id_record = lambda uid: None
        bot.db.is_banned = lambda uid: False
        bot.db.is_temp_banned = lambda uid: False
        text, kb = bot._build_profile_view(user)
        assert "Henüz kayıtlı değil" in text
        assert kb.inline_keyboard[0][0].callback_data == "profil_siteid"
        assert "Gir" in kb.inline_keyboard[0][0].text

        record = {"site_id": "SITE_4001", "updated_at": "2026-07-14T15:16:00"}
        bot.db.get_site_id_record = lambda uid: record
        text, kb = bot._build_profile_view(user)
        assert "SITE_4001" in text
        assert "Değiştir" in kb.inline_keyboard[0][0].text
    finally:
        bot.db = original_db


async def test_admin_password_toggle():
    original_flag = bot.ADMIN_REQUIRE_PASSWORD
    original_ids = list(bot.ADMIN_IDS)
    original_db = bot.db
    try:
        bot.db = SecurityDB()
        bot.db.admin_log = lambda *a, **k: None
        bot.ADMIN_IDS[:] = [555]
        bot.authenticated_admins.pop(555, None)

        # Parola zorunlu: oturum ACILMADAN parola sorulur
        bot.ADMIN_REQUIRE_PASSWORD = True
        message = SilentMessage()
        update = SimpleNamespace(
            effective_user=SimpleNamespace(id=555, username="admin"),
            effective_chat=SimpleNamespace(id=555),
            message=message,
        )
        state = await bot.admin_login_cmd(update, SimpleNamespace())
        assert state == bot.WAITING_PASSWORD
        assert 555 not in bot.authenticated_admins
        assert message.replies
    finally:
        bot.ADMIN_REQUIRE_PASSWORD = original_flag
        bot.ADMIN_IDS[:] = original_ids
        bot.authenticated_admins.pop(555, None)
        bot.db = original_db


async def test_site_id_requires_explicit_confirmation():
    original_path = bot.DB_PATH
    original_db = bot.db
    original_csv = bot.SITE_ID_CSV
    original_xlsx = bot.SITE_ID_XLSX
    original_guard = bot.guard_flood
    original_registered_start = bot._send_registered_start

    async def allow_message(*args, **kwargs):
        return False

    async def skip_registered_start(*args, **kwargs):
        return None

    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "flow_test.db")
            bot.SITE_ID_CSV = os.path.join(temp_dir, "site_id_kayitlari.csv")
            bot.SITE_ID_XLSX = os.path.join(temp_dir, "site_id_kayitlari.xlsx")
            bot.db = bot.Database()
            bot.guard_flood = allow_message
            bot._send_registered_start = skip_registered_start

            user = SimpleNamespace(id=2001, username="test_user", first_name="Test", last_name="Kullanici")
            chat = SimpleNamespace(id=2001)
            context = SimpleNamespace(user_data={}, bot=FakeTelegramBot())

            input_update = SimpleNamespace(
                effective_user=user,
                effective_chat=chat,
                message=FlowMessage("SITE_2001"),
            )
            state = await bot.site_id_receive(input_update, context)
            assert state == bot.SITE_ID_WAIT_CONFIRM
            assert context.user_data["pending_site_id"] == "SITE_2001"
            assert bot.db.get_site_id_record(user.id) is None
            assert not os.path.exists(bot.SITE_ID_CSV)

            no_query = FlowQuery("site_confirm_no")
            no_update = SimpleNamespace(effective_user=user, effective_chat=chat, callback_query=no_query)
            state = await bot.site_id_confirm(no_update, context)
            assert state == bot.SITE_ID_WAIT_INPUT
            assert no_query.answered
            assert bot.db.get_site_id_record(user.id) is None

            await bot.site_id_receive(input_update, context)
            yes_query = FlowQuery("site_confirm_yes")
            yes_update = SimpleNamespace(effective_user=user, effective_chat=chat, callback_query=yes_query)
            state = await bot.site_id_confirm(yes_update, context)
            assert state == bot.ConversationHandler.END
            assert yes_query.answered
            assert bot.db.get_site_id_record(user.id)["site_id"] == "SITE_2001"
            assert os.path.isfile(bot.SITE_ID_CSV)
            assert os.path.isfile(bot.SITE_ID_XLSX)
    finally:
        bot.guard_flood = original_guard
        bot._send_registered_start = original_registered_start
        bot.db = original_db
        bot.DB_PATH = original_path
        bot.SITE_ID_CSV = original_csv
        bot.SITE_ID_XLSX = original_xlsx



async def test_kitle_gonder_hata_raporu():
    """Toplu gonderimde hatalar ture gore sayilmali ve ozetlenmeli."""
    from telegram.error import Forbidden, BadRequest
    users = [{"user_id": 1}, {"user_id": 2}, {"user_id": 3}, {"user_id": 4}]

    async def send_one(uid):
        if uid == 1:
            return
        if uid in (2, 3):
            raise Forbidden("bot was blocked by the user")
        raise BadRequest("can't parse entities")

    ok, fail, ozet = await bot._kitle_gonder(users, send_one)
    assert ok == 1 and fail == 3, (ok, fail)
    assert "2 kisi botu engellemis" in ozet, ozet
    assert "icerik hatasi" in ozet, ozet

    # Hepsi basarili -> ozet bos olmali (admin mesajina ek satir gitmez)
    async def hep_ok(uid):
        return
    ok, fail, ozet = await bot._kitle_gonder(users, hep_ok)
    assert ok == 4 and fail == 0 and ozet == "", (ok, fail, ozet)



async def test_deliver_payload_uzun_caption_kisaltilir():
    """Medyali gonderimde 1024+ karakter aciklama Telegram'a takilmamali."""
    gonderilen = []

    class FakeBot:
        async def send_photo(self, chat_id, photo, caption=None, **kw):
            gonderilen.append(("photo", chat_id, caption))
        async def send_video(self, chat_id, video, caption=None, **kw):
            gonderilen.append(("video", chat_id, caption))
        async def send_message(self, chat_id, text, **kw):
            gonderilen.append(("text", chat_id, text))

    ctx = SimpleNamespace(bot=FakeBot())
    original_path, original_db = bot.DB_PATH, bot.db
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "caption_test.db")
            bot.db = bot.Database()
            bot.db.add_user(777001, "capt_test", "Cap", "Tan")
            uzun = "A" * 1400
            payload = {"type": "photo", "title": "Baslik", "text": uzun,
                       "media_id": "FILEID123", "media_type": "photo"}
            ok, fail, total, ozet = await bot.deliver_payload(ctx, payload, "all")
            assert fail == 0, ozet
            assert gonderilen, "hic gonderim yapilmadi"
            for tur, _cid, cap in gonderilen:
                if tur in ("photo", "video"):
                    assert cap is not None and len(cap) <= 1024, len(cap or "")
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path



async def test_sihirbaz_girdileri_kaybolmuyor():
    """'get(...) or {}' hatasi regresyonu: bos sozlukle baslayan sihirbazlarda
    yazilan soru/baslik/metin user_data icinde KALMALI (v8.4 ve oncesinde kayboluyordu)."""
    class FakeMsg:
        def __init__(self, text=""):
            self.text = text; self.video = None; self.photo = None
        async def reply_text(self, *a, **k): pass
        async def delete(self): pass
    class FakeUser:
        id = 999888; username = "wiz"; first_name = "W"; last_name = ""; language_code = ""
    def upd(text):
        return SimpleNamespace(message=FakeMsg(text), effective_user=FakeUser(),
                               effective_chat=SimpleNamespace(id=1, type="private"))

    original_path, original_db = bot.DB_PATH, bot.db
    original_auth = bot.is_authenticated
    temp_ctx = tempfile.TemporaryDirectory()
    bot.DB_PATH = os.path.join(temp_ctx.name, "wizard_test.db")
    bot.db = bot.Database()
    # Sihirbaz finalleri artik oturum dogruluyor; sahte admin oturumu ac.
    bot.is_authenticated = lambda uid: uid == 999888
    try:
        await _sihirbaz_senaryosu(upd)
    finally:
        bot.is_authenticated = original_auth
        bot.db = original_db
        bot.DB_PATH = original_path
        temp_ctx.cleanup()


async def _sihirbaz_senaryosu(upd):
    # SSS: soru adimi
    ctx = SimpleNamespace(user_data={"faq": {}})
    r = await bot.faq_q_step(upd("Nasil para cekerim?"), ctx)
    assert r == bot.FAQ_A
    assert (ctx.user_data["faq"].get("q") or "").startswith("Nasil"), ctx.user_data

    # SSS: soru kaybolduysa cevap adimi BOS kayit olusturmamali
    ctx2 = SimpleNamespace(user_data={"faq": {}})
    onceki = len(bot.db.list_faqs(only_active=False))
    r2 = await bot.faq_a_step(upd("Cevap metni"), ctx2)
    assert r2 == bot.FAQ_Q, r2
    assert len(bot.db.list_faqs(only_active=False)) == onceki, "bos soru kaydedilmemeli"

    # Duyuru: baslik + metin adimlari user_data'da birikmeli, kayit dolu olmali
    ctx3 = SimpleNamespace(user_data={"ann": {}, "wizard_lock": "ann_new"})
    await bot.ann_title_step(upd("Kampanya Duyurusu"), ctx3)
    await bot.ann_body_step(upd("Detayli aciklama metni"), ctx3)
    assert ctx3.user_data["ann"].get("title") and ctx3.user_data["ann"].get("body"), ctx3.user_data
    await bot.ann_media_skip(upd(""), ctx3)
    kayitlar = bot.db.list_announcements(only_active=False)
    son = kayitlar[0] if kayitlar else {}
    assert (son.get("title") or "").startswith("Kampanya"), son

    # Bos icerikli eski kayit gonderime CIKMAMALI, anlasilir uyari donmeli
    ctx4 = SimpleNamespace(bot=None)
    ok, fail, total, ozet = await bot.deliver_payload(ctx4, {"type": "text", "title": "", "text": ""}, "all")
    assert ok == 0 and fail == 0 and "BOS" in ozet, (ok, fail, ozet)



def test_v85_duzeltmeleri():
    """v8.5: atomik zamanli kilit, ban upsert, etiket icerme eslesmesi, gecersiz link butonlari."""
    original_path, original_db = bot.DB_PATH, bot.db
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot.DB_PATH = os.path.join(temp_dir, "v85_test.db")
            d = bot.Database()
            bot.db = d

            # claim_scheduled: yalnizca ILK sahiplenme basarili olmali
            sid = d.add_scheduled('{"type":"text","text":"x"}', "all",
                                  "2000-01-01T00:00:00")
            assert d.claim_scheduled(sid) is True
            assert d.claim_scheduled(sid) is False, "ikinci sahiplenme reddedilmeli"
            d.set_scheduled_status(sid, "iptal")
            assert d.claim_scheduled(sid) is False, "iptal kayit sahiplenilememeli"

            # ban_user: tabloda olmayan kullanici da banlanabilmeli
            d.ban_user(555444333, "test", by_admin=1)
            assert d.is_banned(555444333), "tabloda olmayan ID banlanamadi"

            # _match_menu_label: ozellestirilmis etikette icerme eslesmesi
            assert bot._match_menu_label("🎁 kampanyalar") == "kampanya"
            assert bot._match_menu_label("💬 destek hattı") == "destek"
            assert bot._match_menu_label("para çek") == "paracek"
            assert bot._match_menu_label("merhaba nasılsın") is None or True  # serbest metin patlamamali

            # Gecersiz linkler buton uretmemeli (ekran olmesin)
            d.set_setting("bonus_link", "gecersiz")
            assert bot.build_campaign_keyboard() is None
            d.set_setting("register_link", "gecersiz-link")
            d.set_setting("giris_link", "https://ornek.com/giris")
            kb = bot._build_start_inline()
            assert kb is not None and len(kb.inline_keyboard[0]) == 1, "yalnizca gecerli giris butonu kalmali"
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path



def test_handler_kayit_sirasi():
    """Desensiz son-durak yakalayicisi TUM callback handler'lardan SONRA kayitli
    olmali; onde olursa butun admin paneli oluyor (v8.5'te yasandi)."""
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "bot.py"), encoding="utf-8") as f:
        kaynak = f.read()
    son_durak = kaynak.rindex("app.add_handler(CallbackQueryHandler(stale_button_cb))")
    admin_kayit = kaynak.rindex("CallbackQueryHandler(admin_cb, pattern=")
    um_kayit = kaynak.rindex("CallbackQueryHandler(user_menu_cb, pattern=")
    assert son_durak > admin_kayit, "stale_button_cb admin_cb'den ONCE kayitli — panel olur!"
    assert son_durak > um_kayit, "stale_button_cb kullanici menusunden ONCE kayitli!"



async def test_yerlesik_asistan():
    """Cevrimdisi asistan: yazim hatasina dayanikli niyet analizi + SSS eslesme."""
    # 1) Niyet analizi (saf fonksiyon)
    ornekler = [
        ("param gelmedii ne zaman yatacak", "cekim"),
        ("cekimm yapamiyorum yardim", "cekim"),
        ("YATIRAMIYORUM kart gecmiyor", "yatirim"),
        ("bonusum hala tanimlanmadi", "bonus"),
        ("siteye giremiyorumm adres nedir", "giris"),
        ("nasil uye olabilirim", "kayit"),
        ("turnuvaya nasil katilirim odul havuzu ne", "turnuva"),
        ("site id nasil girilir", "siteid"),
        ("merhaba nasilsin", "selam"),
        ("sikayet etmek istiyorum dolandirildim", "sikayet"),
    ]
    for metin, beklenen in ornekler:
        niyet, puan = bot._analyze_intent(metin)
        assert niyet is not None and niyet["key"] == beklenen, (metin, niyet and niyet["key"], puan)

    # Alakasiz metin niyet uretmemeli
    niyet, _ = bot._analyze_intent("bugun hava cok guzel purkiie")
    assert niyet is None, niyet

    # 2) SSS eslesmesi + tam akis (gercek gecici DB ile)
    original_path, original_db = bot.DB_PATH, bot.db
    original_guard = bot.guard_flood
    original_ids = list(bot.ADMIN_IDS)

    async def allow(*a, **k):
        return False

    try:
        with tempfile.TemporaryDirectory() as td:
            bot.DB_PATH = os.path.join(td, "asistan.db")
            bot.db = bot.Database()
            bot.guard_flood = allow
            bot.ADMIN_IDS[:] = [999999]
            bot.db.add_faq("Cevrim sarti nedir?", "Bonus cevrimi 10 kattir; detaylar sitede.")
            bot.db.add_faq("Para yatirma limiti kac?", "Minimum 100 TL, maksimum 50.000 TL.")

            faq, skor = bot._faq_best_match("cevrim sarti ne kadar")
            assert faq is not None and "Cevrim" in faq["question"], (faq, skor)

            class FlowMsg:
                def __init__(self, text): self.text = text; self.replies = []
                async def reply_text(self, *a, **k): self.replies.append((a, k))

            def yap(text, uid):
                m = FlowMsg(text)
                return m, SimpleNamespace(effective_user=SimpleNamespace(id=uid, username="u", first_name="A", last_name=None, language_code="tr"),
                                          effective_chat=SimpleNamespace(id=uid, type="private"), message=m)

            # Niyet cevabi + dogru buton
            bot._cb_last_press.clear()
            m, u = yap("param gelmedi hala", 7101)
            await bot.handle_text(u, SimpleNamespace(user_data={}, bot=None))
            assert m.replies and "PARA ÇEKİM" in m.replies[0][0][0]

            # SSS cevabi (niyete takilmayan soru)
            bot._cb_last_press.clear()
            m, u = yap("cevrim sarti nedir acaba", 7102)
            await bot.handle_text(u, SimpleNamespace(user_data={}, bot=None))
            assert m.replies and "10 kat" in m.replies[0][0][0], m.replies[0][0][0][:120]

            # Hicbiri: kibar yonlendirme + destek merkezi
            bot._cb_last_press.clear()
            m, u = yap("xyzqw prkt lorem", 7103)
            await bot.handle_text(u, SimpleNamespace(user_data={}, bot=None))
            assert m.replies and "ANLAYAMADIM" in m.replies[0][0][0]
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path
        bot.guard_flood = original_guard
        bot.ADMIN_IDS[:] = original_ids



def test_bilgi_bankasi_ve_yardim():
    """Hazir SSS seti yuklenir (bir kez) ve panel yardim konulari eksiksizdir."""
    original_path, original_db = bot.DB_PATH, bot.db
    try:
        with tempfile.TemporaryDirectory() as td:
            bot.DB_PATH = os.path.join(td, "seed.db")
            bot.db = bot.Database()
            n = bot.seed_bilgi_bankasi()
            assert n == len(bot._VARSAYILAN_SSS) and n >= 15, n
            assert len(bot.db.list_faqs()) == n
            # Ikinci cagri tekrar yuklememeli
            assert bot.seed_bilgi_bankasi() == 0
            # Admin hepsini silse bile geri gelmemeli (bayrak)
            for f in bot.db.list_faqs(only_active=False):
                bot.db.delete_faq(f["id"])
            assert bot.seed_bilgi_bankasi() == 0
            # Asistan seed'lenen bilgi bankasindan cevap bulabilmeli
            bot.db.set_setting("faq_seeded", "")
            bot.seed_bilgi_bankasi()
            faq, skor = bot._faq_best_match("cevrim sarti nedir")
            assert faq is not None and "Cevrim" in faq["question"], (faq, skor)
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path

    # Yardim: her konunun metni var ve bos degil
    for key, _ad in bot._YARDIM_KONULARI:
        assert (bot._YARDIM_METINLERI.get(key) or "").strip(), key
    kb = bot._yardim_liste_kb()
    assert kb is not None and len(kb.inline_keyboard) >= 6



def test_admin_klavyesi():
    """Admin alt klavyesi: duzen ve etiket yonlendirmesi."""
    kb = bot._build_admin_reply_menu()
    assert [len(r) for r in kb.keyboard] == [2, 2, 2, 2, 2]
    for etiket, beklenen in [("🛠 Panel", "panel"), ("📊 İstatistik", "istatistik"),
                             ("🎯 Kampanyalar", "kampanyalar"), ("📢 Duyurular", "duyurular"),
                             ("📣 Bildirim Gönder", "bildirim"), ("📅 Zamanlı Mesaj", "zamanli"),
                             ("❓ SSS Yönetimi", "sssyonetim"), ("📖 Yardım", "yardim"),
                             ("👤 Kullanıcı Menüsü", "kullanici"), ("🚪 Çıkış", "cikis")]:
        assert bot._match_admin_label(etiket.lower()) == beklenen, etiket
    # Kullanici etiketleri admin haritasina TAKILMAZ
    assert bot._match_admin_label("🏠 ana menü") is None
    assert bot._match_admin_label("🤴 profil") is None



def test_siteid_hedefleme():
    """Site ID tabanli hedefleme: kitle cozumu, ID listesi, toplu segment."""
    original_path, original_db = bot.DB_PATH, bot.db
    try:
        with tempfile.TemporaryDirectory() as td:
            bot.DB_PATH = os.path.join(td, "sid.db")
            d = bot.Database(); bot.db = d
            # 3 kullanici: 2'si Site ID baglamis, 1'i baglamamis
            for uid in (201, 202, 203):
                d.add_user(uid, f"u{uid}", "Ad", "")
            d.upsert_site_id(201, "u201", "Ad", "", "11111")
            d.upsert_site_id(202, "u202", "Ad", "", "22222")

            # siteid hedefi: yalniz baglayanlar
            kitle = [u["user_id"] for u in bot.resolve_audience("siteid")]
            assert sorted(kitle) == [201, 202], kitle
            # ID listesi hedefi
            kitle = [u["user_id"] for u in bot.resolve_audience("ids:11111")]
            assert kitle == [201], kitle
            # olmayan ID bos doner
            assert bot.resolve_audience("ids:99999") == []
            # banli uye siteid kitlesine girmez
            d.ban_user(202, "test")
            kitle = [u["user_id"] for u in bot.resolve_audience("siteid")]
            assert kitle == [201], kitle
            # etiketler
            assert "Site ID" in bot.target_label("siteid")
            assert "(1 ID)" in bot.target_label("ids:11111")
            # /kim sorgusu
            r = d.find_by_site_id("11111")
            assert r and r[0]["telegram_id"] == 201
            # toplu segment (siteseg mantigi): eslesenlere kademe
            for u in d.get_users_by_site_ids(["11111"]):
                d.set_segment(u["user_id"], "platin")
            assert d.get_user(201)["segment"] == "platin"
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path



def test_gonderim_gecmisi_ve_duyuru_duzenleme():
    """Gecmis kaydi + tekrar icin payload saklama + duyuru alan duzenleme."""
    original_path, original_db = bot.DB_PATH, bot.db
    try:
        with tempfile.TemporaryDirectory() as td:
            bot.DB_PATH = os.path.join(td, "hist.db")
            d = bot.Database(); bot.db = d

            bot._log_broadcast("📢 Duyuru #1", "siteid", 5, 1, 6,
                               {"type": "text", "title": "T", "text": "M"}, 42)
            kayitlar = d.list_broadcast_logs()
            assert len(kayitlar) == 1
            b = kayitlar[0]
            assert b["kind"].endswith("#1") and b["target"] == "siteid"
            assert (b["ok"], b["fail"], b["total"]) == (5, 1, 6)
            import json as _json
            p = _json.loads(b["payload"])
            assert bot.build_payload_text(p).startswith("<b>T</b>")
            assert d.get_broadcast_log(b["id"])["admin_id"] == 42

            aid = d.add_announcement("Eski Baslik", "Metin", "", "")
            d.update_announcement_field(aid, "title", "Yeni Baslik")
            d.update_announcement_field(aid, "media_id", "FILE1")
            d.update_announcement_field(aid, "media_type", "photo")
            a = d.get_announcement(aid)
            assert a["title"] == "Yeni Baslik" and a["media_id"] == "FILE1"
            try:
                d.update_announcement_field(aid, "id", "9")
                assert False, "beyaz liste calismadi"
            except ValueError:
                pass
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path



async def test_denetim_duzeltmeleri():
    """Son denetim: kisa anahtar eslesme, NOCASE Site ID, klavye korumasi."""
    # 3) 'id' gibi kisa anahtarlar artik tam eslesmeyle puanlanir
    assert bot._fuzzy_token("id", "id") is True
    niyet, _ = bot._analyze_intent("id yazmak istiyorum kimlik nerede")
    assert niyet is not None and niyet["key"] == "siteid", niyet and niyet["key"]

    # 4) Site ID sorgulari buyuk/kucuk harf duyarsiz
    original_path, original_db = bot.DB_PATH, bot.db
    try:
        with tempfile.TemporaryDirectory() as td:
            bot.DB_PATH = os.path.join(td, "nocase.db")
            d = bot.Database(); bot.db = d
            d.add_user(301, "u", "A", "")
            d.upsert_site_id(301, "u", "A", "", "AbC123")
            assert d.find_by_site_id("abc123"), "kucuk harfle bulunamadi"
            assert [u["user_id"] for u in d.get_users_by_site_ids(["ABC123"])] == [301]
    finally:
        bot.db = original_db
        bot.DB_PATH = original_path

    # 2) Sihirbaz metin adimi menu tusunu ICERIK sanmaz
    class FakeMsg:
        def __init__(self, text=""):
            self.text = text; self.video = None; self.photo = None
            self.replies = []
        async def reply_text(self, *a, **k): self.replies.append(a[0] if a else "")
        async def delete(self): pass
    msg = FakeMsg("🛠 Panel")
    upd = SimpleNamespace(message=msg,
                          effective_user=SimpleNamespace(id=888777, username="w", first_name="W", last_name="", language_code=""),
                          effective_chat=SimpleNamespace(id=1, type="private"))
    ctx = SimpleNamespace(user_data={"faq": {}})
    durum = await bot.faq_q_step(upd, ctx)
    assert durum == bot.FAQ_Q, durum
    assert msg.replies and "ALINMADI" in msg.replies[0]
    assert not (ctx.user_data["faq"].get("q")), "menu tusu soru olarak kaydedildi!"
    # Normal metin ayni adimda calismaya devam eder
    msg2 = FakeMsg("Cevrim sarti nedir?")
    upd2 = SimpleNamespace(message=msg2, effective_user=upd.effective_user,
                           effective_chat=upd.effective_chat)
    durum2 = await bot.faq_q_step(upd2, ctx)
    assert durum2 == bot.FAQ_A, durum2


# =====================================================================
# EKRAN GORUNTUSU KAYDI TESTLERI
# =====================================================================
from datetime import datetime, timezone
import json as _json


class SSTgFile:
    def __init__(self, size=4096):
        self.file_size = size

    async def download_to_drive(self, path):
        with open(path, "wb") as f:
            f.write(b"\xff\xd8fake-jpeg-data")


class SSBot(FakeTelegramBot):
    def __init__(self, size=4096):
        super().__init__()
        self.size = size

    async def get_file(self, file_id):
        self.calls.append(("get_file", {"file_id": file_id}))
        return SSTgFile(self.size)

    async def send_document(self, **kwargs):
        self.calls.append(("document", kwargs))


class SSMessage:
    def __init__(self, text="", photo=None, caption="", date=None, document=None):
        self.text = text
        self.photo = photo
        self.document = document
        self.caption = caption
        self.video = None
        self.date = date or datetime(2026, 9, 22, 14, 30, 5, tzinfo=timezone.utc)
        self.replies = []

    async def reply_text(self, *args, **kwargs):
        self.replies.append((args, kwargs))


class SSQuery(FlowQuery):
    def __init__(self, data, chat_id=4242):
        super().__init__(data)
        self.message = SimpleNamespace(chat_id=chat_id, chat=SimpleNamespace(id=chat_id), text="x")
        self.markups = []
        self.alerts = []

    async def answer(self, *args, **kwargs):
        self.answered = True
        if args:
            self.alerts.append(args[0])

    async def edit_message_reply_markup(self, reply_markup=None):
        self.markups.append(reply_markup)


def _ss_user(uid=4242, uname="testci"):
    return SimpleNamespace(id=uid, username=uname, first_name="Ali", last_name="", language_code="tr")


def _ss_upd(user, msg=None, query=None):
    return SimpleNamespace(message=msg, callback_query=query, effective_user=user,
                           effective_chat=SimpleNamespace(id=user.id, type="private"))


def _ss_photo(file_id="ph-big", caption=""):
    return SSMessage(photo=[SimpleNamespace(file_id="ph-small"), SimpleNamespace(file_id=file_id)], caption=caption)


class _SSOrtam:
    """Gecici DB + gecici screenshots klasoru + flood/admin yamalari."""
    def __enter__(self):
        self.td = tempfile.TemporaryDirectory()
        self.orig = (bot.DB_PATH, bot.db, bot.SCREENSHOTS_DIR, bot.SCREENSHOTS_CSV, bot.guard_flood,
                     list(bot.ADMIN_IDS), bot.is_authenticated, bot.SCREENSHOT_DAILY_LIMIT)
        self.orig_throttle = bot._cb_throttled
        bot._cb_throttled = lambda *a, **k: False
        bot.DB_PATH = os.path.join(self.td.name, "ss.db")
        bot.db = bot.Database()
        bot.SCREENSHOTS_DIR = os.path.join(self.td.name, "screenshots")
        os.makedirs(bot.SCREENSHOTS_DIR)
        bot.SCREENSHOTS_CSV = os.path.join(bot.SCREENSHOTS_DIR, "ekran_goruntuleri.csv")

        async def allow(*a, **k):
            return False
        bot.guard_flood = allow
        bot.ADMIN_IDS[:] = [777]
        return self

    def __exit__(self, *exc):
        (bot.DB_PATH, bot.db, bot.SCREENSHOTS_DIR, bot.SCREENSHOTS_CSV, bot.guard_flood,
         ids, bot.is_authenticated, bot.SCREENSHOT_DAILY_LIMIT) = self.orig
        bot.ADMIN_IDS[:] = ids
        bot._cb_throttled = self.orig_throttle
        self.td.cleanup()
        return False

    def dosyalar(self):
        return sorted(f for f in os.listdir(bot.SCREENSHOTS_DIR) if not f.endswith(".csv"))


def test_ekran_dosya_adi_ve_kategoriler():
    """Dosya adi: kategori + tarih + girilen bilgiler; Turkce/tehlikeli karakterler temizlenir."""
    cat = {"key": "arkadas", "label": "👥 Arkadaşını Getir", "fields": ["Site ID", "Arkadaş ID"]}
    when = datetime(2026, 9, 22, 14, 30, 5)
    fmt = "{kategori}_{tarih}_{saat}_{alanlar}_{tgid}"
    assert bot.build_screenshot_filename(cat, ["ABC123", "XYZ789"], 4242, when, fmt=fmt) == \
        "Arkadasini-Getir_2026-09-22_14-30-05_ABC123_XYZ789_4242"
    # Alan adi yer tutuculari calisir, bilinmeyen yer tutucu bos kalir
    assert bot.build_screenshot_filename(cat, ["A1", "B2"], 1, when, fmt="{arkadas_id}-{site_id}-{yok}") == "B2-A1"
    assert bot.build_screenshot_filename(cat, ["A1", "B2"], 1, when, fmt="{alan2}_{alan1}_{kategori_kodu}") == "B2_A1_arkadas"
    # Yol saldirisi ve bosluklar dosya adina sizamaz
    tehlikeli = bot.build_screenshot_filename({"key": "cekim", "label": "Çekim", "fields": ["Tutar"]},
                                              ["../../etc/passwd"], 5, when, fmt="{kategori}_{alanlar}")
    assert tehlikeli == "Cekim_etc-passwd", tehlikeli
    assert ".." not in tehlikeli and "/" not in tehlikeli
    assert bot.build_screenshot_filename({"key": "x", "label": "Para Yatırma", "fields": ["Tutar"]},
                                         ["500 TL"], 5, when, fmt="{kategori}_{alanlar}") == "Para-Yatirma_500-TL"
    # Bos alan listesi cift alt cizgi birakmaz; tamamen bos sablon guvenli varsayilana duser
    assert bot.build_screenshot_filename({"key": "x", "label": "Çekim", "fields": []}, [], 5, when,
                                         fmt="{kategori}_{alanlar}_{tgid}") == "Cekim_5"
    assert bot.build_screenshot_filename({"key": "x", "label": "", "fields": []}, [], 5, when, fmt="{yok}") == \
        "Ekran_2026-09-22_14-30-05_5"
    # Windows ayrilmis adlari
    assert bot.build_screenshot_filename({"key": "x", "label": "L", "fields": ["A"]}, ["CON"], 5, when, fmt="{alan1}") == "Ekran_CON"
    # Kategori tanimi ayristirma
    cats = bot.parse_screenshot_categories(
        "arkadas:👥 Arkadaşını Getir:Site ID|Arkadaş ID, bonus:🎁 Bonus:Site ID, bozuk, KOTU KEY:X:Y, arkadas:Tekrar:Z, alansiz:📄 Diğer")
    assert [c["key"] for c in cats] == ["arkadas", "bonus", "alansiz"], cats
    assert cats[0]["fields"] == ["Site ID", "Arkadaş ID"] and cats[2]["fields"] == []
    assert bot.parse_screenshot_categories("") == [] and bot.parse_screenshot_categories("::") == []
    cok_alan = bot.parse_screenshot_categories("k:L:" + "|".join(f"A{i}" for i in range(10)))
    assert len(cok_alan[0]["fields"]) == bot.SCREENSHOT_MAX_FIELDS
    assert len(bot.parse_screenshot_categories(bot.BUILTIN_SCREENSHOT_CATEGORIES)) == 5
    # Alan dogrulama: ID alanlari Site ID kuralina uyar, diger alanlar serbest
    assert bot.validate_screenshot_field("Site ID", "abc 123")[1]
    assert bot.validate_screenshot_field("Arkadaş ID", "ABC.1_2-3") == ("ABC.1_2-3", "")
    assert bot.validate_screenshot_field("Tutar", "  500   TL ") == ("500 TL", "")
    assert bot.validate_screenshot_field("Tutar", "")[1] and bot.validate_screenshot_field("Tutar", "!!!")[1]
    assert bot.validate_screenshot_field("Konu", "x" * 61)[1]
    # Alt klavye etiketi taninir; serbest cumle taninmaz
    assert bot._match_menu_label("📸 ekran görüntüsü") == "ekran"
    assert bot._match_menu_label("ekran görüntüsü gönder") == "ekran"
    assert bot._EKRAN_ETIKET_FILTRESI.filter(SimpleNamespace(text="📸 Ekran Görüntüsü"))
    assert not bot._EKRAN_ETIKET_FILTRESI.filter(SimpleNamespace(text="ekranım donuyor sürekli yardım edin lütfen"))
    # Ayni ad varsa _2 eklenir
    with _SSOrtam() as ortam:
        p1, f1 = bot._unique_screenshot_path("Ad_2026", ".jpg")
        assert f1 == "Ad_2026.jpg" and p1.startswith(bot.SCREENSHOTS_DIR)
        open(p1, "wb").close()
        p2, f2 = bot._unique_screenshot_path("Ad_2026", ".jpg")
        assert f2 == "Ad_2026_2.jpg"
        # Bilinmeyen uzanti jpg'ye duser
        assert bot._unique_screenshot_path("X", ".exe")[1] == "X.jpg"
        # DB metodlari
        d = bot.db
        sid = d.add_screenshot(4242, "u", "Ali", "arkadas", "👥 Arkadaşını Getir",
                               _json.dumps([{"alan": "Site ID", "deger": "ABC"}]), "a.jpg", "not")
        sid2 = d.add_screenshot(4243, "v", "Veli", "bonus", "🎁 Bonus", "[]", "b.jpg")
        assert d.count_screenshots() == 2 and d.count_screenshots(only_new=True) == 2
        assert d.count_screenshots_today(4242) == 1 and d.count_screenshots_today(1) == 0
        assert [r["id"] for r in d.list_screenshots()] == [sid2, sid]
        assert [r["id"] for r in d.list_screenshots(category="bonus")] == [sid2]
        d.set_screenshot_status(sid, "okundu")
        assert d.count_screenshots(only_new=True) == 1 and d.get_screenshot(sid)["status"] == "okundu"
        assert bot._ss_fields_list(d.get_screenshot(sid)) == [("Site ID", "ABC")]
        assert "ABC" in bot._ss_list_lines(d.list_screenshots())
        for i in range(80):
            d.add_screenshot(1000 + i, "u", "Ad", "diger", "📄 Diğer",
                             _json.dumps([{"alan": f"A{j}", "deger": "X" * 20} for j in range(6)]), f"f{i}.jpg")
        uzun = bot._ss_list_lines(d.list_screenshots(limit=100))
        assert len(uzun) <= 3400 and "kisaltildi" in uzun, len(uzun)
        assert len(bot._ss_list_lines(d.list_screenshots(limit=20))) < 3300  # komut varsayilani sinira sigar
        assert "#" + str(sid) in bot._ss_admin_detail_text(d.get_screenshot(sid))


async def test_ekran_foto_once_akisi():
    """Yol 1: once foto -> kategori secenekleri -> bilgiler -> dosya girilen bilgilerle adlanir."""
    with _SSOrtam() as ortam:
        user = _ss_user()
        ctx = SimpleNamespace(user_data={}, bot=SSBot())
        msg = _ss_photo(caption="deneme")
        r = await bot.ss_photo_entry(_ss_upd(user, msg), ctx)
        assert r == bot.SS_CAT, r
        assert msg.replies and "EKRAN GÖRÜNTÜSÜ" in msg.replies[-1][0][0]
        etiketler = [b.text for row in msg.replies[-1][1]["reply_markup"].inline_keyboard for b in row]
        assert any("Arkadaşını Getir" in e for e in etiketler) and any("Vazgeç" in e for e in etiketler), etiketler
        assert ctx.user_data["ss"]["file_id"] == "ph-big" and ctx.user_data["wizard_lock"][0] == "ss"

        # Var olmayan kategori: ayni adimda kal
        q0 = SSQuery("sscat_yok")
        assert await bot.ss_cat_pick(_ss_upd(user, query=q0), ctx) == bot.SS_CAT and q0.alerts

        q = SSQuery("sscat_arkadas")
        r = await bot.ss_cat_pick(_ss_upd(user, query=q), ctx)
        assert r == bot.SS_FIELD and q.answered and "Site ID" in q.edits[-1][0][0]

        # Gecersiz Site ID -> hata, ayni adim
        r = await bot.ss_field_step(_ss_upd(user, SSMessage(text="ab c!")), ctx)
        assert r == bot.SS_FIELD and "HATA" in ctx.bot.calls[-1][1]["text"]
        # Menu tusu icerik sayilmaz
        r = await bot.ss_field_step(_ss_upd(user, SSMessage(text="🏠 Ana Menü")), ctx)
        assert r == bot.SS_FIELD
        r = await bot.ss_field_step(_ss_upd(user, SSMessage(text="ABC123")), ctx)
        assert r == bot.SS_FIELD and "Arkadaş ID" in ctx.bot.calls[-1][1]["text"]
        r = await bot.ss_field_step(_ss_upd(user, SSMessage(text="XYZ789")), ctx)
        assert r == bot.ConversationHandler.END

        yerel = msg.date.astimezone().replace(tzinfo=None)
        beklenen = f"Arkadasini-Getir_{yerel.strftime('%Y-%m-%d_%H-%M-%S')}_ABC123_XYZ789_4242.jpg"
        assert ortam.dosyalar() == [beklenen], ortam.dosyalar()
        with open(os.path.join(bot.SCREENSHOTS_DIR, beklenen), "rb") as f:
            assert f.read().startswith(b"\xff\xd8")
        kayit = bot.db.list_screenshots()[0]
        assert kayit["file_name"] == beklenen and kayit["category"] == "arkadas" and kayit["note"] == "deneme"
        assert [d["deger"] for d in _json.loads(kayit["fields"])] == ["ABC123", "XYZ789"]
        with open(bot.SCREENSHOTS_CSV, encoding="utf-8-sig") as f:
            satirlar = list(csv.reader(f))
        assert satirlar[0] == bot.SCREENSHOT_CSV_HEADERS and satirlar[1][-1] == beklenen and "ABC123" in satirlar[1][6]
        # Kullaniciya onay, admine (777) foto bildirimi
        onay = [c for c in ctx.bot.calls if c[0] == "message" and c[1]["chat_id"] == 4242 and "KAYDEDİLDİ" in c[1]["text"]]
        assert onay and beklenen in onay[-1][1]["text"]
        bildirim = [c for c in ctx.bot.calls if c[0] == "photo" and c[1]["chat_id"] == 777]
        assert bildirim and "ABC123" in bildirim[-1][1]["caption"] and bildirim[-1][1]["photo"] == "ph-big"
        assert "ss" not in ctx.user_data and "wizard_lock" not in ctx.user_data

        # Iptal butonu ve /iptal temizler
        await bot.ss_photo_entry(_ss_upd(user, _ss_photo()), ctx)
        qc = SSQuery("ss_cancel")
        assert await bot.ss_cat_pick(_ss_upd(user, query=qc), ctx) == bot.ConversationHandler.END
        assert "ss" not in ctx.user_data and "VAZGEÇİLDİ" in qc.edits[-1][0][0]
        await bot.ss_photo_entry(_ss_upd(user, _ss_photo()), ctx)
        m = SSMessage(text="/iptal")
        assert await bot.ss_cancel_cmd(_ss_upd(user, m), ctx) == bot.ConversationHandler.END
        assert "ss" not in ctx.user_data and m.replies

        # Indirme hatasi: dosya yazilmaz, kullanici bilgilendirilir, kayit acilmaz
        class KirikBot(SSBot):
            async def get_file(self, file_id):
                raise RuntimeError("telegram down")
        ctx2 = SimpleNamespace(user_data={}, bot=KirikBot())
        await bot.ss_photo_entry(_ss_upd(user, _ss_photo("ph-2")), ctx2)
        await bot.ss_cat_pick(_ss_upd(user, query=SSQuery("sscat_arkadas")), ctx2)
        await bot.ss_field_step(_ss_upd(user, SSMessage(text="A1")), ctx2)
        r = await bot.ss_field_step(_ss_upd(user, SSMessage(text="B2")), ctx2)
        assert r == bot.ConversationHandler.END and "KAYDEDİLEMEDİ" in ctx2.bot.calls[-1][1]["text"]
        assert len(ortam.dosyalar()) == 1 and bot.db.count_screenshots() == 1
        # 20MB ustu dosya reddedilir
        ctx3 = SimpleNamespace(user_data={}, bot=SSBot(size=bot.SCREENSHOT_MAX_BYTES + 1))
        await bot.ss_photo_entry(_ss_upd(user, _ss_photo("ph-3")), ctx3)
        await bot.ss_cat_pick(_ss_upd(user, query=SSQuery("sscat_diger")), ctx3)
        await bot.ss_field_step(_ss_upd(user, SSMessage(text="A1")), ctx3)
        r = await bot.ss_field_step(_ss_upd(user, SSMessage(text="konu")), ctx3)
        assert r == bot.ConversationHandler.END and "KAYDEDİLEMEDİ" in ctx3.bot.calls[-1][1]["text"]
        assert len(ortam.dosyalar()) == 1


async def test_ekran_bilgi_once_akisi():
    """Yol 2: /ekran -> kategori -> bilgiler -> foto gelince DOGRUDAN kaydedilir."""
    with _SSOrtam() as ortam:
        user = _ss_user(uid=5150, uname="")
        ctx = SimpleNamespace(user_data={}, bot=SSBot())
        m = SSMessage(text="/ekran")
        r = await bot.ss_start_cmd(_ss_upd(user, m), ctx)
        assert r == bot.SS_CAT and "EKRAN GÖRÜNTÜSÜ GÖNDER" in m.replies[-1][0][0]
        assert ctx.user_data["ss"]["mode"] == "data_first" and ctx.user_data["ss"]["file_id"] is None
        # Kayitli Site ID varsa tek tus buton cikar
        bot.db.upsert_site_id(5150, "", "Ali", "", "SITE77")
        q = SSQuery("sscat_yatirim", chat_id=5150)
        r = await bot.ss_cat_pick(_ss_upd(user, query=q), ctx)
        assert r == bot.SS_FIELD
        butonlar = [b.text for row in q.edits[-1][1]["reply_markup"].inline_keyboard for b in row]
        assert any("SITE77" in b for b in butonlar), butonlar
        q2 = SSQuery("ss_usesite", chat_id=5150)
        r = await bot.ss_usesite_cb(_ss_upd(user, query=q2), ctx)
        assert r == bot.SS_FIELD and ctx.user_data["ss"]["values"] == ["SITE77"]
        assert "Tutar" in ctx.bot.calls[-1][1]["text"]
        r = await bot.ss_field_step(_ss_upd(user, SSMessage(text="500 TL")), ctx)
        assert r == bot.SS_PHOTO and "EKRAN GÖRÜNTÜSÜNÜ GÖNDER" in ctx.bot.calls[-1][1]["text"]
        # Foto yerine yazi: uyar, bekle
        m2 = SSMessage(text="hey")
        assert await bot.ss_photo_step(_ss_upd(user, m2), ctx) == bot.SS_PHOTO and "HATA" in m2.replies[-1][0][0]
        # Resim DOSYA olarak (png) gelirse de kabul edilir, uzanti korunur.
        # allow_reentry nedeniyle PTB fotoyu GIRIS NOKTASINA (ss_photo_entry) yollar;
        # o da "simdi fotoyu gonder" adiminda oldugumuz icin DOGRUDAN kaydetmeli (regresyon).
        m3 = SSMessage(document=SimpleNamespace(file_id="doc-1", mime_type="image/png", file_name="ss.PNG"))
        r = await bot.ss_photo_entry(_ss_upd(user, m3), ctx)
        assert r == bot.ConversationHandler.END
        assert "ss" not in ctx.user_data
        yerel = m3.date.astimezone().replace(tzinfo=None)
        beklenen = f"Para-Yatirma_{yerel.strftime('%Y-%m-%d_%H-%M-%S')}_SITE77_500-TL_5150.png"
        assert ortam.dosyalar() == [beklenen], ortam.dosyalar()
        # Dosya olarak gelen gorsel admine de dosya olarak iletilir
        assert [c for c in ctx.bot.calls if c[0] == "document" and c[1]["chat_id"] == 777]

        # Ana menu butonu ile giris + bilgi adiminda foto gelirse akis bozulmaz
        q3 = SSQuery("um_ekran", chat_id=5150)
        r = await bot.ss_start_cb(_ss_upd(user, query=q3), ctx)
        assert r == bot.SS_CAT and q3.answered
        await bot.ss_cat_pick(_ss_upd(user, query=SSQuery("sscat_bonus", chat_id=5150)), ctx)
        await bot.ss_field_step(_ss_upd(user, SSMessage(text="ID9")), ctx)
        # Bilgi adiminda foto: giris noktasi akisi SIFIRLAMAZ, girilenler korunur
        m4 = _ss_photo("ph-mid")
        assert await bot.ss_photo_entry(_ss_upd(user, m4), ctx) == bot.SS_FIELD
        assert ctx.user_data["ss"]["file_id"] == "ph-mid" and ctx.user_data["ss"]["values"] == ["ID9"]
        assert "FOTO ALINDI" in m4.replies[0][0][0] and "Bonus Adı" in m4.replies[1][0][0]
        r = await bot.ss_field_step(_ss_upd(user, SSMessage(text="Hoşgeldin Bonusu")), ctx)
        assert r == bot.ConversationHandler.END
        assert any(f.startswith("Bonus-Talebi_") and f.endswith("_ID9_Hosgeldin-Bonusu_5150.jpg") for f in ortam.dosyalar()), ortam.dosyalar()

        # /start ve /iptal akis ortasinda: kayit birakilir (akis /start'i YUTMAZ, Site ID akisi bozulmaz)
        await bot.ss_start_cmd(_ss_upd(user, SSMessage(text="/ekran")), ctx)
        assert "ss" in ctx.user_data
        orig_reg = bot._send_registered_start
        async def sessiz(*a, **k): pass
        bot._send_registered_start = sessiz
        try:
            assert await bot.start_cmd(_ss_upd(user, SSMessage(text="/start")), ctx) == bot.ConversationHandler.END
        finally:
            bot._send_registered_start = orig_reg
        assert "ss" not in ctx.user_data and "wizard_lock" not in ctx.user_data
        # PTB durumu kalmis olabilir: ss yokken gelen metin normal yola (handle_text) verilir, durum kapanir
        orig_ht = bot.handle_text
        gelen = []
        async def sahte_ht(u, c): gelen.append(u.message.text)
        bot.handle_text = sahte_ht
        try:
            assert await bot.ss_field_step(_ss_upd(user, SSMessage(text="🏠 Ana Menü")), ctx) == bot.ConversationHandler.END
            assert await bot.ss_photo_step(_ss_upd(user, SSMessage(text="merhaba")), ctx) == bot.ConversationHandler.END
        finally:
            bot.handle_text = orig_ht
        assert gelen == ["🏠 Ana Menü", "merhaba"], gelen
        # ss yokken SS_PHOTO durumuna foto gelirse yeni akis baslar
        m9 = _ss_photo("ph-new")
        assert await bot.ss_photo_step(_ss_upd(user, m9), ctx) == bot.SS_CAT and ctx.user_data["ss"]["file_id"] == "ph-new"
        mi = SSMessage(text="/iptal")
        await bot.iptal_cmd(_ss_upd(user, mi), ctx)
        assert "ss" not in ctx.user_data and mi.replies
        # Acik sikayet akisi varken foto/ekran yeni akis BASLATMAZ (bilgiler sikayete dusmesin)
        ctx.user_data["cmp"] = {"kind": "text", "text": None, "image_id": None}
        mc = _ss_photo("ph-c")
        assert await bot.ss_photo_entry(_ss_upd(user, mc), ctx) == bot.ConversationHandler.END
        assert "AÇIK İŞLEM" in mc.replies[-1][0][0] and "ss" not in ctx.user_data
        mc2 = SSMessage(text="/ekran")
        assert await bot.ss_start_cmd(_ss_upd(user, mc2), ctx) == bot.ConversationHandler.END and "AÇIK İŞLEM" in mc2.replies[-1][0][0]
        # Sikayet zaman asimi cmp'yi temizler; /iptal da temizler
        assert await bot.sikayet_timeout(_ss_upd(user, SSMessage(text="x")), ctx) == bot.ConversationHandler.END
        assert "cmp" not in ctx.user_data
        ctx.user_data["cmp"] = {"kind": None}
        await bot.iptal_cmd(_ss_upd(user, SSMessage(text="/iptal")), ctx)
        assert "cmp" not in ctx.user_data
        # Album: ayni anda gelen 2. foto sessizce yok sayilir, flood sayacina girmez
        bot._cb_throttled = ortam.orig_throttle
        bot._cb_last_press.clear()
        a1, a2 = _ss_photo("alb-1"), _ss_photo("alb-2")
        assert await bot.ss_photo_entry(_ss_upd(user, a1), ctx) == bot.SS_CAT
        assert await bot.ss_photo_entry(_ss_upd(user, a2), ctx) == bot.ConversationHandler.END and a2.replies == []
        assert ctx.user_data["ss"]["file_id"] == "alb-1"
        bot._cb_throttled = lambda *a, **k: False
        ctx.user_data.pop("ss", None); ctx.user_data.pop("wizard_lock", None)

        # Ana menude buton var; ozellik kapaliyken yok ve /ekran kibarca reddeder
        assert any(b.callback_data == "um_ekran" for row in bot.build_main_menu().inline_keyboard for b in row)
        bot.db.set_setting("screenshots_enabled", "0")
        assert not any(b.callback_data == "um_ekran" for row in bot.build_main_menu().inline_keyboard for b in row)
        m5 = SSMessage(text="/ekran")
        assert await bot.ss_start_cmd(_ss_upd(user, m5), ctx) == bot.ConversationHandler.END and "kapalı" in m5.replies[-1][0][0]
        # Kapaliyken kullanici fotosu SESSIZCE yok sayilir
        m6 = _ss_photo("ph-k")
        assert await bot.ss_photo_entry(_ss_upd(user, m6), ctx) == bot.ConversationHandler.END and m6.replies == []
        bot.db.set_setting("screenshots_enabled", "1")

        # Gunluk limit
        bot.SCREENSHOT_DAILY_LIMIT = 2
        m7 = _ss_photo("ph-l")
        assert await bot.ss_photo_entry(_ss_upd(user, m7), ctx) == bot.ConversationHandler.END
        assert "LİMİT" in ctx.bot.calls[-1][1]["text"] and "ss" not in ctx.user_data
        # Beyaz listedeki admin limite takilmaz
        bot.ADMIN_IDS[:] = [777, 5150]
        m8 = _ss_photo("ph-a")
        assert await bot.ss_photo_entry(_ss_upd(user, m8), ctx) == bot.SS_CAT
        bot.ADMIN_IDS[:] = [777]


async def test_ekran_admin_secenekleri():
    """Admin foto gonderince secenek: kutuphane (eski davranis) ya da ekran goruntusu."""
    with _SSOrtam() as ortam:
        bot.ADMIN_IDS[:] = [9001, 777]
        bot.is_authenticated = lambda uid: uid == 9001
        admin = _ss_user(uid=9001, uname="patron")
        ctx = SimpleNamespace(user_data={}, bot=SSBot())
        m = _ss_photo("lib-1", caption="Yeni Banner")
        r = await bot.ss_photo_entry(_ss_upd(admin, m), ctx)
        assert r == bot.SS_ACTION and "FOTO ALINDI" in m.replies[-1][0][0]
        q = SSQuery("ssact_library", chat_id=9001)
        r = await bot.ss_action_pick(_ss_upd(admin, query=q), ctx)
        assert r == bot.ConversationHandler.END
        rows = bot.db.list_campaign_media()
        assert len(rows) == 1 and rows[0]["media_id"] == "lib-1" and rows[0]["name"] == "Yeni Banner"
        assert bot.db.get_setting("campaign_media_id") == "lib-1"
        assert "KUTUPHANEYE EKLENDI" in ctx.bot.calls[-1][1]["text"] and "ss" not in ctx.user_data

        m2 = _ss_photo("lib-2")
        assert await bot.ss_photo_entry(_ss_upd(admin, m2), ctx) == bot.SS_ACTION
        q2 = SSQuery("ssact_screenshot", chat_id=9001)
        r = await bot.ss_action_pick(_ss_upd(admin, query=q2), ctx)
        assert r == bot.SS_CAT and q2.edits
        await bot.ss_cat_pick(_ss_upd(admin, query=SSQuery("sscat_cekim", chat_id=9001)), ctx)
        await bot.ss_field_step(_ss_upd(admin, SSMessage(text="ADM1")), ctx)
        r = await bot.ss_field_step(_ss_upd(admin, SSMessage(text="1000")), ctx)
        assert r == bot.ConversationHandler.END
        assert any(f.startswith("Para-Cekme_") and f.endswith("_ADM1_1000_9001.jpg") for f in ortam.dosyalar()), ortam.dosyalar()
        # Bildirim yalnizca DIGER adminlere (777) gider, gonderen admine gitmez
        hedefler = [c[1]["chat_id"] for c in ctx.bot.calls if c[0] == "photo"]
        assert hedefler == [777], hedefler
        assert len(bot.db.list_campaign_media()) == 1, "ekran goruntusu kutuphaneye SIZMAMALI"

        # Video eskisi gibi dogrudan kutuphaneye
        mv = SSMessage(); mv.video = SimpleNamespace(file_id="vid-1"); mv.photo = None
        await bot.handle_campaign_media(_ss_upd(admin, mv), ctx)
        assert bot.db.get_setting("campaign_media_id") == "vid-1" and len(bot.db.list_campaign_media()) == 2

        # Ozellik kapaliyken admin fotosu eski davranisla dogrudan kutuphaneye
        bot.db.set_setting("screenshots_enabled", "0")
        m3 = _ss_photo("lib-3")
        assert await bot.ss_photo_entry(_ss_upd(admin, m3), ctx) == bot.ConversationHandler.END
        assert bot.db.get_setting("campaign_media_id") == "lib-3" and "ss" not in ctx.user_data
        bot.db.set_setting("screenshots_enabled", "1")

        # Baska sihirbaz acikken foto islenmez, uyarilir
        ctx.user_data["wizard_lock"] = ("aset", time.time())
        m4 = _ss_photo("lib-4")
        assert await bot.ss_photo_entry(_ss_upd(admin, m4), ctx) == bot.ConversationHandler.END
        assert "SIHIRBAZ ACIK" in m4.replies[-1][0][0] and "ss" not in ctx.user_data
        ctx.user_data.pop("wizard_lock")

        # Oturum akis ortasinda dolarsa kutuphaneye yazilmaz
        m5 = _ss_photo("lib-5")
        await bot.ss_photo_entry(_ss_upd(admin, m5), ctx)
        bot.is_authenticated = lambda uid: False
        q5 = SSQuery("ssact_library", chat_id=9001)
        assert await bot.ss_action_pick(_ss_upd(admin, query=q5), ctx) == bot.ConversationHandler.END
        assert "OTURUM DOLDU" in q5.edits[-1][0][0] and bot.db.get_setting("campaign_media_id") == "lib-3"

        # Panel ayarlari: kategori tanimi dogrulanir, dosya adi sablonu dogrulanir
        bot.is_authenticated = lambda uid: uid == 9001
        class FakeMsg:
            def __init__(self, text): self.text = text; self.replies = []
            async def reply_text(self, *a, **k): self.replies.append((a, k))
        def aupd(text):
            return SimpleNamespace(message=FakeMsg(text), effective_user=admin,
                                   effective_chat=SimpleNamespace(id=9001, type="private"))
        actx = SimpleNamespace(user_data={"aset_key": "screenshot_categories", "wizard_lock": ("aset", time.time())}, bot=ctx.bot)
        assert await bot.aset_save(aupd("bozuk tanim"), actx) == bot.ASET_WAIT
        assert await bot.aset_save(aupd("a:<b>:c"), actx) == bot.ASET_WAIT
        assert await bot.aset_save(aupd("tek:🧾 Fatura:Site ID|Fiş No"), actx) == bot.ConversationHandler.END
        assert [c["key"] for c in bot.screenshot_categories()] == ["tek"]
        assert bot.screenshot_categories()[0]["fields"] == ["Site ID", "Fiş No"]
        actx = SimpleNamespace(user_data={"aset_key": "screenshot_name_format", "wizard_lock": ("aset", time.time())}, bot=ctx.bot)
        assert await bot.aset_save(aupd("sabit_ad"), actx) == bot.ASET_WAIT
        assert await bot.aset_save(aupd("{kategori}-{alanlar}"), actx) == bot.ConversationHandler.END
        assert bot.screenshot_name_format() == "{kategori}-{alanlar}"
        # /temizle varsayilana dondurur
        actx = SimpleNamespace(user_data={"aset_key": "screenshot_categories", "wizard_lock": ("aset", time.time())}, bot=ctx.bot)
        assert await bot.aset_clear(aupd("/temizle"), actx) == bot.ConversationHandler.END
        assert len(bot.screenshot_categories()) == 5


def test_ekran_handler_sirasi():
    """Ekran goruntusu sihirbazi: sikayet sihirbazindan SONRA, /start akisindan,
    serbest foto handler'indan ve um_ callback'inden ONCE kayitli olmali."""
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "bot.py"), encoding="utf-8") as f:
        kaynak = f.read()
    ekran = kaynak.rindex("ss_photo_entry),")
    sikayet = kaynak.rindex('CommandHandler("sikayet", sikayet_start')
    start = kaynak.rindex('CommandHandler("start", start_cmd, filters.ChatType.PRIVATE),\n                      CallbackQueryHandler(profil_siteid_cb')
    foto = kaynak.rindex("(filters.PHOTO | filters.VIDEO), handle_campaign_media)")
    um = kaynak.rindex("CallbackQueryHandler(user_menu_cb, pattern=")
    assert sikayet < start < ekran, "ekran sihirbazi /start + Site ID akisindan SONRA olmali"
    assert ekran < um < foto
    # Akis acikken foto giris noktasina duser (allow_reentry); giris noktasi kaldigi adima yonlendirir
    assert "ss_photo_midway_field" not in kaynak[kaynak.index("SS_CAT:[CallbackQueryHandler(ss_cat_pick"):][:600]
    assert "apekran_" in kaynak[kaynak.rindex("CallbackQueryHandler(admin_cb, pattern="):][:400]
    # Komut menusu ve admin komutlari
    assert any(c.command == "ekran" for c in bot._parse_user_commands())
    assert bot._YARDIM_METINLERI.get("ekran")
    assert any(b.callback_data == "ap_ekran" for row in bot._panel_keyboard() for b in row)


def main():
    test_text_and_media_helpers()
    test_play_menu_uses_current_login_link()
    test_media_library_database()
    test_site_id_validation_and_password_hash()
    test_site_id_storage_and_exports()
    test_profile_view_registered_and_unregistered()
    test_rich_tokens_and_full_payload()
    test_scheduled_prenotify_and_retime()
    asyncio.run(test_campaign_media_variants())
    asyncio.run(test_unauthorized_admin_is_silent())
    asyncio.run(test_start_welcomes_unregistered_before_site_id())
    asyncio.run(test_start_default_flow_without_site_id())
    asyncio.run(test_site_id_requires_explicit_confirmation())
    asyncio.run(test_welcome_media_priority())
    test_reply_menu_build()
    test_campaigns_announcements_faq_db()
    test_main_menu_layout()
    test_auto_segments()
    asyncio.run(test_menu_button_text_routing())
    asyncio.run(test_smart_support_routing())
    asyncio.run(test_site_id_cancel_message())
    asyncio.run(test_play_menu_falls_back_to_commands_on_bad_link())
    asyncio.run(test_admin_password_toggle())
    asyncio.run(test_kitle_gonder_hata_raporu())
    asyncio.run(test_deliver_payload_uzun_caption_kisaltilir())
    asyncio.run(test_sihirbaz_girdileri_kaybolmuyor())
    test_v85_duzeltmeleri()
    asyncio.run(test_yerlesik_asistan())
    test_bilgi_bankasi_ve_yardim()
    test_admin_klavyesi()
    test_siteid_hedefleme()
    test_gonderim_gecmisi_ve_duyuru_duzenleme()
    asyncio.run(test_denetim_duzeltmeleri())
    test_handler_kayit_sirasi()
    test_ekran_dosya_adi_ve_kategoriler()
    asyncio.run(test_ekran_foto_once_akisi())
    asyncio.run(test_ekran_bilgi_once_akisi())
    asyncio.run(test_ekran_admin_secenekleri())
    test_ekran_handler_sirasi()
    print("Tum otomatik testler basarili.")


if __name__ == "__main__":
    main()
